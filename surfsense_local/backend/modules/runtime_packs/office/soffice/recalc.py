"""Recalculate a copy of a workbook with LibreOffice and read back its formulas' values."""

import tempfile
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from modules.runtime_packs.office.runtime import OfficeRuntime, require_runtime
from modules.runtime_packs.office.soffice.convert import convert


@dataclass(frozen=True)
class Recalculated:
    """Formula cells' values after LibreOffice recalculated a copy, by sheet and cell.

    `errors` holds cells LibreOffice computed to an error, such as #DIV/0!.
    """

    values: dict[str, dict[str, Any]]
    errors: dict[str, dict[str, str]]
    version: str


def recalc(
    xlsx: Path, *, deadline: float, runtime: OfficeRuntime | None = None
) -> Recalculated:
    """Every formula cell's value, as LibreOffice computes it on a copy.

    The workbook LibreOffice writes is read and thrown away, never delivered.
    """
    import openpyxl

    office = runtime or require_runtime()
    with tempfile.TemporaryDirectory() as out:
        computed = convert(xlsx, "xlsx", Path(out), deadline=deadline, runtime=office)
        formulas = _formula_cells(openpyxl.load_workbook(xlsx, data_only=False))
        values: dict[str, dict[str, Any]] = {}
        errors: dict[str, dict[str, str]] = {}
        book = openpyxl.load_workbook(computed, data_only=True)
        try:
            for sheet, cells in formulas.items():
                if sheet not in book.sheetnames:
                    continue
                for ref in cells:
                    cell = book[sheet][ref]
                    bucket = errors if cell.data_type == "e" else values
                    bucket.setdefault(sheet, {})[ref] = cell.value
        finally:
            book.close()
    return Recalculated(values=values, errors=errors, version=office.version)


def _formula_cells(book: Any) -> dict[str, list[str]]:
    try:
        return {
            sheet.title: list(dict.fromkeys(_computed_cells(sheet)))
            for sheet in book.worksheets
        }
    finally:
        book.close()


def _computed_cells(sheet: Any) -> Iterator[str]:
    """Each formula cell, and every cell an array formula fills, which holds no formula."""
    from openpyxl.worksheet.formula import ArrayFormula

    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type != "f":
                continue
            yield cell.coordinate
            if isinstance(cell.value, ArrayFormula):
                yield from _cells_in(cell.value.ref)


def _cells_in(ref: str) -> Iterator[str]:
    """The coordinates a range such as B1:C3 covers."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    first_column, first_row, last_column, last_row = range_boundaries(ref)
    for row in range(first_row, last_row + 1):
        for column in range(first_column, last_column + 1):
            yield f"{get_column_letter(column)}{row}"
