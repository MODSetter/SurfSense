"""A workbook as the cells a revision names: each sheet's used range, then each non-empty cell's address and value or formula.

Read from the file, not the indexed text: Docling's tables drop sheet names,
addresses and formulas, and the model then guessed them for set_cell.
"""

import json
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

VALUE_CHARS = 200
# Room for each sheet's heading and the notes after the cells.
_HEADING_BYTES = 120


class SheetNotFoundError(Exception):
    def __init__(self, sheet: str, sheets: list[str]) -> None:
        super().__init__(sheet)
        self.sheet = sheet
        self.sheets = sheets


@dataclass
class SheetMap:
    name: str
    used_range: str | None = None  # None for an empty sheet
    lines: list[str] = field(default_factory=list)
    left_out: int = 0  # non-empty cells past the cut


@dataclass(frozen=True)
class CellMap:
    sheets: list[SheetMap]
    cut_at: tuple[str, int] | None  # the sheet and row the first cell left out is on


def cell_map(path: Path, sheet: str | None, from_row: int, budget: int) -> CellMap:
    """Every sheet's cells from `from_row`, or only `sheet`'s, until `budget` bytes are spent.

    Read-only, with formulas as written: the cached values are only what Excel
    last calculated, and a revision writes formulas.
    """
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = book.worksheets
        if sheet is not None:
            sheets = [s for s in sheets if s.title == sheet]
            if not sheets:
                raise SheetNotFoundError(sheet, [s.title for s in book.worksheets])
        maps: list[SheetMap] = []
        cut_at: tuple[str, int] | None = None
        left = budget
        for worksheet in sheets:
            mapped = SheetMap(worksheet.title)
            left -= _HEADING_BYTES + len(worksheet.title.encode())
            corners: list[int] = []  # min row, min column, max row, max column
            # A file's stored dimension can be wrong, and read-only mode trusts it.
            worksheet.reset_dimensions()
            for row in worksheet.iter_rows():
                for cell in row:
                    shown = _shown(cell)
                    if shown is None:
                        continue
                    corners = _grown(corners, cell.row, cell.column)
                    if cell.row < from_row:
                        continue
                    line = f"{get_column_letter(cell.column)}{cell.row} = {shown}"
                    size = len(line.encode()) + 1
                    if cut_at is None and size <= left:
                        mapped.lines.append(line)
                        left -= size
                        continue
                    if cut_at is None:
                        cut_at = (worksheet.title, cell.row)
                    mapped.left_out += 1
            if corners:
                top, first, bottom, last = corners
                mapped.used_range = (
                    f"{get_column_letter(first)}{top}:{get_column_letter(last)}{bottom}"
                )
            maps.append(mapped)
        return CellMap(maps, cut_at)
    finally:
        book.close()


def _grown(corners: list[int], row: int, column: int) -> list[int]:
    if not corners:
        return [row, column, row, column]
    top, first, bottom, last = corners
    return [min(top, row), min(first, column), max(bottom, row), max(last, column)]


def _shown(cell: object) -> str | None:
    """The cell as a formula bar has it, text in quotes; None for an empty cell."""
    value = getattr(cell, "value", None)
    if value is None or value == "":
        return None
    if getattr(cell, "data_type", None) == "f":
        # An array formula is an object holding its text.
        return str(getattr(value, "text", None) or value)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, str):
        if getattr(cell, "data_type", None) == "e":
            return value  # an error such as #DIV/0!
        if len(value) > VALUE_CHARS:
            return (
                f"{json.dumps(value[:VALUE_CHARS], ensure_ascii=False)} "
                f"(cut, {len(value):,} characters)"
            )
        return json.dumps(value, ensure_ascii=False)
    return str(value)
