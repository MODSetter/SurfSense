"""Studio's Word and PDF take two paths: a small model writes Markdown that a
builder renders, a strong one writes a script the runner runs. Both keep a spec."""

from io import BytesIO
from typing import Any

import docx
import pytest
from sqlalchemy.orm import Session

from modules.artifacts.models import Artifact
from modules.artifacts.schemas import ArtifactRead, StudioJobCreate
from modules.artifacts.service import create_artifact_job, regenerate_artifact
from modules.documents.models import Document, DocumentStatus
from modules.llm.model_type import ModelType
from modules.llm.models import SelectedModel
from modules.llm.profile import Fingerprint, Tier
from modules.llm.resolution import ResolvedGeneration
from modules.workspaces.models import Workspace
from shared.config import get_storage_settings
from tests.integration.worker.conftest import stub_model  # noqa: F401
from worker.studio import run
from worker.studio.office.docx import docx as word
from worker.studio.office.pdf import pdf

pytestmark = [pytest.mark.integration, pytest.mark.usefixtures("stub_model")]

MARKDOWN = """\
# Halvorsen Freight proposal

We propose a **two-phase** rollout.

![Halvorsen logo](image:{figure})

| Phase | Cost |
|---|---|
| Pilot | 12,000 |
"""

SCRIPT = """\
# title: Halvorsen Freight proposal
import os
import docx
from docx.shared import Cm

document = docx.Document()
document.add_heading("Halvorsen Freight proposal", level=1)
document.add_picture(os.path.join(os.environ["IMAGES_DIR"], "{figure}.png"), width=Cm(4))
document.add_paragraph("We propose a two-phase rollout.")
document.save(os.environ["OUTPUT_PATH"])
"""

PDF_SCRIPT = """\
# title: Halvorsen pricing
import os
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

drawing = canvas.Canvas(os.environ["OUTPUT_PATH"], pagesize=A4)
drawing.drawString(72, 760, "Halvorsen Freight pricing")
drawing.showPage()
drawing.save()
"""


def _choose(session: Session, monkeypatch: pytest.MonkeyPatch, provider: str) -> None:
    """A chat model on this computer (llamacpp) or on a server (anything else)."""
    session.add(
        SelectedModel(model_type=ModelType.TEXT_GEN, provider="llamacpp", name="m")
    )
    session.commit()
    selection = type(
        "Selection",
        (),
        {
            "provider": provider,
            "name": "m",
            "tier": Tier.CAPABLE,
            "fingerprint": Fingerprint(provider, "m"),
        },
    )()
    monkeypatch.setattr(
        "worker.studio.job.resolve_generation",
        lambda _session: ResolvedGeneration(selection, None),
    )


def _model(monkeypatch: pytest.MonkeyPatch, *replies: str) -> list[dict[str, Any]]:
    """Answer each model call in turn, recording what it was sent."""
    calls: list[dict[str, Any]] = []

    def fake(_model: object, system: str, sources: object, **kwargs: Any) -> str:
        calls.append({"system": system, "sources": sources, **kwargs})
        return replies[min(len(calls), len(replies)) - 1]

    monkeypatch.setattr("worker.studio.shared.generate.run_model", fake)
    return calls


def _no_exec(monkeypatch: pytest.MonkeyPatch) -> None:
    def refuse(_code: str) -> dict:
        raise AssertionError("Word and PDF must not run code with exec()")

    monkeypatch.setattr("worker.studio.office.runner.execute", refuse)


def _job(
    session: Session, workspace: Workspace, source: Document, fmt: str
) -> Artifact:
    return create_artifact_job(
        session,
        workspace,
        StudioJobCreate(format=fmt, document_ids=[source.id], prompt="a proposal"),
    )


def _ready(session: Session, artifact: Artifact) -> None:
    session.expire_all()
    assert artifact.document.status is DocumentStatus.READY, (
        artifact.document.error_message
    )


def _primary(artifact: Artifact) -> bytes:
    (file,) = artifact.files
    return (get_storage_settings().data_dir / file.storage_key).read_bytes()


def test_a_small_model_writes_markdown_that_becomes_a_word_file_and_its_spec(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    source_figure: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Markdown with a source figure becomes a Word file, its spec and v1."""
    _choose(session, monkeypatch, "llamacpp")
    _no_exec(monkeypatch)
    markdown = MARKDOWN.format(figure=source_figure)
    calls = _model(monkeypatch, markdown)
    artifact = _job(session, workspace, logo_source, "docx")

    run(artifact.id)

    _ready(session, artifact)
    assert len(calls) == 1
    # The figure list names each figure the model may place, with its caption.
    assert f"image:{source_figure}" in calls[0]["system"]
    assert artifact.document.title == "Halvorsen Freight proposal"
    assert artifact.files[0].mime_type == word.mime
    opened = docx.Document(BytesIO(_primary(artifact)))
    assert len(opened.inline_shapes) == 1
    assert artifact.artifact_metadata["spec"] == {
        "kind": "markdown",
        "text": markdown.strip(),
        "format": "docx",
        "images": [source_figure],
    }
    assert artifact.artifact_metadata["version"] == {
        "root": artifact.id,
        "number": 1,
        "parent": None,
    }
    read = ArtifactRead.of(artifact)
    assert read.spec_kind == "markdown"
    assert read.version is not None and read.version.root_id == artifact.id
    assert "two-phase" in artifact.document.content


def test_a_small_model_writes_markdown_that_becomes_a_pdf(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The PDF builder renders the same kind of spec."""
    _choose(session, monkeypatch, "llamacpp")
    _no_exec(monkeypatch)
    _model(monkeypatch, "# Pricing\n\nThe pilot costs 12,000.\n")
    artifact = _job(session, workspace, logo_source, "pdf")

    run(artifact.id)

    _ready(session, artifact)
    assert artifact.files[0].mime_type == pdf.mime
    assert _primary(artifact).startswith(b"%PDF")
    assert artifact.artifact_metadata["spec"]["kind"] == "markdown"
    assert artifact.artifact_metadata["spec"]["format"] == "pdf"


def test_a_strong_model_writes_a_script_that_runs_in_the_runner_and_is_kept(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    source_figure: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The script runs in its own process with the figure it names, and is kept."""
    _choose(session, monkeypatch, "openai_compatible")
    _no_exec(monkeypatch)
    script = SCRIPT.format(figure=source_figure)
    calls = _model(monkeypatch, f"```python\n{script}```")
    artifact = _job(session, workspace, logo_source, "docx")

    run(artifact.id)

    _ready(session, artifact)
    assert len(calls) == 1
    assert "OUTPUT_PATH" in calls[0]["system"]
    assert source_figure in calls[0]["system"]
    assert artifact.document.title == "Halvorsen Freight proposal"
    opened = docx.Document(BytesIO(_primary(artifact)))
    assert len(opened.inline_shapes) == 1
    assert artifact.artifact_metadata["spec"] == {
        "kind": "python",
        "text": script.strip(),
        "format": "docx",
        "images": [source_figure],
    }
    assert artifact.artifact_metadata["version"]["number"] == 1
    assert ArtifactRead.of(artifact).spec_kind == "python"


def test_a_strong_models_pdf_script_runs_in_the_runner(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A ReportLab script is titled by its first comment and indexed by its text."""
    _choose(session, monkeypatch, "openai_compatible")
    _no_exec(monkeypatch)
    _model(monkeypatch, PDF_SCRIPT)
    artifact = _job(session, workspace, logo_source, "pdf")

    run(artifact.id)

    _ready(session, artifact)
    assert artifact.document.title == "Halvorsen pricing"
    assert _primary(artifact).startswith(b"%PDF")
    assert artifact.document.content == "Halvorsen Freight pricing"
    assert artifact.artifact_metadata["spec"]["kind"] == "python"


def test_a_failing_script_is_shown_its_error_and_the_fix_is_kept(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The model sees what failed, and the script that ran is the spec."""
    _choose(session, monkeypatch, "openai_compatible")
    calls = _model(monkeypatch, "raise ValueError('no pricing table')", PDF_SCRIPT)
    artifact = _job(session, workspace, logo_source, "pdf")

    run(artifact.id)

    _ready(session, artifact)
    assert len(calls) == 2
    repair = calls[1]["repair"]
    assert "no pricing table" in repair.instruction
    assert artifact.artifact_metadata["spec"]["text"] == PDF_SCRIPT.strip()


def test_a_script_that_keeps_failing_fails_with_a_reason_studio_can_retry(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Not a "Script error:": Retry asks the model again, which is the fix here."""
    _choose(session, monkeypatch, "openai_compatible")
    calls = _model(monkeypatch, "raise ValueError('still broken')")
    artifact = _job(session, workspace, logo_source, "docx")

    with pytest.raises(RuntimeError):
        run(artifact.id)

    session.expire_all()
    assert len(calls) == 3
    assert artifact.document.status is DocumentStatus.FAILED
    reason = artifact.document.error_message or ""
    assert "still broken" in reason
    assert not reason.startswith("Script error: ")
    assert "spec" not in artifact.artifact_metadata


def test_regenerate_drafts_a_studio_document_again(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Studio's draft is made from the sources, so a regenerate asks the model again."""
    _choose(session, monkeypatch, "llamacpp")
    _model(monkeypatch, "# First draft\n", "# Second draft\n")
    artifact = _job(session, workspace, logo_source, "docx")
    run(artifact.id)

    regenerate_artifact(session, artifact)
    run(artifact.id)

    _ready(session, artifact)
    assert artifact.document.title == "Second draft"
    assert artifact.artifact_metadata["spec"]["text"] == "# Second draft"
    assert artifact.artifact_metadata["version"]["number"] == 1
