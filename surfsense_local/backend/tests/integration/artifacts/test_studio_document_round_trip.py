"""Studio's Word and PDF over the routes the panel calls: draft v1, Refine it into v2.

A scripted model stands in for the selected one, so each path runs end to end:
a local model's Markdown through the builders, a remote model's script through
the runner.
"""

from io import BytesIO
from typing import Any

import docx
import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_c
import pytest
from httpx import AsyncClient
from huey import Huey
from sqlalchemy.orm import Session

from modules.artifacts.models import Artifact
from modules.documents.models import Document
from modules.llm.model_type import ModelType
from modules.llm.models import SelectedModel
from modules.llm.profile import Fingerprint, Tier
from modules.llm.providers.types import Message
from modules.llm.resolution import ResolvedGeneration
from modules.workspaces.models import Workspace
from shared.config import get_storage_settings
from shared.queue import studio_queue
from tests.integration.worker.conftest import stub_model  # noqa: F401

pytestmark = [pytest.mark.integration, pytest.mark.usefixtures("stub_model")]

LOCAL = "llamacpp"
REMOTE = "openai_compatible"

CHART = (
    '{"type": "bar", "title": "Cost by phase", "labels": ["Pilot", "Rollout"], '
    '"series": [{"name": "Cost", "values": [12000, 30000]}]}'
)

MARKDOWN = """\
# Halvorsen Freight proposal

We propose a two-phase rollout.

![Halvorsen logo](image:FIGURE)

| Phase | Cost |
|---|---|
| Pilot | 12,000 |
| Rollout | 30,000 |

```chart
CHART
```
"""


def _markdown(figure: str) -> str:
    """The Markdown a local model drafts: a source figure, a table and a chart."""
    return MARKDOWN.replace("FIGURE", figure).replace("CHART", CHART)


WORD_SCRIPT = """\
# title: Halvorsen Freight proposal
import os
import docx
from docx.shared import Cm

document = docx.Document()
document.add_heading("Halvorsen Freight proposal", level=1)
document.add_picture(os.path.join(os.environ["IMAGES_DIR"], "{figure}.png"), width=Cm(4))
table = document.add_table(rows=2, cols=2)
for row, cells in enumerate([("Phase", "Cost"), ("Pilot", "{cost}")]):
    for column, text in enumerate(cells):
        table.cell(row, column).text = text
document.save(os.environ["OUTPUT_PATH"])
"""

PDF_SCRIPT = """\
# title: Halvorsen pricing
import os
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

drawing = canvas.Canvas(os.environ["OUTPUT_PATH"], pagesize=A4)
drawing.drawString(72, 760, "{line}")
drawing.drawImage(os.path.join(os.environ["IMAGES_DIR"], "{figure}.png"), 72, 600)
drawing.showPage()
drawing.save()
"""


def _choose(session: Session, monkeypatch: pytest.MonkeyPatch, provider: str) -> None:
    """The chat model the worker resolves: on llama.cpp, or on a server.

    The row only makes the format available; the worker's resolution is what
    the strength rule reads.
    """
    session.add(SelectedModel(model_type=ModelType.TEXT_GEN, provider=LOCAL, name="m"))
    session.commit()
    selection = type(
        "Selection",
        (),
        {
            "provider": provider,
            "name": "m",
            "tier": Tier.CAPABLE,
            "fingerprint": Fingerprint(provider, "m"),
        },
    )()
    monkeypatch.setattr(
        "worker.studio.job.resolve_generation",
        lambda _session: ResolvedGeneration(selection, None),
    )

    async def window(_session: Session) -> tuple[str, int]:
        return "m", 32_768

    monkeypatch.setattr(
        "modules.artifacts.studio_documents.fits.selected_model_window", window
    )


def _scripted_model(
    monkeypatch: pytest.MonkeyPatch, *replies: str
) -> list[list[Message]]:
    """Answer each model call with the next reply, recording what it was sent."""
    calls: list[list[Message]] = []

    def complete(_model: object, messages: list[Message], **_kwargs: Any) -> str:
        calls.append(messages)
        return replies[len(calls) - 1]

    monkeypatch.setattr("worker.studio.shared.generate.complete", complete)
    return calls


def _work_off(queue: Huey) -> None:
    while (task := queue.dequeue()) is not None:
        queue.execute(task)


async def _draft(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    source: Document,
    fmt: str,
) -> dict[str, Any]:
    workspace_id, source_id = workspace.id, source.id
    _end_read(session)
    created = await client.post(
        f"/workspaces/{workspace_id}/studio/jobs",
        json={"format": fmt, "document_ids": [source_id], "prompt": "a proposal"},
    )
    assert created.status_code == 201, created.text
    _work_off(studio_queue)
    return await _listed(client, workspace_id, created.json()["id"])


async def _refine(
    client: AsyncClient, session: Session, workspace: Workspace, artifact_id: int
) -> dict[str, Any]:
    workspace_id = workspace.id
    _end_read(session)
    response = await client.post(
        f"/artifacts/{artifact_id}/refine", json={"instruction": "Raise the pilot cost"}
    )
    assert response.status_code == 202, response.text
    assert response.json()["status"] == "pending"
    _work_off(studio_queue)
    return await _listed(client, workspace_id, response.json()["id"])


async def _listed(
    client: AsyncClient, workspace_id: int, artifact_id: int
) -> dict[str, Any]:
    """The artifact as the Studio list shows it, which is what the panel reads."""
    listed = (await client.get(f"/workspaces/{workspace_id}/artifacts")).json()
    (found,) = [artifact for artifact in listed if artifact["id"] == artifact_id]
    assert found["status"] == "ready", found["error_message"]
    return found


def _end_read(session: Session) -> None:
    """Every transaction takes SQLite's one write lock, so the test's must end first."""
    session.commit()


def _primary(session: Session, artifact_id: int) -> bytes:
    session.expire_all()
    (file,) = session.get(Artifact, artifact_id).files
    path = get_storage_settings().data_dir / file.storage_key
    _end_read(session)
    return path.read_bytes()


def _spec(session: Session, artifact_id: int) -> dict[str, Any]:
    session.expire_all()
    spec = session.get(Artifact, artifact_id).artifact_metadata["spec"]
    _end_read(session)
    return spec


def _pdf_text_and_images(data: bytes) -> tuple[str, int]:
    pdf = pdfium.PdfDocument(data)
    try:
        text = ""
        images = 0
        for page in pdf:
            text += page.get_textpage().get_text_range()
            images += len(list(page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE])))
        return text, images
    finally:
        pdf.close()


async def test_a_local_model_drafts_word_as_markdown_and_refines_it(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    source_figure: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A table, a source figure and a chart in v1; v2 is the rewrite, on the same root."""
    _choose(session, monkeypatch, LOCAL)
    first = _markdown(source_figure)
    revised = first.replace("12,000", "15,000")
    calls = _scripted_model(monkeypatch, first, f"```markdown\n{revised}```")

    v1 = await _draft(client, session, workspace, logo_source, "docx")

    assert v1["spec_kind"] == "markdown"
    assert v1["version"] == {"root_id": v1["id"], "number": 1, "parent_id": None}
    assert v1["title"] == "Halvorsen Freight proposal"
    assert _spec(session, v1["id"])["images"] == [source_figure]
    word = docx.Document(BytesIO(_primary(session, v1["id"])))
    (table,) = word.tables
    assert [cell.text for cell in table.rows[1].cells] == ["Pilot", "12,000"]
    # The source figure and the drawn chart.
    assert len(word.inline_shapes) == 2

    v2 = await _refine(client, session, workspace, v1["id"])

    assert v2["spec_kind"] == "markdown"
    assert v2["version"] == {"root_id": v1["id"], "number": 2, "parent_id": v1["id"]}
    assert _spec(session, v2["id"])["text"] == revised.strip()
    revised_word = docx.Document(BytesIO(_primary(session, v2["id"])))
    assert revised_word.tables[0].rows[1].cells[1].text == "15,000"
    assert len(revised_word.inline_shapes) == 2
    # The refine read v1's whole Markdown, chart block and all.
    assert CHART in calls[1][-1].content
    assert _spec(session, v1["id"])["text"] == first.strip()


async def test_a_remote_model_drafts_word_as_a_script_and_refines_it(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    source_figure: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The script runs in the runner with the figure it names; v2 runs its rewrite."""
    _choose(session, monkeypatch, REMOTE)

    def refuse(_code: str) -> dict:
        raise AssertionError("Word must not run code with exec()")

    monkeypatch.setattr("worker.studio.office.runner.execute", refuse)
    first = WORD_SCRIPT.format(figure=source_figure, cost="12,000")
    revised = WORD_SCRIPT.format(figure=source_figure, cost="15,000")
    calls = _scripted_model(monkeypatch, first, f"```python\n{revised}```")

    v1 = await _draft(client, session, workspace, logo_source, "docx")

    assert v1["spec_kind"] == "python"
    assert v1["version"] == {"root_id": v1["id"], "number": 1, "parent_id": None}
    assert _spec(session, v1["id"])["images"] == [source_figure]
    word = docx.Document(BytesIO(_primary(session, v1["id"])))
    assert word.tables[0].rows[1].cells[1].text == "12,000"
    assert len(word.inline_shapes) == 1

    v2 = await _refine(client, session, workspace, v1["id"])

    assert v2["spec_kind"] == "python"
    assert v2["version"] == {"root_id": v1["id"], "number": 2, "parent_id": v1["id"]}
    assert _spec(session, v2["id"])["text"] == revised.strip()
    revised_word = docx.Document(BytesIO(_primary(session, v2["id"])))
    assert revised_word.tables[0].rows[1].cells[1].text == "15,000"
    assert len(revised_word.inline_shapes) == 1
    assert first.strip() in calls[1][-1].content


@pytest.mark.parametrize("provider", [LOCAL, REMOTE])
async def test_a_pdf_of_each_kind_renders_and_refines(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    source_figure: str,
    monkeypatch: pytest.MonkeyPatch,
    provider: str,
) -> None:
    """Markdown through the ReportLab builder, or a ReportLab script in the runner."""
    _choose(session, monkeypatch, provider)
    if provider == LOCAL:
        first = _markdown(source_figure)
        revised = first.replace("12,000", "15,000")
        figures = 2  # the source figure and the chart
    else:
        first = PDF_SCRIPT.format(figure=source_figure, line="Pilot 12,000")
        revised = PDF_SCRIPT.format(figure=source_figure, line="Pilot 15,000")
        figures = 1
    _scripted_model(monkeypatch, first, revised)

    v1 = await _draft(client, session, workspace, logo_source, "pdf")

    assert v1["spec_kind"] == ("markdown" if provider == LOCAL else "python")
    data = _primary(session, v1["id"])
    assert data.startswith(b"%PDF")
    text, images = _pdf_text_and_images(data)
    assert "12,000" in text
    assert images == figures

    v2 = await _refine(client, session, workspace, v1["id"])

    assert v2["version"] == {"root_id": v1["id"], "number": 2, "parent_id": v1["id"]}
    assert v2["spec_kind"] == v1["spec_kind"]
    text, images = _pdf_text_and_images(_primary(session, v2["id"]))
    assert "15,000" in text
    assert images == figures
