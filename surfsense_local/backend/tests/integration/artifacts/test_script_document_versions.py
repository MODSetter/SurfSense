"""A document made with a source's image, rendered, then edited into its next version.

Everything but the embedder is real: the figures pass keeps the image, the
service stores each script, and the queued job runs it through the runner.
"""

from io import BytesIO

import docx
import pytest
from huey import Huey
from PIL import Image
from sqlalchemy.orm import Session

from modules.artifacts.models import Artifact
from modules.artifacts.script_documents.service import create_script_document
from modules.documents.models import Document, DocumentStatus
from modules.documents.source_figures import FiguresPending, list_figures
from modules.workspaces.models import Workspace
from shared.config import get_storage_settings
from shared.queue import ingest_queue, studio_queue
from tests.integration.worker.conftest import stub_model  # noqa: F401

# The job indexes what the script wrote; the stub stands in for the embedder.
pytestmark = [pytest.mark.integration, pytest.mark.usefixtures("stub_model")]

PROPOSAL = """\
import os
import docx
from docx.shared import Cm

document = docx.Document()
document.add_picture(os.path.join(os.environ["IMAGES_DIR"], "{logo}.png"), width=Cm(4))
document.add_heading("{heading}", level=1)
document.add_paragraph("We propose a two-phase rollout for Halvorsen Freight.")
document.save(os.environ["OUTPUT_PATH"])
"""


def _work_off(queue: Huey) -> None:
    """Run what the queue holds, as the worker's consumer would."""
    while (task := queue.dequeue()) is not None:
        queue.execute(task)


def _word_file(artifact: Artifact) -> docx.document.Document:
    (file,) = artifact.files
    path = get_storage_settings().data_dir / file.storage_key
    return docx.Document(BytesIO(path.read_bytes()))


def _pictures(document: docx.document.Document) -> list[tuple[int, int]]:
    """The pixel size of each picture the Word file holds."""
    return [
        Image.open(BytesIO(part.blob)).size
        for part in document.part.package.image_parts
    ]


def test_a_document_made_with_a_source_image_is_edited_into_its_second_version(
    session: Session, workspace: Workspace, logo_source: Document
) -> None:
    """v1 places the source's logo; v2 is made from v1 and leaves v1 as it was."""
    # The logo was ingested before figures were kept: asking queues the pass.
    with pytest.raises(FiguresPending):
        list_figures(session, workspace.id, logo_source.id)
    _work_off(ingest_queue)
    (logo,) = list_figures(session, workspace.id, logo_source.id)

    first = create_script_document(
        session,
        workspace,
        title="Client proposal",
        format="docx",
        script=PROPOSAL.format(logo=logo.name, heading="Client proposal"),
        base_artifact_id=None,
        image_names=[logo.name],
    )
    _work_off(studio_queue)
    session.expire_all()
    assert first.document.status is DocumentStatus.READY, first.document.error_message

    second = create_script_document(
        session,
        workspace,
        title="Client proposal",
        format="docx",
        script=PROPOSAL.format(logo=logo.name, heading="Revised proposal"),
        base_artifact_id=first.id,
        image_names=[logo.name],
    )
    _work_off(studio_queue)

    session.expire_all()
    assert second.document.status is DocumentStatus.READY, second.document.error_message
    assert second.artifact_metadata["version"] == {
        "root": first.id,
        "number": 2,
        "parent": first.id,
    }
    assert second.artifact_metadata["source_document_ids"] == [logo_source.id]
    assert first.document.content.startswith("# Client proposal\n")
    assert second.document.content.startswith("# Revised proposal\n")
    for version in (first, second):
        assert _pictures(_word_file(version)) == [(logo.width, logo.height)]
