"""Refine: one model call rewrites a version's whole spec into the next version."""

from io import BytesIO
from typing import Any

import docx
import pytest
from httpx import AsyncClient, Response
from huey import Huey
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from modules.artifacts.models import Artifact
from modules.artifacts.schemas import StudioJobCreate
from modules.artifacts.script_documents.service import create_script_document
from modules.artifacts.service import create_artifact_job
from modules.documents.models import Document, DocumentStatus
from modules.documents.source_figures.layout import figures_dir, write_index
from modules.llm.model_type import ModelType
from modules.llm.models import SelectedModel
from modules.llm.profile import Fingerprint, Tier
from modules.llm.providers.types import Message
from modules.llm.resolution import ResolvedGeneration
from modules.workspaces.models import Workspace
from shared.config import get_storage_settings
from shared.queue import studio_queue
from tests.integration.worker.conftest import stub_model  # noqa: F401
from worker.studio import run

pytestmark = [pytest.mark.integration, pytest.mark.usefixtures("stub_model")]

DRAFT = (
    "# Halvorsen Freight proposal\n\nWe propose a two-phase rollout over six months."
)
REVISED = "# Halvorsen Freight proposal\n\nA two-phase rollout."

SCRIPT = """\
# title: Client proposal
import os
import docx

document = docx.Document()
document.add_heading("Client proposal", level=1)
document.add_paragraph("{line}")
document.save(os.environ["OUTPUT_PATH"])
"""


@pytest.fixture
def local_model(session: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    """A model on llama.cpp, whose window the refine route reads as 32k tokens."""
    session.add(
        SelectedModel(
            model_type=ModelType.TEXT_GEN, provider="llamacpp", name="Qwen3-4B"
        )
    )
    session.commit()
    _window(monkeypatch, 32_768)
    _resolve(monkeypatch, "llamacpp")


def _resolve(monkeypatch: pytest.MonkeyPatch, provider: str) -> None:
    """The model the worker resolves; one on a server drafts Word and PDF as a script."""
    selection = type(
        "Selection",
        (),
        {
            "provider": provider,
            "name": "Qwen3-4B",
            "tier": Tier.CAPABLE,
            "fingerprint": Fingerprint(provider, "Qwen3-4B"),
        },
    )()
    monkeypatch.setattr(
        "worker.studio.job.resolve_generation",
        lambda _session: ResolvedGeneration(selection, None),
    )


def _window(monkeypatch: pytest.MonkeyPatch, tokens: int) -> None:
    async def window(_session: Session) -> tuple[str, int]:
        return "Qwen3-4B", tokens

    monkeypatch.setattr(
        "modules.artifacts.studio_documents.fits.selected_model_window", window
    )


def _draft_reply(monkeypatch: pytest.MonkeyPatch, reply: str) -> None:
    monkeypatch.setattr(
        "worker.studio.shared.generate.run_model", lambda *_a, **_k: reply
    )


def _refine_reply(monkeypatch: pytest.MonkeyPatch, reply: str) -> list[list[Message]]:
    calls: list[list[Message]] = []

    def fake(_model: object, messages: list[Message], **_kwargs: Any) -> str:
        calls.append(messages)
        return reply

    monkeypatch.setattr("worker.studio.shared.generate.complete", fake)
    return calls


def _work_off(queue: Huey) -> None:
    while (task := queue.dequeue()) is not None:
        queue.execute(task)


async def _refine(
    client: AsyncClient, session: Session, artifact_id: int, instruction: str
) -> Response:
    """Ask over HTTP with the test's own read transaction ended: SQLite has one writer."""
    session.commit()
    return await client.post(
        f"/artifacts/{artifact_id}/refine", json={"instruction": instruction}
    )


def _markdown_v1(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    monkeypatch: pytest.MonkeyPatch,
) -> Artifact:
    return _studio_v1(session, workspace, logo_source, monkeypatch, DRAFT)


def _script_v1(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    monkeypatch: pytest.MonkeyPatch,
) -> Artifact:
    """Studio's v1 drafted as a script by a model on a server."""
    _resolve(monkeypatch, "openai_compatible")
    reply = f"```python\n{SCRIPT.format(line='Six months.')}```"
    return _studio_v1(session, workspace, logo_source, monkeypatch, reply)


def _studio_v1(
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    monkeypatch: pytest.MonkeyPatch,
    reply: str,
) -> Artifact:
    _draft_reply(monkeypatch, reply)
    artifact = create_artifact_job(
        session,
        workspace,
        StudioJobCreate(format="docx", document_ids=[logo_source.id]),
    )
    _work_off(studio_queue)
    session.expire_all()
    assert artifact.document.status is DocumentStatus.READY, (
        artifact.document.error_message
    )
    return artifact


def _agent_script(session: Session, workspace: Workspace) -> Artifact:
    """A Word document the agent wrote as a script in chat."""
    artifact = create_script_document(
        session,
        workspace,
        title="Client proposal",
        format="docx",
        script=SCRIPT.format(line="Six months."),
        base_artifact_id=None,
        image_names=[],
    )
    _work_off(studio_queue)
    session.expire_all()
    assert artifact.document.status is DocumentStatus.READY
    return artifact


def _primary(artifact: Artifact) -> bytes:
    (file,) = artifact.files
    return (get_storage_settings().data_dir / file.storage_key).read_bytes()


def _count(session: Session) -> int:
    return session.scalar(select(func.count()).select_from(Artifact))


async def test_refining_markdown_makes_the_next_version_and_leaves_the_first_alone(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The rewritten Markdown is v2's spec and file; v1 keeps its own."""
    first = _markdown_v1(session, workspace, logo_source, monkeypatch)
    first_bytes = _primary(first)
    calls = _refine_reply(monkeypatch, f"```markdown\n{REVISED}\n```")

    response = await _refine(client, session, first.id, "Make it shorter")

    assert response.status_code == 202, response.text
    body = response.json()
    assert body["status"] == "pending"
    assert body["spec_kind"] == "markdown"
    assert body["version"] == {"root_id": first.id, "number": 2, "parent_id": first.id}
    assert body["title"] == "Halvorsen Freight proposal"

    _work_off(studio_queue)

    session.expire_all()
    second = session.get(Artifact, body["id"])
    assert second.document.status is DocumentStatus.READY, second.document.error_message
    assert second.artifact_metadata["spec"]["text"] == REVISED
    assert second.artifact_metadata["spec"]["kind"] == "markdown"
    assert (
        "A two-phase rollout."
        in docx.Document(BytesIO(_primary(second))).paragraphs[1].text
    )
    # One call, holding the whole current Markdown and the instruction.
    (messages,) = calls
    assert DRAFT in messages[-1].content
    assert "Make it shorter" in messages[-1].content
    assert first.artifact_metadata["spec"]["text"] == DRAFT
    assert _primary(first) == first_bytes


async def test_refining_a_script_runs_the_rewritten_script_in_the_runner(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The rewritten script runs as v2 through the runner; v1 keeps its script and file."""
    first = _script_v1(session, workspace, logo_source, monkeypatch)
    first_script = first.artifact_metadata["spec"]["text"]
    revised = SCRIPT.format(line="Three months.")
    _refine_reply(monkeypatch, f"```python\n{revised}```")

    response = await _refine(client, session, first.id, "Halve the timeline")
    assert response.status_code == 202, response.text
    assert response.json()["spec_kind"] == "python"
    _work_off(studio_queue)

    session.expire_all()
    second = session.get(Artifact, response.json()["id"])
    assert second.document.status is DocumentStatus.READY, second.document.error_message
    assert second.artifact_metadata["spec"]["text"] == revised.strip()
    assert second.artifact_metadata["version"]["number"] == 2
    assert (
        docx.Document(BytesIO(_primary(second))).paragraphs[1].text == "Three months."
    )
    assert first.artifact_metadata["spec"]["text"] == first_script
    assert docx.Document(BytesIO(_primary(first))).paragraphs[1].text == "Six months."


async def test_a_failed_refine_leaves_earlier_versions_untouched(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A rewrite that fails fails only its own version, which keeps no spec."""
    first = _script_v1(session, workspace, logo_source, monkeypatch)
    first_script = first.artifact_metadata["spec"]["text"]
    _refine_reply(monkeypatch, "raise ValueError('lost the table')")

    response = await _refine(client, session, first.id, "Add a table")
    _work_off(studio_queue)

    session.expire_all()
    second = session.get(Artifact, response.json()["id"])
    assert second.document.status is DocumentStatus.FAILED
    assert "lost the table" in (second.document.error_message or "")
    assert "spec" not in second.artifact_metadata
    assert first.document.status is DocumentStatus.READY
    assert first.artifact_metadata["spec"]["text"] == first_script


async def test_the_agents_documents_are_refined_in_its_chat_not_here(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    local_model: None,
) -> None:
    """Decision 8: the agent edits its own documents, so Refine refuses them and says so."""
    script = _agent_script(session, workspace)
    before = _count(session)

    response = await _refine(client, session, script.id, "Halve the timeline")

    assert response.status_code == 409
    assert "chat" in response.json()["detail"]
    assert _count(session) == before
    assert studio_queue.pending_count() == 0


async def test_only_a_ready_studio_word_or_pdf_version_is_listed_as_refinable(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The panel offers Refine from this flag: a ready Studio draft or refine, nothing else."""
    first = _markdown_v1(session, workspace, logo_source, monkeypatch)
    _draft_reply(monkeypatch, "# Summary\n\nRings.")
    summary = create_artifact_job(
        session,
        workspace,
        StudioJobCreate(format="summary", document_ids=[logo_source.id]),
    )
    _work_off(studio_queue)
    script = _agent_script(session, workspace)
    first_id, summary_id, script_id = first.id, summary.id, script.id
    workspace_id = workspace.id
    _refine_reply(monkeypatch, f"```markdown\n{REVISED}\n```")
    pending = (await _refine(client, session, first_id, "Shorter")).json()

    listed = await client.get(f"/workspaces/{workspace_id}/artifacts")
    refinable = {row["id"]: row["refinable"] for row in listed.json()}

    assert pending["refinable"] is False
    assert refinable == {
        first_id: True,
        pending["id"]: False,
        summary_id: False,
        script_id: False,
    }
    _work_off(studio_queue)
    ready = await client.get(f"/artifacts/{pending['id']}")
    assert ready.json()["refinable"] is True


async def test_a_document_longer_than_the_models_window_is_refused_with_a_reason(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Nothing is created or queued when the rewrite cannot fit the window."""
    first = _markdown_v1(session, workspace, logo_source, monkeypatch)
    _window(monkeypatch, 1_000)
    before = _count(session)

    response = await _refine(client, session, first.id, "Make it shorter")

    assert response.status_code == 409
    assert "too long for Qwen3-4B" in response.json()["detail"]
    assert _count(session) == before
    assert studio_queue.pending_count() == 0


async def test_an_artifact_without_a_spec_cannot_be_refined(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Only a Word document or PDF keeps a spec to rewrite."""
    _draft_reply(monkeypatch, "# Summary\n\nRings.")
    summary = create_artifact_job(
        session,
        workspace,
        StudioJobCreate(format="summary", document_ids=[logo_source.id]),
    )
    _work_off(studio_queue)

    response = await _refine(client, session, summary.id, "Shorter")

    assert response.status_code == 409
    assert "Word document or PDF" in response.json()["detail"]


async def test_a_version_still_being_made_cannot_be_refined(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A pending version has no spec yet to rewrite."""
    first = _markdown_v1(session, workspace, logo_source, monkeypatch)
    pending = await _refine(client, session, first.id, "Shorter")

    response = await _refine(client, session, pending.json()["id"], "Shorter still")

    assert response.status_code == 409


@pytest.mark.parametrize("instruction", ["", "   ", "x" * 2001])
async def test_an_instruction_must_say_something_and_fit(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
    instruction: str,
) -> None:
    """An empty or overlong instruction is refused before anything is read."""
    first = _markdown_v1(session, workspace, logo_source, monkeypatch)

    response = await _refine(client, session, first.id, instruction)

    assert response.status_code == 422


async def test_refining_an_unknown_artifact_is_not_found(
    client: AsyncClient, local_model: None
) -> None:
    """An id no artifact has is a 404."""
    response = await client.post(
        "/artifacts/999/refine", json={"instruction": "Shorter"}
    )

    assert response.status_code == 404


async def test_the_figures_listed_in_the_prompt_count_against_the_window(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A source with many captioned figures lengthens the prompt the model reads."""
    first = _markdown_v1(session, workspace, logo_source, monkeypatch)
    folder = figures_dir(
        get_storage_settings().document_dir(workspace.id, logo_source.id)
    )
    folder.mkdir(parents=True, exist_ok=True)
    caption = "Figure: freight volume by depot and quarter, " * 5
    entries = [
        {"n": n, "width": 40, "height": 20, "page": 1, "caption": caption}
        for n in range(1, 101)
    ]
    write_index(folder, entries, None)
    # Room for the short spec and the prompt, not for 100 figures' lines.
    _window(monkeypatch, 4_000)
    before = _count(session)

    response = await _refine(client, session, first.id, "Make it shorter")

    assert response.status_code == 409, response.text
    assert "too long for Qwen3-4B" in response.json()["detail"]
    assert _count(session) == before


def _refine_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    def unreachable(*_args: Any, **_kwargs: Any) -> str:
        raise RuntimeError("The model could not be reached")

    monkeypatch.setattr("worker.studio.shared.generate.complete", unreachable)


async def _refined_then_failed(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    monkeypatch: pytest.MonkeyPatch,
) -> int:
    """v2, a refine whose rewrite failed, so Retry is what the user reaches for."""
    first = _markdown_v1(session, workspace, logo_source, monkeypatch)
    _refine_fails(monkeypatch)
    response = await _refine(client, session, first.id, "Make it shorter")
    _work_off(studio_queue)
    second_id = response.json()["id"]
    session.expire_all()
    assert session.get(Artifact, second_id).document.status is DocumentStatus.FAILED
    session.commit()
    return second_id


async def test_regenerating_a_refined_version_is_refused_when_it_no_longer_fits(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Retry asks for the same rewrite, so the model selected now must hold it too."""
    second_id = await _refined_then_failed(
        client, session, workspace, logo_source, monkeypatch
    )
    _window(monkeypatch, 1_000)

    retried = await client.post(f"/artifacts/{second_id}/regenerate")

    assert retried.status_code == 409
    assert "too long for Qwen3-4B" in retried.json()["detail"]
    session.expire_all()
    assert session.get(Artifact, second_id).document.status is DocumentStatus.FAILED
    assert studio_queue.pending_count() == 0


async def test_regenerating_a_refined_version_that_fits_asks_again(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Within the window, Retry queues the same rewrite."""
    second_id = await _refined_then_failed(
        client, session, workspace, logo_source, monkeypatch
    )
    _refine_reply(monkeypatch, f"```markdown\n{REVISED}\n```")

    retried = await client.post(f"/artifacts/{second_id}/regenerate")
    _work_off(studio_queue)

    assert retried.status_code == 202, retried.text
    session.expire_all()
    second = session.get(Artifact, second_id)
    assert second.document.status is DocumentStatus.READY, second.document.error_message
    assert second.artifact_metadata["spec"]["text"] == REVISED


async def test_regenerating_a_draft_is_not_priced_against_the_window(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A v1's Retry drafts from the sources again; only a refine's rewrite must fit."""
    first = _markdown_v1(session, workspace, logo_source, monkeypatch)
    _window(monkeypatch, 1_000)
    session.commit()

    retried = await client.post(f"/artifacts/{first.id}/regenerate")
    _work_off(studio_queue)

    assert retried.status_code == 202, retried.text
    session.expire_all()
    assert first.document.status is DocumentStatus.READY


@pytest.mark.parametrize("ending", ["failed", "cancelled"])
async def test_a_version_whose_retry_did_not_finish_is_not_refined(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    logo_source: Document,
    local_model: None,
    monkeypatch: pytest.MonkeyPatch,
    ending: str,
) -> None:
    """It still keeps its last spec, but the panel offers no Refine, so the route agrees."""
    first = _markdown_v1(session, workspace, logo_source, monkeypatch)
    first_id = first.id
    session.commit()

    def unreachable(*_args: Any, **_kwargs: Any) -> str:
        raise RuntimeError("The model could not be reached")

    monkeypatch.setattr("worker.studio.shared.generate.run_model", unreachable)
    assert (await client.post(f"/artifacts/{first_id}/regenerate")).status_code == 202
    if ending == "cancelled":
        assert (await client.post(f"/artifacts/{first_id}/cancel")).status_code == 200
    else:
        studio_queue.dequeue()
        with pytest.raises(RuntimeError):
            run(first_id)
    session.expire_all()
    assert first.document.status is DocumentStatus(ending)
    assert "spec" in first.artifact_metadata
    before = _count(session)
    session.commit()

    shown = await client.get(f"/artifacts/{first_id}")
    response = await _refine(client, session, first_id, "Make it shorter")

    assert shown.json()["refinable"] is False
    assert response.status_code == 409, response.text
    assert "did not finish" in response.json()["detail"]
    assert _count(session) == before


DECK = """\
import os
from pptx import Presentation

deck = Presentation()
deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = "Rollout"
deck.save(os.environ["OUTPUT_PATH"])
"""

WORKBOOK = """\
import os
from openpyxl import Workbook

book = Workbook()
book.active.append(["Phase", "Months"])
book.save(os.environ["OUTPUT_PATH"])
"""


@pytest.mark.parametrize(("format", "script"), [("pptx", DECK), ("xlsx", WORKBOOK)])
async def test_the_agents_decks_and_workbooks_are_not_refinable_either(
    client: AsyncClient,
    session: Session,
    workspace: Workspace,
    local_model: None,
    format: str,
    script: str,
) -> None:
    """They keep a python spec as its Word documents do, and are edited in its chat too."""
    made = create_script_document(
        session,
        workspace,
        title="Rollout",
        format=format,
        script=script,
        base_artifact_id=None,
        image_names=[],
    )
    _work_off(studio_queue)
    made_id = made.id
    session.expire_all()
    assert made.document.status is DocumentStatus.READY, made.document.error_message
    before = _count(session)
    session.commit()

    shown = await client.get(f"/artifacts/{made_id}")
    response = await _refine(client, session, made_id, "Add a slide on costs")

    assert shown.json()["spec_kind"] == "python"
    assert shown.json()["refinable"] is False
    assert response.status_code == 409
    assert "chat" in response.json()["detail"]
    assert _count(session) == before
