"""Studio's job makes a revised copy's version with an engine and never writes the user's file."""

from io import BytesIO
from typing import Any

import docx
import openpyxl
import pptx
import pytest
from sqlalchemy.orm import Session

from modules.artifacts.models import Artifact
from modules.artifacts.revised_copies.decide_all import decide_all
from modules.artifacts.revised_copies.service import (
    create_next_version,
    create_revised_copy,
)
from modules.artifacts.revised_copies.source_file import sha256_of, source_file
from modules.artifacts.revised_copies.versions import revision_base
from modules.artifacts.script_documents.version import version_of
from modules.documents.models import DocumentStatus
from modules.workspaces.models import Workspace
from shared.config import get_storage_settings
from tests.integration.artifacts.revised_copies.source_files import (
    add_source,
    contract_docx,
    has_part,
    pitch_pptx,
    pricing_xlsx,
    source_path,
    word_xml,
)
from tests.integration.worker.conftest import stub_model  # noqa: F401
from worker.studio import run

pytestmark = [pytest.mark.integration, pytest.mark.usefixtures("stub_model")]

NAME = "MSA_Acme.docx"
REPLACE = {
    "op": "replace_text",
    "quote": "within 30 days",
    "text": "within 45 days",
}


def _v1(
    session: Session, workspace: Workspace, source_id: int, operations: list[dict]
) -> Artifact:
    source = source_file(session, workspace.id, source_id)
    return create_revised_copy(
        session,
        workspace.id,
        source,
        source_sha256=sha256_of(source.path),
        operations=operations,
        chat_thread_id=None,
    )


def _next(
    session: Session, workspace: Workspace, artifact: Artifact, operations: list[dict]
) -> Artifact:
    base = revision_base(session, workspace.id, artifact.id)
    return create_next_version(
        session,
        workspace.id,
        base,
        action="edit",
        operations=operations,
        chat_thread_id=None,
    )


def _made(session: Session, artifact: Artifact) -> Artifact:
    run(artifact.id)
    session.expire_all()
    assert artifact.document.status is DocumentStatus.READY, (
        artifact.document.error_message
    )
    return artifact


def _primary(artifact: Artifact) -> bytes:
    (file,) = artifact.files
    return (get_storage_settings().data_dir / file.storage_key).read_bytes()


def _kept(artifact: Artifact, key: str) -> bytes:
    record = artifact.artifact_metadata["revision"][key]
    return (get_storage_settings().data_dir / record["storage_key"]).read_bytes()


def _text(data: bytes) -> str:
    return "\n".join(p.text for p in docx.Document(BytesIO(data)).paragraphs)


def _revision(artifact: Artifact) -> dict[str, Any]:
    return artifact.artifact_metadata["revision"]


def test_version_1_is_the_word_file_with_tracked_changes_and_the_file_is_untouched(
    session: Session, workspace: Workspace
) -> None:
    """Decisions 11 and 12: a new artifact derived from the source; every edit tracked."""
    original = contract_docx()
    source_id = add_source(session, workspace.id, NAME, original)
    path = source_path(workspace.id, source_id, NAME)
    modified = path.stat().st_mtime_ns

    artifact = _made(session, _v1(session, workspace, source_id, [REPLACE]))

    assert path.read_bytes() == original
    assert path.stat().st_mtime_ns == modified
    assert artifact.document.title == "MSA_Acme (revised)"
    assert artifact.format == "docx"
    (file,) = artifact.files
    assert file.original_filename == "MSA_Acme (revised).docx"
    body = word_xml(_primary(artifact))
    assert "<w:ins " in body and "<w:del " in body
    assert 'w:author="SurfSense"' in body
    revision = _revision(artifact)
    assert revision["derived_from_document_id"] == source_id
    assert revision["source_name"] == NAME
    assert revision["counts"]["changes"] >= 2
    assert revision["counts"]["comments"] == 0
    assert revision["report"]["applied"] == 1
    assert artifact.artifact_metadata["source_document_ids"] == [source_id]
    assert version_of(artifact.artifact_metadata).number == 1


def test_the_body_and_the_clean_file_read_as_the_document_with_every_change_made(
    session: Session, workspace: Workspace
) -> None:
    """Search finds the revised wording; Clean is the file to send without markup."""
    source_id = add_source(session, workspace.id, NAME, contract_docx())

    artifact = _made(session, _v1(session, workspace, source_id, [REPLACE]))

    assert "within 45 days" in artifact.document.content
    assert "30 days" not in artifact.document.content
    clean = _kept(artifact, "clean")
    assert "<w:ins " not in word_xml(clean) and "<w:del " not in word_xml(clean)
    assert "within 45 days" in _text(clean)
    assert _revision(artifact)["external"] is None


def test_the_next_version_continues_from_the_newest_and_keeps_its_changes(
    session: Session, workspace: Workspace
) -> None:
    """A comment added later sits beside the first version's tracked changes."""
    source_id = add_source(session, workspace.id, NAME, contract_docx())
    first = _made(session, _v1(session, workspace, source_id, [REPLACE]))

    second = _made(
        session,
        _next(
            session,
            workspace,
            first,
            [
                {
                    "op": "add_comment",
                    "quote": "60 days notice",
                    "text": "Their template says 90.",
                }
            ],
        ),
    )

    version = version_of(second.artifact_metadata)
    assert (version.root, version.number, version.parent) == (first.id, 2, first.id)
    revision = _revision(second)
    assert revision["base_number"] == 1
    assert revision["counts"]["comments"] == 1
    assert revision["counts"]["changes"] == _revision(first)["counts"]["changes"]
    assert has_part(_primary(second), "word/comments.xml")


def test_an_internal_comment_stays_out_of_the_file_the_user_sends(
    session: Session, workspace: Workspace
) -> None:
    """Decision 18: "With changes" serves the external file; Clean has no comments."""
    source_id = add_source(session, workspace.id, NAME, contract_docx())

    artifact = _made(
        session,
        _v1(
            session,
            workspace,
            source_id,
            [
                {**REPLACE, "comment": "Matches our policy."},
                {
                    "op": "add_comment",
                    "quote": "60 days notice",
                    "text": "Push back here.",
                    "internal": True,
                },
            ],
        ),
    )

    revision = _revision(artifact)
    assert len(revision["internal_comment_ids"]) == 1
    assert "Push back here." in word_xml(_primary(artifact), "word/comments.xml")
    external = _kept(artifact, "external")
    assert "Matches our policy." in word_xml(external, "word/comments.xml")
    assert "Push back here." not in word_xml(external, "word/comments.xml")
    assert "Push back here." not in word_xml(_kept(artifact, "clean"))


INTERNAL = {
    "op": "add_comment",
    "quote": "60 days notice",
    "text": "Push back here.",
    "internal": True,
}


def test_a_later_version_keeps_earlier_internal_comments_out_of_the_file_sent(
    session: Session, workspace: Workspace
) -> None:
    """An internal comment stays internal in every version after the one that added it."""
    source_id = add_source(session, workspace.id, NAME, contract_docx())
    first = _made(session, _v1(session, workspace, source_id, [INTERNAL]))

    second = _made(session, _next(session, workspace, first, [REPLACE]))
    accepted = _made(session, decide_all(session, second, "accept_all"))

    for version in (second, accepted):
        assert "Push back here." in word_xml(_primary(version), "word/comments.xml")
        external = _kept(version, "external")
        assert "Push back here." not in word_xml(external, "word/comments.xml")
        assert "Push back here." not in word_xml(_kept(version, "clean"))


def test_accepting_all_makes_a_version_with_no_changes_left(
    session: Session, workspace: Workspace
) -> None:
    """Accept all is a version of its own; the one before keeps its changes."""
    source_id = add_source(session, workspace.id, NAME, contract_docx())
    first = _made(session, _v1(session, workspace, source_id, [REPLACE]))

    accepted = _made(session, decide_all(session, first, "accept_all"))

    revision = _revision(accepted)
    assert revision["action"] == "accept_all"
    assert revision["counts"]["changes"] == 0
    assert revision["report"] is None
    assert "within 45 days" in accepted.document.content
    assert "<w:ins " not in word_xml(_primary(accepted))


def test_rejecting_all_brings_the_original_wording_back(
    session: Session, workspace: Workspace
) -> None:
    """Reject all restores what the user's file said."""
    source_id = add_source(session, workspace.id, NAME, contract_docx())
    first = _made(session, _v1(session, workspace, source_id, [REPLACE]))

    rejected = _made(session, decide_all(session, first, "reject_all"))

    assert _revision(rejected)["counts"]["changes"] == 0
    assert "within 30 days" in rejected.document.content


def test_a_source_changed_after_the_check_fails_the_version_without_a_retry(
    session: Session, workspace: Workspace
) -> None:
    """Edits checked against other bytes are not applied to these."""
    source_id = add_source(session, workspace.id, NAME, contract_docx())
    artifact = _v1(session, workspace, source_id, [REPLACE])
    path = source_path(workspace.id, source_id, NAME)
    path.write_bytes(path.read_bytes() + b"\0")

    run(artifact.id)  # returning, not raising, is what spares a Huey retry

    session.expire_all()
    assert artifact.document.status is DocumentStatus.FAILED
    assert "changed on disk" in (artifact.document.error_message or "")
    assert artifact.files == []


def test_an_edit_the_engine_refuses_fails_the_version_and_saves_nothing(
    session: Session, workspace: Workspace
) -> None:
    """The tool checks first; the job holds the same rule for a file that changed meanwhile."""
    source_id = add_source(session, workspace.id, NAME, contract_docx())
    artifact = _v1(
        session,
        workspace,
        source_id,
        [{"op": "replace_text", "quote": "not in the file", "text": "x"}],
    )

    run(artifact.id)

    session.expire_all()
    assert artifact.document.status is DocumentStatus.FAILED
    assert (artifact.document.error_message or "").startswith(
        "Nothing was saved: operation #0 was refused."
    )


def test_an_excel_source_gets_its_cells_changed_and_keeps_its_suffix(
    session: Session, workspace: Workspace
) -> None:
    """The cell changes, the formula stays, and the copy is still a workbook."""
    source_id = add_source(session, workspace.id, "Pricing.xlsx", pricing_xlsx())

    artifact = _made(
        session,
        _v1(
            session,
            workspace,
            source_id,
            [{"op": "set_cell", "sheet": "Pricing", "cell": "B2", "value": 15000}],
        ),
    )

    assert artifact.format == "xlsx"
    assert artifact.files[0].original_filename == "Pricing (revised).xlsx"
    book = openpyxl.load_workbook(BytesIO(_primary(artifact)))
    assert book["Pricing"]["B2"].value == 15000
    assert book["Pricing"]["B4"].value == "=SUM(B2:B3)"
    assert 'Sheet "Pricing"' in artifact.document.content
    revision = _revision(artifact)
    assert revision["counts"] is None and revision["clean"] is None


def test_a_powerpoint_source_gets_its_text_replaced_in_place(
    session: Session, workspace: Workspace
) -> None:
    """Only the quoted text changes; the other slide is as it was."""
    source_id = add_source(session, workspace.id, "Pitch.pptx", pitch_pptx())

    artifact = _made(
        session,
        _v1(
            session,
            workspace,
            source_id,
            [{"op": "replace_text", "slide": 2, "quote": "Pricing", "text": "Costs"}],
        ),
    )

    deck = pptx.Presentation(BytesIO(_primary(artifact)))
    assert [s.shapes.title.text for s in deck.slides] == ["Halvorsen Freight", "Costs"]
    assert "Slide 2: Costs" in artifact.document.content
