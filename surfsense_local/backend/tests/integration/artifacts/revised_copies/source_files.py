"""Source files as the user uploads them, for a revised copy to start from."""

import zipfile
from io import BytesIO
from pathlib import Path

import docx
import openpyxl
import pptx
from sqlalchemy.orm import Session

from modules.documents.models import Document, DocumentStatus, DocumentType
from shared.config import get_storage_settings

CONTRACT_PARAGRAPHS = (
    "Master Services Agreement",
    "Payment is due within 30 days of the invoice date.",
    "Either party may terminate this agreement with 60 days notice.",
)


def contract_docx() -> bytes:
    """A three-paragraph agreement, as the user wrote it in Word."""
    document = docx.Document()
    for text in CONTRACT_PARAGRAPHS:
        document.add_paragraph(text)
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def pricing_xlsx() -> bytes:
    """A sheet of costs with a total formula."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Pricing"
    sheet.append(["Item", "Cost"])
    sheet.append(["Pilot", 12000])
    sheet.append(["Rollout", 30000])
    sheet["B4"] = "=SUM(B2:B3)"
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def pitch_pptx() -> bytes:
    """A two-slide deck whose slides have titles."""
    deck = pptx.Presentation()
    for title in ("Halvorsen Freight", "Pricing"):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


def add_source(session: Session, workspace_id: int, name: str, data: bytes) -> int:
    """A ready FILE source whose original is `data`, stored as an upload is."""
    source = Document(
        workspace_id=workspace_id,
        title=name,
        document_type=DocumentType.FILE,
        status=DocumentStatus.READY,
        content=name,
    )
    session.add(source)
    session.commit()
    folder = get_storage_settings().document_dir(workspace_id, source.id)
    folder.mkdir(parents=True)
    (folder / name).write_bytes(data)
    return source.id


def source_path(workspace_id: int, source_id: int, name: str) -> Path:
    """Where an upload's original is kept."""
    return get_storage_settings().document_dir(workspace_id, source_id) / name


def word_xml(data: bytes, part: str = "word/document.xml") -> str:
    """One part of a Word package, as text."""
    with zipfile.ZipFile(BytesIO(data)) as package:
        return package.read(part).decode("utf-8")


def has_part(data: bytes, part: str) -> bool:
    """Whether the package holds the part."""
    with zipfile.ZipFile(BytesIO(data)) as package:
        return part in package.namelist()
