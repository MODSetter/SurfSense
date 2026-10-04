"""Studio's job runs a document script as stored and keeps what it wrote."""

import sqlite3
import threading
import time
from io import BytesIO
from pathlib import Path
from typing import Any

import docx
import pytest
from sqlalchemy import text
from sqlalchemy.orm import Session

from modules.artifacts.models import Artifact
from modules.artifacts.script_documents.service import create_script_document
from modules.artifacts.service import regenerate_artifact
from modules.documents.models import Document, DocumentStatus
from modules.embedding import encoder
from modules.llm.model_type import ModelType
from modules.llm.models import SelectedModel
from modules.workspaces.models import Workspace
from shared.config import get_storage_settings
from shared.db import create_session_factory
from tests.integration.worker.conftest import stub_model  # noqa: F401
from worker.studio import run
from worker.studio.office.docx import docx as word
from worker.studio.office.pdf import pdf

# The job indexes what the script wrote; the stub stands in for the embedder.
pytestmark = [pytest.mark.integration, pytest.mark.usefixtures("stub_model")]

WORD_SCRIPT = """\
import os
import docx

document = docx.Document()
document.add_heading("Client proposal", level=1)
document.add_paragraph("We propose a two-phase rollout for Halvorsen Freight.")
table = document.add_table(rows=2, cols=2)
table.style = "Table Grid"
table.cell(0, 0).text = "Phase"
table.cell(0, 1).text = "Cost"
table.cell(1, 0).text = "Pilot"
table.cell(1, 1).text = "12,000"
document.save(os.environ["OUTPUT_PATH"])
"""

PDF_SCRIPT = """\
import os
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

drawing = canvas.Canvas(os.environ["OUTPUT_PATH"], pagesize=A4)
drawing.drawString(72, 760, "Halvorsen Freight pricing")
drawing.showPage()
drawing.save()
"""

LOGO_SCRIPT = """\
import os
import docx
from docx.shared import Cm

name = "{name}"
document = docx.Document()
document.add_picture(os.path.join(os.environ["IMAGES_DIR"], name + ".png"), width=Cm(4))
document.add_paragraph("Logo placed.")
document.save(os.environ["OUTPUT_PATH"])
"""


def _create(session: Session, workspace: Workspace, **overrides: Any) -> Artifact:
    request: dict[str, Any] = {
        "title": "Client proposal",
        "format": "docx",
        "script": WORD_SCRIPT,
        "base_artifact_id": None,
        "image_names": [],
    }
    return create_script_document(session, workspace, **{**request, **overrides})


def _primary(artifact: Artifact) -> bytes:
    (file,) = artifact.files
    return (get_storage_settings().data_dir / file.storage_key).read_bytes()


def _matches(session: Session, term: str) -> int:
    return session.scalar(
        text("SELECT count(*) FROM chunks_fts WHERE chunks_fts MATCH :query"),
        {"query": term},
    )


def test_a_python_docx_script_becomes_a_ready_word_file_whose_text_is_searchable(
    session: Session, workspace: Workspace
) -> None:
    """Word: the file is kept as written, and its text becomes the searchable body."""
    artifact = _create(session, workspace, title="Proposal: Halvorsen?")

    run(artifact.id)

    session.expire_all()
    assert artifact.document.status is DocumentStatus.READY, (
        artifact.document.error_message
    )
    assert artifact.document.title == "Proposal: Halvorsen?"
    (file,) = artifact.files
    assert file.mime_type == word.mime
    assert file.original_filename == "Proposal Halvorsen.docx"
    opened = docx.Document(BytesIO(_primary(artifact)))
    assert opened.paragraphs[1].text.startswith("We propose a two-phase rollout")
    assert "Phase | Cost" in artifact.document.content
    assert _matches(session, "Halvorsen") == 1


def test_a_reportlab_script_becomes_a_ready_pdf(
    session: Session, workspace: Workspace
) -> None:
    """PDF: ReportLab's file is kept, and its text layer is the body."""
    artifact = _create(session, workspace, format="pdf", script=PDF_SCRIPT)

    run(artifact.id)

    session.expire_all()
    assert artifact.document.status is DocumentStatus.READY, (
        artifact.document.error_message
    )
    assert artifact.files[0].mime_type == pdf.mime
    assert _primary(artifact).startswith(b"%PDF")
    assert artifact.document.content == "Halvorsen Freight pricing"


def test_a_source_image_the_spec_names_is_where_the_script_looks(
    session: Session, workspace: Workspace, source_figure: str
) -> None:
    """A figure the spec names is copied to IMAGES_DIR under its name."""
    artifact = _create(
        session,
        workspace,
        script=LOGO_SCRIPT.format(name=source_figure),
        image_names=[source_figure],
    )

    run(artifact.id)

    session.expire_all()
    assert artifact.document.status is DocumentStatus.READY, (
        artifact.document.error_message
    )
    opened = docx.Document(BytesIO(_primary(artifact)))
    assert len(opened.inline_shapes) == 1


def test_a_failing_script_fails_the_job_with_its_error_and_is_not_retried(
    session: Session, workspace: Workspace
) -> None:
    """The error and the traceback's tail are what the agent fixes its script from."""
    artifact = _create(
        session,
        workspace,
        script="rows = []\nraise ValueError('the pricing table is empty')\n",
    )

    run(artifact.id)  # returning, not raising, is what spares a Huey retry

    session.expire_all()
    assert artifact.document.status is DocumentStatus.FAILED
    reason = artifact.document.error_message or ""
    assert "ValueError: the pricing table is empty" in reason.splitlines()[0]
    assert len(reason.splitlines()) > 1  # the traceback's tail follows the error
    assert len(reason) <= 500
    assert artifact.files == []
    assert session.scalar(text("SELECT count(*) FROM chunks")) == 0


def test_a_file_that_is_not_a_word_file_fails_the_job_without_a_retry(
    session: Session, workspace: Workspace
) -> None:
    """Bytes that do not open as the format are a failed run, not a ready file."""
    artifact = _create(
        session,
        workspace,
        script=(
            "import os\nopen(os.environ['OUTPUT_PATH'], 'w').write('Dear client,')\n"
        ),
    )

    run(artifact.id)

    session.expire_all()
    assert artifact.document.status is DocumentStatus.FAILED
    assert (
        artifact.document.error_message
        == "Script error: the script wrote a file that is not a valid .docx"
    )


def test_a_script_document_that_fails_after_its_run_is_not_retried_either(
    session: Session, workspace: Workspace, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The agent already read FAILED and moved on; a retry turning this version
    READY behind it would leave the user a version the chat never announced."""

    def embedder_down(*_args: Any) -> list[list[float]]:
        raise RuntimeError("the embedder could not be loaded")

    monkeypatch.setattr("modules.embedding.encoder.embed", embedder_down)
    artifact = _create(session, workspace)

    run(artifact.id)  # returning, not raising, is what spares a Huey retry

    session.expire_all()
    assert artifact.document.status is DocumentStatus.FAILED
    assert artifact.document.error_message == "the embedder could not be loaded"


def test_regenerate_runs_the_stored_script_again_with_no_chat_model_chosen(
    session: Session, workspace: Workspace
) -> None:
    """The stored script is the document's source: nothing is redrafted, so the
    Word format's chat model is not needed, as the agent's own model wrote it."""
    artifact = _create(session, workspace)
    run(artifact.id)
    assert session.get(SelectedModel, ModelType.TEXT_GEN) is None

    regenerate_artifact(session, artifact)
    run(artifact.id)

    session.expire_all()
    assert artifact.generation == 2
    assert artifact.document.status is DocumentStatus.READY, (
        artifact.document.error_message
    )
    assert artifact.artifact_metadata["spec"]["text"] == WORD_SCRIPT
    assert docx.Document(BytesIO(_primary(artifact))).paragraphs[0].text == (
        "Client proposal"
    )


def test_a_cancel_stops_a_running_script_without_waiting_out_its_limit(
    session: Session, workspace: Workspace, tmp_path: Path
) -> None:
    """The Studio slot is freed within a poll of the cancel, not after two minutes."""
    started = tmp_path / "started"
    artifact = _create(
        session,
        workspace,
        script=f"import pathlib, time\npathlib.Path({str(started)!r}).touch()\n"
        "time.sleep(60)\n",
    )

    def cancel_once_running() -> None:
        deadline = time.monotonic() + 30
        while not started.exists() and time.monotonic() < deadline:
            time.sleep(0.05)
        with create_session_factory(session.get_bind())() as other:
            other.get(Document, artifact.document_id).status = DocumentStatus.CANCELLED
            other.commit()

    canceller = threading.Thread(target=cancel_once_running)
    canceller.start()
    begun = time.monotonic()
    run(artifact.id)
    elapsed = time.monotonic() - begun
    canceller.join()

    session.expire_all()
    assert artifact.document.status is DocumentStatus.CANCELLED
    assert elapsed < 20
    assert artifact.files == []


# Longer than SQLite's 5 s busy wait (shared.db), as another job embedding a long body.
LOCK_SECONDS = 6.5


def _hold_the_write_lock_once(session: Session, started: Path, released: Path) -> None:
    """Take the write lock once the script runs, keep it past the busy wait, then say so."""
    deadline = time.monotonic() + 30
    while not started.exists() and time.monotonic() < deadline:
        time.sleep(0.05)
    with create_session_factory(session.get_bind())() as other:
        other.execute(text("SELECT 1"))  # every transaction begins with the write lock
        time.sleep(LOCK_SECONDS)
        other.commit()
    released.touch()


def test_a_write_lock_held_elsewhere_does_not_fail_a_running_script(
    session: Session, workspace: Workspace, tmp_path: Path
) -> None:
    """The cancel check skips a look it cannot get; a correct script is never failed for it."""
    started, released = tmp_path / "started", tmp_path / "released"
    script = (
        "import os, pathlib, time\n"
        f"pathlib.Path({str(started)!r}).touch()\n"
        f"while not pathlib.Path({str(released)!r}).exists():\n"
        "    time.sleep(0.05)\n" + WORD_SCRIPT
    )
    artifact = _create(session, workspace, script=script)
    holder = threading.Thread(
        target=_hold_the_write_lock_once, args=(session, started, released)
    )
    holder.start()

    run(artifact.id)
    holder.join()

    session.expire_all()
    assert artifact.document.status is DocumentStatus.READY, (
        artifact.document.error_message
    )


def test_the_write_lock_is_free_while_the_body_is_embedded(
    session: Session, workspace: Workspace, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Embedding a long body takes as long as the body; other writers must not wait on it."""
    embed = encoder.embed
    others_could_write: list[bool] = []

    def embed_while_another_writes(*args: Any) -> list[list[float]]:
        other = sqlite3.connect(
            get_storage_settings().database_path, timeout=0.1, isolation_level=None
        )
        try:
            other.execute("BEGIN IMMEDIATE")
            other.execute("ROLLBACK")
            others_could_write.append(True)
        except sqlite3.OperationalError:
            others_could_write.append(False)
        finally:
            other.close()
        return embed(*args)

    monkeypatch.setattr("modules.embedding.encoder.embed", embed_while_another_writes)
    artifact = _create(session, workspace)

    run(artifact.id)

    session.expire_all()
    assert artifact.document.status is DocumentStatus.READY
    assert others_could_write == [True]
