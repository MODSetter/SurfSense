"""Figures real Docling keeps from generated PDF, Word and PowerPoint files, offline.

Ingest reads the born-digital PDF from its text layer and queues the figures-only
pass, which then reads it with Docling once nothing else is queued.
"""

import io
import json
import os
import subprocess
import sys

import docx
import pptx
import pytest
from docx.shared import Inches
from PIL import Image, ImageDraw
from pptx.util import Inches as SlideInches
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

from modules.documents.models import Document, DocumentStatus, DocumentType
from modules.documents.tasks import PRIORITY_IDLE
from modules.workspaces.models import Workspace
from shared.config import StorageSettings
from shared.db import create_db_engine, create_session_factory
from shared.migrations import upgrade_to_head
from tests.integration.worker.test_parsing import BACKEND, PACK_ROOT

pytestmark = pytest.mark.integration

PHOTO_CAPTION = "Figure 1: The harbour at dawn."
CHART_CAPTION = "Figure 2: Yearly costs by region."

# Every file converts in one child, which pays Docling's import and model load once.
CHILD = """
import socket

def refuse(*_args, **_kwargs):
    raise OSError("this test refuses outbound network")

socket.socket.connect = refuse
socket.socket.connect_ex = refuse

from sqlalchemy import select

from modules.documents.models import Document, DocumentStatus
from shared.config import get_storage_settings
from shared.db import create_db_engine, create_session_factory, import_models
from shared.queue import ingest_queue
from worker.ingestion.parsing import markdown_for

import_models()
engine = create_db_engine(get_storage_settings().database_path)
with create_session_factory(engine)() as session:
    for document in session.scalars(select(Document).order_by(Document.id)):
        markdown_for(document)
        # As ingest finishes a source: ready, with its figures kept or queued.
        document.status = DocumentStatus.READY
        session.commit()
engine.dispose()

# The ingest worker, once no upload waits: the figures-only pass, as queued.
while ingest_queue.pending_count():
    task = ingest_queue.dequeue()
    print(f"queued: {task.name} {task.args} {task.priority}")
    ingest_queue.execute(task)
"""


def photo(size: tuple[int, int], tint: int) -> bytes:
    """A raster picture with enough shape that layout reads it as one."""
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    for x in range(0, size[0], 20):
        draw.rectangle([x, 0, x + 10, size[1]], fill=(x % 255, tint, 200))
    draw.ellipse(
        [size[0] // 4, size[1] // 6, size[0] * 3 // 4, size[1] * 5 // 6], "orange"
    )
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def report_pdf() -> bytes:
    """A raster photo on page one and a vector bar chart on page two, each captioned."""
    styles = getSampleStyleSheet()
    bar_chart = Drawing(400, 220)
    chart = VerticalBarChart()
    chart.x, chart.y, chart.width, chart.height = 40, 30, 340, 170
    chart.data = [(13, 5, 20, 22, 37), (14, 6, 21, 23, 38)]
    chart.categoryAxis.categoryNames = ["2021", "2022", "2023", "2024", "2025"]
    bar_chart.add(chart)
    story = [
        Paragraph("Quarterly report", styles["Title"]),
        Paragraph("The quarter in figures. " * 12, styles["BodyText"]),
        PdfImage(io.BytesIO(photo((400, 300), 100)), width=300, height=225),
        Paragraph(PHOTO_CAPTION, styles["Italic"]),
        PageBreak(),
        Paragraph("Costs by region and year. " * 12, styles["BodyText"]),
        bar_chart,
        Paragraph(CHART_CAPTION, styles["Italic"]),
    ]
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4).build(story)
    return out.getvalue()


def pasted_docx() -> bytes:
    """A Word memo with a picture pasted under its heading, captioned."""
    document = docx.Document()
    document.add_heading("A memo with a pasted picture", 1)
    document.add_picture(io.BytesIO(photo((400, 300), 100)), width=Inches(3))
    document.add_paragraph("Figure 1: A pasted picture.", style="Caption")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def deck_pptx() -> bytes:
    """Two slides, each with a different picture."""
    deck = pptx.Presentation()
    for tint in (100, 30):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "A slide with a picture"
        picture = io.BytesIO(photo((400, 300), tint))
        slide.shapes.add_picture(picture, SlideInches(1), SlideInches(2))
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


@pytest.fixture(scope="module")
def kept(tmp_path_factory: pytest.TempPathFactory) -> dict[str, dict]:
    """Each generated file's figures index, keyed by its file name."""
    if PACK_ROOT is None:
        pytest.skip("run scripts/fetch_docling_models.py to exercise the real parser")

    root = tmp_path_factory.mktemp("figures")
    files = {
        "report.pdf": report_pdf(),
        "memo.docx": pasted_docx(),
        "deck.pptx": deck_pptx(),
    }
    storage = StorageSettings(data_dir=root / "data")
    engine = create_db_engine(storage.database_path)
    upgrade_to_head(engine)
    folders, ids = {}, {}
    with create_session_factory(engine)() as session:
        workspace = Workspace(name="Figures")
        session.add(workspace)
        session.flush()
        for name, content in files.items():
            source = Document(
                workspace_id=workspace.id,
                title=name,
                document_type=DocumentType.FILE,
                status=DocumentStatus.PROCESSING,
            )
            session.add(source)
            session.flush()
            folder = storage.document_dir(workspace.id, source.id)
            folder.mkdir(parents=True)
            (folder / name).write_bytes(content)
            folders[name], ids[name] = folder, source.id
        session.commit()
    engine.dispose()

    empty_cache = root / "hf"
    empty_cache.mkdir()
    child = subprocess.run(
        [sys.executable, "-c", CHILD],
        cwd=BACKEND,
        env={
            **os.environ,
            "SURFSENSE_LOCAL_DATA_DIR": str(storage.data_dir),
            "SURFSENSE_LOCAL_MODELS_DIR": str(PACK_ROOT),
            "HF_HOME": str(empty_cache),
            "HF_HUB_OFFLINE": "1",
            # Docling's own reading of a Word file's pictures. Ingest reads Word
            # without Docling and leaves them to the figures pass, which asks
            # Docling: test_ingest_word.py.
            "SURFSENSE_LOCAL_DOCX_READER": "docling",
        },
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=900,
    )
    assert child.returncode == 0, child.stderr[-4000:]
    # Only the born-digital PDF is left to the figures pass, below every ingest.
    queued = [line for line in child.stdout.splitlines() if line.startswith("queued:")]
    assert queued == [
        f"queued: extract_figures ({ids['report.pdf']},) {PRIORITY_IDLE}"
    ], child.stdout[-4000:]
    return {
        name: {
            "index": json.loads(
                (folder / "figures" / "figures.json").read_text("utf-8")
            ),
            "folder": folder / "figures",
        }
        for name, folder in folders.items()
    }


def test_a_pdfs_raster_photo_and_drawn_chart_are_kept_with_captions_and_pages(
    kept: dict[str, dict],
) -> None:
    """The chart is vector drawing, not an embedded image: layout finds it anyway."""
    figures = kept["report.pdf"]["index"]["figures"]

    assert [(f["n"], f["caption"], f["page"]) for f in figures] == [
        (1, PHOTO_CAPTION, 1),
        (2, CHART_CAPTION, 2),
    ]
    # Cropped from the page at 2x: the photo was placed at 300 x 225 points.
    assert 540 <= figures[0]["width"] <= 660
    assert 400 <= figures[0]["height"] <= 500
    for figure in figures:
        with Image.open(kept["report.pdf"]["folder"] / f"{figure['n']}.png") as image:
            assert image.size == (figure["width"], figure["height"])


def test_a_word_files_pasted_picture_is_kept_at_its_own_pixels(
    kept: dict[str, dict],
) -> None:
    """Docling's Word reader links no caption and knows no page."""
    assert kept["memo.docx"]["index"]["figures"] == [
        {"n": 1, "caption": None, "page": None, "width": 400, "height": 300}
    ]


def test_a_decks_pictures_are_kept_with_their_slide_numbers(
    kept: dict[str, dict],
) -> None:
    """Docling reports a slide's number as its page."""
    figures = kept["deck.pptx"]["index"]["figures"]

    assert [(f["page"], f["width"], f["height"]) for f in figures] == [
        (1, 400, 300),
        (2, 400, 300),
    ]
