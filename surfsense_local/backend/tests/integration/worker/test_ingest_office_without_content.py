"""Office files with nothing to read, such as a template, ingested with real Docling."""

import io
import zipfile
from collections.abc import Iterator

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import Engine, text
from sqlalchemy.orm import Session

from modules.documents.models import Document, DocumentStatus, DocumentType
from modules.documents.source_figures.layout import figures_dir, read_index
from modules.documents.tasks import ingest_document
from modules.workspaces.models import Workspace
from shared.config import get_storage_settings
from shared.db import create_session_factory
from shared.queue import ingest_queue
from worker.ingestion import run

pytestmark = pytest.mark.integration


@pytest.fixture
def session(engine: Engine) -> Iterator[Session]:
    """A session on the migrated database ingest opens again by path."""
    with create_session_factory(engine)() as opened:
        yield opened


def upload(session: Session, name: str, content: bytes) -> Document:
    """A pending uploaded file with its original on disk, as the upload route leaves it."""
    workspace = Workspace(name="Brand")
    session.add(workspace)
    session.flush()
    document = Document(
        workspace_id=workspace.id,
        title=name,
        document_type=DocumentType.FILE,
        dedup_key=name,
        document_metadata={"suffix": name[name.rindex(".") :]},
    )
    session.add(document)
    session.commit()
    folder = get_storage_settings().document_dir(workspace.id, document.id)
    folder.mkdir(parents=True)
    (folder / name).write_bytes(content)
    return document


def saved(office_file: object) -> bytes:
    """A python-docx, python-pptx or openpyxl file's bytes."""
    buffer = io.BytesIO()
    office_file.save(buffer)  # type: ignore[attr-defined]
    return buffer.getvalue()


def template_deck() -> bytes:
    """A widescreen template with renamed layouts and no slides, as PowerPoint saves one."""
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.slide_layouts[0].name = "Brand_Title"
    deck.slide_layouts[1].name = "Brand Content"
    return saved(deck)


def test_a_deck_with_no_slides_is_a_ready_template(
    session: Session, stub_model: None
) -> None:
    """Docling refuses a deck with no pages; a template must still be usable."""
    document = upload(session, "brand.pptx", template_deck())

    run(document.id)

    session.expire_all()
    assert document.status is DocumentStatus.READY
    assert document.error_message is None
    content = document.content or ""
    assert content.startswith("PowerPoint template with no slides.")
    assert "13.33 x 7.5 in (16:9)" in content
    # Exact names, in order, so the agent can pick a layout by name.
    assert content.index("- Brand_Title") < content.index("- Brand Content")
    assert "- Title Only" in content

    found = session.scalar(
        text("SELECT count(*) FROM chunks_fts WHERE chunks_fts MATCH 'template'")
    )
    assert found == 1
    # No pictures to keep, recorded so the figures-only pass never reads it again.
    folder = figures_dir(
        get_storage_settings().document_dir(document.workspace_id, document.id)
    )
    assert read_index(folder) == []


def test_a_deck_with_slides_is_still_read_by_docling(
    session: Session, stub_model: None
) -> None:
    """Only a deck with no slides at all is described rather than read."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Harbour budget"
    slide.placeholders[1].text = "Ferry subsidy rises 4 percent."
    document = upload(session, "budget.pptx", saved(deck))

    run(document.id)

    session.expire_all()
    assert document.status is DocumentStatus.READY
    assert "Ferry subsidy rises 4 percent." in (document.content or "")
    assert "template" not in (document.content or "")


@pytest.mark.parametrize(
    ("name", "content"),
    [
        ("blank.docx", lambda: saved(WordDocument())),
        ("blank.xlsx", lambda: saved(Workbook())),
    ],
)
def test_a_blank_word_or_excel_file_is_ready(
    session: Session, stub_model: None, name: str, content: object
) -> None:
    """Docling reads these as empty rather than refusing them."""
    document = upload(session, name, content())  # type: ignore[operator]

    run(document.id)

    session.expire_all()
    assert document.status is DocumentStatus.READY


def test_a_file_docling_refuses_fails_without_retries(
    session: Session, stub_model: None
) -> None:
    """The same bytes are refused again, so a retry would only delay the failed state."""
    broken = io.BytesIO()
    with zipfile.ZipFile(broken, "w") as package:
        package.writestr("word/document.xml", "<not word")
    document = upload(session, "broken.docx", broken.getvalue())

    ingest_document(document.id)
    ingest_queue.execute(ingest_queue.dequeue())

    session.expire_all()
    assert document.status is DocumentStatus.FAILED
    assert "could not load document" in (document.error_message or "")
    assert ingest_queue.pending_count() == 0
    assert ingest_queue.scheduled_count() == 0
