"""Word files ingested by the reader in worker/ingestion/word_markdown, Docling
left for the files it cannot read and for their figures."""

import io
import json
from collections.abc import Iterator
from types import SimpleNamespace
from typing import Any

import docx
import pytest
from docling_core.types.doc import ImageRef
from docx.oxml import parse_xml
from docx.shared import Inches
from sqlalchemy import Engine, text
from sqlalchemy.orm import Session

from modules.documents.models import DocumentStatus
from modules.documents.source_figures.layout import read_index
from modules.documents.source_figures.listing import FiguresPending, list_figures
from modules.documents.tasks import (
    PRIORITY_ASKED,
    PRIORITY_BULK,
    extract_figures,
)
from shared.db import create_session_factory
from shared.queue import ingest_queue
from tests.integration.worker.test_figures import (
    add_source,
    figures_dir,
    png,
    reading,
)
from tests.integration.worker.test_figures_docling import photo
from worker.ingestion import run

pytestmark = pytest.mark.integration

TEXT_BOX_NAMESPACES = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
)


@pytest.fixture
def docling(monkeypatch: pytest.MonkeyPatch) -> list[Any]:
    """Docling reads each file as the next reading here, and fails the test
    when it is asked with none left."""
    readings: list[Any] = []

    def convert(_source: object) -> SimpleNamespace:
        if not readings:
            raise AssertionError("Docling was asked to read a file")
        return SimpleNamespace(document=readings.pop(0))

    monkeypatch.setattr(
        "worker.ingestion.parsing._converter",
        lambda: SimpleNamespace(convert=convert),
    )
    return readings


@pytest.fixture
def session(engine: Engine) -> Iterator[Session]:
    """A session on the migrated database ingest opens again by path."""
    with create_session_factory(engine)() as opened:
        yield opened


def memo(picture: bool = False, header_logo: bool = False) -> bytes:
    """A short memo: a heading, a numbered list, a table and, if asked, a picture
    in the body or a logo in the header."""
    document = docx.Document()
    document.add_heading("Harbour budget", 1)
    document.add_paragraph("Ferry subsidy rises 4 percent.", style="List Number")
    document.add_paragraph("Pier repairs move to spring.", style="List Number")
    table = document.add_table(rows=2, cols=2)
    for r, row in enumerate([["Line", "EUR"], ["Ferry", "120,000"]]):
        for c, cell_text in enumerate(row):
            table.cell(r, c).text = cell_text
    if picture:
        document.add_picture(io.BytesIO(photo((400, 300), 100)), width=Inches(3))
    if header_logo:
        logo = io.BytesIO(photo((120, 60), 30))
        header = document.sections[0].header.paragraphs[0]
        header.add_run().add_picture(logo, width=Inches(1))
    return saved(document)


def saved(document: Any) -> bytes:
    """The document's bytes, as an upload carries them."""
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def figures_queued() -> list[tuple[tuple, int]]:
    """The figures passes waiting, as their arguments and priority."""
    return [
        (task.args, task.priority)
        for task in ingest_queue.pending()
        if task.name == extract_figures.s(0).name
    ]


def test_a_word_file_is_ready_with_the_readers_text_and_no_docling(
    session: Session, stub_model: None, docling: list[Any]
) -> None:
    """The docling fixture fails the test if Docling is asked to read anything."""
    source = add_source(session, "budget.docx", memo())

    run(source.id)

    session.expire_all()
    assert source.status is DocumentStatus.READY
    assert source.content == (
        "## Harbour budget\n\n"
        "1. Ferry subsidy rises 4 percent.\n"
        "2. Pier repairs move to spring.\n\n"
        "| Line | EUR |\n|---|---|\n| Ferry | 120,000 |"
    )
    found = session.scalar(
        text("SELECT count(*) FROM chunks_fts WHERE chunks_fts MATCH 'ferry'")
    )
    assert found == 1


def test_the_switch_hands_word_files_back_to_docling(
    session: Session,
    stub_model: None,
    docling: list[Any],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The escape hatch, should this reader misread a file."""
    monkeypatch.setenv("SURFSENSE_LOCAL_DOCX_READER", "docling")
    docling.append(reading())
    source = add_source(session, "budget.docx", memo())

    run(source.id)

    session.expire_all()
    assert source.content == "Quarterly results."
    assert not docling  # Docling read it


def test_a_word_file_the_reader_cannot_open_is_read_by_docling(
    session: Session, stub_model: None, docling: list[Any]
) -> None:
    """Docling may still read what this reader refuses, such as a Strict OOXML file."""
    docling.append(reading())
    source = add_source(session, "strict.docx", b"PK not quite a Word file")

    run(source.id)

    session.expire_all()
    assert source.status is DocumentStatus.READY
    assert source.content == "Quarterly results."


@pytest.mark.parametrize("picture", [False, True])
def test_a_reread_word_file_forgets_the_figures_of_its_last_reading(
    session: Session, stub_model: None, docling: list[Any], picture: bool
) -> None:
    """A changed file's old figures must not be offered as its figures."""
    source = add_source(session, "budget.docx", memo(picture=picture))
    folder = figures_dir(source)
    folder.mkdir()
    (folder / "1.png").write_bytes(b"old")
    (folder / "figures.json").write_text(
        json.dumps({"version": 1, "figures": [{"n": 1, "width": 1, "height": 1}]}),
        encoding="utf-8",
    )

    run(source.id)

    assert not (folder / "1.png").exists()
    # With a picture the figures pass reads it; without one there is none to read.
    assert read_index(folder) == (None if picture else [])


def test_a_word_file_without_pictures_lists_none_and_queues_no_pass(
    session: Session, stub_model: None, docling: list[Any]
) -> None:
    """Docling would find none, so nobody waits on a second reading of the file."""
    source = add_source(session, "budget.docx", memo())

    run(source.id)

    assert list_figures(session, source.workspace_id, source.id) == []
    assert figures_queued() == []


@pytest.mark.parametrize(
    "where", [{"picture": True}, {"header_logo": True}], ids=["body", "header"]
)
def test_a_word_file_with_pictures_queues_their_pass_as_its_ingest_ends(
    session: Session, stub_model: None, docling: list[Any], where: dict[str, bool]
) -> None:
    """So a Studio draft from it can place them; behind other files' text."""
    source = add_source(session, "memo.docx", memo(**where))

    run(source.id)

    session.expire_all()
    assert source.status is DocumentStatus.READY
    assert figures_queued() == [((source.id,), PRIORITY_BULK)]


def test_a_word_file_with_only_shapes_leaves_its_figures_to_an_ask(
    session: Session, stub_model: None, docling: list[Any]
) -> None:
    """Docling draws shapes only when LibreOffice is on PATH: no pass for every
    file with a text box, and no index saying it has none either."""
    document = docx.Document()
    document.add_paragraph("Intro.")._p.append(
        parse_xml(
            f"<w:r {TEXT_BOX_NAMESPACES}><w:drawing><wp:inline>"
            '<wp:extent cx="914400" cy="457200"/><wp:docPr id="1" name="Box"/>'
            "<a:graphic><a:graphicData"
            ' uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            "<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>Boxed</w:t></w:r></w:p>"
            "</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic>"
            "</wp:inline></w:drawing></w:r>"
        )
    )
    source = add_source(session, "boxed.docx", saved(document))

    run(source.id)

    session.expire_all()
    assert source.content == "Boxed\n\nIntro."
    assert read_index(figures_dir(source)) is None
    assert figures_queued() == []


def test_a_picture_only_word_file_keeps_its_figures_at_ingest(
    session: Session, stub_model: None, docling: list[Any]
) -> None:
    """No text for the reader, so Docling reads the file, and its pictures with it."""
    document = docx.Document()
    document.add_picture(io.BytesIO(photo((400, 300), 100)), width=Inches(3))
    docling.append(reading({"image": ImageRef.from_pil(png((120, 80)), dpi=72)}))
    source = add_source(session, "scan.docx", saved(document))

    run(source.id)

    assert not docling  # Docling read it
    assert read_index(figures_dir(source)) == [
        {"n": 1, "caption": None, "page": None, "width": 120, "height": 80}
    ]
    assert figures_queued() == []


def test_a_word_files_pictures_come_from_the_figures_pass(
    session: Session, stub_model: None
) -> None:
    """Real Docling. Someone asking goes ahead of the bulk copy ingest queued,
    which then finds the figures kept and leaves them."""
    source = add_source(session, "memo.docx", memo(picture=True))
    run(source.id)
    session.expire_all()
    assert source.status is DocumentStatus.READY
    folder = figures_dir(source)
    assert read_index(folder) is None

    with pytest.raises(FiguresPending):
        list_figures(session, source.workspace_id, source.id)
    assert figures_queued() == [
        ((source.id,), PRIORITY_ASKED),
        ((source.id,), PRIORITY_BULK),
    ]
    session.close()  # the pass reads the row in a transaction of its own
    first = ingest_queue.dequeue()
    assert (first.args, first.priority) == ((source.id,), PRIORITY_ASKED)
    extract_figures.call_local(source.id)
    kept = [{"n": 1, "caption": None, "page": None, "width": 400, "height": 300}]
    assert read_index(folder) == kept

    extract_figures.call_local(source.id)  # the bulk copy, later

    assert read_index(folder) == kept
