"""Real Docling on a generated Word file, with and without the style cache."""

import random
from contextlib import nullcontext
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.parts.document import DocumentPart

from worker.ingestion import parsing

pytestmark = pytest.mark.integration

# The distinct (style_id, style_type) pairs Docling asks for in `contract`. A
# change means a Docling or python-docx bump reads styles another way: check the
# cache still applies before updating it, or a long contract reads 3.5x slower.
STYLE_LOOKUPS = 20
WORDS = ["party", "notice", "payment", "term", "clause", "supplier", "buyer", "fee"]


def contract(folder: Path) -> Path:
    """Many styled paragraphs of many runs, as a long contract has, with headings,
    lists, numbering, a table, character styles, a header and broken style references."""
    rng = random.Random(7)
    document = Document()
    clauses = [
        document.styles.add_style(f"Clause {n}", WD_STYLE_TYPE.PARAGRAPH)
        for n in range(8)
    ]
    defined = document.styles.add_style("Defined Term", WD_STYLE_TYPE.CHARACTER)
    defined.font.bold = True
    numbered = document.styles.add_style("Numbered Clause", WD_STYLE_TYPE.PARAGRAPH)
    numbered.element.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="1"/></w:numPr>')
    )
    document.styles.element.append(
        parse_xml(
            f'<w:style {nsdecls("w")} w:type="paragraph"><w:name w:val="No Id"/></w:style>'
        )
    )
    document.sections[0].header.paragraphs[0].text = "Short-form contract"
    document.sections[0].header.paragraphs[0].style = "Header"

    document.add_paragraph("Services contract", style="Title")
    for section in range(6):
        document.add_heading(f"Part {section + 1}", level=1 + section % 3)
        for n in range(20):
            style = [*clauses, "List Bullet", "List Number", numbered][n % 11]
            paragraph = document.add_paragraph(style=style)
            for _ in range(6):
                run = paragraph.add_run(" ".join(rng.sample(WORDS, 3)) + " ")
                if rng.random() < 0.3:
                    run.style = rng.choice([defined, "Strong", "Emphasis"])
        for missing in ("DoesNotExist", "Strong", ""):
            paragraph = document.add_paragraph("A clause with a broken style.")
            reference = OxmlElement("w:pStyle")
            reference.set(qn("w:val"), missing)
            paragraph._p.get_or_add_pPr().insert(0, reference)

    table = document.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.paragraphs[0].add_run(f"{rng.choice(WORDS)} {r}.{c}")
            cell.paragraphs[0].style = "List Bullet" if c == 2 else "Normal"
    table.cell(0, 0).merge(table.cell(0, 1))

    path = folder / "contract.docx"
    document.save(str(path))
    return path


def reading(path: Path) -> tuple[str, dict[str, Any]]:
    """What ingest keeps of the file: its markdown, and the whole DoclingDocument."""
    converted = parsing.convert(path)
    return converted.export_to_markdown(), converted.export_to_dict()


@pytest.fixture
def lookups(monkeypatch: pytest.MonkeyPatch) -> list[Any]:
    """Each lookup that reaches python-docx's own get_style."""
    seen: list[Any] = []
    real = DocumentPart.get_style

    def counting(part: DocumentPart, style_id: Any, style_type: Any) -> Any:
        seen.append((style_id, style_type))
        return real(part, style_id, style_type)

    monkeypatch.setattr(DocumentPart, "get_style", counting)
    return seen


def test_the_cache_leaves_docling_reading_unchanged(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Ingest, figures, citations and Studio all build on this reading."""
    path = contract(tmp_path)
    cached = reading(path)

    monkeypatch.setattr(parsing, "cached_style_lookups", nullcontext)
    assert reading(path) == cached


def test_docling_reaches_python_docx_once_per_style(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, lookups: list[Any]
) -> None:
    """The cache only pays while Docling looks styles up through DocumentPart."""
    path = contract(tmp_path)
    parsing.convert(path)
    cached = len(lookups)

    lookups.clear()
    monkeypatch.setattr(parsing, "cached_style_lookups", nullcontext)
    parsing.convert(path)

    assert cached == STYLE_LOOKUPS
    assert len(lookups) > 50 * cached
