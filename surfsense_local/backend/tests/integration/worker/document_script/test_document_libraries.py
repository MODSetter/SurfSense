"""The libraries a document script is told it has work in the child process."""

import io
from pathlib import Path

import openpyxl
import pypdfium2
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

from worker.document_script.run import run_document_script

pytestmark = pytest.mark.integration


def test_a_python_docx_script_produces_a_word_file() -> None:
    """Word: a heading and a real table, read back by python-docx."""
    script = """\
import os
from docx import Document

document = Document()
document.add_heading("Client proposal", level=1)
table = document.add_table(rows=2, cols=2)
table.rows[0].cells[0].text = "Item"
table.rows[1].cells[0].text = "Design"
document.save(os.environ["OUTPUT_PATH"])
"""

    result = run_document_script(script, output_name="document.docx", images={})

    assert result.ok and result.output is not None
    document = Document(io.BytesIO(result.output))
    assert document.paragraphs[0].text == "Client proposal"
    assert document.tables[0].rows[1].cells[0].text == "Design"


def test_a_reportlab_script_produces_a_pdf() -> None:
    """PDF: ReportLab's platypus, read back by pypdfium2."""
    script = """\
import os
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate

styles = getSampleStyleSheet()
SimpleDocTemplate(os.environ["OUTPUT_PATH"], pagesize=A4).build(
    [Paragraph("Quarterly report", styles["Title"])]
)
"""

    result = run_document_script(script, output_name="document.pdf", images={})

    assert result.ok and result.output is not None
    pdf = pypdfium2.PdfDocument(result.output)
    try:
        assert len(pdf) == 1
        assert "Quarterly report" in pdf[0].get_textpage().get_text_range()
    finally:
        pdf.close()


def test_a_matplotlib_chart_is_drawn_without_a_display_and_placed_in_word() -> None:
    """A chart is saved as a PNG beside OUTPUT_PATH, then placed in the document."""
    script = """\
import os
import matplotlib
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm

assert matplotlib.get_backend().lower() == "agg"
chart = os.path.join(os.path.dirname(os.environ["OUTPUT_PATH"]), "costs.png")
figure, axes = plt.subplots()
axes.bar(["2024", "2025", "2026"], [120, 140, 165])
axes.set_title("Yearly costs")
figure.savefig(chart, dpi=100)

document = Document()
document.add_picture(chart, width=Cm(15))
document.save(os.environ["OUTPUT_PATH"])
"""

    result = run_document_script(script, output_name="document.docx", images={})

    assert result.ok, result.traceback_tail
    assert result.output is not None
    assert len(Document(io.BytesIO(result.output)).inline_shapes) == 1


def test_an_image_passed_in_is_readable_at_images_dir(tmp_path: Path) -> None:
    """A source figure named in the call is a PNG at IMAGES_DIR/<name>.png."""
    figure = tmp_path / "figure.png"
    Image.new("RGB", (40, 30), "red").save(figure)
    script = """\
import os
from PIL import Image

with Image.open(os.path.join(os.environ["IMAGES_DIR"], "12-1.png")) as image:
    size = f"{image.width}x{image.height}"
with open(os.environ["OUTPUT_PATH"], "w") as out:
    out.write(size)
"""

    result = run_document_script(
        script, output_name="size.txt", images={"12-1": figure}
    )

    assert result.ok
    assert result.output == b"40x30"


@pytest.mark.parametrize("name", ["../logo", "a/b", ""])
def test_an_image_name_must_be_a_plain_file_name(tmp_path: Path, name: str) -> None:
    """Images land inside IMAGES_DIR, never beside it."""
    figure = tmp_path / "figure.png"
    Image.new("RGB", (4, 4)).save(figure)

    with pytest.raises(ValueError):
        run_document_script("x = 1\n", output_name="a.pdf", images={name: figure})


def test_a_python_pptx_script_produces_a_deck() -> None:
    """PowerPoint: a title slide and a bullet slide, read back by python-pptx."""
    script = """\
import os
from pptx import Presentation
from pptx.util import Inches

deck = Presentation()
deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
cover = deck.slides.add_slide(deck.slide_layouts[0])
cover.shapes.title.text = "Quarterly review"
body = deck.slides.add_slide(deck.slide_layouts[1])
body.shapes.title.text = "Costs"
body.placeholders[1].text = "Costs rose in the north."
deck.save(os.environ["OUTPUT_PATH"])
"""

    result = run_document_script(script, output_name="document.pptx", images={})

    assert result.ok, result.traceback_tail
    deck = Presentation(io.BytesIO(result.output))
    assert [slide.shapes.title.text for slide in deck.slides] == [
        "Quarterly review",
        "Costs",
    ]


@pytest.mark.parametrize("library", ["xlsxwriter", "openpyxl"])
def test_an_excel_script_produces_a_workbook_with_a_formula(library: str) -> None:
    """Excel: both libraries the skill may name write a formula openpyxl reads back."""
    script = {
        "xlsxwriter": """\
import os
import xlsxwriter

book = xlsxwriter.Workbook(os.environ["OUTPUT_PATH"])
sheet = book.add_worksheet("Costs")
sheet.write_column(0, 0, [1, 2])
sheet.write_formula(2, 0, "=SUM(A1:A2)")
book.close()
""",
        "openpyxl": """\
import os
import openpyxl

book = openpyxl.Workbook()
sheet = book.active
sheet.title = "Costs"
sheet.append([1])
sheet.append([2])
sheet["A3"] = "=SUM(A1:A2)"
book.save(os.environ["OUTPUT_PATH"])
""",
    }[library]

    result = run_document_script(script, output_name="document.xlsx", images={})

    assert result.ok, result.traceback_tail
    book = openpyxl.load_workbook(io.BytesIO(result.output))
    assert book["Costs"]["A3"].value == "=SUM(A1:A2)"


def _template_deck(path: Path) -> bytes:
    """A one-slide deck to stand in for a source the user uploaded."""
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Old cover"
    deck.save(path)
    return path.read_bytes()


def test_a_template_is_a_copy_at_template_path(tmp_path: Path) -> None:
    """The script opens the source's file as TEMPLATE_PATH, and what it does to it stays in the run."""
    source = tmp_path / "Brand.pptx"
    original = _template_deck(source)
    script = """\
import os
from pptx import Presentation

deck = Presentation(os.environ["TEMPLATE_PATH"])
deck.slides[0].shapes.title.text = "New cover"
deck.save(os.environ["TEMPLATE_PATH"])
deck.save(os.environ["OUTPUT_PATH"])
"""

    result = run_document_script(
        script, output_name="document.pptx", images={}, template=source
    )

    assert result.ok, result.traceback_tail
    made = Presentation(io.BytesIO(result.output))
    assert made.slides[0].shapes.title.text == "New cover"
    assert source.read_bytes() == original


def test_without_a_template_there_is_no_template_path() -> None:
    """A script made without a template is not handed a path to a file that is not there."""
    script = """\
import os

with open(os.environ["OUTPUT_PATH"], "w") as out:
    out.write(str("TEMPLATE_PATH" in os.environ))
"""

    result = run_document_script(script, output_name="seen.txt", images={})

    assert result.ok, result.traceback_tail
    assert result.output == b"False"
