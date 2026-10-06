"""SurfSense's features on a real LibreOffice: a PDF of a Word file, and a workbook's real totals.

Runs the LibreOffice SURFSENSE_TEST_LIBREOFFICE names, else the one at the
standard install path, through an explicit runtime: never as Office support's
own choice, since a developer's install may be on an ended branch.
"""

import os
import sys
import time
from io import BytesIO
from pathlib import Path

import docx
import openpyxl
import pytest
import xlsxwriter

from modules.agent.previews.page_images import page_count
from modules.office_support.libreoffice import PackedLibreOffice
from modules.runtime_packs.office.program import program_in
from modules.runtime_packs.office.runtime import OfficeRuntime
from worker.studio.script_document.cached_values import with_cached_values

pytestmark = pytest.mark.office

# Generous: a fresh profile's first start takes about 5 s (measured), more on a busy machine.
DEADLINE_SECONDS = 120


def _installed() -> Path | None:
    named = os.environ.get("SURFSENSE_TEST_LIBREOFFICE")
    if named:
        return Path(named)
    if sys.platform == "win32":
        return Path(os.environ.get("PROGRAMW6432", r"C:\Program Files")) / "LibreOffice"
    if sys.platform == "darwin":
        return Path("/Applications/LibreOffice.app")
    return Path("/usr/lib/libreoffice")


@pytest.fixture
def office() -> PackedLibreOffice:
    """This machine's LibreOffice, named explicitly, as Settings would never pick an ended branch."""
    root = _installed()
    if root is None or not program_in(root).is_file():
        pytest.skip("no LibreOffice installed; set SURFSENSE_TEST_LIBREOFFICE")
    return PackedLibreOffice(OfficeRuntime(program_in(root), "test", "installed"))


def _deadline() -> float:
    return time.monotonic() + DEADLINE_SECONDS


def test_a_word_file_becomes_a_pdf_and_stays_as_it_was(
    office: PackedLibreOffice, tmp_path: Path
) -> None:
    """Every page is laid out, and LibreOffice only ever opens a copy."""
    document = docx.Document()
    document.add_heading("Client proposal", level=1)
    document.add_page_break()
    document.add_paragraph("Page two.")
    word = tmp_path / "proposal.docx"
    document.save(word)
    before = word.read_bytes()

    pdf = office.pdf_of(word, deadline=_deadline())

    assert pdf.startswith(b"%PDF-")
    assert page_count(pdf) == 2
    assert word.read_bytes() == before


def test_a_scripts_zero_totals_become_libreoffices_values_in_the_delivered_file(
    office: PackedLibreOffice, tmp_path: Path
) -> None:
    """xlsxwriter caches 0 for both; 04 measured 6 and 110 after LibreOffice."""
    buffer = BytesIO()
    book = xlsxwriter.Workbook(buffer, {"in_memory": True})
    sheet = book.add_worksheet("Costs")
    sheet.write_column("A1", [2, 3, 50])
    sheet.write_formula("B1", "=A1*A2")
    sheet.write_formula("B2", "=SUM(A1:A3)*2")
    book.close()
    data = buffer.getvalue()
    workbook = tmp_path / "costs.xlsx"
    workbook.write_bytes(data)

    values = office.values_of(workbook, deadline=_deadline())
    delivered = with_cached_values(data, values)

    cached = openpyxl.load_workbook(BytesIO(delivered.data), data_only=True)["Costs"]
    assert (cached["B1"].value, cached["B2"].value) == (6, 110)
    formulas = openpyxl.load_workbook(BytesIO(delivered.data))["Costs"]
    assert formulas["B2"].value == "=SUM(A1:A3)*2"
    assert (delivered.written, delivered.left_blank) == (2, 0)


def test_every_cell_an_array_formula_fills_gets_libreoffices_value(
    office: PackedLibreOffice, tmp_path: Path
) -> None:
    """xlsxwriter caches 0 in each cell of the range, not only the first."""
    buffer = BytesIO()
    book = xlsxwriter.Workbook(buffer, {"in_memory": True})
    sheet = book.add_worksheet("Costs")
    sheet.write_column("A1", [2, 3, 50])
    sheet.write_array_formula("B1:B3", "{=A1:A3*2}")
    book.close()
    data = buffer.getvalue()
    workbook = tmp_path / "costs.xlsx"
    workbook.write_bytes(data)

    values = office.values_of(workbook, deadline=_deadline())
    delivered = with_cached_values(data, values)

    cached = openpyxl.load_workbook(BytesIO(delivered.data), data_only=True)["Costs"]
    assert [cached[f"B{row}"].value for row in (1, 2, 3)] == [4, 6, 100]
