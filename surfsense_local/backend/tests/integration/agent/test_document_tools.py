"""The document tools as opencode's MCP client calls them: render, read, and list images."""

import re
import sqlite3
import threading
import time
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from sqlalchemy import Engine, select, update

from modules.agent.previews import Previews
from modules.agent.tool_endpoint import list_images, render_document
from modules.artifacts.models import Artifact
from modules.artifacts.script_documents.service import create_script_document
from modules.artifacts.script_documents.spec import DocumentSpec
from modules.artifacts.script_documents.version import ArtifactVersion
from modules.artifacts.studio_documents.recipe import RECIPE_KEY, drafted
from modules.artifacts.studio_documents.service import create_refine_version
from modules.documents.models import Document, DocumentStatus, DocumentType
from modules.documents.source_figures.layout import figures_dir, write_index
from modules.documents.tasks import extract_figures
from modules.workspaces.models import Workspace
from shared.config import get_storage_settings
from shared.db import create_session_factory
from shared.queue import ingest_queue
from tests.integration.agent.conftest import declare_image_input
from tests.integration.agent.tool_endpoint_client import ToolEndpoint
from tests.integration.worker.conftest import stub_model  # noqa: F401

# The job indexes what the script wrote; the stub stands in for the embedder.
pytestmark = [
    pytest.mark.integration,
    pytest.mark.usefixtures("stub_model", "model_reads_images"),
]

STOP = "If this is your third failed run for this request, stop and tell the user what failed."

WORD = """\
import os
import docx

document = docx.Document()
document.add_heading("Client proposal", level=1)
document.add_paragraph("We propose a two-phase rollout for Halvorsen Freight.")
document.add_paragraph("{closing}")
document.save(os.environ["OUTPUT_PATH"])
"""

PDF = """\
import os
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

drawing = canvas.Canvas(os.environ["OUTPUT_PATH"], pagesize=A4)
for page in ("Cover", "Pricing"):
    drawing.drawString(72, 760, page)
    drawing.showPage()
drawing.save()
"""

LOGO = """\
import os
import docx

document = docx.Document()
document.add_picture(os.path.join(os.environ["IMAGES_DIR"], "{name}.png"))
document.save(os.environ["OUTPUT_PATH"])
"""

FAILING = "rows = []\nraise ValueError('the pricing table is empty')\n"

# A page 2 pt wide and 200 inches tall, as a script steered by a source might make.
STRIP = """\
import os
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

drawing = canvas.Canvas(os.environ["OUTPUT_PATH"], pagesize=A4)
drawing.drawString(72, 760, "Cover")
drawing.showPage()
drawing.setPageSize((2, 14400))
drawing.showPage()
drawing.save()
"""


def render(**arguments: object) -> dict[str, object]:
    """A render call's arguments: a Word document unless told otherwise."""
    return {
        "title": "Client proposal",
        "format": "docx",
        "script": WORD.format(closing="Signed."),
        **arguments,
    }


def _artifact_id(text: str) -> int:
    """The artifact a successful render names on its first line."""
    return int(text.split(",")[0].removeprefix("Rendered artifact "))


async def _listed(tools: ToolEndpoint, workspace_id: int) -> list[dict]:
    reply = await tools.client.get(f"/workspaces/{workspace_id}/artifacts")
    return reply.json()


def _source(
    engine: Engine,
    workspace_id: int,
    status: DocumentStatus = DocumentStatus.READY,
    document_type: DocumentType = DocumentType.FILE,
) -> int:
    with create_session_factory(engine)() as session:
        source = Document(
            workspace_id=workspace_id,
            title="Logo.png",
            document_type=document_type,
            status=status,
            content="Halvorsen Freight logo",
        )
        session.add(source)
        session.commit()
        return source.id


def _logo_source(engine: Engine, workspace_id: int) -> int:
    """A logo uploaded as a source, its one figure kept as ingest keeps it."""
    source_id = _source(engine, workspace_id)
    folder = get_storage_settings().document_dir(workspace_id, source_id)
    folder.mkdir(parents=True)
    logo = BytesIO()
    Image.new("RGB", (40, 20), "red").save(logo, format="PNG")
    (folder / "Logo.png").write_bytes(logo.getvalue())
    extract_figures.call_local(source_id)
    return source_id


def _report_source(engine: Engine, workspace_id: int) -> int:
    """A report whose ingest kept two figures, one with a caption."""
    source_id = _source(engine, workspace_id)
    folder = figures_dir(get_storage_settings().document_dir(workspace_id, source_id))
    folder.mkdir(parents=True)
    for n in (1, 2):
        Image.new("RGB", (1200, 800), "blue").save(folder / f"{n}.png")
    write_index(
        folder,
        [
            {
                "n": 1,
                "caption": "Figure 2. Yearly costs",
                "page": 3,
                "width": 1200,
                "height": 800,
            },
            {"n": 2, "caption": None, "page": 5, "width": 1200, "height": 800},
        ],
        None,
    )
    return source_id


async def test_a_render_waits_for_the_ready_document_and_says_what_it_made(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """The model learns the artifact, its version, its size and its opening text."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False, text
    artifact_id = _artifact_id(text)
    assert text.splitlines()[0] == (
        f"Rendered artifact {artifact_id}, version 1: Client proposal"
    )
    assert "3 paragraphs" in text
    assert "We propose a two-phase rollout for Halvorsen Freight." in text
    (listed,) = await _listed(tools, workspace_id)
    assert (listed["id"], listed["status"], listed["spec_kind"]) == (
        artifact_id,
        "ready",
        "python",
    )
    assert listed["version"] == {"root_id": artifact_id, "number": 1, "parent_id": None}


async def test_a_render_naming_its_artifact_makes_the_next_version(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """An edit keeps the document: same root, the next number, the old version kept."""
    workspace_id = await tools.workspace()
    first, _ = await tools.call(workspace_id, "render_document", render())
    first_id = _artifact_id(first)

    second, is_error = await tools.call(
        workspace_id,
        "render_document",
        render(script=WORD.format(closing="Kind regards."), artifact_id=first_id),
    )

    assert is_error is False, second
    second_id = _artifact_id(second)
    assert second.startswith(f"Rendered artifact {second_id}, version 2:")
    assert "Kind regards." in second
    versions = {a["id"]: a["version"] for a in await _listed(tools, workspace_id)}
    assert versions[second_id] == {
        "root_id": first_id,
        "number": 2,
        "parent_id": first_id,
    }


async def test_a_pdf_render_counts_its_pages_and_lists_previews_to_open(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """The agent opens each preview with `read`, relative to its own folder."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(
        workspace_id, "render_document", render(format="pdf", script=PDF)
    )

    assert is_error is False, text
    artifact_id = _artifact_id(text)
    assert "2 pages" in text
    folder = tools.folder(workspace_id)
    previews = [
        f"outputs/previews/{artifact_id}-v1/page-1.png",
        f"outputs/previews/{artifact_id}-v1/page-2.png",
    ]
    for preview in previews:
        assert f"- {preview}" in text
        assert (folder / preview).is_file()
    assert "headers and footers" not in text


async def test_a_model_that_cannot_read_images_is_drawn_no_previews(
    tools: ToolEndpoint, studio_worker: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    """It would be shown none of them, and a Word preview waits up to 30 s on Electron."""
    declare_image_input(False)
    drawn: list[object] = []
    monkeypatch.setattr(
        render_document, "previews_for", lambda *args, **kwargs: drawn.append(args)
    )
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False, text
    assert drawn == []
    assert "No page previews: the selected model cannot read images" in text


async def test_a_word_render_without_the_desktop_app_says_why_it_has_no_previews(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """Word pages are drawn by Electron, which no test and no Docker stack runs."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False, text
    assert "No page previews: Word pages are drawn by the SurfSense desktop app" in text


async def test_word_previews_say_they_leave_out_headers_and_footers(
    tools: ToolEndpoint, studio_worker: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A header logo the preview cannot show is not a missing image to render again for."""
    workspace_id = await tools.workspace()

    def previews_for(artifact: Artifact, folder: Path, time_left: float) -> Previews:
        return Previews(
            [folder / "outputs" / "previews" / f"{artifact.id}-v1" / "page-1.png"]
        )

    monkeypatch.setattr(render_document, "previews_for", previews_for)

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False, text
    assert "-v1/page-1.png" in text
    assert "Word previews leave out headers and footers" in text


async def test_a_page_too_long_and_thin_to_draw_is_named_beside_the_others(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """The pages that could be drawn are listed, and the model learns which one was not."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(
        workspace_id, "render_document", render(format="pdf", script=STRIP)
    )

    assert is_error is False, text
    artifact_id = _artifact_id(text)
    assert f"- outputs/previews/{artifact_id}-v1/page-1.png" in text
    assert "page-2.png" not in text
    assert "Page 2 was not drawn" in text


async def test_the_word_preview_gets_only_the_time_the_call_has_left(
    tools: ToolEndpoint, studio_worker: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    """opencode gives up on the whole call, so the preview cannot have its full 30 s late in it."""
    given: list[float] = []

    def previews_for(_artifact: Artifact, _folder: Path, time_left: float) -> Previews:
        given.append(time_left)
        return Previews([], "not drawn in this test")

    monkeypatch.setattr(render_document, "previews_for", previews_for)
    monkeypatch.setattr(render_document, "CALL_SECONDS", 20)
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False, text
    (time_left,) = given
    assert 0 < time_left < 20


async def test_a_render_stops_waiting_when_the_call_must_answer(
    tools: ToolEndpoint, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Time spent before the wait comes out of it: the call answers before opencode gives up."""
    monkeypatch.setattr(render_document, "CALL_SECONDS", 1)
    workspace_id = await tools.workspace()
    started = time.monotonic()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert time.monotonic() - started < render_document.WAIT_SECONDS / 10
    assert is_error is False
    assert "appears in Studio when it is ready" in text


def _fail_while_holding_the_write_lock(engine: Engine, seconds: float) -> None:
    """Fail the first artifact's job as the worker would, holding the lock past the busy wait."""
    with create_session_factory(engine)() as session:
        while (document_id := session.scalar(select(Artifact.document_id))) is None:
            session.commit()
            time.sleep(0.05)
        session.execute(
            update(Document)
            .where(Document.id == document_id)
            .values(status=DocumentStatus.FAILED, error_message="the disk is full")
        )
        time.sleep(seconds)
        session.commit()


async def test_a_render_waits_out_a_write_that_holds_the_database_past_the_busy_wait(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """Studio holds the write lock while it embeds a long document; the wait looks again."""
    workspace_id = await tools.workspace()
    # SQLite's busy wait gives up after 5 s (shared.db).
    worker = threading.Thread(
        target=_fail_while_holding_the_write_lock, args=(engine, 6.0), daemon=True
    )
    worker.start()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    worker.join()
    assert is_error is True
    assert "the disk is full" in text
    assert text.endswith(STOP)


async def test_a_named_source_image_reaches_the_script(
    tools: ToolEndpoint, engine: Engine, studio_worker: None
) -> None:
    """The figure the call names is in IMAGES_DIR under that name."""
    workspace_id = await tools.workspace()
    name = f"{_logo_source(engine, workspace_id)}-1"

    text, is_error = await tools.call(
        workspace_id,
        "render_document",
        render(script=LOGO.format(name=name), images=[name]),
    )

    assert is_error is False, text


async def test_a_failing_script_returns_its_error_and_the_stop_rule(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """The model fixes its script from the traceback, and stops after a third failure."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(
        workspace_id, "render_document", render(script=FAILING)
    )

    assert is_error is True
    (failed,) = await _listed(tools, workspace_id)
    assert failed["status"] == "failed"
    assert "ValueError: the pricing table is empty" in text
    assert "line 2, in <module>" in text  # the traceback points into the script
    assert f"artifact_id {failed['id']}" in text
    assert "Fix the script" in text
    assert text.endswith(STOP)


async def test_a_version_that_failed_outside_its_script_is_not_blamed_on_it(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """A full disk or a busy database is no reason to rewrite a script that was right."""
    workspace_id = await tools.workspace()
    worker = threading.Thread(
        target=_fail_while_holding_the_write_lock, args=(engine, 0.0), daemon=True
    )
    worker.start()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    worker.join()
    (failed,) = await _listed(tools, workspace_id)
    assert is_error is True
    assert "the disk is full" in text
    assert "Fix the script" not in text
    assert f"same script again with artifact_id {failed['id']}" in text
    assert text.endswith(STOP)


async def test_a_ready_version_is_reported_though_the_database_was_busy(
    tools: ToolEndpoint,
    engine: Engine,
    studio_worker: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Another writer holding the lock past the busy wait must not hide a version that was made."""
    settled = render_document.wait_for_outcome

    def settle_then_lock(*args: object) -> object:
        outcome = settled(*args)
        held = threading.Event()
        threading.Thread(
            target=_hold_the_write_lock, args=(engine, held, 6.0), daemon=True
        ).start()
        held.wait(10)
        return outcome

    monkeypatch.setattr(render_document, "wait_for_outcome", settle_then_lock)
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False, text
    assert text.startswith("Rendered artifact ")
    assert "Halvorsen Freight" in text


def _hold_the_write_lock(engine: Engine, held: threading.Event, seconds: float) -> None:
    """Hold the write lock past SQLite's 5 s busy wait (shared.db)."""
    with create_session_factory(engine)() as session:
        session.execute(select(1))
        held.set()
        time.sleep(seconds)
        session.commit()


async def test_a_render_still_running_at_the_limit_says_where_it_will_appear(
    tools: ToolEndpoint, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Studio may be busy with other jobs; the model must not render it again."""
    monkeypatch.setattr(render_document, "WAIT_SECONDS", 0.5)
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False
    (pending,) = await _listed(tools, workspace_id)
    assert pending["status"] == "pending"
    assert f"artifact {pending['id']}" in text
    assert "appears in Studio when it is ready" in text


@pytest.mark.parametrize(
    ("arguments", "says"),
    [
        ({"title": "", "format": "docx", "script": "x"}, "title"),
        (
            {"title": "Proposal", "format": "odt", "script": "x"},
            "docx, pdf, pptx or xlsx",
        ),
        ({"title": "Proposal", "format": "docx"}, "script"),
        (
            {"title": "Proposal", "format": "docx", "script": "x", "artifact_id": "1"},
            "artifact_id",
        ),
        (
            {"title": "Proposal", "format": "docx", "script": "x", "images": "7-1"},
            "images",
        ),
    ],
)
async def test_a_render_missing_what_it_needs_says_what(
    tools: ToolEndpoint, arguments: dict, says: str
) -> None:
    """A refusal names the argument to fix, and nothing is made."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", arguments)

    assert is_error is True
    assert says in text
    assert await _listed(tools, workspace_id) == []


async def test_optional_arguments_sent_as_null_are_left_out(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """Many models fill every optional field, with null where they mean none."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(
        workspace_id, "render_document", render(artifact_id=None, images=None)
    )

    assert is_error is False, text
    assert text.startswith("Rendered artifact ")


async def test_a_render_cannot_continue_another_workspaces_document(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """A workspace's tools edit that workspace's documents only."""
    workspace_id, elsewhere = await tools.workspace(), await tools.workspace()
    other, _ = await tools.call(elsewhere, "render_document", render())

    text, is_error = await tools.call(
        workspace_id, "render_document", render(artifact_id=_artifact_id(other))
    )

    assert is_error is True
    assert "no artifact" in text
    assert await _listed(tools, workspace_id) == []


async def test_a_render_naming_an_image_no_source_has_is_refused(
    tools: ToolEndpoint,
) -> None:
    """A script can only be handed images the workspace's sources hold."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(
        workspace_id, "render_document", render(images=["999-1"])
    )

    assert is_error is True
    assert '"999-1"' in text


async def test_reading_any_version_returns_the_newest_script(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """An edit starts from the newest version, whichever one the model remembers."""
    workspace_id = await tools.workspace()
    first, _ = await tools.call(workspace_id, "render_document", render())
    first_id = _artifact_id(first)
    newest_script = WORD.format(closing="Kind regards.")
    second, _ = await tools.call(
        workspace_id,
        "render_document",
        render(script=newest_script, artifact_id=first_id),
    )

    text, is_error = await tools.call(
        workspace_id, "read_document", {"artifact_id": first_id}
    )

    assert is_error is False, text
    second_id = _artifact_id(second)
    assert f"version 2, artifact {second_id}" in text
    assert "Word document" in text
    assert newest_script in text
    lines = newest_script.count("\n")
    assert f"lines 1-{lines} of {lines}" in text
    assert "offset" not in text


# opencode cuts a tool result past either and keeps the whole in a folder every
# workspace's agent can read (its Truncate).
OPENCODE_MAX_LINES = 2000
OPENCODE_MAX_BYTES = 50 * 1024
_PAGE = re.compile(r"Script, lines ([0-9]+)-([0-9]+) of ([0-9]+):\n")
# After the page's own last newline, a blank line and where the next page starts.
_NEXT = re.compile(r"(?<=\n)\nThe script goes on .* offset ([0-9]+) .*\Z", re.DOTALL)


async def _read_in_pages(
    tools: ToolEndpoint, workspace_id: int, artifact_id: int
) -> tuple[str, list[str]]:
    """The script, put back together from every page the model would read; and the pages."""
    pages: list[str] = []
    script = ""
    arguments: dict[str, object] = {"artifact_id": artifact_id}
    while True:
        text, is_error = await tools.call(workspace_id, "read_document", arguments)
        assert is_error is False, text
        pages.append(text)
        first, last, total = map(int, _PAGE.search(text).groups())
        assert first == arguments.get("offset", 1)
        body = text[_PAGE.search(text).end() :]
        following = _NEXT.search(body)
        if following is None:
            assert last == total
            return script + body, pages
        assert int(following[1]) == last + 1
        script += body[: following.start()]
        arguments = {"artifact_id": artifact_id, "offset": last + 1}


@pytest.mark.parametrize(
    "padding",
    [
        pytest.param([f"# note {n}\n" for n in range(2500)], id="many-lines"),
        pytest.param([f"# {'x' * 300}\n" for _ in range(400)], id="long-lines"),
    ],
)
async def test_a_long_script_is_read_in_pages_opencode_keeps_whole(
    tools: ToolEndpoint, studio_worker: None, padding: list[str]
) -> None:
    """Each page stays under opencode's cut, and the pages make up the script exactly."""
    workspace_id = await tools.workspace()
    script = WORD.format(closing="Signed.") + "".join(padding)
    rendered, _ = await tools.call(
        workspace_id, "render_document", render(script=script)
    )

    read, pages = await _read_in_pages(tools, workspace_id, _artifact_id(rendered))

    assert read == script
    assert len(pages) > 1
    for page in pages:
        assert page.count("\n") < OPENCODE_MAX_LINES
        assert len(page.encode()) < OPENCODE_MAX_BYTES
        assert "Its newest is version 1" in page
    assert "surfsense_read_document" in pages[0]


@pytest.mark.parametrize("offset", [0, 99, "2", True])
async def test_reading_from_a_line_the_script_does_not_have_is_refused(
    tools: ToolEndpoint, studio_worker: None, offset: object
) -> None:
    """The model learns which offsets there are rather than reading nothing."""
    workspace_id = await tools.workspace()
    rendered, _ = await tools.call(workspace_id, "render_document", render())

    text, is_error = await tools.call(
        workspace_id,
        "read_document",
        {"artifact_id": _artifact_id(rendered), "offset": offset},
    )

    assert is_error is True
    lines = WORD.format(closing="Signed.").count("\n")
    assert f"from 1 to {lines}" in text


async def test_an_offset_sent_as_null_reads_from_the_top(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """Null is how many models leave an optional field out."""
    workspace_id = await tools.workspace()
    rendered, _ = await tools.call(workspace_id, "render_document", render())

    text, is_error = await tools.call(
        workspace_id,
        "read_document",
        {"artifact_id": _artifact_id(rendered), "offset": None},
    )

    assert is_error is False, text
    assert WORD.format(closing="Signed.") in text


async def test_reading_a_document_studio_drafted_says_it_has_no_script(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """Studio's own buttons keep no script, so there is nothing to edit here."""
    workspace_id = await tools.workspace()
    with create_session_factory(engine)() as session:
        document = Document(
            workspace_id=workspace_id,
            title="Word",
            document_type=DocumentType.ARTIFACT,
            status=DocumentStatus.READY,
        )
        session.add(document)
        session.flush()
        drafted = Artifact(
            document_id=document.id, workspace_id=workspace_id, format="docx"
        )
        session.add(drafted)
        session.commit()
        drafted_id = drafted.id

    text, is_error = await tools.call(
        workspace_id, "read_document", {"artifact_id": drafted_id}
    )

    assert is_error is True
    assert "no script" in text


async def test_a_document_studio_drafted_in_markdown_cannot_be_continued_here(
    tools: ToolEndpoint, engine: Engine, studio_worker: None
) -> None:
    """Studio's Markdown is Refine's to rewrite; a script written blind would join its versions."""
    workspace_id = await tools.workspace()
    with create_session_factory(engine)() as session:
        document = Document(
            workspace_id=workspace_id,
            title="Report",
            document_type=DocumentType.ARTIFACT,
            status=DocumentStatus.READY,
        )
        session.add(document)
        session.flush()
        drafted_in_studio = Artifact(
            document_id=document.id, workspace_id=workspace_id, format="docx"
        )
        session.add(drafted_in_studio)
        session.flush()
        drafted_in_studio.artifact_metadata = {
            "spec": DocumentSpec(
                "markdown", "# Report\n\nText.", "docx", ()
            ).as_metadata(),
            "version": ArtifactVersion(drafted_in_studio.id, 1, None).as_metadata(),
            RECIPE_KEY: drafted(),
            "source_document_ids": [],
            "prompt": None,
        }
        session.commit()
        drafted_id = drafted_in_studio.id

    text, is_error = await tools.call(
        workspace_id, "render_document", render(artifact_id=drafted_id)
    )

    assert is_error is True, text
    assert "no script" in text
    assert [a["id"] for a in await _listed(tools, workspace_id)] == [drafted_id]


async def test_reading_another_workspaces_document_is_refused(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """Another workspace is someone else's research."""
    workspace_id, elsewhere = await tools.workspace(), await tools.workspace()
    other, _ = await tools.call(elsewhere, "render_document", render())

    text, is_error = await tools.call(
        workspace_id, "read_document", {"artifact_id": _artifact_id(other)}
    )

    assert is_error is True
    assert "Client proposal" not in text


async def test_a_sources_images_are_listed_with_their_captions_and_sizes(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """The model picks an image by its caption and page, and places it by name."""
    workspace_id = await tools.workspace()
    report = _report_source(engine, workspace_id)

    text, is_error = await tools.call(
        workspace_id, "list_images", {"source_ids": [report]}
    )

    assert is_error is False, text
    assert (
        f'- {report}-1: 1200x800 px, page 3, caption "Figure 2. Yearly costs"' in text
    )
    assert f"- {report}-2: 1200x800 px, page 5, no caption" in text


async def test_listed_images_are_copied_with_the_write_lock_free(
    tools: ToolEndpoint, engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A report can have hundreds of figures; Studio and the API must not wait on their copies."""
    workspace_id = await tools.workspace()
    report = _report_source(engine, workspace_id)
    others_could_write: list[bool] = []
    show = list_images.show_figure

    def show_while_another_writes(*args: object) -> str:
        other = sqlite3.connect(
            get_storage_settings().database_path, timeout=0.1, isolation_level=None
        )
        try:
            other.execute("BEGIN IMMEDIATE")
            other.execute("ROLLBACK")
            others_could_write.append(True)
        except sqlite3.OperationalError:
            others_could_write.append(False)
        finally:
            other.close()
        return show(*args)

    monkeypatch.setattr(list_images, "show_figure", show_while_another_writes)

    text, is_error = await tools.call(
        workspace_id, "list_images", {"source_ids": [report]}
    )

    assert is_error is False, text
    assert others_could_write == [True, True]
    assert f"at sources/figures/{report}-2.png" in text


async def test_a_listed_image_can_be_opened_from_the_sources_folder(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """The agent looks at a chart, to read its values, before it draws a new one."""
    workspace_id = await tools.workspace()
    report = _report_source(engine, workspace_id)
    storage = get_storage_settings()

    text, _ = await tools.call(workspace_id, "list_images", {"source_ids": [report]})

    kept = figures_dir(storage.document_dir(workspace_id, report))
    shown = tools.folder(workspace_id) / "sources" / "figures"
    for n in (1, 2):
        assert f"sources/figures/{report}-{n}.png" in text
        copied = (shown / f"{report}-{n}.png").read_bytes()
        assert copied == (kept / f"{n}.png").read_bytes()


async def test_an_image_source_is_its_own_image(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """A logo uploaded as a source is the one image it holds."""
    workspace_id = await tools.workspace()
    logo = _logo_source(engine, workspace_id)

    text, _ = await tools.call(workspace_id, "list_images", {"source_ids": [logo]})

    assert f"- {logo}-1: 40x20 px" in text


async def test_a_source_ingested_before_images_were_kept_says_they_are_coming(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """The figures-only pass is queued once, and the model is told to ask again."""
    workspace_id = await tools.workspace()
    older = _source(engine, workspace_id)
    folder = get_storage_settings().document_dir(workspace_id, older)
    folder.mkdir(parents=True)
    Image.new("RGB", (40, 20), "red").save(folder / "Logo.png")

    text, is_error = await tools.call(
        workspace_id, "list_images", {"source_ids": [older]}
    )

    assert is_error is False
    assert f"Source {older}" in text
    assert "being extracted" in text
    assert [task.args for task in ingest_queue.pending()] == [(older,)]


async def test_listing_says_which_sources_have_no_images_or_are_not_here(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """Each source the call names gets an answer, so none is mistaken for empty."""
    workspace_id, elsewhere = await tools.workspace(), await tools.workspace()
    note = _source(engine, workspace_id, document_type=DocumentType.NOTE)
    foreign = _logo_source(engine, elsewhere)

    text, is_error = await tools.call(
        workspace_id, "list_images", {"source_ids": [note]}
    )
    theirs, refused = await tools.call(
        workspace_id, "list_images", {"source_ids": [note, foreign]}
    )

    assert is_error is False
    assert f"Source {note}" in text and "no images" in text
    # Another workspace's source is never among this thread's.
    assert refused is True
    assert f"Source {foreign} is not selected" in theirs
    assert f"{foreign}-1" not in theirs
    shown = tools.folder(workspace_id) / "sources"
    assert not (shown / "figures").exists()


async def test_listing_without_sources_says_what_to_name(tools: ToolEndpoint) -> None:
    """A small model may leave the argument out; the refusal says how to fill it."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "list_images", {})

    assert is_error is True
    assert "sources/" in text


async def test_a_preview_that_breaks_still_reports_the_ready_version(
    tools: ToolEndpoint, studio_worker: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The version is made; an error result would send the model to make it again."""

    def broken(_artifact: Artifact, _folder: Path, time_left: float) -> None:
        raise OSError("disk full")

    monkeypatch.setattr(render_document, "previews_for", broken)
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False
    assert text.startswith("Rendered artifact ")
    assert "No page previews: drawing them failed." in text


async def test_reading_a_document_refined_in_studio_skips_the_version_with_no_script_yet(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """A Refine's version keeps its script only once it renders; until then, and
    if it fails, the newest script is the one before it. Refine is Studio's, so
    the first version is a script Studio drafted."""
    workspace_id = await tools.workspace()
    with create_session_factory(engine)() as session:
        first = create_script_document(
            session,
            session.get(Workspace, workspace_id),
            title="Client proposal",
            format="docx",
            script=WORD.format(closing="Signed."),
            base_artifact_id=None,
            image_names=[],
        )
        first_id = first.id
        first.document.status = DocumentStatus.READY  # as its run left it
        first.artifact_metadata = {**first.artifact_metadata, RECIPE_KEY: drafted()}
        session.commit()
        # As the refine route leaves it: pending, its rewrite not written yet.
        create_refine_version(session, first_id, "Shorter")

    text, is_error = await tools.call(
        workspace_id, "read_document", {"artifact_id": first_id}
    )

    assert is_error is False, text
    assert f"Its newest is version 1, artifact {first_id}" in text
    assert "We propose a two-phase rollout" in text
