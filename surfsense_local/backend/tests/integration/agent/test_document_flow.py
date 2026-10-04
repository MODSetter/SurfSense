"""The demo's document flow through the tool endpoint, the model's calls played in order.

A logo source, a Word proposal with that logo and a chart, an edit to it, and a
PDF version the agent can look at: the calls a model makes across the three turns
of 07-create-and-edit-mvp, with no model and no network.
"""

from io import BytesIO

import docx
import pytest
from PIL import Image
from sqlalchemy import Engine

from modules.documents.models import Document, DocumentStatus, DocumentType
from modules.documents.tasks import extract_figures
from shared.config import get_storage_settings
from shared.db import create_session_factory
from tests.integration.agent.tool_endpoint_client import ToolEndpoint
from tests.integration.worker.conftest import stub_model  # noqa: F401

# The job indexes what the script wrote; the stub stands in for the embedder.
pytestmark = [
    pytest.mark.integration,
    pytest.mark.usefixtures("stub_model", "model_reads_images"),
]

CHART = """\
import os
import matplotlib.pyplot as plt

figure, axes = plt.subplots(figsize=(6, 3))
axes.bar(["2025", "2026", "2027"], [120, 95, 80])
axes.set_title("Yearly costs")
chart = os.path.join(os.path.dirname(os.environ["OUTPUT_PATH"]), "costs.png")
figure.savefig(chart, dpi=100)
logo = os.path.join(os.environ["IMAGES_DIR"], "{logo}.png")
"""

WORD = (
    CHART
    + """\
import docx
from docx.shared import Inches

document = docx.Document()
document.add_picture(logo, width=Inches(1))
document.add_heading("{heading}", level=1)
document.add_paragraph("A two-phase rollout for Halvorsen Freight.")
document.add_picture(chart, width=Inches(6))
document.save(os.environ["OUTPUT_PATH"])
"""
)

PDF = (
    CHART
    + """\
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate

styles = getSampleStyleSheet()
SimpleDocTemplate(os.environ["OUTPUT_PATH"], pagesize=A4).build(
    [
        Image(logo, width=72, height=36),
        Paragraph("Client proposal", styles["Title"]),
        PageBreak(),
        Image(chart, width=432, height=216),
    ]
)
"""
)


def _logo_source(engine: Engine, workspace_id: int) -> int:
    """A logo uploaded as a source, its figure kept by the real figures pass."""
    with create_session_factory(engine)() as session:
        source = Document(
            workspace_id=workspace_id,
            title="Logo.png",
            document_type=DocumentType.FILE,
            status=DocumentStatus.READY,
            content="Halvorsen Freight logo",
        )
        session.add(source)
        session.commit()
        source_id = source.id
    folder = get_storage_settings().document_dir(workspace_id, source_id)
    folder.mkdir(parents=True)
    logo = BytesIO()
    Image.new("RGB", (200, 100), "red").save(logo, format="PNG")
    (folder / "Logo.png").write_bytes(logo.getvalue())
    extract_figures.call_local(source_id)
    return source_id


def _made(text: str) -> tuple[int, int]:
    """The artifact and version a successful render names on its first line."""
    first = text.splitlines()[0].removeprefix("Rendered artifact ")
    artifact, version = first.split(":")[0].split(", version ")
    return int(artifact), int(version)


async def _pictures_in_word(tools: ToolEndpoint, artifact_id: int) -> int:
    reply = await tools.client.get(f"/artifacts/{artifact_id}/files/primary")
    reply.raise_for_status()
    return len(docx.Document(BytesIO(reply.content)).inline_shapes)


async def test_the_agent_drafts_edits_and_checks_a_document_from_its_sources(
    tools: ToolEndpoint, engine: Engine, studio_worker: None
) -> None:
    """Each version keeps its script, edits share a root, and the PDF has pages to open."""
    workspace_id = await tools.workspace()
    source_id = _logo_source(engine, workspace_id)

    # Turn 1: find the logo, then draft the Word document with it and a chart.
    images, is_error = await tools.call(
        workspace_id, "list_images", {"source_ids": [source_id]}
    )
    assert is_error is False, images
    logo = f"{source_id}-1"
    assert f"- {logo}: 200x100 px" in images

    first_script = WORD.format(logo=logo, heading="Client proposal")
    first, is_error = await tools.call(
        workspace_id,
        "render_document",
        {
            "title": "Client proposal",
            "format": "docx",
            "script": first_script,
            "images": [logo],
        },
    )
    assert is_error is False, first
    first_id, first_version = _made(first)
    assert first_version == 1
    assert "A two-phase rollout for Halvorsen Freight." in first
    assert await _pictures_in_word(tools, first_id) == 2

    # Turn 2: read the script back, and render the edit as the next version.
    script, is_error = await tools.call(
        workspace_id, "read_document", {"artifact_id": first_id}
    )
    assert is_error is False, script
    assert f"version 1, artifact {first_id}" in script
    assert f"Source images it places: {logo}" in script
    assert first_script in script

    second, is_error = await tools.call(
        workspace_id,
        "render_document",
        {
            "title": "Client proposal",
            "format": "docx",
            "script": WORD.format(logo=logo, heading="Proposal for Halvorsen"),
            "images": [logo],
            "artifact_id": first_id,
        },
    )
    assert is_error is False, second
    second_id, second_version = _made(second)
    assert second_version == 2
    assert "Proposal for Halvorsen" in second
    assert await _pictures_in_word(tools, second_id) == 2

    # Turn 3: a PDF for the client, a document of its own, with pages to check.
    pdf, is_error = await tools.call(
        workspace_id,
        "render_document",
        {
            "title": "Client proposal (PDF)",
            "format": "pdf",
            "script": PDF.format(logo=logo),
            "images": [logo],
        },
    )
    assert is_error is False, pdf
    pdf_id, pdf_version = _made(pdf)
    assert pdf_version == 1
    assert "2 pages" in pdf
    agent_folder = get_storage_settings().agent_working_dir(workspace_id)
    previews = [
        line.removeprefix("- ") for line in pdf.splitlines() if "previews/" in line
    ]
    assert previews == [
        f"outputs/previews/{pdf_id}-v1/page-1.png",
        f"outputs/previews/{pdf_id}-v1/page-2.png",
    ]
    for preview in previews:
        with Image.open(agent_folder / preview) as page:
            assert page.width == 1000

    listed = await tools.client.get(f"/workspaces/{workspace_id}/artifacts")
    versions = {a["id"]: a["version"] for a in listed.json()}
    assert versions == {
        first_id: {"root_id": first_id, "number": 1, "parent_id": None},
        second_id: {"root_id": first_id, "number": 2, "parent_id": first_id},
        pdf_id: {"root_id": pdf_id, "number": 1, "parent_id": None},
    }
