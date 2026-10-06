"""Converting a Word document or deck to PDF, as opencode's MCP client calls the tool.

LibreOffice is played by a stand-in; Studio's worker runs on a thread.
"""

from io import BytesIO
from pathlib import Path

import docx
import pytest
from sqlalchemy import Engine

from modules.artifacts.models import Artifact, ArtifactFileRole
from modules.documents.models import DocumentStatus
from modules.source_scope.schemas import SourceScope
from shared.config import get_storage_settings
from shared.db import create_session_factory
from tests.integration.agent.test_office_documents import _uploaded
from tests.integration.agent.tool_endpoint_client import ToolEndpoint
from tests.integration.worker.conftest import stub_model  # noqa: F401
from tests.office_stand_in import office_on

# The job indexes the PDF's text; the stub stands in for the embedder.
pytestmark = [
    pytest.mark.integration,
    pytest.mark.usefixtures("stub_model", "model_reads_images", "studio_worker"),
]

WORD = """\
import os
import docx

document = docx.Document()
document.add_heading("Client proposal", level=1)
document.add_paragraph("We propose a two-phase rollout.")
document.save(os.environ["OUTPUT_PATH"])
"""


def _docx() -> bytes:
    document = docx.Document()
    document.add_paragraph("Board pack")
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def _artifact_id(text: str) -> int:
    return int(text.split(",")[0].removeprefix("Rendered artifact "))


async def _rendered(tools: ToolEndpoint, workspace_id: int) -> int:
    text, is_error = await tools.call(
        workspace_id,
        "render_document",
        {"title": "Client proposal", "format": "docx", "script": WORD},
    )
    assert is_error is False, text
    return _artifact_id(text)


async def test_a_rendered_word_version_becomes_a_new_pdf_with_its_pages_shown(
    tools: ToolEndpoint, engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The PDF is its own document in Studio, and the Word version stays as it was."""
    workspace_id = await tools.workspace()
    word_id = await _rendered(tools, workspace_id)
    office = office_on(monkeypatch, pages=3)
    with create_session_factory(engine)() as session:
        word_file = _primary(session.get(Artifact, word_id))
        before = word_file.read_bytes()

    text, images, is_error = await tools.call_content(
        workspace_id, "convert_document", {"artifact_id": word_id, "format": "pdf"}
    )

    assert is_error is False, text
    pdf_id = _artifact_id(text)
    assert pdf_id != word_id
    assert (
        text.splitlines()[0]
        == f"Rendered artifact {pdf_id}, version 1: Client proposal"
    )
    assert (
        f"A PDF of 3 pages, converted by LibreOffice 26.8.1 from artifact {word_id}, "
        "which is unchanged."
    ) in text
    assert "LibreOffice page 1" in text
    assert "in order: page 1, page 2, page 3." in text
    assert len(images) == 3
    assert word_file.read_bytes() == before
    assert office.read[-1] == (".docx", before)
    with create_session_factory(engine)() as session:
        pdf = session.get(Artifact, pdf_id)
        assert pdf is not None
        assert (pdf.format, pdf.document.status) == ("pdf", DocumentStatus.READY)
        assert _primary(pdf).read_bytes().startswith(b"%PDF-")


async def test_a_selected_word_source_converts_and_its_file_is_left_alone(
    tools: ToolEndpoint, engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """LibreOffice is handed a copy of the original's bytes; the original is never opened for writing."""
    office = office_on(monkeypatch, pages=2)
    workspace_id = await tools.workspace()
    original = _docx()
    source_id = _uploaded(engine, workspace_id, "Board pack.docx", original)

    text, is_error = await tools.call(
        workspace_id, "convert_document", {"document_id": source_id, "format": "pdf"}
    )

    assert is_error is False, text
    assert f"from source {source_id}, which is unchanged." in text
    assert office.read == [(".docx", original)]


async def test_a_source_the_thread_may_not_use_is_refused_before_anything_converts(
    tools: ToolEndpoint, engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A source is its content: the turn's ticks hold here as in every tool."""
    office = office_on(monkeypatch)
    workspace_id = await tools.workspace()
    ticked = _uploaded(engine, workspace_id, "Ticked.docx", _docx())
    unticked = _uploaded(engine, workspace_id, "Unticked.docx", _docx())
    thread = tools.thread(workspace_id, SourceScope(document_ids=[ticked]))

    text, is_error = await tools.call(
        workspace_id,
        "convert_document",
        {"document_id": unticked, "format": "pdf"},
        thread=thread,
    )

    assert is_error is True
    assert f"Source {unticked} is not selected for this request." in text
    assert office.read == []


async def test_with_office_support_off_converting_is_refused_in_a_sentence(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """The tool stays listed, and the model learns where the user turns it on."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Board pack.docx", _docx())

    text, is_error = await tools.call(
        workspace_id, "convert_document", {"document_id": source_id, "format": "pdf"}
    )

    assert is_error is True
    assert "needs Office support, which is off on this computer" in text
    assert "Settings > Office support" in text


@pytest.mark.parametrize(
    ("arguments", "refusal"),
    [
        ({"format": "pdf"}, "Give exactly one of artifact_id"),
        ({"artifact_id": 1, "document_id": 2, "format": "pdf"}, "Give exactly one"),
        ({"artifact_id": 1, "format": "docx"}, "format must be pdf"),
    ],
)
async def test_a_call_naming_the_wrong_things_is_refused(
    tools: ToolEndpoint,
    monkeypatch: pytest.MonkeyPatch,
    arguments: dict[str, object],
    refusal: str,
) -> None:
    """Each refusal names the argument to fix."""
    office_on(monkeypatch)
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "convert_document", arguments)

    assert is_error is True
    assert refusal in text


async def test_a_source_that_is_not_word_or_powerpoint_is_refused(
    tools: ToolEndpoint, engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A workbook or a PDF has nothing to convert to PDF here."""
    office_on(monkeypatch)
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Figures.xlsx", b"PK not read")

    text, is_error = await tools.call(
        workspace_id, "convert_document", {"document_id": source_id, "format": "pdf"}
    )

    assert is_error is True
    assert "is not a Word (.docx) or PowerPoint (.pptx) file" in text


async def test_a_conversion_libreoffice_fails_says_why_and_leaves_a_failed_pdf(
    tools: ToolEndpoint, engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """LibreOffice's own sentence reaches the model, not a bare failure."""
    office_on(monkeypatch, failure="LibreOffice could not convert it: a broken file")
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Board pack.docx", _docx())

    text, is_error = await tools.call(
        workspace_id, "convert_document", {"document_id": source_id, "format": "pdf"}
    )

    assert is_error is True
    assert text == (
        "Converting to PDF failed: LibreOffice could not convert it: a broken file"
    )


def _primary(artifact: Artifact | None) -> Path:
    assert artifact is not None
    primary = next(f for f in artifact.files if f.role is ArtifactFileRole.PRIMARY)
    return get_storage_settings().data_dir / primary.storage_key
