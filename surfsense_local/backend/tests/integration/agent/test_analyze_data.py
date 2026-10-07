"""The analysis tool as opencode's MCP client calls it: a pandas script over ticked spreadsheets, run by the Studio worker."""

import base64
import re
import time
from io import BytesIO

import docx
import openpyxl
import pytest
from PIL import Image
from sqlalchemy import Engine

from modules.agent.data_analysis import task
from modules.agent.tool_endpoint import analyze_data
from modules.source_scope.schemas import SourceScope
from shared.config import get_storage_settings
from tests.integration.agent.conftest import declare_image_input
from tests.integration.agent.test_document_tools import _artifact_id
from tests.integration.agent.test_office_documents import _primary, _uploaded
from tests.integration.agent.test_source_pages import _pdf
from tests.integration.agent.tool_endpoint_client import ToolEndpoint
from tests.integration.worker.conftest import stub_model  # noqa: F401

pytestmark = [
    pytest.mark.integration,
    pytest.mark.usefixtures("stub_model", "model_reads_images", "studio_worker"),
]

STOP = "If this is your third failed run for this request, stop and tell the user what failed."

SALES = (
    "month,region,revenue\n"
    "Jan,North,120\nJan,South,80\nFeb,North,150\nFeb,South,95\nMar,North,170\n"
)

REVENUE_BY_MONTH = """\
import os
import pandas as pd
import matplotlib.pyplot as plt

df = pd.read_csv(os.path.join(os.environ["INPUT_DIR"], "Sales.csv"))
totals = df.groupby("month", sort=False)["revenue"].sum()
print("Total revenue:", int(df["revenue"].sum()))
out = os.environ["OUTPUT_DIR"]
totals.to_csv(os.path.join(out, "totals.csv"))
totals.plot(kind="bar")
plt.savefig(os.path.join(out, "revenue by month.png"))
"""

FAILING = "import pandas as pd\nrows = []\nraise KeyError('revenu')\n"

CHART_IN_WORD = """\
import os
import docx

document = docx.Document()
document.add_heading("Revenue", level=1)
document.add_picture(os.path.join(os.environ["IMAGES_DIR"], "{name}.png"))
document.save(os.environ["OUTPUT_PATH"])
"""


def analyse(source_ids: list[int], **arguments: object) -> dict[str, object]:
    """An analysis call's arguments: revenue by month unless told otherwise."""
    return {
        "title": "Revenue by month",
        "document_ids": source_ids,
        "script": REVENUE_BY_MONTH,
        **arguments,
    }


def _sales(engine: Engine, workspace_id: int) -> int:
    return _uploaded(engine, workspace_id, "Sales.csv", SALES.encode())


def _run(text: str) -> str:
    match = re.match(r"Analysis run ([0-9a-f]{8}): ", text)
    assert match, text
    return match[1]


def _chart_name(text: str) -> str:
    match = re.search(r"analysis-[0-9a-f]{8}-[A-Za-z0-9_-]+", text)
    assert match, text
    return match[0]


async def test_an_analysis_returns_what_it_printed_its_tables_and_its_charts(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """The model reads the result it computed, sees the chart, and learns its name for a document."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    original = (
        get_storage_settings().document_dir(workspace_id, source_id) / "Sales.csv"
    )
    before = original.read_bytes(), original.stat().st_mtime_ns

    text, images, is_error = await tools.call_content(
        workspace_id, "analyze_data", analyse([source_id])
    )

    assert is_error is False, text
    run = _run(text)
    assert text.splitlines()[0] == f"Analysis run {run}: Revenue by month"
    assert "Total revenue: 615" in text
    assert "totals.csv, 3 rows and 2 columns:" in text
    assert "| month | revenue |\n| --- | --- |\n| Jan | 200 |" in text
    assert f"analysis-{run}-revenue_by_month" in text
    assert f"outputs/analysis/{run}/" in text
    (chart,) = images
    assert chart["mimeType"] == "image/jpeg"
    with Image.open(BytesIO(base64.b64decode(chart["data"]))) as image:
        assert image.width > 100
    kept = tools.folder(workspace_id) / "outputs" / "analysis" / run
    assert sorted(path.name for path in kept.iterdir()) == [
        "revenue_by_month.png",
        "totals.csv",
    ]
    # The source is only ever copied.
    assert (original.read_bytes(), original.stat().st_mtime_ns) == before


async def test_a_workbook_source_is_read_with_pandas(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """An .xlsx source arrives under its own name; the manifest names its source."""
    workspace_id = await tools.workspace()
    book = openpyxl.Workbook()
    book.active.append(["product", "units"])
    book.active.append(["Crate", 12])
    book.active.append(["Pallet", 30])
    data = BytesIO()
    book.save(data)
    source_id = _uploaded(engine, workspace_id, "Stock.xlsx", data.getvalue())
    script = (
        "import json, os\n"
        "import pandas as pd\n"
        "folder = os.environ['INPUT_DIR']\n"
        "print(json.load(open(os.path.join(folder, 'manifest.json'))))\n"
        "print(pd.read_excel(os.path.join(folder, 'Stock.xlsx'))['units'].sum())\n"
    )

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id], script=script)
    )

    assert is_error is False, text
    assert (
        f"[{{'file': 'Stock.xlsx', 'document_id': {source_id}, 'title': 'Stock.xlsx'}}]"
        in text
    )
    assert "\n42\n" in text


async def test_a_model_that_reads_no_images_gets_the_charts_by_name(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """A text-only model still places the chart in a document by its name."""
    declare_image_input(False)
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)

    text, images, is_error = await tools.call_content(
        workspace_id, "analyze_data", analyse([source_id])
    )

    assert is_error is False, text
    assert images == []
    assert "the selected model cannot read images" in text
    assert _chart_name(text).endswith("-revenue_by_month")


@pytest.mark.parametrize(
    ("name", "data"),
    [("Plan.pdf", _pdf((595, 842))), ("Notes.md", b"# Notes\n")],
)
async def test_a_source_that_is_not_a_spreadsheet_is_refused(
    tools: ToolEndpoint, engine: Engine, name: str, data: bytes
) -> None:
    """Analysis reads tables; a document's text is read in sources/ instead."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, name, data)

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id])
    )

    assert is_error is True
    assert text == (
        f'Source {source_id} ("{name}") is not a spreadsheet or CSV file '
        "(.xlsx, .csv or .tsv): read its text in sources/ instead."
    )


async def test_a_source_not_ticked_for_the_turn_is_refused(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """The analysis reads a source's whole file, so it keeps to the ticked sources."""
    workspace_id = await tools.workspace()
    ticked = _sales(engine, workspace_id)
    unticked = _uploaded(engine, workspace_id, "Salaries.csv", b"name,pay\nA,1\n")
    thread = tools.thread(workspace_id, SourceScope(document_ids=[ticked]))

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([ticked, unticked]), thread=thread
    )

    assert is_error is True
    assert f"Source {unticked} is not selected for this request." in text
    assert not (tools.folder(workspace_id, thread) / "outputs" / "analysis").exists()


async def test_a_failing_script_returns_its_error_to_fix(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """The model fixes the script from the error line and the frames in it."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)

    text, images, is_error = await tools.call_content(
        workspace_id, "analyze_data", analyse([source_id], script=FAILING)
    )

    assert is_error is True
    assert images == []
    assert text.startswith("Analysis \"Revenue by month\" failed:\nKeyError: 'revenu'")
    assert "line 3, in <module>" in text
    assert "Fix the script and run it again." in text
    assert text.endswith(STOP)


async def test_three_failed_analyses_stop_analyses_but_not_renders(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """Analyses keep their own count: a turn that fixed its data can still render."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    failing = analyse([source_id], script=FAILING)
    for _ in range(3):
        _, is_error = await tools.call(workspace_id, "analyze_data", failing)
        assert is_error is True

    stopped, stopped_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id])
    )
    render, render_error = await tools.call(
        workspace_id, "render_document", {"format": "docx", "script": "x = 1"}
    )

    assert stopped_error is True
    assert stopped.startswith(
        "3 analysis runs failed in this request, so no more analysis runs run"
    )
    assert render_error is True
    assert render == "Give the document a title."


async def test_long_printed_output_is_cut_with_a_note(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """A whole frame printed would fill the model's window; it learns to print less."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    script = "for n in range(5000):\n    print('row', n, 'x' * 60)\n"

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id], script=script)
    )

    assert is_error is False, text
    assert "row 0 x" in text
    assert "row 4999" not in text
    assert "Printed output cut after 6,000 characters" in text
    assert len(text) < 8000


async def test_a_big_table_is_shown_in_part(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """A saved table shows its first rows and columns; the file keeps all of it."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    script = (
        "import os\n"
        "import pandas as pd\n"
        "frame = pd.DataFrame({f'c{n}': range(50) for n in range(12)})\n"
        "frame.to_csv(os.path.join(os.environ['OUTPUT_DIR'], 'wide.csv'), index=False)\n"
    )

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id], script=script)
    )

    assert is_error is False, text
    assert "wide.csv, 50 rows and 12 columns" in text
    assert "| c0 | c1 | c2 | c3 | c4 | c5 | c6 | c7 |\n" in text
    assert "| 19 |" in text and "| 20 |" not in text
    assert "First 20 rows and 8 columns shown." in text


async def test_a_chart_from_an_analysis_can_be_placed_in_a_document(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """The chart's name is an image name for render_document, at IMAGES_DIR like a source figure."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    text, _ = await tools.call(workspace_id, "analyze_data", analyse([source_id]))
    name = _chart_name(text)

    rendered, is_error = await tools.call(
        workspace_id,
        "render_document",
        {
            "title": "Revenue report",
            "format": "docx",
            "script": CHART_IN_WORD.format(name=name),
            "images": [name],
        },
    )

    assert is_error is False, rendered
    made = docx.Document(BytesIO(_primary(engine, _artifact_id(rendered))))
    assert len(made.inline_shapes) == 1


async def test_a_chart_from_another_threads_analysis_is_refused(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """A chart is its thread's output, made from that thread's sources."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    text, _ = await tools.call(workspace_id, "analyze_data", analyse([source_id]))
    name = _chart_name(text)
    other = tools.thread(workspace_id)

    rendered, is_error = await tools.call(
        workspace_id,
        "render_document",
        {
            "title": "Revenue report",
            "format": "docx",
            "script": CHART_IN_WORD.format(name=name),
            "images": [name],
        },
        thread=other,
    )

    assert is_error is True
    assert rendered == (
        f'This chat has no analysis chart named "{name}": use a name '
        "surfsense_analyze_data gave in this chat."
    )


async def test_a_run_the_worker_cannot_carry_out_is_not_blamed_on_the_script(
    tools: ToolEndpoint, engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A full disk is no reason to rewrite a script that was right, nor one of the three."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)

    def disk_full(*args: object, **kwargs: object) -> None:
        raise OSError("No space left on device")

    monkeypatch.setattr(task, "run_analysis_script", disk_full)

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id])
    )

    assert is_error is True
    assert text == (
        'Analysis "Revenue by month" could not run: OSError(\'No space left on '
        "device'). Tell the user what failed."
    )


async def test_a_run_the_worker_has_not_finished_by_the_deadline_says_so(
    tools: ToolEndpoint, engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The call answers inside opencode's timeout, and the model does not run it again."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    monkeypatch.setattr(analyze_data, "CALL_SECONDS", 1)

    def slow(*args: object, **kwargs: object) -> None:
        time.sleep(3)

    monkeypatch.setattr(task, "run_analysis_script", slow)

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id])
    )

    assert is_error is True
    assert text.startswith('Analysis "Revenue by month" did not finish within 1 s')
    assert "do not run it again" in text


async def test_printed_output_of_many_short_lines_is_cut_by_lines_too(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """opencode moves a result past 2,000 lines into a folder every thread can read."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    script = "print('.\\n' * 3000)\n"

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id], script=script)
    )

    assert is_error is False, text
    assert len(text.splitlines()) < 400
    assert "Printed output cut after 200 lines" in text


async def test_files_left_out_are_named_only_in_part(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """A script saving thousands of files gets a count, not every name."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    script = (
        "import os\n"
        "for n in range(3000):\n"
        "    open(os.path.join(os.environ['OUTPUT_DIR'], f'note-{n:04}.md'), 'w')"
        ".close()\n"
    )

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id], script=script)
    )

    assert is_error is False, text
    assert "Not kept: note-0000.md, " in text
    assert "and 2,980 more" in text
    assert len(text.encode()) < 10_000


async def test_a_failure_with_a_huge_message_is_cut_to_what_fixes_it(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """An exception can carry a whole frame; the model gets its start and the failing line."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    script = "import sys\nprint('x' * 60000, file=sys.stderr)\nraise ValueError('y' * 60000)\n"

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id], script=script)
    )

    assert is_error is True
    assert text.startswith('Analysis "Revenue by month" failed:\nValueError: yyy')
    assert "line 3, in <module>" in text
    assert text.endswith(STOP)
    assert len(text.encode()) < 12_000


async def test_tables_of_wide_characters_stay_within_what_the_model_reads_whole(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """opencode counts bytes, and a Chinese cell takes three per character."""
    workspace_id = await tools.workspace()
    source_id = _sales(engine, workspace_id)
    script = (
        "import os\n"
        "import pandas as pd\n"
        "cell = '\u6536\u5165' * 30\n"
        "print(cell * 100)\n"
        "for n in range(8):\n"
        "    frame = pd.DataFrame({f'{cell}{c}': [cell] * 30 for c in range(10)})\n"
        "    frame.to_csv(os.path.join(os.environ['OUTPUT_DIR'], f't{n}.csv'), "
        "index=False)\n"
    )

    text, is_error = await tools.call(
        workspace_id, "analyze_data", analyse([source_id], script=script)
    )

    assert is_error is False, text
    assert "t0.csv, 30 rows and 10 columns:" in text
    assert "more tables are kept; open one with read." in text
    assert len(text.encode()) < 45_000
