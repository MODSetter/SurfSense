"""PowerPoint decks and Excel workbooks through the render tool, and documents made from a source template.

Each script runs in the real runner behind Studio's job, as opencode's MCP client calls it.
"""

import threading
from collections.abc import Iterator
from io import BytesIO

import docx
import pptx
import pytest
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm
from pptx.util import Inches
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen import canvas
from sqlalchemy import Engine

from modules.agent.previews import docx_snapshots
from modules.agent.previews.docx_snapshots import SnapshotRequest
from modules.artifacts.models import Artifact
from modules.documents.models import Document, DocumentStatus, DocumentType
from shared.config import get_storage_settings
from shared.db import create_session_factory
from tests.integration.agent.test_document_tools import _artifact_id
from tests.integration.agent.tool_endpoint_client import ToolEndpoint
from tests.integration.worker.conftest import stub_model  # noqa: F401
from worker.ingestion import run as run_ingest

pytestmark = [
    pytest.mark.integration,
    pytest.mark.usefixtures("stub_model", "model_reads_images"),
]

DECK = """\
import os
from pptx import Presentation
from pptx.util import Inches

deck = Presentation()
deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
cover = deck.slides.add_slide(deck.slide_layouts[0])
cover.shapes.title.text = "Quarterly review"
body = deck.slides.add_slide(deck.slide_layouts[1])
body.shapes.title.text = "Costs"
body.placeholders[1].text = "Costs rose in the north."
deck.save(os.environ["OUTPUT_PATH"])
"""

WORKBOOK = """\
import os
import xlsxwriter

book = xlsxwriter.Workbook(os.environ["OUTPUT_PATH"])
sheet = book.add_worksheet("Costs")
bold = book.add_format({"bold": True})
sheet.write_row(0, 0, ["Year", "Cost"], bold)
for row, (year, cost) in enumerate([("2024", 120), ("2025", 135)], start=1):
    sheet.write_row(row, 0, [year, cost])
sheet.write(3, 0, "Total", bold)
sheet.write_formula(3, 1, "=SUM(B2:B3)")
book.close()
"""

# The skill's pattern: keep the template's styles, section and header; empty its body.
LETTER_FROM_TEMPLATE = """\
import os
import docx

document = docx.Document(os.environ["TEMPLATE_PATH"])
body = document.element.body
for block in list(body):
    if block.tag.endswith("}sectPr"):
        continue
    body.remove(block)
document.add_heading("Proposal for Halvorsen Freight", level=1)
document.add_paragraph("We propose a two-phase rollout.")
document.save(os.environ["OUTPUT_PATH"])
"""

DECK_FROM_TEMPLATE = """\
import os
from pptx import Presentation

deck = Presentation(os.environ["TEMPLATE_PATH"])
slides = deck.slides._sldIdLst
for slide in list(slides):
    deck.part.drop_rel(slide.rId)
    slides.remove(slide)
layout = next(l for l in deck.slide_layouts if l.name == "Brand Title Only")
deck.slides.add_slide(layout).shapes.title.text = "Quarterly review"
deck.save(os.environ["OUTPUT_PATH"])
"""


def render(**arguments: object) -> dict[str, object]:
    """A render call's arguments: a deck unless told otherwise."""
    return {"title": "Quarterly review", "format": "pptx", "script": DECK, **arguments}


def _uploaded(engine: Engine, workspace_id: int, name: str, data: bytes) -> int:
    """A source the user uploaded, its original kept where ingest keeps it."""
    with create_session_factory(engine)() as session:
        source = Document(
            workspace_id=workspace_id,
            title=name,
            document_type=DocumentType.FILE,
            status=DocumentStatus.READY,
            content=f"The text of {name}",
        )
        session.add(source)
        session.commit()
        source_id = source.id
    folder = get_storage_settings().document_dir(workspace_id, source_id)
    folder.mkdir(parents=True)
    (folder / name).write_bytes(data)
    return source_id


def _letterhead() -> bytes:
    """A Word file with its own style, a header, wide margins and its own body text."""
    document = docx.Document()
    document.styles.add_style("Halvorsen Body", WD_STYLE_TYPE.PARAGRAPH)
    section = document.sections[0]
    section.left_margin = Cm(4)
    section.header.paragraphs[0].text = "Halvorsen Freight letterhead"
    document.add_paragraph("Old body text to remove.", style="Halvorsen Body")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _brand_deck() -> bytes:
    """A 16:9 deck with one slide of its own and a layout named for the brand."""
    deck = pptx.Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.slide_layouts[5]._element.cSld.set("name", "Brand Title Only")
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Brand cover"
    buffer = BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _primary(engine: Engine, artifact_id: int) -> bytes:
    with create_session_factory(engine)() as session:
        (file,) = session.get(Artifact, artifact_id).files
        return (get_storage_settings().data_dir / file.storage_key).read_bytes()


async def test_a_deck_render_counts_its_slides_and_shows_their_text(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """The model learns its slides' titles and text, in order, from the result."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False, text
    artifact_id = _artifact_id(text)
    assert text.splitlines()[0] == (
        f"Rendered artifact {artifact_id}, version 1: Quarterly review"
    )
    assert "A PowerPoint deck of 2 slides." in text
    assert "## Slide 1: Quarterly review" in text
    assert "## Slide 2: Costs\nCosts rose in the north." in text
    listed = (await tools.client.get(f"/workspaces/{workspace_id}/artifacts")).json()
    assert [(a["format"], a["spec_kind"]) for a in listed] == [("pptx", "python")]


async def test_a_deck_without_the_desktop_app_says_why_it_has_no_slide_previews(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """Slides are printed by Electron, as Word pages are; no test runs it."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(workspace_id, "render_document", render())

    assert is_error is False, text
    assert (
        "No page previews: PowerPoint slides are drawn by the SurfSense desktop app"
        in text
    )


@pytest.fixture
def electron(monkeypatch: pytest.MonkeyPatch) -> Iterator[list[SnapshotRequest]]:
    """A stand-in for Electron on the snapshot queue: it prints every request as two landscape pages."""
    queue = docx_snapshots.DocxSnapshots()
    monkeypatch.setattr(docx_snapshots, "snapshots", queue)
    served: list[SnapshotRequest] = []
    stop = threading.Event()
    queue.next_request()  # polling from the start

    def serve() -> None:
        while not stop.is_set():
            request = queue.next_request()
            if request is None:
                stop.wait(0.02)
                continue
            served.append(request)
            queue.deliver(request.id, _slides_pdf(2))

    thread = threading.Thread(target=serve, daemon=True)
    thread.start()
    yield served
    stop.set()
    thread.join(timeout=5)


def _slides_pdf(pages: int) -> bytes:
    out = BytesIO()
    drawing = canvas.Canvas(out, pagesize=landscape(letter))
    for number in range(1, pages + 1):
        drawing.drawString(72, 72, f"Slide {number}")
        drawing.showPage()
    drawing.save()
    return out.getvalue()


async def test_a_deck_render_is_a_studio_version_whose_slides_electron_prints_as_a_deck(
    tools: ToolEndpoint, studio_worker: None, electron: list[SnapshotRequest]
) -> None:
    """End to end: the version Studio lists, the pptx print request, and the slides the model is shown."""
    workspace_id = await tools.workspace()

    text, images, is_error = await tools.call_content(
        workspace_id, "render_document", render()
    )

    assert is_error is False, text
    artifact_id = _artifact_id(text)
    assert [(r.format, r.artifact_id, r.source_file, r.pages) for r in electron] == [
        ("pptx", artifact_id, None, "1-4")
    ]
    assert (
        f"The slide previews of version 1 of artifact {artifact_id} come with this "
        "result as images, in order: slide 1, slide 2." in text
    )
    assert [item["mimeType"] for item in images] == ["image/jpeg", "image/jpeg"]
    assert "differ a little from PowerPoint" in text
    listed = (await tools.client.get(f"/workspaces/{workspace_id}/artifacts")).json()
    assert [
        (a["id"], a["format"], a["status"], a["version"]["number"]) for a in listed
    ] == [(artifact_id, "pptx", "ready", 1)]


async def test_a_workbook_render_returns_its_summary_to_check_in_place_of_pages(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """A workbook has no pages; its sheets, rows and formulas are what the model checks."""
    workspace_id = await tools.workspace()

    text, images, is_error = await tools.call_content(
        workspace_id,
        "render_document",
        render(title="Costs", format="xlsx", script=WORKBOOK),
    )

    assert is_error is False, text
    assert images == []
    assert "An Excel workbook of 1 sheet." in text
    assert 'Sheet "Costs": A1:B4\nYear | Cost\n2024 | 120' in text
    assert "- Costs!B4: =SUM(B2:B3)" in text
    assert "No page previews: a workbook has no pages" in text
    assert "outputs/previews" not in text


async def test_a_word_document_made_from_a_template_keeps_its_header_and_margins(
    tools: ToolEndpoint, engine: Engine, studio_worker: None
) -> None:
    """The template's section and header carry over; its body does not; the source is untouched."""
    workspace_id = await tools.workspace()
    letterhead = _letterhead()
    source_id = _uploaded(engine, workspace_id, "Letterhead.docx", letterhead)

    text, is_error = await tools.call(
        workspace_id,
        "render_document",
        {
            "title": "Proposal",
            "format": "docx",
            "script": LETTER_FROM_TEMPLATE,
            "template_source_id": source_id,
        },
    )

    assert is_error is False, text
    made = docx.Document(BytesIO(_primary(engine, _artifact_id(text))))
    assert "Halvorsen Body" in [style.name for style in made.styles]
    section = made.sections[0]
    assert section.header.paragraphs[0].text == "Halvorsen Freight letterhead"
    assert round(section.left_margin.cm) == 4
    assert [p.text for p in made.paragraphs] == [
        "Proposal for Halvorsen Freight",
        "We propose a two-phase rollout.",
    ]
    original = get_storage_settings().document_dir(workspace_id, source_id)
    assert (original / "Letterhead.docx").read_bytes() == letterhead


async def test_a_deck_made_from_a_template_and_its_next_version_both_start_from_it(
    tools: ToolEndpoint, engine: Engine, studio_worker: None
) -> None:
    """The next version leaves template_source_id out and still gets the template."""
    workspace_id = await tools.workspace()
    brand = _brand_deck()
    source_id = _uploaded(engine, workspace_id, "Brand.pptx", brand)
    first, is_error = await tools.call(
        workspace_id,
        "render_document",
        render(script=DECK_FROM_TEMPLATE, template_source_id=source_id),
    )
    assert is_error is False, first

    second, is_error = await tools.call(
        workspace_id,
        "render_document",
        render(script=DECK_FROM_TEMPLATE, artifact_id=_artifact_id(first)),
    )

    assert is_error is False, second
    for result in (first, second):
        made = pptx.Presentation(BytesIO(_primary(engine, _artifact_id(result))))
        assert round(made.slide_width / made.slide_height, 2) == 1.78
        assert [s.shapes.title.text for s in made.slides] == ["Quarterly review"]
        assert made.slides[0].slide_layout.name == "Brand Title Only"
    original = get_storage_settings().document_dir(workspace_id, source_id)
    assert (original / "Brand.pptx").read_bytes() == brand
    script, _ = await tools.call(
        workspace_id, "read_document", {"artifact_id": _artifact_id(second)}
    )
    assert f"Template source it starts from: {source_id}" in script


async def test_a_deck_starts_from_an_uploaded_template_with_no_slides(
    tools: ToolEndpoint, engine: Engine, studio_worker: None
) -> None:
    """A template ingests as a ready source, so a deck can start from it."""
    workspace_id = await tools.workspace()
    template = pptx.Presentation()
    template.slide_width, template.slide_height = Inches(13.333), Inches(7.5)
    template.slide_layouts[5]._element.cSld.set("name", "Brand Title Only")
    buffer = BytesIO()
    template.save(buffer)
    with create_session_factory(engine)() as session:
        source = Document(
            workspace_id=workspace_id,
            title="Brand.pptx",
            document_type=DocumentType.FILE,
            document_metadata={"suffix": ".pptx"},
        )
        session.add(source)
        session.commit()
        source_id = source.id
    folder = get_storage_settings().document_dir(workspace_id, source_id)
    folder.mkdir(parents=True)
    (folder / "Brand.pptx").write_bytes(buffer.getvalue())
    run_ingest(source_id)

    text, is_error = await tools.call(
        workspace_id,
        "render_document",
        render(script=DECK_FROM_TEMPLATE, template_source_id=source_id),
    )

    assert is_error is False, text
    made = pptx.Presentation(BytesIO(_primary(engine, _artifact_id(text))))
    assert round(made.slide_width / made.slide_height, 2) == 1.78
    assert [s.shapes.title.text for s in made.slides] == ["Quarterly review"]
    assert made.slides[0].slide_layout.name == "Brand Title Only"


@pytest.mark.parametrize(
    ("format", "name", "says"),
    [
        ("pptx", "Letterhead.docx", "is not a PowerPoint file (.pptx)"),
        (
            "xlsx",
            "Budget.xlsx",
            "A template applies to a Word document or a PowerPoint deck only",
        ),
    ],
)
async def test_a_template_the_format_cannot_start_from_is_refused_in_a_sentence(
    tools: ToolEndpoint, engine: Engine, format: str, name: str, says: str
) -> None:
    """Excel takes no template, and a deck starts only from a deck."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, name, b"PK\x03\x04 office file")

    text, is_error = await tools.call(
        workspace_id,
        "render_document",
        render(format=format, template_source_id=source_id),
    )

    assert is_error is True
    assert says in text
    listed = (await tools.client.get(f"/workspaces/{workspace_id}/artifacts")).json()
    assert listed == []


@pytest.mark.parametrize("value", ["12", 1.5, True])
async def test_a_template_source_that_is_no_number_is_refused(
    tools: ToolEndpoint, value: object
) -> None:
    """The refusal names the argument and what it takes."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(
        workspace_id, "render_document", render(template_source_id=value)
    )

    assert is_error is True
    assert "template_source_id" in text
