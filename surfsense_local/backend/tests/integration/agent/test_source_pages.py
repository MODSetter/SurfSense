"""Pages of a source drawn for the agent to look at, as opencode's MCP client calls the tool.

PDF pages are drawn in the API. Word and PowerPoint pages come back as a PDF
from Electron, which a thread plays here on the snapshot queue itself.
"""

import threading
from collections.abc import Callable, Iterator
from io import BytesIO

import docx
import pptx
import pytest
from lxml import etree
from PIL import Image
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from sqlalchemy import Engine

from modules.agent.previews import docx_snapshots
from modules.agent.previews.docx_snapshots import SnapshotRequest
from modules.source_scope.schemas import SourceScope
from shared.config import get_storage_settings
from tests.integration.agent.conftest import MAX_PATH, declare_image_input
from tests.integration.agent.test_office_documents import _uploaded
from tests.integration.agent.tool_endpoint_client import ToolEndpoint

pytestmark = [pytest.mark.integration, pytest.mark.usefixtures("model_reads_images")]


@pytest.fixture(autouse=True)
def fresh_snapshots(monkeypatch: pytest.MonkeyPatch) -> docx_snapshots.DocxSnapshots:
    """No Electron has polled yet."""
    fresh = docx_snapshots.DocxSnapshots()
    monkeypatch.setattr(docx_snapshots, "snapshots", fresh)
    return fresh


def _pdf(*sizes: tuple[float, float]) -> bytes:
    out = BytesIO()
    drawing = canvas.Canvas(out)
    for number, size in enumerate(sizes, 1):
        drawing.setPageSize(size)
        drawing.drawString(72, 72, f"Page {number}")
        drawing.showPage()
    drawing.save()
    return out.getvalue()


def _docx() -> bytes:
    document = docx.Document()
    document.add_heading("Letterhead", level=1)
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def _pptx(slides: int) -> bytes:
    deck = pptx.Presentation()
    for number in range(1, slides + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Slide {number}"
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


def _call(document_id: int, **arguments: object) -> dict[str, object]:
    return {"document_id": document_id, **arguments}


@pytest.fixture
def electron(
    fresh_snapshots: docx_snapshots.DocxSnapshots,
) -> Iterator[Callable[[Callable[[SnapshotRequest], bytes]], list[SnapshotRequest]]]:
    """Start an Electron stand-in that prints each request with the function it is given."""
    stop = threading.Event()
    threads: list[threading.Thread] = []

    def start(print_: Callable[[SnapshotRequest], bytes]) -> list[SnapshotRequest]:
        served: list[SnapshotRequest] = []
        fresh_snapshots.next_request()  # polling from the start

        def serve() -> None:
            while not stop.is_set():
                request = fresh_snapshots.next_request()
                if request is None:
                    stop.wait(0.02)
                    continue
                served.append(request)
                assert fresh_snapshots.source_file(request.id) == request.source_file
                fresh_snapshots.deliver(request.id, print_(request))

        thread = threading.Thread(target=serve, daemon=True)
        thread.start()
        threads.append(thread)
        return served

    yield start
    stop.set()
    for thread in threads:
        thread.join(timeout=5)


async def test_a_pdfs_first_pages_are_drawn_where_the_agent_reads_them(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """By default the first four pages, each at most 1,000 px on its long side."""
    workspace_id = await tools.workspace()
    original = _pdf(*[A4] * 6)
    source_id = _uploaded(engine, workspace_id, "Brand guide.pdf", original)

    text, is_error = await tools.call(workspace_id, "source_pages", _call(source_id))

    assert is_error is False, text
    assert f'Source {source_id} ("Brand guide.pdf") has 6 pages.' in text
    for page in (1, 2, 3, 4):
        relative = f"sources/pages/{source_id}-p{page}.png"
        assert f"- {relative}" in text
        with Image.open(tools.folder(workspace_id) / relative) as image:
            assert max(image.size) == 1000
            assert image.height > image.width
    assert f"{source_id}-p5.png" not in text
    folder = get_storage_settings().document_dir(workspace_id, source_id)
    assert (folder / "Brand guide.pdf").read_bytes() == original


async def test_pages_whose_paths_cannot_fit_are_refused_in_a_sentence(
    tools: ToolEndpoint, engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Under a data folder this deep, `sources/pages/` has no room for a page's name."""
    storage = get_storage_settings()
    workspace_id = await tools.workspace()
    pages = tools.folder(workspace_id) / "sources" / "pages"
    padding = MAX_PATH - len(str(pages / "1-p1.png"))
    monkeypatch.setattr(storage, "data_dir", storage.data_dir / ("d" * padding))
    declare_image_input(True)
    source_id = _uploaded(engine, workspace_id, "Brand guide.pdf", _pdf(A4))

    text, is_error = await tools.call(workspace_id, "source_pages", _call(source_id))

    assert is_error is True
    assert "too deep on this computer" in text
    assert not any((tools.folder(workspace_id) / "sources" / "pages").glob("*.png"))


async def test_the_pages_asked_for_are_the_pages_drawn(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """Page 5 is landscape here, so its image shows it is page 5 that was drawn."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(
        engine, workspace_id, "Brand.pdf", _pdf(A4, A4, A4, A4, landscape(A4))
    )

    text, is_error = await tools.call(
        workspace_id, "source_pages", _call(source_id, pages=[5, 2])
    )

    assert is_error is False, text
    pages = tools.folder(workspace_id) / "sources" / "pages"
    with Image.open(pages / f"{source_id}-p5.png") as five:
        assert five.width == 1000 and five.height < 1000
    with Image.open(pages / f"{source_id}-p2.png") as two:
        assert two.height == 1000
    assert text.index("-p2.png") < text.index("-p5.png")


@pytest.mark.parametrize(
    ("pages", "says"),
    [
        ([4], "has 3 pages: ask for pages from 1 to 3"),
        ([0], "page numbers from 1"),
        ([1, 2, 3, 4, 5], "at most 4 pages"),
        ("1-2", "page numbers from 1"),
    ],
)
async def test_pages_it_cannot_draw_are_refused_with_what_it_can(
    tools: ToolEndpoint, engine: Engine, pages: object, says: str
) -> None:
    """The model learns which pages there are and how many it may ask for."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Brand.pdf", _pdf(A4, A4, A4))

    text, is_error = await tools.call(
        workspace_id, "source_pages", _call(source_id, pages=pages)
    )

    assert is_error is True
    assert says in text


async def test_a_word_sources_pages_are_printed_by_electron_from_its_original(
    tools: ToolEndpoint, engine: Engine, electron
) -> None:
    """Electron is handed the source's own file to lay out as Word; a short print gives the count."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Letterhead.docx", _docx())
    served = electron(lambda request: _pdf(A4, A4))

    text, is_error = await tools.call(workspace_id, "source_pages", _call(source_id))

    assert is_error is False, text
    (request,) = served
    assert (request.format, request.pages) == ("docx", "1-4")
    assert request.source_file is not None
    assert request.source_file.name == "Letterhead.docx"
    assert "has 2 pages" in text
    assert f"- sources/pages/{source_id}-p2.png" in text
    assert "headers and footers" in text


async def test_a_decks_slides_are_counted_and_printed_as_asked(
    tools: ToolEndpoint, engine: Engine, electron
) -> None:
    """A deck's slide count is known before printing; the slides asked for are the ones printed."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Brand.pptx", _pptx(6))
    served = electron(lambda request: _pdf(landscape(A4), landscape(A4)))

    text, is_error = await tools.call(
        workspace_id, "source_pages", _call(source_id, pages=[6, 3])
    )

    assert is_error is False, text
    assert [(r.format, r.pages) for r in served] == [("pptx", "3,6")]
    assert f'Source {source_id} ("Brand.pptx") has 6 slides.' in text
    assert f"- sources/pages/{source_id}-p3.png" in text
    assert f"- sources/pages/{source_id}-p6.png" in text


async def test_a_16_9_slide_is_drawn_no_wider_than_1000_px(
    tools: ToolEndpoint, engine: Engine, electron
) -> None:
    """A widescreen slide prints at 1280 x 720 CSS px, 960 x 540 pt, where the
    scale to 1,000 px comes to a hair over it and the pixels round up."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Brand.pptx", _pptx(1))
    electron(lambda request: _pdf((960, 540)))

    text, is_error = await tools.call(workspace_id, "source_pages", _call(source_id))

    assert is_error is False, text
    page = tools.folder(workspace_id) / "sources" / "pages" / f"{source_id}-p1.png"
    with Image.open(page) as image:
        assert image.width == 1000


def _with_sections(deck_file: bytes) -> bytes:
    """The deck as PowerPoint saves one with sections, which list its slides by id again."""
    deck = pptx.Presentation(BytesIO(deck_file))
    presentation = deck.part._element
    p = presentation.nsmap["p"]
    p14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    extension = etree.SubElement(
        etree.SubElement(presentation, f"{{{p}}}extLst"),
        f"{{{p}}}ext",
        uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}",
    )
    section = etree.SubElement(
        etree.SubElement(extension, f"{{{p14}}}sectionLst"),
        f"{{{p14}}}section",
        name="Default Section",
        id="{0E5B2E3A-6C0F-4B7E-9C51-1D2B3C4D5E6F}",
    )
    listed = etree.SubElement(section, f"{{{p14}}}sldIdLst")
    for slide in deck.slides:
        etree.SubElement(listed, f"{{{p14}}}sldId", id=str(slide.slide_id))
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


async def test_a_deck_with_sections_counts_each_slide_once(
    tools: ToolEndpoint, engine: Engine, electron
) -> None:
    """A section's list of slide ids is not more slides."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Brand.pptx", _with_sections(_pptx(2)))
    served = electron(lambda request: _pdf(landscape(A4), landscape(A4)))

    text, is_error = await tools.call(workspace_id, "source_pages", _call(source_id))

    assert is_error is False, text
    assert f'Source {source_id} ("Brand.pptx") has 2 slides.' in text
    assert [r.pages for r in served] == ["1-2"]


async def test_a_deck_asked_for_slides_it_lacks_is_refused_before_printing(
    tools: ToolEndpoint, engine: Engine, electron
) -> None:
    """No print is spent on a slide that is not there."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Brand.pptx", _pptx(2))
    served = electron(lambda request: _pdf(A4))

    text, is_error = await tools.call(
        workspace_id, "source_pages", _call(source_id, pages=[3])
    )

    assert is_error is True
    assert "has 2 slides: ask for slides from 1 to 2" in text
    assert served == []


async def test_a_print_with_other_pages_than_asked_draws_none(
    tools: ToolEndpoint, engine: Engine, electron
) -> None:
    """A page labelled as one it is not would mislead the model."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Brand.pptx", _pptx(6))
    electron(lambda request: _pdf(*[A4] * 4))

    text, is_error = await tools.call(
        workspace_id, "source_pages", _call(source_id, pages=[5, 6])
    )

    assert is_error is False, text
    assert "No pages were drawn" in text
    assert ".png" not in text


async def test_without_electron_a_word_source_gets_no_pages_and_says_why(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """Without the desktop app, Word and PowerPoint pages cannot be drawn."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Letterhead.docx", _docx())

    text, is_error = await tools.call(workspace_id, "source_pages", _call(source_id))

    assert is_error is False, text
    assert "Word pages are drawn by the SurfSense desktop app" in text


@pytest.mark.parametrize(
    ("name", "says"),
    [
        ("Logo.png", "surfsense_list_images"),
        ("Budget.xlsx", "PDF, Word or PowerPoint"),
        ("Notes.txt", "PDF, Word or PowerPoint"),
    ],
)
async def test_a_source_with_no_pages_to_draw_is_refused(
    tools: ToolEndpoint, engine: Engine, name: str, says: str
) -> None:
    """Only a PDF, a Word file or a deck has pages; the refusal says what to do instead."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, name, b"some bytes")

    text, is_error = await tools.call(workspace_id, "source_pages", _call(source_id))

    assert is_error is True
    assert says in text


async def test_a_source_outside_the_workspace_or_the_turn_is_refused(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """A source the turn may not use is refused, naming what it may."""
    workspace_id, elsewhere = await tools.workspace(), await tools.workspace()
    ticked = _uploaded(engine, workspace_id, "Plan.pdf", _pdf(A4))
    unticked = _uploaded(engine, workspace_id, "Brand.pdf", _pdf(A4))
    theirs = _uploaded(engine, elsewhere, "Theirs.pdf", _pdf(A4))
    thread = tools.thread(workspace_id, SourceScope(document_ids=[ticked]))

    outside, outside_error = await tools.call(
        workspace_id, "source_pages", _call(unticked), thread=thread
    )
    foreign, foreign_error = await tools.call(
        workspace_id, "source_pages", _call(theirs), thread=thread
    )

    assert outside_error is True
    assert f"Source {unticked} is not selected" in outside
    assert foreign_error is True
    assert f"Source {theirs} is not selected" in foreign
    pages = tools.folder(workspace_id, thread) / "sources" / "pages"
    assert not pages.exists() or not list(pages.iterdir())


async def test_a_model_that_reads_no_images_is_not_offered_the_tool(
    tools: ToolEndpoint, engine: Engine
) -> None:
    """Page images would reach it as errors, so the tool is neither listed nor run."""
    declare_image_input(False)
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Brand.pdf", _pdf(A4))

    listed = await tools.request(workspace_id, "tools/list")
    text, is_error = await tools.call(workspace_id, "source_pages", _call(source_id))

    assert "source_pages" not in [tool["name"] for tool in listed["result"]["tools"]]
    assert is_error is True
    assert "cannot read images" in text


async def test_a_deck_with_no_slides_says_so(
    tools: ToolEndpoint, engine: Engine, electron
) -> None:
    """Nothing is sent to print, and the model learns why."""
    workspace_id = await tools.workspace()
    source_id = _uploaded(engine, workspace_id, "Empty.pptx", _pptx(0))
    served = electron(lambda request: _pdf(A4))

    text, is_error = await tools.call(workspace_id, "source_pages", _call(source_id))

    assert is_error is True
    assert "has no slides to draw" in text
    assert served == []
