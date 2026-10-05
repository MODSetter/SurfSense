"""The documents skill as the agent reads it: its examples run in the script runner."""

import re
from io import BytesIO
from pathlib import Path

import docx
import openpyxl
import pptx
import pypdfium2
import pytest
from docx.shared import Cm
from lxml import etree
from pptx.util import Inches

from modules.agent.opencode_config import DOCUMENTS_SKILL, skills_folder
from worker.document_script.run import run_document_script

pytestmark = pytest.mark.integration

_EXAMPLE = re.compile(r"^## Example: ([\w ]+)\n\n```python\n(.*?)^```", re.M | re.S)


def _examples() -> dict[str, str]:
    """Each example script in the skill, by the heading it is under."""
    text = (skills_folder() / DOCUMENTS_SKILL / "SKILL.md").read_text(encoding="utf-8")
    return dict(_EXAMPLE.findall(text))


def test_the_skill_names_itself_as_opencode_finds_it() -> None:
    """opencode lists a skill by its frontmatter name, and the config allows only that one."""
    text = (skills_folder() / DOCUMENTS_SKILL / "SKILL.md").read_text(encoding="utf-8")

    frontmatter = text.split("---")[1]
    assert f"name: {DOCUMENTS_SKILL}\n" in frontmatter
    assert "description: " in frontmatter


def test_the_skill_says_what_the_agent_can_look_at_and_what_previews_leave_out() -> (
    None
):
    """A source's figures can be opened, and a Word preview hides headers and footers."""
    text = (skills_folder() / DOCUMENTS_SKILL / "SKILL.md").read_text(encoding="utf-8")

    assert "sources/figures/" in text
    assert "headers and footers" in text
    # Most local models read no images: opencode answers their `read` with an error.
    assert (
        "If `read` cannot show you the image, place that figure with its caption "
        "and tell the user" in text
    )
    # The pages come with the render; a text-only model is told so in the result.
    assert "When the result says there are no previews" in text
    assert "Then look at every page the result shows before you go on." in text
    # Opening a preview again sends the page twice, and both copies stay.
    assert (
        "Open a preview file with `read` only when you need a closer look at one "
        "page: every image stays in the conversation." in text
    )


def test_the_word_example_renders_a_document_with_its_table_and_chart() -> None:
    """A model copies the example's habits, so the example itself must run."""
    result = run_document_script(
        _examples()["Word"], output_name="document.docx", images={}
    )

    assert result.ok, f"{result.error}\n{result.traceback_tail}"
    document = docx.Document(BytesIO(result.output))
    assert document.paragraphs[0].style.name == "Title"
    assert [
        p.style.name for p in document.paragraphs if p.text.startswith("A pilot")
    ] == ["List Bullet"]
    assert document.tables[0].rows[0].cells[0].text == "Phase"
    assert len(document.inline_shapes) == 1
    assert round(document.sections[0].page_width.cm, 1) == 21.0


def test_the_pdf_example_renders_an_a4_page_with_its_table() -> None:
    """The PDF example's text, table and list all reach the page."""
    result = run_document_script(
        _examples()["PDF"], output_name="document.pdf", images={}
    )

    assert result.ok, f"{result.error}\n{result.traceback_tail}"
    pdf = pypdfium2.PdfDocument(result.output)
    try:
        page = pdf[0]
        text = page.get_textpage().get_text_range()
        assert (round(page.get_width()), round(page.get_height())) == (595, 842)
    finally:
        pdf.close()
    for words in ("Client proposal", "Costs & timeline", "Rollout", "48,000 EUR"):
        assert words in text


def test_the_powerpoint_example_renders_a_16_9_deck_with_a_table_chart_and_notes() -> (
    None
):
    """The deck example uses layouts by name, a real table, a native chart and notes."""
    result = run_document_script(
        _examples()["PowerPoint"], output_name="document.pptx", images={}
    )

    assert result.ok, f"{result.error}\n{result.traceback_tail}"
    deck = pptx.Presentation(BytesIO(result.output))
    assert round(deck.slide_width / deck.slide_height, 2) == 1.78
    titles = [slide.shapes.title.text for slide in deck.slides]
    assert titles == ["Rollout plan", "Why now", "Pricing"]
    pricing = deck.slides[2]
    assert any(shape.has_table for shape in pricing.shapes)
    assert any(shape.has_chart for shape in pricing.shapes)
    assert deck.slides[1].notes_slide.notes_text_frame.text == (
        "Start with the cost rise."
    )


def test_the_excel_example_writes_formulas_for_its_totals() -> None:
    """Totals are formulas the user's edits keep right, not numbers worked out in the script."""
    result = run_document_script(
        _examples()["Excel"], output_name="document.xlsx", images={}
    )

    assert result.ok, f"{result.error}\n{result.traceback_tail}"
    sheet = openpyxl.load_workbook(BytesIO(result.output))["Costs"]
    assert sheet["B5"].value == "=SUM(B2:B4)"
    assert sheet["D2"].value == "=B2*(1+C2)"
    assert sheet["B2"].value == 12000
    assert sheet.freeze_panes == "A2"
    assert len(sheet._charts) == 1


def test_the_excel_examples_formulas_carry_the_values_they_come_to() -> None:
    """Studio's viewer shows a formula's stored value, which xlsxwriter leaves at 0 unless given one."""
    result = run_document_script(
        _examples()["Excel"], output_name="document.xlsx", images={}
    )

    assert result.ok, f"{result.error}\n{result.traceback_tail}"
    shown = openpyxl.load_workbook(BytesIO(result.output), data_only=True)["Costs"]
    assert shown["D2"].value == pytest.approx(13200)
    assert shown["B5"].value == 66000
    assert shown["D5"].value == pytest.approx(13200 + 55200 + 6300)


def test_the_word_template_example_keeps_the_letterhead_and_drops_its_body(
    tmp_path: Path,
) -> None:
    """The header and margins stay; the template's own paragraphs go."""
    template = tmp_path / "Letterhead.docx"
    letterhead = docx.Document()
    letterhead.sections[0].left_margin = Cm(4)
    letterhead.sections[0].header.paragraphs[0].text = "Halvorsen Freight"
    letterhead.add_paragraph("Old body text.")
    letterhead.add_table(rows=1, cols=1).cell(0, 0).text = "Old table."
    letterhead.save(template)

    result = run_document_script(
        _examples()["Word from a template"],
        output_name="document.docx",
        images={},
        template=template,
    )

    assert result.ok, f"{result.error}\n{result.traceback_tail}"
    made = docx.Document(BytesIO(result.output))
    assert made.sections[0].header.paragraphs[0].text == "Halvorsen Freight"
    assert round(made.sections[0].left_margin.cm) == 4
    assert [p.text for p in made.paragraphs] == [
        "Client proposal",
        "We propose a two-phase rollout.",
    ]
    assert made.tables == []


def test_the_powerpoint_template_example_replaces_the_slides_and_keeps_the_size(
    tmp_path: Path,
) -> None:
    """The template's slides go; its slide size and layouts make the new ones."""
    template = tmp_path / "Brand.pptx"
    brand = pptx.Presentation()
    brand.slide_width, brand.slide_height = Inches(13.333), Inches(7.5)
    for title in ("Old cover", "Old agenda"):
        brand.slides.add_slide(brand.slide_layouts[0]).shapes.title.text = title
    brand.save(template)

    result = run_document_script(
        _examples()["PowerPoint from a template"],
        output_name="document.pptx",
        images={},
        template=template,
    )

    assert result.ok, f"{result.error}\n{result.traceback_tail}"
    deck = pptx.Presentation(BytesIO(result.output))
    assert round(deck.slide_width / deck.slide_height, 2) == 1.78
    assert [s.shapes.title.text for s in deck.slides] == ["Rollout plan", "Why now"]


_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def test_the_powerpoint_template_example_leaves_no_section_naming_a_removed_slide(
    tmp_path: Path,
) -> None:
    """A brand deck saved by PowerPoint often has sections, which list its slides
    by id; one naming a slide the script removed makes PowerPoint repair the file."""
    template = tmp_path / "Brand.pptx"
    brand = pptx.Presentation()
    # One more slide than the example makes, so a section cannot name only new ones.
    for title in ("Old cover", "Old agenda", "Old close"):
        brand.slides.add_slide(brand.slide_layouts[0]).shapes.title.text = title
    ids = [slide.slide_id for slide in brand.slides]
    presentation = brand.part._element
    extensions = presentation.find(f"{{{presentation.nsmap['p']}}}extLst")
    if extensions is None:
        extensions = etree.SubElement(
            presentation, f"{{{presentation.nsmap['p']}}}extLst"
        )
    extension = etree.SubElement(
        extensions,
        f"{{{presentation.nsmap['p']}}}ext",
        uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}",
    )
    sections = etree.SubElement(extension, f"{{{_P14}}}sectionLst")
    section = etree.SubElement(
        sections,
        f"{{{_P14}}}section",
        name="Default Section",
        id="{0E5B2E3A-6C0F-4B7E-9C51-1D2B3C4D5E6F}",
    )
    listed = etree.SubElement(section, f"{{{_P14}}}sldIdLst")
    for slide_id in ids:
        etree.SubElement(listed, f"{{{_P14}}}sldId", id=str(slide_id))
    brand.save(template)

    result = run_document_script(
        _examples()["PowerPoint from a template"],
        output_name="document.pptx",
        images={},
        template=template,
    )

    assert result.ok, f"{result.error}\n{result.traceback_tail}"
    deck = pptx.Presentation(BytesIO(result.output))
    kept = {str(slide.slide_id) for slide in deck.slides}
    named = {
        listed_id.get("id") for listed_id in deck.part._element.iter(f"{{{_P14}}}sldId")
    }
    # Sections, when any are left, must hold exactly the deck's slides.
    assert not named or named == kept, f"sections name {named}, the deck has {kept}"
