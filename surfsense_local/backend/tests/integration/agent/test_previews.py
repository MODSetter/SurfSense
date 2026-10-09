"""Page images of a rendered document, for the agent to check what it made.

PDF pages are drawn in the API. Word pages come back as a PDF from Electron,
which a stand-in plays here over the real routes on a real port, holding the
key Electron hands the API at launch.
"""

import asyncio
import threading
import time
from collections.abc import Callable, Iterator
from io import BytesIO
from pathlib import Path
from typing import Any

import docx
import httpx
import pptx
import pytest
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlalchemy import Engine
from sqlalchemy.orm import Session

from api.main import create_app
from modules.agent.previews import Previews, docx_snapshots, previews_for, snapshot_key
from modules.artifacts.models import Artifact, ArtifactFile, ArtifactFileRole
from modules.documents.models import Document, DocumentStatus, DocumentType
from modules.workspaces.models import Workspace
from shared.config import get_agent_settings, get_storage_settings
from shared.db import create_session_factory
from tests.integration.agent.conftest import MAX_PATH

pytestmark = pytest.mark.integration

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PREVIEWS = "/agent/previews/docx-snapshots"
SNAPSHOT_KEY = "the-key-electron-minted"
ELECTRON = {"Authorization": f"Bearer {SNAPSHOT_KEY}"}


@pytest.fixture(autouse=True)
def fresh_snapshots(monkeypatch: pytest.MonkeyPatch) -> docx_snapshots.DocxSnapshots:
    """No Electron has polled this test's API yet."""
    fresh = docx_snapshots.DocxSnapshots()
    monkeypatch.setattr(docx_snapshots, "snapshots", fresh)
    return fresh


@pytest.fixture(autouse=True)
def electron_launched_the_api(monkeypatch: pytest.MonkeyPatch) -> None:
    """The API was started by Electron, which chose the snapshot key."""
    monkeypatch.setattr(
        snapshot_key.get_snapshot_key_settings(), "docx_snapshot_key", SNAPSHOT_KEY
    )


@pytest.fixture
def ready_artifact(engine: Engine) -> Callable[..., Artifact]:
    """Make a ready, versioned artifact whose primary file holds the given bytes.

    Returned detached with its files loaded, as a caller that waits must hold
    no transaction.
    """

    def make(format: str, data: bytes, mime: str, number: int = 1) -> Artifact:
        with create_session_factory(engine)() as session:
            artifact = _artifact(session, format, data, mime, number)
            session.commit()
            assert artifact.files
        return artifact

    return make


def _artifact(
    session: Session, format: str, data: bytes, mime: str, number: int
) -> Artifact:
    workspace = Workspace(name="Proposals")
    session.add(workspace)
    session.flush()
    document = Document(
        workspace_id=workspace.id,
        title="Client proposal",
        document_type=DocumentType.ARTIFACT,
        status=DocumentStatus.READY,
    )
    session.add(document)
    session.flush()
    artifact = Artifact(
        document_id=document.id, workspace_id=workspace.id, format=format
    )
    session.add(artifact)
    session.flush()
    artifact.artifact_metadata = {
        "version": {"root": artifact.id, "number": number, "parent": None}
    }
    storage = get_storage_settings()
    path = storage.artifact_dir(workspace.id, artifact.id) / f"primary.{format}"
    path.parent.mkdir(parents=True)
    path.write_bytes(data)
    artifact.files.append(
        ArtifactFile(
            role=ArtifactFileRole.PRIMARY,
            storage_key=str(path.relative_to(storage.data_dir)),
            original_filename=f"Client proposal.{format}",
            mime_type=mime,
            size_bytes=len(data),
            checksum_sha256="0" * 64,
        )
    )
    return artifact


def _pdf(pages: int) -> bytes:
    """An A4 PDF with a line of text on each page."""
    out = BytesIO()
    drawing = canvas.Canvas(out, pagesize=A4)
    for number in range(1, pages + 1):
        drawing.drawString(72, 760, f"Page {number}")
        drawing.showPage()
    drawing.save()
    return out.getvalue()


def _pdf_of_sizes(*sizes: tuple[float, float]) -> bytes:
    """A PDF with one page of each size, in points."""
    out = BytesIO()
    drawing = canvas.Canvas(out)
    for size in sizes:
        drawing.setPageSize(size)
        drawing.showPage()
    drawing.save()
    return out.getvalue()


def _docx() -> bytes:
    out = BytesIO()
    document = docx.Document()
    document.add_heading("Client proposal", level=1)
    document.save(out)
    return out.getvalue()


def _thread_folder(artifact: Artifact) -> Path:
    """The folder of the thread whose render asked for the previews."""
    return get_storage_settings().thread_working_dir(artifact.workspace_id, 1)


def _preview_folder(artifact: Artifact, number: int) -> Path:
    return (
        _thread_folder(artifact) / "outputs" / "previews" / f"{artifact.id}-v{number}"
    )


async def test_a_pdf_gets_its_pages_as_1000_pixel_wide_images_in_the_agent_folder(
    ready_artifact: Callable[..., Artifact],
) -> None:
    """A PDF version's pages land in outputs/previews/<id>-v<n>/ as PNGs the agent can read."""
    artifact = ready_artifact("pdf", _pdf(2), "application/pdf", number=3)

    previews = await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    folder = _preview_folder(artifact, 3)
    assert previews == Previews(
        pages=[folder / "page-1.png", folder / "page-2.png"], reason=None
    )
    for page in previews.pages:
        with Image.open(page) as image:
            assert image.format == "PNG"
            assert image.width == 1000


async def test_only_the_first_four_pages_are_drawn(
    ready_artifact: Callable[..., Artifact],
) -> None:
    """A long document costs the model at most four images."""
    artifact = ready_artifact("pdf", _pdf(6), "application/pdf")

    previews = await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    assert [page.name for page in previews.pages] == [
        "page-1.png",
        "page-2.png",
        "page-3.png",
        "page-4.png",
    ]


async def test_drawing_a_version_again_replaces_its_old_pages(
    ready_artifact: Callable[..., Artifact],
) -> None:
    """No page from an earlier drawing of the same version is left to mislead the agent."""
    artifact = ready_artifact("pdf", _pdf(1), "application/pdf")
    stale = _preview_folder(artifact, 1) / "page-2.png"
    stale.parent.mkdir(parents=True)
    stale.write_bytes(b"from an earlier run")

    previews = await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    assert [page.name for page in previews.pages] == ["page-1.png"]
    assert not stale.exists()


async def test_a_page_is_drawn_no_larger_than_an_image_the_model_can_take(
    ready_artifact: Callable[..., Artifact],
) -> None:
    """A tall page is narrowed to fit the height; one too thin to read is skipped and named.

    At 1,000 px wide, a 2 x 14,400 pt page would be a 28 GB bitmap in the API.
    """
    sizes = [A4, (2, 14400), (595, 4000)]
    artifact = ready_artifact("pdf", _pdf_of_sizes(*sizes), "application/pdf")

    previews = await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    assert [page.name for page in previews.pages] == ["page-1.png", "page-3.png"]
    with Image.open(previews.pages[0]) as cover:
        assert cover.width == 1000
    with Image.open(previews.pages[1]) as tall:
        assert tall.height <= 4000
        assert tall.width < 1000
    assert previews.reason is not None
    assert "Page 2 was not drawn" in previews.reason
    assert "2 x 14400 pt" in previews.reason


async def test_previews_whose_paths_cannot_fit_are_not_drawn_and_say_why(
    ready_artifact: Callable[..., Artifact], tmp_path: Path
) -> None:
    """Under a thread folder this deep, `outputs/previews/` has no room for a page's name."""
    artifact = ready_artifact("pdf", _pdf(2), "application/pdf")
    folder = tmp_path / ("t" * (MAX_PATH - len(str(tmp_path)) - 20))

    previews = previews_for(artifact, folder)

    assert previews.pages == []
    assert "too deep on this computer" in (previews.reason or "")
    assert not (folder / "outputs").exists()


async def test_a_file_that_is_not_a_pdf_gives_no_pages_and_says_why(
    ready_artifact: Callable[..., Artifact],
) -> None:
    """A file pdfium cannot open is a reason, not a failed tool call."""
    artifact = ready_artifact("pdf", b"not a pdf at all", "application/pdf")

    previews = await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    assert previews.pages == []
    assert previews.reason is not None
    assert previews.reason.startswith("The pages could not be drawn")


async def test_a_version_whose_file_is_gone_gives_no_pages_and_says_why(
    ready_artifact: Callable[..., Artifact],
) -> None:
    """A version whose file left the disk is a reason, not a failed tool call."""
    artifact = ready_artifact("pdf", _pdf(1), "application/pdf")
    storage = get_storage_settings()
    (storage.data_dir / artifact.files[0].storage_key).unlink()

    previews = await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    assert previews == Previews(
        [], "The document's file is missing, so no pages were drawn."
    )


async def test_an_artifact_that_is_no_document_version_is_refused(
    ready_artifact: Callable[..., Artifact],
) -> None:
    """Previews are drawn for document versions only, which name their folder."""
    artifact = ready_artifact("pdf", _pdf(1), "application/pdf")
    artifact.artifact_metadata = None

    with pytest.raises(ValueError, match="not a document version"):
        previews_for(artifact, _thread_folder(artifact))


async def test_without_electron_a_word_document_gets_no_pages_and_a_reason(
    ready_artifact: Callable[..., Artifact],
) -> None:
    """Without the desktop app (Docker, tests) Word gets no pages at once, and says why."""
    artifact = ready_artifact("docx", _docx(), DOCX_MIME)

    previews = await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    assert previews.pages == []
    assert previews.reason is not None
    assert "desktop app" in previews.reason


async def test_a_word_snapshot_waits_no_longer_than_the_caller_has_left(
    ready_artifact: Callable[..., Artifact],
    fresh_snapshots: docx_snapshots.DocxSnapshots,
) -> None:
    """The render tool's call has a deadline; a late preview must not run past it."""
    artifact = ready_artifact("docx", _docx(), DOCX_MIME)
    fresh_snapshots.next_request()  # Electron polls, but never prints this one
    started = time.monotonic()

    previews = await asyncio.to_thread(
        previews_for, artifact, _thread_folder(artifact), time_left=3.5
    )

    assert time.monotonic() - started < 10
    assert previews == Previews([], "The Word preview did not finish within 3.5 s.")


async def test_with_too_little_time_left_a_word_document_is_not_sent_to_print(
    ready_artifact: Callable[..., Artifact],
    fresh_snapshots: docx_snapshots.DocxSnapshots,
) -> None:
    """Electron could not print in the time left, so it is not asked to, and the model learns why."""
    artifact = ready_artifact("docx", _docx(), DOCX_MIME)
    fresh_snapshots.next_request()

    previews = await asyncio.to_thread(
        previews_for, artifact, _thread_folder(artifact), time_left=1
    )

    assert previews.pages == []
    assert previews.reason is not None
    assert "no time left" in previews.reason


class StandInForElectron:
    """Polls the API as Electron does and answers each snapshot request it gets."""

    def __init__(self, base_url: str, answer: Callable[[httpx.Client, dict], None]):
        self.base_url = base_url
        self.answer = answer
        self.polled = threading.Event()
        self.served: list[dict[str, Any]] = []
        self._stop = threading.Event()
        self._thread = threading.Thread(target=self._run, daemon=True)

    def __enter__(self) -> "StandInForElectron":
        self._thread.start()
        assert self.polled.wait(5), "the stand-in never reached the API"
        return self

    def __exit__(self, *exc: object) -> None:
        self._stop.set()
        self._thread.join(timeout=5)

    def _run(self) -> None:
        with httpx.Client(base_url=self.base_url, headers=ELECTRON, timeout=10) as http:
            while not self._stop.is_set():
                reply = http.get(f"{PREVIEWS}/next")
                self.polled.set()
                if reply.status_code == 200:
                    self.served.append(reply.json())
                    self.answer(http, reply.json())
                self._stop.wait(0.05)


def _print_two_pages(http: httpx.Client, request: dict) -> None:
    """Fetch the Word file the request names and post a two-page PDF back."""
    word = http.get(request["file_url"])
    assert word.status_code == 200
    assert word.content.startswith(b"PK")  # a docx is a zip
    posted = http.post(
        f"{PREVIEWS}/{request['id']}/pdf",
        content=_pdf(2),
        headers={"content-type": "application/pdf"},
    )
    assert posted.status_code == 204


async def test_a_word_document_is_printed_by_electron_and_drawn_as_pages(
    base_url: str, ready_artifact: Callable[..., Artifact]
) -> None:
    """Electron fetches the Word file the request names, and the PDF it posts becomes pages."""
    artifact = ready_artifact("docx", _docx(), DOCX_MIME, number=2)

    with StandInForElectron(base_url, _print_two_pages) as electron:
        previews = await asyncio.to_thread(
            previews_for, artifact, _thread_folder(artifact)
        )

    assert electron.served[0]["file_url"] == f"/artifacts/{artifact.id}/files/primary"
    folder = _preview_folder(artifact, 2)
    assert previews == Previews(
        pages=[folder / "page-1.png", folder / "page-2.png"], reason=None
    )


async def test_a_failure_electron_reports_becomes_the_reason(
    base_url: str, ready_artifact: Callable[..., Artifact]
) -> None:
    """What went wrong in Electron reaches the model in the reason."""
    artifact = ready_artifact("docx", _docx(), DOCX_MIME)

    def report_failure(http: httpx.Client, request: dict) -> None:
        http.post(
            f"{PREVIEWS}/{request['id']}/failure",
            json={"reason": "docx-preview could not read the file"},
        ).raise_for_status()

    with StandInForElectron(base_url, report_failure):
        previews = await asyncio.to_thread(
            previews_for, artifact, _thread_folder(artifact)
        )

    assert previews.pages == []
    assert previews.reason == (
        "The Word preview failed: docx-preview could not read the file"
    )


@pytest.fixture
def http(base_url: str) -> Iterator[httpx.Client]:
    """A client on the real API's port, holding Electron's key."""
    with httpx.Client(base_url=base_url, headers=ELECTRON, timeout=10) as client:
        yield client


async def test_with_nothing_waiting_the_poll_is_empty(http: httpx.Client) -> None:
    """Electron's poll costs nothing when no Word document waits."""
    reply = http.get(f"{PREVIEWS}/next")

    assert reply.status_code == 204


@pytest.mark.parametrize(
    ("method", "route", "body"),
    [
        ("GET", "next", None),
        ("POST", "some-request/pdf", b"%PDF-1.7"),
        ("POST", "some-request/failure", b'{"reason": "no"}'),
    ],
)
async def test_a_web_page_may_not_call_the_snapshot_routes(
    http: httpx.Client, method: str, route: str, body: bytes | None
) -> None:
    """A page in the user's browser can neither take a request nor answer one."""
    reply = http.request(
        method,
        f"{PREVIEWS}/{route}",
        content=body,
        headers={"origin": "https://example.com"},
    )

    assert reply.status_code == 403


@pytest.mark.parametrize(
    ("method", "route", "body"),
    [
        ("GET", "next", None),
        ("POST", "some-request/pdf", b"%PDF-1.7"),
        ("POST", "some-request/failure", b'{"reason": "no"}'),
    ],
)
@pytest.mark.parametrize(
    "authorization",
    [None, "Bearer a-guess", SNAPSHOT_KEY],
    ids=["no key", "another key", "the key without its scheme"],
)
async def test_a_caller_without_electrons_key_may_not_call_the_snapshot_routes(
    base_url: str,
    method: str,
    route: str,
    body: bytes | None,
    authorization: str | None,
) -> None:
    """Another process on this machine sends no Origin either, so the key is what tells."""
    headers = {} if authorization is None else {"Authorization": authorization}
    with httpx.Client(base_url=base_url, timeout=10) as http:
        reply = http.request(
            method, f"{PREVIEWS}/{route}", content=body, headers=headers
        )

    assert reply.status_code == 401


async def test_a_refused_poll_does_not_count_as_electron_running(
    base_url: str, ready_artifact: Callable[..., Artifact]
) -> None:
    """A process polling without the key gets no request, and the API still says the app is not running."""
    with httpx.Client(base_url=base_url, timeout=10) as http:
        assert http.get(f"{PREVIEWS}/next").status_code == 401

    artifact = ready_artifact("docx", _docx(), DOCX_MIME)
    previews = await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    assert previews.pages == []
    assert "desktop app" in (previews.reason or "")


@pytest.mark.parametrize("no_key", [None, ""])
async def test_an_api_electron_gave_no_key_refuses_every_caller(
    http: httpx.Client, monkeypatch: pytest.MonkeyPatch, no_key: str | None
) -> None:
    """Started by hand beside an opencode, the API has no key to check, so nothing is let in."""
    monkeypatch.setattr(
        snapshot_key.get_snapshot_key_settings(), "docx_snapshot_key", no_key
    )

    assert http.get(f"{PREVIEWS}/next").status_code == 401


async def test_an_answer_to_no_waiting_request_is_not_found(http: httpx.Client) -> None:
    """An answer that comes after its waiter gave up, or for no request, is refused."""
    pdf = http.post(
        f"{PREVIEWS}/no-such-request/pdf",
        content=_pdf(1),
        headers={"content-type": "application/pdf"},
    )
    failure = http.post(
        f"{PREVIEWS}/no-such-request/failure", json={"reason": "nothing"}
    )

    assert pdf.status_code == 404
    assert failure.status_code == 404


async def test_a_body_that_is_not_a_pdf_is_refused(http: httpx.Client) -> None:
    """Only a PDF is handed to pdfium as a snapshot."""
    reply = http.post(
        f"{PREVIEWS}/no-such-request/pdf",
        content=b"<html>not a pdf</html>",
        headers={"content-type": "application/pdf"},
    )

    assert reply.status_code == 422


async def test_an_api_without_opencode_has_no_snapshot_routes(
    engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Previews exist for the agent alone, so Electron's poll finds nothing to call."""
    monkeypatch.setattr(get_agent_settings(), "opencode_url", None)
    monkeypatch.setattr(get_agent_settings(), "opencode_password", None)
    app = create_app()
    app.state.session_factory = create_session_factory(engine)

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        reply = await client.get(f"{PREVIEWS}/next")

    assert reply.status_code == 404


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _pptx() -> bytes:
    deck = pptx.Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Quarterly review"
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


async def test_a_deck_is_printed_by_electron_as_a_pptx_and_drawn_as_slides(
    base_url: str, ready_artifact: Callable[..., Artifact]
) -> None:
    """Electron learns the file is a deck, and prints its first four slides."""
    artifact = ready_artifact("pptx", _pptx(), PPTX_MIME)

    def print_slides(http: httpx.Client, request: dict) -> None:
        deck = http.get(request["file_url"])
        assert deck.content.startswith(b"PK")
        http.post(
            f"{PREVIEWS}/{request['id']}/pdf",
            content=_pdf(1),
            headers={"content-type": "application/pdf"},
        ).raise_for_status()

    with StandInForElectron(base_url, print_slides) as electron:
        previews = await asyncio.to_thread(
            previews_for, artifact, _thread_folder(artifact)
        )

    assert electron.served == [
        {
            "id": electron.served[0]["id"],
            "file_url": f"/artifacts/{artifact.id}/files/primary",
            "format": "pptx",
            "pages": "1-4",
        }
    ]
    assert previews == Previews([_preview_folder(artifact, 1) / "page-1.png"], None)


async def test_without_electron_a_deck_gets_no_slides_and_a_reason(
    ready_artifact: Callable[..., Artifact],
) -> None:
    """Slides are Electron's to print, as Word pages are."""
    artifact = ready_artifact("pptx", _pptx(), PPTX_MIME)

    previews = await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    assert previews.pages == []
    assert "PowerPoint slides are drawn by the SurfSense desktop app" in (
        previews.reason or ""
    )


async def test_a_word_request_names_its_format_and_pages(
    base_url: str, ready_artifact: Callable[..., Artifact]
) -> None:
    """Every request says how to lay the file out and which pages to print."""
    artifact = ready_artifact("docx", _docx(), DOCX_MIME)

    with StandInForElectron(base_url, _print_two_pages) as electron:
        await asyncio.to_thread(previews_for, artifact, _thread_folder(artifact))

    assert (electron.served[0]["format"], electron.served[0]["pages"]) == (
        "docx",
        "1-4",
    )


async def test_a_sources_file_is_served_to_the_print_window_only_while_it_prints(
    base_url: str,
    tmp_path: Path,
    fresh_snapshots: docx_snapshots.DocxSnapshots,
) -> None:
    """The window fetches it with neither key nor a missing Origin, as it fetches an artifact's file."""
    original = tmp_path / "Brand.pptx"
    original.write_bytes(_pptx())
    fetched: list[httpx.Response] = []

    def fetch_as_the_window(http: httpx.Client, request: dict) -> None:
        with httpx.Client(base_url=base_url, timeout=10) as window:
            fetched.append(window.get(request["file_url"], headers={"origin": "null"}))
        http.post(
            f"{PREVIEWS}/{request['id']}/pdf",
            content=_pdf(2),
            headers={"content-type": "application/pdf"},
        ).raise_for_status()
        with httpx.Client(base_url=base_url, timeout=10) as window:
            fetched.append(window.get(request["file_url"]))

    with StandInForElectron(base_url, fetch_as_the_window) as electron:
        pdf = await asyncio.to_thread(
            fresh_snapshots.snapshot,
            format="pptx",
            source_file=original,
            pages="2,5",
            timeout=10,
        )

    (request,) = electron.served
    assert request["file_url"] == f"{PREVIEWS}/{request['id']}/file"
    assert (request["format"], request["pages"]) == ("pptx", "2,5")
    assert fetched[0].status_code == 200
    assert fetched[0].content == original.read_bytes()
    assert fetched[1].status_code == 404  # answered: no longer served
    assert pdf.startswith(b"%PDF")


async def test_no_file_is_served_for_a_request_no_one_made(http: httpx.Client) -> None:
    """The file route serves nothing but a source being printed."""
    reply = http.get(f"{PREVIEWS}/no-such-request/file")

    assert reply.status_code == 404
