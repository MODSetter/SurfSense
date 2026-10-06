"""The read tool over a workbook: the cells, addresses and formulas a revision names."""

from io import BytesIO

import openpyxl
import pytest

from modules.source_scope.schemas import SourceScope
from tests.integration.agent.tool_endpoint_client import ToolEndpoint
from tests.integration.artifacts.revised_copies.source_files import (
    add_source,
    contract_docx,
    pricing_xlsx,
)
from tests.integration.worker.conftest import stub_model  # noqa: F401

pytestmark = [
    pytest.mark.integration,
    pytest.mark.usefixtures("stub_model", "model_reads_images"),
]

TOOL = "read_document"


def _source(tools: ToolEndpoint, workspace_id: int, name: str, data: bytes) -> int:
    with tools.sessions() as session:
        return add_source(session, workspace_id, name, data)


def _budget_xlsx() -> bytes:
    """Two sheets: a text that looks like a formula, and a total across sheets."""
    book = openpyxl.Workbook()
    costs = book.active
    costs.title = "Costs"
    costs["B2"] = "Travel"
    costs["C2"] = 1250.5
    costs["C3"] = "=SUM(C2:C2)"
    costs["D2"] = '="not a formula"'
    costs["D2"].data_type = "s"
    summary = book.create_sheet("Summary")
    summary["A1"] = "=Costs!C3*2"
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _wide_xlsx(rows: int) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Ledger"
    for row in range(1, rows + 1):
        sheet.append([f"Entry {row}", row, f"=B{row}*2"])
    book.create_sheet("Notes")["A1"] = "kept"
    out = BytesIO()
    book.save(out)
    return out.getvalue()


async def test_a_workbook_source_reads_as_its_sheets_cells_and_formulas(
    tools: ToolEndpoint,
) -> None:
    """Docling's tables drop sheet names, addresses and formulas, so the model guessed them."""
    workspace_id = await tools.workspace()
    source_id = _source(tools, workspace_id, "Budget.xlsx", _budget_xlsx())

    text, is_error = await tools.call(workspace_id, TOOL, {"document_id": source_id})

    assert is_error is False, text
    lines = text.splitlines()
    assert lines[0].startswith(f'Source {source_id} ("Budget.xlsx") is a workbook')
    assert 'Sheet "Costs", used range B2:D3:' in lines
    assert 'Sheet "Summary", used range A1:A1:' in lines
    assert 'B2 = "Travel"' in lines
    assert "C2 = 1250.5" in lines
    assert "C3 = =SUM(C2:C2)" in lines
    assert 'D2 = "=\\"not a formula\\""' in lines
    assert "A1 = =Costs!C3*2" in lines
    assert "cut" not in text


async def test_one_sheet_can_be_read_on_its_own(tools: ToolEndpoint) -> None:
    """A big first sheet must not hide the one the user asked about."""
    workspace_id = await tools.workspace()
    source_id = _source(tools, workspace_id, "Budget.xlsx", _budget_xlsx())

    text, is_error = await tools.call(
        workspace_id, TOOL, {"document_id": source_id, "sheet": "Summary"}
    )

    assert is_error is False, text
    assert "A1 = =Costs!C3*2" in text
    assert "Travel" not in text


async def test_a_sheet_the_workbook_lacks_is_refused_naming_its_sheets(
    tools: ToolEndpoint,
) -> None:
    """The model's next call can name a sheet that is there."""
    workspace_id = await tools.workspace()
    source_id = _source(tools, workspace_id, "Budget.xlsx", _budget_xlsx())

    text, is_error = await tools.call(
        workspace_id, TOOL, {"document_id": source_id, "sheet": "Totals"}
    )

    assert is_error is True
    assert 'There is no sheet "Totals". Its sheets: "Costs", "Summary".' in text


async def test_a_long_map_is_cut_and_says_where(tools: ToolEndpoint) -> None:
    """A result over opencode's 50 KB is cut by opencode, silently mid-cell."""
    workspace_id = await tools.workspace()
    source_id = _source(tools, workspace_id, "Ledger.xlsx", _wide_xlsx(5000))

    text, is_error = await tools.call(workspace_id, TOOL, {"document_id": source_id})

    assert is_error is False, text
    assert len(text.encode()) < 50_000
    assert 'Sheet "Ledger", used range A1:C5000:' in text
    assert 'Sheet "Notes", used range A1:A1:' in text  # every sheet is named first
    assert 'A1 = "Entry 1"' in text
    assert "C5000 = " not in text
    cut = text.splitlines()[-1]
    assert cut.startswith("The map was cut at row ")
    row = int(cut.removeprefix("The map was cut at row ").split()[0])
    assert f'Call again with sheet "Ledger" and offset {row}' in cut


async def test_an_offset_reads_on_from_that_row(tools: ToolEndpoint) -> None:
    """The call a cut map names reads the rest of the sheet."""
    workspace_id = await tools.workspace()
    source_id = _source(tools, workspace_id, "Ledger.xlsx", _wide_xlsx(5000))

    text, is_error = await tools.call(
        workspace_id,
        TOOL,
        {"document_id": source_id, "sheet": "Ledger", "offset": 4990},
    )

    assert is_error is False, text
    assert 'A4990 = "Entry 4990"' in text
    assert "C5000 = =B5000*2" in text
    assert "A4989 = " not in text
    assert "The map was cut" not in text


async def test_a_source_the_user_did_not_tick_is_refused(tools: ToolEndpoint) -> None:
    """The original file is a source, so it keeps to the turn's ticks."""
    workspace_id = await tools.workspace()
    ticked = _source(tools, workspace_id, "Other.xlsx", pricing_xlsx())
    unticked = _source(tools, workspace_id, "Pricing.xlsx", pricing_xlsx())

    text, is_error = await tools.call(
        workspace_id,
        TOOL,
        {"document_id": unticked},
        thread=tools.thread(workspace_id, SourceScope(document_ids=[ticked])),
    )

    assert is_error is True
    assert f"Source {unticked} is not selected for this request." in text


async def test_a_source_that_is_no_workbook_is_refused(tools: ToolEndpoint) -> None:
    """Other sources' text is already in sources/."""
    workspace_id = await tools.workspace()
    source_id = _source(tools, workspace_id, "MSA.docx", contract_docx())

    text, is_error = await tools.call(workspace_id, TOOL, {"document_id": source_id})

    assert is_error is True
    assert "reads a workbook source's cells (.xlsx, .xlsm)" in text


async def test_one_id_is_needed(tools: ToolEndpoint) -> None:
    """A source and a document at once is a call the model must fix."""
    workspace_id = await tools.workspace()

    text, is_error = await tools.call(
        workspace_id, TOOL, {"document_id": 1, "artifact_id": 2}
    )

    assert is_error is True
    assert "one of them, not both" in text


async def test_a_revised_workbook_reads_as_its_newest_version(
    tools: ToolEndpoint, studio_worker: None
) -> None:
    """The next edit is made on the copy's newest version, so that is the map the model needs."""
    workspace_id = await tools.workspace()
    source_id = _source(tools, workspace_id, "Pricing.xlsx", pricing_xlsx())
    revised, _ = await tools.call(
        workspace_id,
        "revise_document",
        {
            "document_id": source_id,
            "operations": [
                {"op": "set_cell", "sheet": "Pricing", "cell": "B2", "value": 15000}
            ],
        },
    )
    artifact_id = int(revised.split(",")[0].removeprefix("Rendered artifact "))

    text, is_error = await tools.call(workspace_id, TOOL, {"artifact_id": artifact_id})

    assert is_error is False, text
    assert text.splitlines()[0].startswith(
        f'Artifact {artifact_id}, version 1 of the revised copy of "Pricing.xlsx", '
        "is a workbook"
    )
    assert "B2 = 15000" in text
    assert "B4 = =SUM(B2:B3)" in text


def _wide_row_xlsx(columns: int) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Prices"
    sheet.append([f"Price {column}" for column in range(columns)])
    sheet.append(["after the wide row"])
    out = BytesIO()
    book.save(out)
    return out.getvalue()


async def test_a_row_too_wide_for_one_map_is_read_past(tools: ToolEndpoint) -> None:
    """Calling again from the cut row cut it at the same cell, so the rest was never read."""
    workspace_id = await tools.workspace()
    source_id = _source(tools, workspace_id, "Prices.xlsx", _wide_row_xlsx(5000))
    arguments: dict = {"document_id": source_id, "sheet": "Prices"}

    first, _ = await tools.call(workspace_id, TOOL, arguments)
    cut = first.splitlines()[-1]
    offset = int(cut.split(" and offset ")[1].split()[0])
    after, is_error = await tools.call(
        workspace_id, TOOL, {**arguments, "offset": offset}
    )

    assert is_error is False, after
    assert 'Row 1 of sheet "Prices" holds more cells than one map shows' in cut
    assert 'A2 = "after the wide row"' in after


def _many_sheets_xlsx(sheets: int) -> bytes:
    book = openpyxl.Workbook()
    book.active.title = "Sheet 0000"
    for number in range(1, sheets):
        book.create_sheet(f"Quarterly sheet {number:04d}")["A1"] = number
    out = BytesIO()
    book.save(out)
    return out.getvalue()


async def test_a_workbook_of_many_sheets_stays_within_one_map(
    tools: ToolEndpoint,
) -> None:
    """Every sheet's heading was added past the cut, and opencode cut the result mid-line."""
    workspace_id = await tools.workspace()
    source_id = _source(tools, workspace_id, "Quarters.xlsx", _many_sheets_xlsx(1200))

    text, is_error = await tools.call(workspace_id, TOOL, {"document_id": source_id})

    assert is_error is False, text
    assert len(text.encode()) < 42_000
    assert 'Sheet "Sheet 0000"' in text
    assert "more sheets" in text
