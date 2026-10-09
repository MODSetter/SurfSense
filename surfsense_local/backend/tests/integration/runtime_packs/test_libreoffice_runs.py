"""The runner against a real LibreOffice, named through the test seam, never found on its own.

SURFSENSE_TEST_LIBREOFFICE names an install folder (the one holding program/, or
LibreOffice.app); without it a standard install path is tried, and the tests skip
when there is none. The development machine's 25.2 is an ended branch the app
would refuse; the seam runs it anyway, because only the runner is under test.
"""

import os
import re
import sys
import time
from pathlib import Path

import docx
import psutil
import pypdfium2
import pytest
import xlsxwriter
from pptx import Presentation
from pptx.util import Inches

from modules.runtime_packs.office import layout
from modules.runtime_packs.office.program import program_in
from modules.runtime_packs.office.runtime import OfficeRuntime
from modules.runtime_packs.office.soffice import convert as convert_module
from modules.runtime_packs.office.soffice.convert import convert
from modules.runtime_packs.office.soffice.errors import OfficeTimeout
from modules.runtime_packs.office.soffice.recalc import recalc
from modules.runtime_packs.office.soffice.version import report_version

pytestmark = pytest.mark.office

_STANDARD = {
    "win32": [Path(r"C:\Program Files\LibreOffice")],
    "darwin": [Path("/Applications/LibreOffice.app")],
}.get(sys.platform, [Path("/usr/lib/libreoffice")])


@pytest.fixture
def office() -> OfficeRuntime:
    """The LibreOffice the seam names, as an install the runner is handed."""
    named = os.environ.get("SURFSENSE_TEST_LIBREOFFICE")
    for root in [Path(named)] if named else _STANDARD:
        if program_in(root).is_file():
            return OfficeRuntime(program_in(root), "test", "installed")
    pytest.skip("no LibreOffice; set SURFSENSE_TEST_LIBREOFFICE to an install folder")


def _soon(seconds: float = 120) -> float:
    return time.monotonic() + seconds


def _pdf_text(path: Path) -> str:
    pdf = pypdfium2.PdfDocument(path)
    try:
        return "".join(page.get_textpage().get_text_range() for page in pdf)
    finally:
        pdf.close()


def test_a_docx_becomes_a_pdf_and_the_original_is_untouched(
    office: OfficeRuntime, tmp_path: Path
) -> None:
    """The PDF carries the document's words; the file handed in keeps its bytes."""
    source = tmp_path / "memo.docx"
    document = docx.Document()
    document.add_heading("Quarterly memo", level=1)
    document.add_paragraph("Revenue rose in every region.")
    document.save(source)
    before = source.read_bytes()

    pdf = convert(source, "pdf", tmp_path / "out", deadline=_soon(), runtime=office)

    assert pdf == tmp_path / "out" / "memo.pdf"
    assert pdf.read_bytes().startswith(b"%PDF-")
    assert "Revenue rose in every region." in _pdf_text(pdf)
    assert source.read_bytes() == before


def test_a_pptx_becomes_a_pdf_with_a_page_per_slide(
    office: OfficeRuntime, tmp_path: Path
) -> None:
    """Two slides, two pages."""
    source = tmp_path / "deck.pptx"
    deck = Presentation()
    for title in ("Opening slide", "Closing slide"):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    deck.save(source)

    pdf = convert(source, "pdf", tmp_path / "out", deadline=_soon(), runtime=office)

    assert len(pypdfium2.PdfDocument(pdf)) == 2
    assert "Closing slide" in _pdf_text(pdf)


def test_recalculation_replaces_the_zeros_xlsxwriter_cached(
    office: OfficeRuntime, tmp_path: Path
) -> None:
    """=A1*A2 and =SUM(A1:A3)*10 read 6 and 110; a #DIV/0! is reported apart."""
    source = tmp_path / "book.xlsx"
    book = xlsxwriter.Workbook(source)
    sheet = book.add_worksheet("Data")
    sheet.write_column("A1", [2, 3, 6])
    sheet.write_formula("B1", "=A1*A2")
    sheet.write_formula("B2", "=SUM(A1:A3)*10")
    sheet.write_formula("B3", "=1/0")
    book.close()

    result = recalc(source, deadline=_soon(), runtime=office)

    assert result.values == {"Data": {"B1": 6, "B2": 110}}
    assert result.errors == {"Data": {"B3": "#DIV/0!"}}
    assert result.version == "test"


def test_the_profile_keeps_updates_macros_links_and_proxies_off_after_a_run(
    office: OfficeRuntime, tmp_path: Path
) -> None:
    """LibreOffice keeps what it was seeded with when it writes its profile back."""
    source = tmp_path / "a.docx"
    docx.Document().save(source)
    convert(source, "pdf", tmp_path / "out", deadline=_soon(), runtime=office)

    settings = (layout.profile_slot() / "user" / "registrymodifications.xcu").read_text(
        encoding="utf-8"
    )
    jobs = "/org.openoffice.Office.Jobs/Jobs/org.openoffice.Office.Jobs:Job['UpdateCheck']/Arguments"
    for path, name, value in (
        ("/org.openoffice.Office.Update/Update", "Enabled", "false"),
        (jobs, "AutoCheckEnabled", "false"),
        (jobs, "AutoDownloadEnabled", "false"),
        ("/org.openoffice.Office.Calc/Formula/Load", "OOXMLRecalcMode", "0"),
        (
            "/org.openoffice.Office.Common/Security/Scripting",
            "DisableMacrosExecution",
            "true",
        ),
        ("/org.openoffice.Office.Calc/Content/Update", "Link", "1"),
        ("/org.openoffice.Office.Writer/Content/Update", "Link", "2"),
        ("/org.openoffice.Inet/Settings", "ooInetHTTPSProxyPort", "9"),
    ):
        assert _setting(settings, path, name) == value, name


def _setting(settings: str, path: str, name: str) -> str | None:
    found = re.search(
        rf'<item oor:path="{re.escape(path)}"><prop oor:name="{name}"[^>]*>'
        r"<value>([^<]*)</value>",
        settings,
    )
    return found.group(1) if found else None


def test_a_run_past_its_limit_is_killed_with_everything_it_started(
    office: OfficeRuntime, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """No soffice process outlives the timeout; too little time is left to retry."""
    monkeypatch.setattr(convert_module, "RUN_LIMIT_SECONDS", 0.3)
    source = tmp_path / "a.docx"
    docx.Document().save(source)

    with pytest.raises(OfficeTimeout):
        convert(source, "pdf", tmp_path / "out", deadline=_soon(25), runtime=office)

    profile = str(layout.profile_slot().resolve().as_uri())
    left = [
        process
        for process in psutil.process_iter(["cmdline"])
        if any(profile in part for part in process.info["cmdline"] or ())
    ]
    assert left == []


def test_an_install_reports_its_version(office: OfficeRuntime) -> None:
    """Such as 25.2.7.2: what Settings shows for a confirmed install."""
    version = report_version(office, deadline=_soon())
    assert version.split(".")[0].isdigit()
    assert version.count(".") == 3
