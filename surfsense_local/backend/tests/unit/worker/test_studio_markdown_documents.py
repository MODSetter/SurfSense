"""Studio's Markdown spec becomes a Word file or a PDF through committed builders."""

import json
from io import BytesIO
from pathlib import Path

import docx
import pypdfium2
import pytest
from PIL import Image

from shared import pdfium
from worker.studio.office.markdown import markdown_to_pdf, markdown_to_word
from worker.studio.script_document.extracted_text import pdf_text, word_text

pytestmark = pytest.mark.unit

REPORT = """\
# Halvorsen Freight proposal

We propose a **two-phase** rollout, *starting in March*, as the [brief](https://example.com/brief) asks.

## Phases

- Pilot at the Oslo depot
  - Two trucks
  - One dispatcher
- Rollout to every depot

1. Sign the contract
2. Start the pilot

## Pricing

| Phase | Cost |
|---|---|
| Pilot | 12,000 |
| Rollout | 48,000 |
"""

HEADINGS_AND_CELLS = [
    "Halvorsen Freight proposal",
    "Phases",
    "Pricing",
    "Pilot at the Oslo depot",
    "Two trucks",
    "Sign the contract",
    "12,000",
    "48,000",
    "Rollout",
]

CHART = {
    "type": "bar",
    "title": "Yearly costs",
    "labels": ["2025", "2026"],
    "series": [{"name": "Cost", "values": [12000, 48000]}],
}


def _figure(tmp_path: Path) -> dict[str, Path]:
    path = tmp_path / "12-1.png"
    Image.new("RGB", (400, 200), "red").save(path)
    return {"12-1": path}


def _word(data: bytes) -> docx.document.Document:
    return docx.Document(BytesIO(data))


def _pictures(document: docx.document.Document) -> int:
    return len(document.part.package.image_parts)


def _pdf_images(data: bytes) -> int:
    with pdfium.lock:
        pdf = pypdfium2.PdfDocument(data)
        try:
            return sum(
                1
                for page in pdf
                for obj in page.get_objects()
                if obj.type == pypdfium2.raw.FPDF_PAGEOBJ_IMAGE
            )
        finally:
            pdf.close()


def _chart_block(chart: object) -> str:
    body = chart if isinstance(chart, str) else json.dumps(chart)
    return f"# Costs\n\n```chart\n{body}\n```\n"


def test_a_report_becomes_a_word_file_holding_every_heading_and_cell() -> None:
    """Headings stay headings and table cells stay cells."""
    text = word_text(markdown_to_word(REPORT, {}))

    for expected in HEADINGS_AND_CELLS:
        assert expected in text
    assert "# Halvorsen Freight proposal" in text
    assert "## Pricing" in text
    assert "Phase | Cost" in text


def test_a_report_becomes_a_pdf_holding_every_heading_and_cell() -> None:
    """The PDF's text layer holds every heading and cell."""
    data = markdown_to_pdf(REPORT, {})

    assert data.startswith(b"%PDF")
    text = pdf_text(data)
    for expected in HEADINGS_AND_CELLS:
        assert expected in text


def test_word_keeps_emphasis_lists_and_links_as_word_structure() -> None:
    """Marks become run properties, lists Word's list styles, links hyperlinks."""
    document = _word(markdown_to_word(REPORT, {}))

    runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
    assert any(run.bold and run.text == "two-phase" for run in runs)
    assert any(run.italic and run.text == "starting in March" for run in runs)
    styles = {paragraph.text: paragraph.style.name for paragraph in document.paragraphs}
    assert styles["Pilot at the Oslo depot"] == "List Bullet"
    assert styles["Two trucks"] == "List Bullet 2"
    assert styles["Sign the contract"] == "List Number"
    assert "https://example.com/brief" in {
        rel.target_ref for rel in document.part.rels.values() if rel.is_external
    }
    (table,) = document.tables
    assert table.style.name == "Table Grid"


def test_a_known_source_figure_is_placed_in_both_formats(tmp_path: Path) -> None:
    """An image: reference to a figure on the shelf places its PNG."""
    markdown = "# Logo\n\n![Halvorsen logo](image:12-1)\n"
    figures = _figure(tmp_path)

    assert _pictures(_word(markdown_to_word(markdown, figures))) == 1
    assert _pdf_images(markdown_to_pdf(markdown, figures)) == 1


def test_an_unknown_figure_becomes_its_caption_in_italics(tmp_path: Path) -> None:
    """A figure no source holds, or any other target, leaves its caption."""
    markdown = "# Logo\n\n![Halvorsen logo](image:99-4)\n\n![Remote](https://example.com/x.png)\n"
    figures = _figure(tmp_path)

    document = _word(markdown_to_word(markdown, figures))
    assert _pictures(document) == 0
    italic = [r.text for p in document.paragraphs for r in p.runs if r.italic]
    assert "Halvorsen logo" in italic
    assert "Remote" in italic

    data = markdown_to_pdf(markdown, figures)
    assert _pdf_images(data) == 0
    assert "Halvorsen logo" in pdf_text(data)


def test_a_valid_chart_block_is_drawn_as_a_picture() -> None:
    """A chart block that matches its shape is drawn, not tabulated."""
    markdown = _chart_block(CHART)

    document = _word(markdown_to_word(markdown, {}))
    assert _pictures(document) == 1
    assert not document.tables
    assert _pdf_images(markdown_to_pdf(markdown, {})) == 1


@pytest.mark.parametrize("kind", ["line", "pie"])
def test_line_and_pie_charts_are_drawn_too(kind: str) -> None:
    """Each of the three chart types draws."""
    assert (
        _pictures(_word(markdown_to_word(_chart_block({**CHART, "type": kind}), {})))
        == 1
    )


def test_a_chart_that_does_not_match_becomes_a_table_of_its_data() -> None:
    """A chart that cannot be drawn keeps its data, with a note why."""
    mismatched = {**CHART, "type": "radar"}
    markdown = _chart_block(mismatched)

    document = _word(markdown_to_word(markdown, {}))
    assert _pictures(document) == 0
    text = word_text(markdown_to_word(markdown, {}))
    assert "2025 | 12000" in text
    assert "shown as a table" in text

    pdf = pdf_text(markdown_to_pdf(markdown, {}))
    assert "48000" in pdf
    assert "shown as a table" in pdf


def test_a_chart_that_does_not_parse_still_renders_with_a_note() -> None:
    """Broken JSON does not fail the document."""
    markdown = _chart_block("{not json")

    text = word_text(markdown_to_word(markdown, {}))
    assert "could not be drawn" in text
    assert "could not be drawn" in pdf_text(markdown_to_pdf(markdown, {}))


def test_text_that_reportlab_would_read_as_markup_is_kept_as_written() -> None:
    """ReportLab's markup characters in the text are escaped."""
    markdown = "# A & B\n\nKeep <b>this</b> & that.\n"

    text = pdf_text(markdown_to_pdf(markdown, {}))

    assert "A & B" in text
    assert "Keep <b>this</b> & that." in text


def test_control_characters_a_model_wrote_do_not_break_either_builder() -> None:
    """Characters XML cannot hold are dropped; the text around them stays."""
    markdown = "# Report\x00\n\nBefore\x01 \x0bafter \ufffe end.\n"

    assert "Before after  end." in word_text(markdown_to_word(markdown, {}))
    assert "Before after" in pdf_text(markdown_to_pdf(markdown, {}))


def test_a_quote_inside_a_list_renders_in_both_formats() -> None:
    """A quote nested in a list item keeps its text in Word and in the PDF."""
    markdown = "- Point\n  > Quoted in the point\n  > - with a list\n- Next\n"

    assert "Quoted in the point" in word_text(markdown_to_word(markdown, {}))
    text = pdf_text(markdown_to_pdf(markdown, {}))
    assert "Quoted in the point" in text
    assert "with a list" in text


def test_a_table_cell_longer_than_a_page_runs_onto_the_next_page() -> None:
    """A row taller than the page is split rather than failing the PDF."""
    words = " ".join(f"word{index}" for index in range(5000))
    markdown = f"| Phase | Notes |\n|---|---|\n| Pilot | {words} |\n"

    text = pdf_text(markdown_to_pdf(markdown, {}))

    assert "word0" in text
    assert "word4999" in text


@pytest.mark.parametrize("columns", [30, 60])
def test_a_table_wider_than_the_page_keeps_every_column(columns: int) -> None:
    """Too many columns for one A4 width are set in several tables, none lost."""
    header = "| " + " | ".join(f"Col{index}" for index in range(columns)) + " |"
    rule = "|" + "---|" * columns
    row = "| " + " | ".join(f"v{index}" for index in range(columns)) + " |"

    text = pdf_text(markdown_to_pdf(f"{header}\n{rule}\n{row}\n", {}))

    for index in range(columns):
        assert f"Col{index}" in text
        assert f"v{index}" in text


def test_a_chart_with_many_series_falls_back_to_a_table_that_fits() -> None:
    """A broken chart's table with more series than the page is wide still renders."""
    chart = {
        "type": "pie",
        "labels": ["2025"],
        "series": [{"name": f"S{index}", "values": [index]} for index in range(40)],
    }

    text = pdf_text(markdown_to_pdf(_chart_block(chart), {}))

    assert "shown as a table" in text
    assert "S39" in text


@pytest.mark.parametrize(
    ("kind", "values"),
    [
        ("pie", [-1, 2]),
        ("pie", [0, 0]),
        ("bar", [1, float("nan")]),
        ("line", [1, float("inf")]),
    ],
)
def test_a_chart_matplotlib_cannot_draw_becomes_a_table(
    kind: str, values: list[float]
) -> None:
    """Negative or all-zero pie slices and non-finite numbers are tabulated."""
    chart = {
        "type": kind,
        "labels": ["a", "b"],
        "series": [{"name": "S", "values": values}],
    }
    markdown = _chart_block(json.dumps(chart))

    document = _word(markdown_to_word(markdown, {}))
    assert _pictures(document) == 0
    assert "shown as a table" in word_text(markdown_to_word(markdown, {}))
    assert "shown as a table" in pdf_text(markdown_to_pdf(markdown, {}))


def test_letters_outside_latin_1_keep_their_shapes_in_the_pdf() -> None:
    """Cyrillic, Greek and CJK text is set in a font that has its letters.

    Arabic and Hebrew get glyphs too, though still in logical order (no bidi).
    """
    lines = ["Привет мир", "Καλημέρα", "中文报告", "日本語のレポート", "한국어 보고서"]
    markdown = (
        "# Отчёт\n\n"
        + "\n\n".join([*lines, "مرحبا", "שלום"])
        + "\n\n| Ключ | 值 |\n|---|---|\n| **Да** | *是* |\n\n```\nШлюз\n```\n"
    )

    text = pdf_text(markdown_to_pdf(markdown, {}))

    assert "■" not in text
    for line in ["Отчёт", *lines, "Ключ", "值", "Да", "是", "Шлюз"]:
        assert line in text


def test_a_tabulated_chart_keeps_each_number_as_written() -> None:
    """The fallback table does not round 1234567 to 1.23457e+06."""
    chart = {
        "type": "radar",
        "labels": ["2025"],
        "series": [{"name": "Cost", "values": [1234567]}],
    }

    assert "2025 | 1234567" in word_text(markdown_to_word(_chart_block(chart), {}))


@pytest.mark.parametrize(
    "markdown",
    [
        "\n".join(">" * level + f" deepest{level}" for level in range(1, 31)),
        "\n".join("  " * level + f"- deepest{level + 1}" for level in range(30)),
        "\n".join("> " + "  " * level + f"- deepest{level + 1}" for level in range(30)),
    ],
    ids=["quotes", "lists", "lists-in-a-quote"],
)
def test_quotes_and_lists_nested_thirty_deep_keep_their_text(markdown: str) -> None:
    """Nothing deep is dropped, and PDF indents stop before the text has no room."""
    assert "deepest30" in word_text(markdown_to_word(markdown, {}))
    assert "deepest30" in pdf_text(markdown_to_pdf(markdown, {}))


def test_a_chart_whose_labels_the_chart_font_lacks_becomes_a_table() -> None:
    """matplotlib's font has no Chinese, so the data is tabulated, not drawn as boxes."""
    chart = {
        "type": "bar",
        "title": "销售",
        "labels": ["一月", "二月"],
        "series": [{"name": "收入", "values": [1, 2]}],
    }

    document = _word(markdown_to_word(_chart_block(chart), {}))

    assert _pictures(document) == 0
    assert "一月 | 1" in word_text(markdown_to_word(_chart_block(chart), {}))


def test_a_chart_in_cyrillic_is_still_drawn() -> None:
    """Letters matplotlib's font has draw as usual."""
    chart = {**CHART, "title": "Расходы", "labels": ["Январь", "Февраль"]}

    assert _pictures(_word(markdown_to_word(_chart_block(chart), {}))) == 1
