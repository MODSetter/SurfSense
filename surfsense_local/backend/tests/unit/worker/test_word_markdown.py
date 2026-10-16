"""Word files read as markdown without Docling, on files python-docx generates."""

import logging
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.document import Document as WordDocument
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from worker.ingestion.word_markdown import markdown_without_docling
from worker.ingestion.word_markdown.document import markdown_of

pytestmark = pytest.mark.unit

W = nsdecls("w")
WR = nsdecls("w", "r")


def saved(document: WordDocument, folder: Path, name: str = "memo.docx") -> Path:
    """The document written to a file, as an upload leaves it on disk."""
    path = folder / name
    document.save(str(path))
    return path


def level(ilvl: int, num_fmt: str, text: str, start: int = 1) -> str:
    """One w:lvl of an abstract list."""
    return (
        f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="{start}"/>'
        f'<w:numFmt w:val="{num_fmt}"/><w:lvlText w:val="{text}"/></w:lvl>'
    )


def with_numbering(document: WordDocument, abstracts: str, nums: str) -> None:
    """Replace the template's numbering.xml with these definitions."""
    root = document.part.numbering_part.element
    for child in list(root):
        root.remove(child)
    for xml in (abstracts, nums):
        for element in parse_xml(f"<w:numbering {W}>{xml}</w:numbering>"):
            root.append(element)


def numbered(document: WordDocument, text: str, num_id: int, ilvl: int = 0) -> None:
    """A paragraph numbered directly, as Word writes one after Restart Numbering."""
    paragraph = document.add_paragraph(text)
    paragraph._p.get_or_add_pPr().append(
        parse_xml(
            f'<w:numPr {W}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
        )
    )


CLAUSES = (
    '<w:abstractNum w:abstractNumId="0">'
    + level(0, "decimal", "%1.")
    + level(1, "decimal", "%1.%2.")
    + level(2, "lowerLetter", "%3)")
    + level(3, "lowerRoman", "(%4)")
    + "</w:abstractNum>"
)


def test_headings_take_their_level_from_the_style(tmp_path: Path) -> None:
    """Markdown levels as Docling gives them: a title is #, Heading 1 is ##."""
    document = Document()
    document.add_paragraph("Services agreement", style="Title")
    document.add_heading("Scope", 1)
    document.add_heading("Hosting", 2)
    outlined = document.styles.add_style("Clause Head", 1)
    outlined.element.get_or_add_pPr().append(
        parse_xml(f'<w:outlineLvl {W} w:val="2"/>')
    )
    document.add_paragraph("Uptime", style=outlined)
    document.add_paragraph("Body text.")

    assert markdown_of(saved(document, tmp_path)) == (
        "# Services agreement\n\n## Scope\n\n### Hosting\n\n#### Uptime\n\nBody text."
    )


def test_nested_lists_show_the_numbers_word_shows(tmp_path: Path) -> None:
    """Rendered from numbering.xml as Word does: 1.1. and a), not 1.1.1."""
    document = Document()
    with_numbering(
        document, CLAUSES, '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
    )
    numbered(document, "Scope", 1, 0)
    numbered(document, "Services", 1, 1)
    numbered(document, "Hosting", 1, 2)
    numbered(document, "Backups", 1, 2)
    numbered(document, "Nightly", 1, 3)
    numbered(document, "Fees", 1, 0)
    numbered(document, "Invoices", 1, 1)
    numbered(document, "Late fees", 1, 2)

    assert markdown_of(saved(document, tmp_path)) == (
        "1. Scope\n"
        "    - 1.1. Services\n"
        "        - a) Hosting\n"
        "        - b) Backups\n"
        "            - (i) Nightly\n"
        "2. Fees\n"
        "    - 2.1. Invoices\n"
        "        - a) Late fees"
    )


def test_lists_on_one_definition_share_their_count_until_one_restarts(
    tmp_path: Path,
) -> None:
    """Word continues a second num on the same abstract list; a startOverride
    restarts it, and start values count from where they say."""
    document = Document()
    with_numbering(
        document,
        CLAUSES
        + '<w:abstractNum w:abstractNumId="1">'
        + level(0, "upperRoman", "Article %1", start=4)
        + "</w:abstractNum>",
        '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
        '<w:num w:numId="2"><w:abstractNumId w:val="0"/></w:num>'
        '<w:num w:numId="3"><w:abstractNumId w:val="0"/>'
        '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'
        '<w:num w:numId="4"><w:abstractNumId w:val="1"/></w:num>',
    )
    numbered(document, "One", 1)
    numbered(document, "Two", 1)
    document.add_paragraph("Between the lists.")
    numbered(document, "Three", 2)
    numbered(document, "Again one", 3)
    numbered(document, "Again two", 3)
    numbered(document, "Fourth article", 4)

    assert markdown_of(saved(document, tmp_path)).split("\n") == [
        "1. One",
        "2. Two",
        "",
        "Between the lists.",
        "",
        "3. Three",
        "1. Again one",
        "2. Again two",
        "- Article IV Fourth article",
    ]


def test_a_heading_numbered_by_its_style_keeps_its_number(tmp_path: Path) -> None:
    """Heading styles linked to a multilevel list: the number is part of the heading."""
    document = Document()
    with_numbering(
        document, CLAUSES, '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
    )
    for style, ilvl in (("Heading 1", 0), ("Heading 2", 1)):
        document.styles[style].element.get_or_add_pPr().append(
            parse_xml(
                f'<w:numPr {W}><w:ilvl w:val="{ilvl}"/><w:numId w:val="1"/></w:numPr>'
            )
        )
    document.add_heading("Definitions", 1)
    document.add_heading("Interpretation", 2)
    document.add_heading("Services", 1)

    assert markdown_of(saved(document, tmp_path)) == (
        "## 1. Definitions\n\n### 1.1. Interpretation\n\n## 2. Services"
    )


def test_bullets_are_dashes(tmp_path: Path) -> None:
    """Whatever glyph the bullet level draws."""
    document = Document()
    document.add_paragraph("Hosting", style="List Bullet")
    document.add_paragraph("Support", style="List Bullet")
    document.add_paragraph("After the list.")

    assert markdown_of(saved(document, tmp_path)) == (
        "- Hosting\n- Support\n\nAfter the list."
    )


def test_a_table_is_one_compact_line_per_row_with_merged_cells_once(
    tmp_path: Path,
) -> None:
    """Docling pads columns and repeats a merged cell in every column it spans."""
    document = Document()
    table = document.add_table(rows=3, cols=3)
    for r, row in enumerate(
        [["Item", "", "Cost"], ["Hosting", "EU", "40"], ["", "US", "35"]]
    ):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 2).text = "40 | 45"

    assert markdown_of(saved(document, tmp_path)) == (
        "| Item |  | Cost |\n"
        "|---|---|---|\n"
        "| Hosting | EU | 40 \\| 45 |\n"
        "|  | US | 35 |"
    )


def test_a_table_inside_a_cell_flattens_into_that_cell(tmp_path: Path) -> None:
    """Markdown cannot nest tables; Docling flattens them the same way."""
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Prices"
    table.cell(1, 0).text = "EU"
    inner = table.cell(1, 1).add_table(rows=2, cols=2)
    for r, row in enumerate([["Tier", "Fee"], ["Basic", "10"]]):
        for c, text in enumerate(row):
            inner.cell(r, c).text = text

    assert markdown_of(saved(document, tmp_path)) == (
        "| Region | Prices |\n|---|---|\n| EU | Tier Fee Basic 10 |"
    )


def test_a_one_cell_table_reads_as_body_text(tmp_path: Path) -> None:
    """A framed box, as Docling reads it."""
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Note: prices exclude VAT."

    assert markdown_of(saved(document, tmp_path)) == "Note: prices exclude VAT."


def test_a_link_keeps_its_words_and_its_address(tmp_path: Path) -> None:
    """An internal anchor has no address to keep."""
    document = Document()
    paragraph = document.add_paragraph("See the ")
    r_id = document.part.relate_to(
        "https://example.com/terms", RT.HYPERLINK, is_external=True
    )
    paragraph._p.append(
        parse_xml(
            f'<w:hyperlink {WR} r:id="{r_id}"><w:r><w:t>full terms</w:t></w:r></w:hyperlink>'
        )
    )
    paragraph._p.append(
        parse_xml(
            f'<w:hyperlink {WR} w:anchor="fees"><w:r><w:t xml:space="preserve"> and fees</w:t></w:r></w:hyperlink>'
        )
    )
    paragraph.add_run(".")

    assert markdown_of(saved(document, tmp_path)) == (
        "See the [full terms](https://example.com/terms) and fees."
    )


def test_tabs_and_breaks_keep_words_apart(tmp_path: Path) -> None:
    """A tab between a clause number and its title must not glue them: 1Definitions."""
    document = Document()
    paragraph = document.add_paragraph()
    paragraph._p.append(
        parse_xml(
            f"<w:r {W}><w:t>1</w:t><w:tab/><w:t>Definitions</w:t><w:br/>"
            '<w:t>Page</w:t><w:br w:type="page"/><w:t>two</w:t></w:r>'
        )
    )

    assert markdown_of(saved(document, tmp_path)) == "1 Definitions\nPage two"


def test_tracked_changes_read_as_if_accepted(tmp_path: Path) -> None:
    """Inserted words are read, deleted ones are not."""
    document = Document()
    paragraph = document.add_paragraph("Fees are ")
    paragraph._p.append(
        parse_xml(
            f'<w:del {W} w:id="1" w:author="A"><w:r><w:delText>forty</w:delText></w:r></w:del>'
        )
    )
    paragraph._p.append(
        parse_xml(
            f'<w:ins {W} w:id="2" w:author="A"><w:r><w:t>fifty</w:t></w:r></w:ins>'
        )
    )
    paragraph.add_run(" euros.")

    assert markdown_of(saved(document, tmp_path)) == "Fees are fifty euros."


def test_footnotes_and_endnotes_follow_the_body(tmp_path: Path) -> None:
    """Docling drops them; they are cited where the text cites them."""
    document = Document()
    for kind, content_type, relationship in (
        ("footnote", CT.WML_FOOTNOTES, RT.FOOTNOTES),
        ("endnote", CT.WML_ENDNOTES, RT.ENDNOTES),
    ):
        notes = (
            f'<w:{kind}s {W}><w:{kind} w:type="separator" w:id="-1"><w:p><w:r>'
            f"<w:separator/></w:r></w:p></w:{kind}>"
            f'<w:{kind} w:id="5"><w:p><w:r><w:{kind}Ref/></w:r>'
            f"<w:r><w:t>The {kind} text.</w:t></w:r></w:p></w:{kind}></w:{kind}s>"
        )
        part = Part(
            PackURI(f"/word/{kind}s.xml"),
            content_type,
            notes.encode(),
            document.part.package,
        )
        document.part.relate_to(part, relationship)
    paragraph = document.add_paragraph("Fees apply")
    paragraph._p.append(parse_xml(f'<w:r {W}><w:footnoteReference w:id="5"/></w:r>'))
    paragraph.add_run(" unless waived")
    paragraph._p.append(parse_xml(f'<w:r {W}><w:endnoteReference w:id="5"/></w:r>'))
    paragraph.add_run(".")

    assert markdown_of(saved(document, tmp_path)) == (
        "Fees apply[^1] unless waived[^i].\n\n"
        "[^1]: The footnote text.\n\n"
        "[^i]: The endnote text."
    )


def notes_part(document: WordDocument, kind: str, notes: str) -> Part:
    """A footnotes or endnotes part holding these w:footnote or w:endnote elements."""
    content_type, relationship = {
        "footnote": (CT.WML_FOOTNOTES, RT.FOOTNOTES),
        "endnote": (CT.WML_ENDNOTES, RT.ENDNOTES),
    }[kind]
    part = Part(
        PackURI(f"/word/{kind}s.xml"),
        content_type,
        f"<w:{kind}s {WR}>{notes}</w:{kind}s>".encode(),
        document.part.package,
    )
    document.part.relate_to(part, relationship)
    return part


def cite(paragraph: Any, kind: str, note_id: int) -> None:
    """A footnote or endnote reference at the paragraph's end."""
    paragraph._p.append(
        parse_xml(f'<w:r {W}><w:{kind}Reference w:id="{note_id}"/></w:r>')
    )


def test_notes_are_labelled_per_kind_in_the_order_first_cited(
    tmp_path: Path,
) -> None:
    """A note cited again keeps its label; footnotes and endnotes count apart."""
    document = Document()
    for kind in ("footnote", "endnote"):
        notes_part(
            document,
            kind,
            "".join(
                f'<w:{kind} w:id="{n}"><w:p><w:r><w:t>{kind} {n}</w:t></w:r></w:p>'
                f"</w:{kind}>"
                for n in (7, 8)
            ),
        )
    paragraph = document.add_paragraph("Text")
    for kind, note_id in (
        ("footnote", 8), ("endnote", 7), ("footnote", 7), ("endnote", 8),
        ("footnote", 8),
    ):  # fmt: skip
        cite(paragraph, kind, note_id)

    assert markdown_of(saved(document, tmp_path)) == (
        "Text[^1][^i][^2][^ii][^1]\n\n"
        "[^1]: footnote 8\n\n[^i]: endnote 7\n\n[^2]: footnote 7\n\n[^ii]: endnote 8"
    )


def test_a_link_in_a_footnote_resolves_against_the_notes_part(tmp_path: Path) -> None:
    """Its r:id names a relationship of footnotes.xml, not of the body."""
    document = Document()
    body_ids = set(document.part.rels)
    part = Part(
        PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, b"", document.part.package
    )
    r_id = part.relate_to("https://example.com/source", RT.HYPERLINK, is_external=True)
    assert r_id in body_ids  # the same id names something else in the body
    part._blob = (
        f'<w:footnotes {WR}><w:footnote w:id="1"><w:p><w:r><w:t xml:space="preserve">'
        f'See </w:t></w:r><w:hyperlink r:id="{r_id}"><w:r><w:t>the source</w:t></w:r>'
        "</w:hyperlink></w:p></w:footnote></w:footnotes>"
    ).encode()
    document.part.relate_to(part, RT.FOOTNOTES)
    cite(document.add_paragraph("Claim"), "footnote", 1)

    assert markdown_of(saved(document, tmp_path)) == (
        "Claim[^1]\n\n[^1]: See [the source](https://example.com/source)"
    )


@pytest.mark.parametrize(
    ("address", "written"),
    [
        (
            "https://en.wikipedia.org/wiki/Foo_(bar) baz",
            "https://en.wikipedia.org/wiki/Foo_(bar)%20baz",
        ),
        ("https://example.com/a)b", "https://example.com/a%29b"),
    ],
)
def test_a_link_address_stays_one_markdown_destination(
    tmp_path: Path, address: str, written: str
) -> None:
    """A space or an unmatched parenthesis would end it: spaces as %20, as
    Docling writes them, and parentheses escaped only when unmatched."""
    document = Document()
    paragraph = document.add_paragraph("See ")
    r_id = document.part.relate_to(address, RT.HYPERLINK, is_external=True)
    paragraph._p.append(
        parse_xml(
            f'<w:hyperlink {WR} r:id="{r_id}"><w:r><w:t>the page</w:t></w:r></w:hyperlink>'
        )
    )

    assert markdown_of(saved(document, tmp_path)) == f"See [the page]({written})"


def test_headers_footers_and_comments_are_left_out(tmp_path: Path) -> None:
    """As Docling's markdown leaves them: repeated page furniture and review talk."""
    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Confidential draft"
    section.footer.paragraphs[0].text = "Page 1"
    paragraph = document.add_paragraph("Fees are fifty euros.")
    document.add_comment(paragraph.runs, text="Check with finance", author="A")

    assert markdown_of(saved(document, tmp_path)) == "Fees are fifty euros."


def test_content_controls_show_their_content(tmp_path: Path) -> None:
    """A control around paragraphs and one around runs alike: a filled-in form."""
    document = Document()
    document.add_paragraph("Before.")._p.addnext(
        parse_xml(
            f"<w:sdt {W}><w:sdtPr/><w:sdtContent><w:p><w:r><w:t>Block control</w:t>"
            "</w:r></w:p></w:sdtContent></w:sdt>"
        )
    )
    paragraph = document.add_paragraph("Client: ")
    paragraph._p.append(
        parse_xml(
            f"<w:sdt {W}><w:sdtPr/><w:sdtContent><w:r><w:t>Halvorsen Freight</w:t></w:r>"
            "</w:sdtContent></w:sdt>"
        )
    )

    assert markdown_of(saved(document, tmp_path)) == (
        "Before.\n\nBlock control\n\nClient: Halvorsen Freight"
    )


def test_moved_text_reads_where_it_was_moved_to(tmp_path: Path) -> None:
    """A tracked move, accepted: the words once, at their new place."""
    document = Document()
    first = document.add_paragraph("Payment ")
    first._p.append(
        parse_xml(
            f'<w:moveFrom {W} w:id="1" w:author="A"><w:r><w:t>within 30 days</w:t>'
            "</w:r></w:moveFrom>"
        )
    )
    second = document.add_paragraph("Invoices are due ")
    second._p.append(
        parse_xml(
            f'<w:moveTo {W} w:id="2" w:author="A"><w:r><w:t>within 30 days</w:t>'
            "</w:r></w:moveTo>"
        )
    )

    assert markdown_of(saved(document, tmp_path)) == (
        "Payment\n\nInvoices are due within 30 days"
    )


def test_a_row_deleted_with_tracked_changes_is_gone(tmp_path: Path) -> None:
    """Not an empty row, which is how Docling leaves it."""
    document = Document()
    table = document.add_table(rows=3, cols=2)
    for r, row in enumerate([["Line", "EUR"], ["Ferry", "120"], ["Pier", "80"]]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    deleted = table.rows[1]._tr
    deleted.get_or_add_trPr().append(
        parse_xml(f'<w:del {W} w:id="3" w:author="A" w:date="2026-01-01T00:00:00Z"/>')
    )

    assert markdown_of(saved(document, tmp_path)) == (
        "| Line | EUR |\n|---|---|\n| Pier | 80 |"
    )


def test_alternate_content_is_read_once_from_its_choice(tmp_path: Path) -> None:
    """Around runs or around paragraphs, Word shows the first choice it reads;
    the fallback is for older readers. Docling drops both."""
    mc = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
    document = Document()
    paragraph = document.add_paragraph("Shown: ")
    paragraph._p.append(
        parse_xml(
            f'<mc:AlternateContent {W} {mc}><mc:Choice Requires="w14"><w:r><w:t>choice'
            "</w:t></w:r></mc:Choice><mc:Fallback><w:r><w:t>fallback</w:t></w:r>"
            "</mc:Fallback></mc:AlternateContent>"
        )
    )
    paragraph._p.addnext(
        parse_xml(
            f'<mc:AlternateContent {W} {mc}><mc:Choice Requires="w14"><w:p><w:r>'
            "<w:t>Body choice</w:t></w:r></w:p></mc:Choice><mc:Fallback><w:p><w:r>"
            "<w:t>Body fallback</w:t></w:r></w:p></mc:Fallback></mc:AlternateContent>"
        )
    )

    assert markdown_of(saved(document, tmp_path)) == "Shown: choice\n\nBody choice"


def test_ruby_reads_as_its_base_text(tmp_path: Path) -> None:
    """The reading guide above the characters is an aid; the base is the text."""
    document = Document()
    document.add_paragraph("Tokyo: ")._p.append(
        parse_xml(
            f"<w:r {W}><w:ruby><w:rubyPr/><w:rt><w:r><w:t>とうきょう"
            "</w:t></w:r></w:rt><w:rubyBase><w:r><w:t>東京</w:t></w:r>"
            "</w:rubyBase></w:ruby></w:r>"
        )
    )

    assert markdown_of(saved(document, tmp_path)) == "Tokyo: 東京"


def test_a_text_box_is_read_once(tmp_path: Path) -> None:
    """Word stores a VML copy of each text box as a fallback."""
    document = Document()
    box = "<w:txbxContent><w:p><w:r><w:t>Boxed quote</w:t></w:r></w:p></w:txbxContent>"
    paragraph = document.add_paragraph("After the box.")
    paragraph._p.insert(
        1,
        parse_xml(
            f'<w:r {W} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
            ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
            ' xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent>'
            f'<mc:Choice Requires="wps"><wps:txbx>{box}</wps:txbx></mc:Choice>'
            f"<mc:Fallback><v:textbox>{box}</v:textbox></mc:Fallback>"
            "</mc:AlternateContent></w:r>"
        ),
    )

    assert markdown_of(saved(document, tmp_path)) == "Boxed quote\n\nAfter the box."


def test_a_blank_document_is_left_to_docling(tmp_path: Path) -> None:
    """No text here may still be text Docling finds, and a blank file stays as before."""
    path = saved(Document(), tmp_path)

    assert markdown_of(path) == ""
    assert markdown_without_docling(path) is None


def test_a_malformed_file_is_left_to_docling(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    """Docling then refuses it with its own reason, as it does today."""
    path = tmp_path / "broken.docx"
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("word/document.xml", "<not word")

    with caplog.at_level(logging.WARNING):
        assert markdown_without_docling(path) is None

    assert "Docling reads broken.docx" in caplog.text


def test_a_file_with_a_native_chart_is_left_to_docling(tmp_path: Path) -> None:
    """Docling writes a chart's data as a table; this reader would drop it."""
    document = Document()
    document.add_paragraph("Costs by region.")
    chart = Part(
        PackURI("/word/charts/chart1.xml"),
        CT.DML_CHART,
        b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>',
        document.part.package,
    )
    document.part.relate_to(chart, RT.CHART)

    assert markdown_without_docling(saved(document, tmp_path)) is None


@pytest.mark.parametrize(
    ("setting", "reads_here"),
    [
        (None, True),
        ("fast", True),
        (" Fast ", True),
        ("docling", False),
        ("DOCLING", False),
    ],
)
def test_the_switch_chooses_the_reader(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    setting: str | None,
    reads_here: bool,
) -> None:
    """SURFSENSE_LOCAL_DOCX_READER: this reader unless it says docling, in any case."""
    document = Document()
    document.add_paragraph("Fees are fifty euros.")
    path = saved(document, tmp_path)
    if setting is None:
        monkeypatch.delenv("SURFSENSE_LOCAL_DOCX_READER", raising=False)
    else:
        monkeypatch.setenv("SURFSENSE_LOCAL_DOCX_READER", setting)

    expected = "Fees are fifty euros." if reads_here else None
    assert markdown_without_docling(path) == expected


def test_an_unknown_switch_value_turns_the_reader_off(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """Whoever set it most likely meant to turn it off."""
    document = Document()
    document.add_paragraph("Fees are fifty euros.")
    monkeypatch.setenv("SURFSENSE_LOCAL_DOCX_READER", "python-docx")

    with caplog.at_level(logging.WARNING):
        assert markdown_without_docling(saved(document, tmp_path)) is None

    assert "SURFSENSE_LOCAL_DOCX_READER" in caplog.text
