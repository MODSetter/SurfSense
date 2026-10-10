"""Studio's PowerPoint and Excel: the model's script runs in the script runner."""

from io import BytesIO

import pptx as python_pptx
import pytest
import xlsxwriter

from modules.llm.profile import Tier
from modules.llm.resolution import ResolvedGeneration
from shared import cancellation
from worker.document_script.run import ScriptResult
from worker.jobs import JobCancelledError
from worker.studio.office import pipeline as office
from worker.studio.office import prompt
from worker.studio.office.docx import docx
from worker.studio.office.pdf import pdf
from worker.studio.office.pptx import pptx
from worker.studio.office.spec import Office
from worker.studio.office.xlsx import xlsx
from worker.studio.script_document import pipeline as script_document
from worker.studio.shared import generate
from worker.studio.shared.artifact import Source

pytestmark = pytest.mark.unit

MODEL = ResolvedGeneration(
    type("Selection", (), {"name": "qwen3:8b", "tier": Tier.CAPABLE})(), None
)


def _model(raw: str, monkeypatch: pytest.MonkeyPatch) -> None:
    """Stand in for the selected model, returning a fixed script."""
    monkeypatch.setattr(generate, "run_model", lambda *a, **k: raw)


def _runner(
    monkeypatch: pytest.MonkeyPatch, *results: ScriptResult
) -> list[dict[str, object]]:
    """Stand in for the script runner, answering each run in turn."""
    calls: list[dict[str, object]] = []

    def run(script: str, **kwargs: object) -> ScriptResult:
        calls.append({"script": script, **kwargs})
        return results[min(len(calls), len(results)) - 1]

    monkeypatch.setattr(script_document, "run_document_script", run)
    return calls


def _ran(output: bytes) -> ScriptResult:
    return ScriptResult(
        ok=True, output=output, error=None, traceback_tail=None, seconds=0.4
    )


def _failed(error: str, tail: str | None = None) -> ScriptResult:
    return ScriptResult(
        ok=False, output=None, error=error, traceback_tail=tail, seconds=0.4
    )


def _deck(title: str) -> bytes:
    deck = python_pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = "Arrival in 2004."
    buffer = BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _workbook(cell: str) -> bytes:
    buffer = BytesIO()
    book = xlsxwriter.Workbook(buffer, {"in_memory": True})
    book.add_worksheet("Missions").write(0, 0, cell)
    book.close()
    return buffer.getvalue()


def test_a_deck_script_runs_in_the_runner_and_keeps_its_file_title_and_text(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The title is the script's `# title:` line; the body is the deck's own text."""
    data = _deck("Cassini")
    _model("```python\n# title: Cassini\nsave()\n```", monkeypatch)
    calls = _runner(monkeypatch, _ran(data))

    built = office.render(pptx, MODEL, [Source(1, "Saturn", "rings")], None)

    assert calls[0]["script"] == "# title: Cassini\nsave()"
    assert calls[0]["output_name"] == "document.pptx"
    assert built.primary == data
    assert built.primary_mime == pptx.mime
    assert built.primary_filename == "Cassini.pptx"
    assert built.title == "Cassini"
    assert "## Slide 1: Cassini" in built.markdown
    assert "Arrival in 2004." in built.markdown


def test_a_script_without_a_title_line_takes_the_prompt_as_its_title(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The user's button fixes the type: an xlsx job stores a .xlsx."""
    _model("save()", monkeypatch)
    _runner(monkeypatch, _ran(_workbook("Huygens")))

    built = office.render(xlsx, MODEL, [], "mission budget")

    assert built.title == "mission budget"
    assert built.primary_filename == "mission budget.xlsx"
    assert built.primary_mime.endswith("spreadsheetml.sheet")
    assert "Huygens" in built.markdown


def test_a_file_that_is_not_its_format_fails_the_job(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Bytes at OUTPUT_PATH that do not open as a workbook are no workbook."""
    _model("save()", monkeypatch)
    _runner(monkeypatch, _ran(b"not a zip"))

    with pytest.raises(RuntimeError, match=r"not a valid \.xlsx"):
        office.render(xlsx, MODEL, [], None)


def test_a_failed_script_goes_back_with_its_traceback_then_the_fix_is_kept(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The repair replays the failed script with the error and the traceback's end."""
    replies = iter(["# title: Cassini\nbroken()", "# title: Cassini\nfixed()"])
    seen: list[generate.Repair | None] = []

    def fake_model(
        _session: object,
        _system: str,
        _sources: object,
        *,
        repair: generate.Repair | None = None,
    ) -> str:
        seen.append(repair)
        return next(replies)

    monkeypatch.setattr(generate, "run_model", fake_model)
    data = _deck("Cassini")
    tail = 'File "script.py", line 2, in <module>\nNameError: broken'
    _runner(monkeypatch, _failed("NameError: broken", tail), _ran(data))

    built = office.render(pptx, MODEL, [], None)

    assert built.primary == data
    assert len(seen) == 2
    assert seen[0] is None
    assert seen[1] is not None
    assert "broken()" in seen[1].reply
    assert "NameError: broken" in seen[1].instruction
    assert 'script.py", line 2' in seen[1].instruction


def test_render_stops_after_three_failed_scripts(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Three broken scripts fail the job; a fourth is never requested."""
    calls = 0

    def fake_model(*_a: object, **_k: object) -> str:
        nonlocal calls
        calls += 1
        return "broken()"

    monkeypatch.setattr(generate, "run_model", fake_model)
    _runner(monkeypatch, _failed("ValueError: still broken"))

    with pytest.raises(RuntimeError, match="still broken"):
        office.render(xlsx, MODEL, [], None)
    assert calls == 3


def test_a_cancel_during_a_failed_attempt_asks_for_no_retry(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Asking the model to fix a script would spend a generation nobody wants."""
    calls = 0
    cancelled = False

    def fake_model(*_a: object, **_k: object) -> str:
        nonlocal calls, cancelled
        calls += 1
        cancelled = True
        return "broken()"

    def check() -> None:
        if cancelled:
            raise JobCancelledError

    monkeypatch.setattr(generate, "run_model", fake_model)
    _runner(monkeypatch, _failed("ValueError: broken"))

    with cancellation.watching(check), pytest.raises(JobCancelledError):
        office.render(xlsx, MODEL, [], None)
    assert calls == 1


@pytest.mark.parametrize("tier", list(Tier))
@pytest.mark.parametrize("spec", [pptx, xlsx], ids=["pptx", "xlsx"])
def test_the_prompt_asks_for_a_file_at_output_path_and_a_title_line(
    tier: Tier, spec: Office
) -> None:
    """The runner's contract, not the old namespace one."""
    text = prompt.build(tier, spec, None)

    assert "OUTPUT_PATH" in text
    assert "# title:" in text
    assert "output_bytes" not in text


def test_each_format_carries_its_skill_and_library() -> None:
    """Every office folder loaded a non-empty SKILL.md and names its library."""
    for spec in (docx, pptx, xlsx, pdf):
        assert spec.skill.strip()
        assert spec.library in {"python-docx", "python-pptx", "xlsxwriter", "reportlab"}
