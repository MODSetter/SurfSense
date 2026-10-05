"""A document script's run becomes a Built: its file, its text, its download name."""

import threading
from io import BytesIO
from pathlib import Path

import docx
import pptx
import pypdfium2
import pytest
import xlsxwriter
from pptx.util import Inches
from reportlab.pdfgen import canvas

from modules.artifacts.script_documents.spec import DocumentScript
from shared import pdfium
from worker.document_script.run import ScriptResult
from worker.studio.office.docx import docx as word
from worker.studio.office.pdf import pdf
from worker.studio.office.pptx import pptx as deck_format
from worker.studio.office.xlsx import xlsx as workbook_format
from worker.studio.script_document import pipeline
from worker.studio.script_document.pipeline import ScriptRunFailedError

pytestmark = pytest.mark.unit


def _word_file() -> bytes:
    document = docx.Document()
    document.add_heading("Client proposal", level=1)
    document.add_paragraph("We propose a two-phase rollout.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Phase"
    table.cell(0, 1).text = "Cost"
    table.cell(1, 0).text = "Pilot"
    table.cell(1, 1).text = "12,000"
    document.add_heading("Timeline", level=2)
    document.add_paragraph("Six weeks.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_file(*pages: str) -> bytes:
    buffer = BytesIO()
    drawing = canvas.Canvas(buffer)
    for text in pages:
        drawing.drawString(72, 720, text)
        drawing.showPage()
    drawing.save()
    return buffer.getvalue()


def _runner(
    monkeypatch: pytest.MonkeyPatch, result: ScriptResult
) -> list[dict[str, object]]:
    """Stand in for the script runner, recording what each run was given."""
    calls: list[dict[str, object]] = []

    def run(script: str, **kwargs: object) -> ScriptResult:
        calls.append({"script": script, **kwargs})
        return result

    monkeypatch.setattr(pipeline, "run_document_script", run)
    return calls


def _ran(output: bytes) -> ScriptResult:
    return ScriptResult(
        ok=True, output=output, error=None, traceback_tail=None, seconds=0.4
    )


def test_a_word_file_keeps_its_bytes_and_reads_back_in_body_order(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Headings stay headings and table cells are searchable, in document order."""
    data = _word_file()
    calls = _runner(monkeypatch, _ran(data))
    script = DocumentScript(text="build()", format="docx", images=())

    built = pipeline.render("Client proposal", script, {})

    assert calls[0]["script"] == "build()"
    assert calls[0]["output_name"] == "document.docx"
    assert built.primary == data
    assert built.primary_mime == word.mime
    assert built.title == "Client proposal"
    assert built.markdown == (
        "# Client proposal\n\n"
        "We propose a two-phase rollout.\n\n"
        "Phase | Cost\nPilot | 12,000\n\n"
        "## Timeline\n\n"
        "Six weeks."
    )


def test_a_merged_cell_reads_once(monkeypatch: pytest.MonkeyPatch) -> None:
    """python-docx repeats a merged cell for every column it spans."""
    document = docx.Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Yearly costs"
    for column, value in enumerate(["2025", "2026", "2027"]):
        table.cell(1, column).text = value
    buffer = BytesIO()
    document.save(buffer)
    _runner(monkeypatch, _ran(buffer.getvalue()))
    script = DocumentScript(text="build()", format="docx", images=())

    built = pipeline.render("Costs", script, {})

    assert built.markdown == "Yearly costs\n2025 | 2026 | 2027"


def test_a_pdf_reads_back_page_by_page(monkeypatch: pytest.MonkeyPatch) -> None:
    """Each page's text layer is its own paragraph, in page order."""
    data = _pdf_file("Executive summary", "Pricing")
    calls = _runner(monkeypatch, _ran(data))
    script = DocumentScript(text="build()", format="pdf", images=())

    built = pipeline.render("Proposal", script, {})

    assert calls[0]["output_name"] == "document.pdf"
    assert built.primary == data
    assert built.primary_mime == pdf.mime
    assert built.markdown == "Executive summary\n\nPricing"


def test_the_named_images_reach_the_runner(monkeypatch: pytest.MonkeyPatch) -> None:
    """The images resolved for the spec are what the runner copies beside it."""
    calls = _runner(monkeypatch, _ran(_pdf_file("Logo")))
    images = {"7-1": Path("figures/1.png")}
    script = DocumentScript(text="build()", format="pdf", images=("7-1",))

    pipeline.render("Proposal", script, images)

    assert calls[0]["images"] == images


def test_the_download_name_is_the_title_without_characters_windows_forbids(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The title the user sees is kept; only what a file system refuses is dropped."""
    _runner(monkeypatch, _ran(_word_file()))
    script = DocumentScript(text="build()", format="docx", images=())

    built = pipeline.render('Q3: plan / "budget"?', script, {})

    assert built.primary_filename == "Q3 plan  budget.docx"


def test_a_file_with_no_text_is_still_searchable_by_its_title(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A page of only drawings has no text layer; the body must not be empty."""
    _runner(monkeypatch, _ran(_pdf_file("")))
    script = DocumentScript(text="build()", format="pdf", images=())

    built = pipeline.render("Org chart", script, {})

    assert built.markdown == "# Org chart"


@pytest.mark.parametrize(
    ("fmt", "output", "error"),
    [
        ("docx", b"not a zip", "the script wrote a file that is not a valid .docx"),
        ("docx", b"%PDF-1.4", "the script wrote a file that is not a valid .docx"),
        ("pdf", b"PK\x03\x04", "the script wrote a file that is not a valid .pdf"),
    ],
)
def test_a_file_that_does_not_open_as_its_format_fails_the_run(
    monkeypatch: pytest.MonkeyPatch, fmt: str, output: bytes, error: str
) -> None:
    """The format the request named is checked by opening the file as it."""
    _runner(monkeypatch, _ran(output))
    script = DocumentScript(text="build()", format=fmt, images=())

    with pytest.raises(ScriptRunFailedError) as failed:
        pipeline.render("Proposal", script, {})

    assert failed.value.reason(500) == error


def test_a_failed_run_carries_the_error_and_the_traceback_tail(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A failed run is reported with the runner's error line and traceback tail."""
    tail = (
        "Traceback (most recent call last):\n"
        '  File "script.py", line 3, in <module>\n'
        "ValueError: the pricing table is empty"
    )
    _runner(
        monkeypatch,
        ScriptResult(
            ok=False,
            output=None,
            error="ValueError: the pricing table is empty",
            traceback_tail=tail,
            seconds=0.2,
        ),
    )
    script = DocumentScript(text="build()", format="docx", images=())

    with pytest.raises(ScriptRunFailedError) as failed:
        pipeline.render("Proposal", script, {})

    assert failed.value.reason(500) == f"ValueError: the pricing table is empty\n{tail}"


def test_a_long_traceback_keeps_its_last_lines_under_the_limit() -> None:
    """The last lines name the failing line of the script; the first are runpy's."""
    frames = [f'  File "runpy.py", line {n}, in _run_code' for n in range(40)]
    tail = "\n".join([*frames, '  File "script.py", line 9', "KeyError: 'total'"])
    failure = ScriptRunFailedError("KeyError: 'total'", tail)

    reason = failure.reason(120)

    assert len(reason) <= 120
    assert reason.startswith("KeyError: 'total'\n")
    assert reason.endswith("  File \"script.py\", line 9\nKeyError: 'total'")
    assert "line 0," not in reason


def test_a_pdf_is_read_under_the_process_wide_pdfium_lock(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """pdfium is not thread-safe, and Studio runs several jobs at once."""
    held_while_opening: list[bool] = []
    opened = pypdfium2.PdfDocument

    def open_checking_the_lock(*args: object, **kwargs: object) -> object:
        held_while_opening.append(pdfium.lock.locked())
        return opened(*args, **kwargs)

    monkeypatch.setattr(pdfium, "lock", threading.Lock())
    monkeypatch.setattr(pypdfium2, "PdfDocument", open_checking_the_lock)
    _runner(monkeypatch, _ran(_pdf_file("Pricing")))
    script = DocumentScript(text="build()", format="pdf", images=())

    built = pipeline.render("Proposal", script, {})

    assert built.markdown == "Pricing"
    assert held_while_opening == [True]
    assert not pdfium.lock.locked()


def _deck_file() -> bytes:
    """A two-slide deck: a title slide, then a slide with bullets, a table and notes."""
    deck = pptx.Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = "Quarterly review"
    cover.placeholders[1].text = "Halvorsen Freight"
    body = deck.slides.add_slide(deck.slide_layouts[5])
    body.shapes.title.text = "Costs"
    box = body.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1))
    box.text_frame.text = "Costs rose in the north."
    box.text_frame.add_paragraph().text = "The south held steady."
    table = body.shapes.add_table(
        2, 2, Inches(1), Inches(3), Inches(4), Inches(1)
    ).table
    for (row, column), value in {
        (0, 0): "Region",
        (0, 1): "Cost",
        (1, 0): "North",
        (1, 1): "120",
    }.items():
        table.cell(row, column).text = value
    body.notes_slide.notes_text_frame.text = "Start with the north."
    buffer = BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_a_deck_reads_back_slide_by_slide_with_titles_tables_and_notes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Each slide's title marks it; its text, table rows and notes follow in order."""
    data = _deck_file()
    calls = _runner(monkeypatch, _ran(data))
    script = DocumentScript(text="build()", format="pptx", images=())

    built = pipeline.render("Quarterly review", script, {})

    assert calls[0]["output_name"] == "document.pptx"
    assert built.primary == data
    assert built.primary_mime == deck_format.mime
    assert built.primary_filename == "Quarterly review.pptx"
    assert built.markdown == (
        "## Slide 1: Quarterly review\n"
        "Halvorsen Freight\n\n"
        "## Slide 2: Costs\n"
        "Costs rose in the north.\n"
        "The south held steady.\n"
        "Region | Cost\n"
        "North | 120\n"
        "Notes: Start with the north."
    )


def _workbook_file(rows: int = 3) -> bytes:
    """A workbook with a header row, values, a formula total and a chart."""
    buffer = BytesIO()
    book = xlsxwriter.Workbook(buffer, {"in_memory": True})
    sheet = book.add_worksheet("Costs")
    sheet.write_row(0, 0, ["Year", "Cost"])
    for row in range(1, rows + 1):
        sheet.write_row(row, 0, [str(2023 + row), 100 + row])
    sheet.write_formula(rows + 1, 1, f"=SUM(B2:B{rows + 1})")
    sheet.write(rows + 1, 0, "Total")
    chart = book.add_chart({"type": "column"})
    chart.add_series({"values": f"=Costs!$B$2:$B${rows + 1}"})
    sheet.insert_chart("D2", chart)
    book.add_worksheet("Notes").write(0, 0, "Figures in thousand EUR")
    book.close()
    return buffer.getvalue()


def test_a_workbook_reads_back_as_a_summary_of_each_sheet_and_its_formulas(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The model checks a workbook from its sheets, first rows and formulas, having no pages to look at."""
    data = _workbook_file()
    calls = _runner(monkeypatch, _ran(data))
    script = DocumentScript(text="build()", format="xlsx", images=())

    built = pipeline.render("Costs", script, {})

    assert calls[0]["output_name"] == "document.xlsx"
    assert built.primary_mime == workbook_format.mime
    assert built.markdown == (
        "A workbook of 2 sheets and 1 chart.\n\n"
        'Sheet "Costs": A1:B5\n'
        "Year | Cost\n"
        "2024 | 101\n"
        "2025 | 102\n"
        "2026 | 103\n"
        "Total | =SUM(B2:B4)\n\n"
        'Sheet "Notes": A1\n'
        "Figures in thousand EUR\n\n"
        "Formulas:\n"
        "- Costs!B5: =SUM(B2:B4)"
    )


def test_a_long_workbook_summary_shows_its_first_rows_and_counts_the_rest(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A sheet of thousands of rows stays a summary the model can read in one result."""
    _runner(monkeypatch, _ran(_workbook_file(rows=500)))
    script = DocumentScript(text="build()", format="xlsx", images=())

    built = pipeline.render("Costs", script, {})

    assert "(492 more rows)" in built.markdown
    assert "2033 | 110" not in built.markdown
    assert len(built.markdown) < 6000


def test_an_array_formula_reads_as_the_formula_it_holds(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """xlsxwriter's array and dynamic-array formulas list with the rest, not as objects."""
    buffer = BytesIO()
    book = xlsxwriter.Workbook(buffer, {"in_memory": True})
    sheet = book.add_worksheet("Costs")
    sheet.write_row(0, 0, [1, 2, 3])
    sheet.write_array_formula("A2:A2", "{=SUM(A1:C1*2)}", None, 12)
    sheet.write_dynamic_array_formula("B2:B2", "=LEN(A1:C1)")
    book.close()
    _runner(monkeypatch, _ran(buffer.getvalue()))
    script = DocumentScript(text="build()", format="xlsx", images=())

    built = pipeline.render("Costs", script, {})

    assert "object" not in built.markdown
    assert "=SUM(A1:C1*2) | =LEN(A1:C1)" in built.markdown
    assert "- Costs!A2: =SUM(A1:C1*2)" in built.markdown
    assert "- Costs!B2: =LEN(A1:C1)" in built.markdown


def test_a_cell_written_at_the_far_end_of_a_sheet_does_not_stall_the_summary(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Every row up to the last one written is read padded to the last column;
    the summary stops reading long before that and says so."""
    buffer = BytesIO()
    book = xlsxwriter.Workbook(buffer, {"in_memory": True})
    sheet = book.add_worksheet("Costs")
    sheet.write_row(0, 0, ["Year", "Cost"])
    sheet.write(0, 16000, "far")
    sheet.write(1_048_575, 0, "bottom")
    book.close()
    _runner(monkeypatch, _ran(buffer.getvalue()))
    script = DocumentScript(text="build()", format="xlsx", images=())
    built: list[str] = []
    summarised = threading.Thread(
        target=lambda: built.append(pipeline.render("Costs", script, {}).markdown),
        daemon=True,
    )

    summarised.start()
    summarised.join(20)

    assert built, "the summary was still reading the sheet after 20 s"
    assert 'Sheet "Costs": A1:WQK1048576' in built[0]
    assert "Year | Cost" in built[0]
    assert "not read" in built[0]


@pytest.mark.parametrize("fmt", ["pptx", "xlsx"])
def test_a_deck_or_workbook_that_does_not_open_fails_the_run(
    monkeypatch: pytest.MonkeyPatch, fmt: str
) -> None:
    """A file that is not the format asked for is the script's fault, said as such."""
    _runner(monkeypatch, _ran(b"not a zip"))
    script = DocumentScript(text="build()", format=fmt, images=())

    with pytest.raises(ScriptRunFailedError) as failed:
        pipeline.render("Proposal", script, {})

    assert failed.value.reason(500) == (
        f"the script wrote a file that is not a valid .{fmt}"
    )


def test_a_template_reaches_the_runner(monkeypatch: pytest.MonkeyPatch) -> None:
    """The source file a template names is what the runner copies beside the script."""
    calls = _runner(monkeypatch, _ran(_deck_file()))
    template = Path("documents/12/Brand.pptx")
    script = DocumentScript(
        text="build()", format="pptx", images=(), template_source_id=12
    )

    pipeline.render("Review", script, {}, template=template)

    assert calls[0]["template"] == template
