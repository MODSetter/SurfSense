import threading
from typing import Any

import pytest
from docx import Document
from docx.document import Document as WordDocument
from docx.parts.document import DocumentPart

from worker.ingestion.docx_style_cache import cached_style_lookups

pytestmark = pytest.mark.unit

STYLES = ["Normal", "Heading 1", "List Bullet"]


def styled_document(paragraphs: int) -> WordDocument:
    """Paragraphs cycling through a few styles, as a long contract repeats them."""
    document = Document()
    for n in range(paragraphs):
        document.add_paragraph(f"Clause {n}", style=STYLES[n % len(STYLES)])
    return document


def style_names(document: WordDocument) -> list[str]:
    """Each paragraph's style, read the way Docling reads it."""
    return [paragraph.style.name for paragraph in document.paragraphs]


@pytest.fixture
def lookups(monkeypatch: pytest.MonkeyPatch) -> list[tuple[Any, ...]]:
    """Each lookup that reaches python-docx's own get_style."""
    seen: list[tuple[Any, ...]] = []
    real = DocumentPart.get_style

    def counting(part: DocumentPart, style_id: Any, style_type: Any) -> Any:
        seen.append((part, style_id, style_type))
        return real(part, style_id, style_type)

    monkeypatch.setattr(DocumentPart, "get_style", counting)
    return seen


def test_each_style_is_looked_up_once_per_document(lookups: list) -> None:
    """Docling reads a paragraph's style once per run; python-docx scans for each."""
    document = styled_document(30)
    expected = style_names(document)
    lookups.clear()

    with cached_style_lookups():
        first, second = style_names(document), style_names(document)

    assert first == second == expected
    assert len(lookups) == len(STYLES)


def test_each_document_keeps_its_own_styles(lookups: list) -> None:
    """Two files read in one block never answer with each other's styles."""
    one, other = styled_document(6), styled_document(6)
    lookups.clear()

    with cached_style_lookups():
        for document in (one, other):
            for paragraph in document.paragraphs:
                assert paragraph.style.element.getparent() is document.styles.element

    assert len(lookups) == 2 * len(STYLES)


def test_python_docx_is_restored_after_the_block(lookups: list) -> None:
    """Writers such as revise_document must see python-docx as it ships."""
    before = DocumentPart.get_style
    document = styled_document(9)

    with cached_style_lookups():
        assert DocumentPart.get_style is not before
    lookups.clear()
    style_names(document)

    assert DocumentPart.get_style is before
    assert len(lookups) == 9


def test_python_docx_is_restored_when_the_read_fails() -> None:
    """A Docling failure mid-read must not leave the cache installed."""
    before = DocumentPart.get_style

    with pytest.raises(RuntimeError), cached_style_lookups():
        raise RuntimeError("Docling refused the file")

    assert DocumentPart.get_style is before


def test_nothing_is_remembered_from_one_block_to_the_next(lookups: list) -> None:
    """The cache lives for one conversion, so a later read sees the file's styles anew."""
    document = styled_document(3)
    lookups.clear()

    for _ in range(2):
        with cached_style_lookups():
            style_names(document)

    assert len(lookups) == 2 * len(STYLES)


def test_another_thread_reads_python_docx_uncached(lookups: list) -> None:
    """A conversion holds the patch process-wide; other threads keep the original."""
    document = styled_document(12)
    lookups.clear()

    with cached_style_lookups():
        reader = threading.Thread(target=style_names, args=(document,))
        reader.start()
        reader.join()

    assert len(lookups) == 12


def test_overlapping_conversions_restore_python_docx_once_both_end() -> None:
    """Conversions on two threads may end in either order."""
    before = DocumentPart.get_style
    entered, release = threading.Event(), threading.Event()

    def other_conversion() -> None:
        with cached_style_lookups():
            entered.set()
            release.wait(timeout=10)

    other = threading.Thread(target=other_conversion)
    with cached_style_lookups():
        other.start()
        assert entered.wait(timeout=10)
    try:
        assert DocumentPart.get_style is not before
    finally:
        release.set()
        other.join()

    assert DocumentPart.get_style is before
