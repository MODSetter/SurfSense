import hashlib
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

from modules.artifacts.revised_copies.engines import powerpoint
from tests.unit.artifacts.revised_copies.office_fixtures import (
    PARSER,
    A,
    build_deck,
    check_deck,
    parts,
    slide_texts,
)

pytestmark = pytest.mark.unit


@pytest.fixture
def deck(tmp_path: Path) -> Path:
    """A real deck with groups, a table, notes, a picture, a link and a chart."""
    return build_deck(tmp_path / "Board.pptx", image=tmp_path / "logo.png")


def run(deck: Path, tmp_path: Path, *operations: dict):
    """Applies the operations to a copy next to the test's files."""
    out = tmp_path / "out.pptx"
    return powerpoint.apply(deck, list(operations), out), out


def runs(path: Path, slide: int, part: str = "slides/slide") -> list[tuple[str, dict]]:
    """Each run's text with its bold, italic and size."""
    root = etree.fromstring(parts(path)[f"ppt/{part}{slide}.xml"], PARSER)
    found = []
    for run_element in root.iter(f"{{{A}}}r"):
        properties = run_element.find(f"{{{A}}}rPr")
        attributes = (
            {}
            if properties is None
            else {k: v for k, v in properties.attrib.items() if k in ("b", "i", "sz")}
        )
        found.append((run_element.findtext(f"{{{A}}}t"), attributes))
    return found


def codes(report) -> set[str]:
    """Every code in the report, from outcomes and notices."""
    return {o.code for o in report.outcomes if o.code} | {
        n.code for n in report.must_tell_user
    }


def test_a_quote_inside_one_run_keeps_that_runs_formatting_and_nothing_else_changes(
    deck, tmp_path
):
    """Only the edited slide's part changes, and the replaced words keep their run's look."""
    before = parts(deck)
    digest = hashlib.sha256(deck.read_bytes()).hexdigest()

    report, out = run(
        deck,
        tmp_path,
        {
            "op": "replace_text",
            "slide": 1,
            "quote": "revenue grew",
            "text": "sales rose",
        },
    )

    assert report.saved and report.applied == 1
    assert ("sales rose", {"b": "0", "i": "1"}) in runs(out, 1)
    assert slide_texts(out)[0].startswith("Quarterly sales rose fast")
    after = parts(out)
    assert [name for name in after if after[name] != before.get(name)] == [
        "ppt/slides/slide1.xml"
    ]
    assert hashlib.sha256(deck.read_bytes()).hexdigest() == digest
    assert "RUN_FORMATTING_MERGED" not in codes(report)
    check_deck(out)


def test_a_quote_across_runs_takes_the_first_runs_formatting_and_says_so(
    deck, tmp_path
):
    """The user is told when differently formatted words were merged into one look."""
    report, out = run(
        deck,
        tmp_path,
        {
            "op": "replace_text",
            "slide": 1,
            "quote": "Quarterly revenue",
            "text": "Annual sales",
        },
    )

    assert report.saved
    assert runs(out, 1)[:3] == [
        ("Annual sales", {"b": "1", "i": "0"}),
        (" grew", {"b": "0", "i": "1"}),
        (" fast", {"b": "0", "i": "0"}),
    ]
    notice = next(n for n in report.must_tell_user if n.code == "RUN_FORMATTING_MERGED")
    assert notice.values == {"slide": 1, "quote": "Quarterly revenue"}
    check_deck(out)


@pytest.mark.parametrize(
    ("quote", "text", "expected"),
    [
        ("Inside the group", "Within the group", "Within the group"),
        ("Due Friday", "Due Monday", "Due Monday"),
        ("Quarterly   revenue", "Annual revenue", "Annual revenue grew fast"),
    ],
)
def test_text_in_groups_tables_and_with_loose_whitespace_is_found(
    deck, tmp_path, quote, text, expected
):
    """Quotes reach grouped shapes and table cells, and whitespace runs match one space."""
    report, out = run(
        deck, tmp_path, {"op": "replace_text", "slide": 1, "quote": quote, "text": text}
    )

    assert report.saved, report.as_text()
    assert expected in slide_texts(out)[0]
    check_deck(out)


def test_notes_are_edited_only_when_asked(deck, tmp_path):
    """where: notes edits the speaker notes and leaves the slide alone."""
    report, out = run(
        deck,
        tmp_path,
        {
            "op": "replace_text",
            "slide": 1,
            "quote": "revenue first",
            "text": "costs first",
            "where": "notes",
        },
    )

    assert report.saved
    edited = Presentation(str(out)).slides[0]
    assert edited.notes_slide.notes_text_frame.text == "Speak about costs first"
    assert "revenue grew" in slide_texts(out)[0]
    check_deck(out)


def test_an_empty_text_deletes_the_quote_and_a_newline_becomes_a_line_break(
    deck, tmp_path
):
    """A newline cannot split the paragraph, so it becomes a line break."""
    report, out = run(
        deck,
        tmp_path,
        {"op": "replace_text", "slide": 1, "quote": " fast", "text": ""},
        {
            "op": "replace_text",
            "slide": 2,
            "quote": "Back to the summary",
            "text": "Back to\nthe summary",
        },
    )

    assert report.saved
    assert slide_texts(out)[0].startswith("Quarterly revenue grew |")
    slide = etree.fromstring(parts(out)["ppt/slides/slide2.xml"], PARSER)
    assert slide.find(f".//{{{A}}}br") is not None
    check_deck(out)


@pytest.mark.parametrize(
    ("operation", "code", "values"),
    [
        (
            {"op": "replace_text", "slide": 1, "quote": "profit", "text": "x"},
            "QUOTE_NOT_FOUND",
            {"quote": "profit"},
        ),
        (
            {"op": "replace_text", "slide": 1, "quote": "e", "text": "x"},
            "QUOTE_AMBIGUOUS",
            {"quote": "e", "count": 13},
        ),
        (
            {"op": "replace_text", "slide": 1, "quote": "group Second", "text": "x"},
            "QUOTE_CROSSES_PARAGRAPHS",
            {"quote": "group Second"},
        ),
        (
            {"op": "replace_text", "slide": 4, "quote": "Page 4", "text": "x"},
            "QUOTE_IN_FIELD",
            {"quote": "Page 4"},
        ),
        (
            {
                "op": "replace_text",
                "slide": 4,
                "quote": "Closing",
                "text": "x",
                "where": "notes",
            },
            "NOTES_MISSING",
            {"slide": 4},
        ),
        (
            {"op": "replace_text", "slide": 9, "quote": "Closing", "text": "x"},
            "SLIDE_NOT_FOUND",
            {"slide": 9, "count": 4},
        ),
        (
            {"op": "delete_slide", "slide": 0},
            "SLIDE_NOT_FOUND",
            {"slide": 0, "count": 4},
        ),
    ],
)
def test_refusals_name_the_problem_and_save_nothing(
    deck, tmp_path, operation, code, values
):
    """One refused operation saves nothing, and the refusal carries a code and the values to act on."""
    report, out = run(deck, tmp_path, {"op": "duplicate_slide", "slide": 2}, operation)

    assert not report.saved
    assert not out.exists()
    assert (
        report.outcomes[1].status,
        report.outcomes[1].code,
        report.outcomes[1].values,
    ) == ("refused", code, values)
    assert report.outcomes[0].code == "BATCH_REFUSED"


@pytest.mark.parametrize(
    ("operation", "field"),
    [
        ({"op": "replace_text", "quote": "x", "text": "y"}, "slide"),
        ({"op": "replace_text", "slide": 1, "text": "y"}, "quote"),
        ({"op": "replace_text", "slide": 1, "quote": "  ", "text": "y"}, "quote"),
        ({"op": "replace_text", "slide": 1, "quote": "x"}, "text"),
        (
            {
                "op": "replace_text",
                "slide": 1,
                "quote": "x",
                "text": "y",
                "where": "footer",
            },
            "where",
        ),
        ({"op": "delete_slide", "slide": "2"}, "slide"),
        ({"op": "duplicate_slide", "slide": True}, "slide"),
    ],
)
def test_malformed_operations_name_the_field(deck, tmp_path, operation, field):
    """The model learns which field to fix."""
    report, _ = run(deck, tmp_path, operation)

    assert (report.outcomes[0].code, report.outcomes[0].values) == (
        "BAD_OPERATION",
        {"field": field},
    )


def test_another_formats_operation_is_unsupported(deck, tmp_path):
    """Operations of another format are refused, not ignored."""
    report, _ = run(
        deck, tmp_path, {"op": "set_cell", "sheet": "A", "cell": "A1", "value": 1}
    )

    assert (report.outcomes[0].code, report.outcomes[0].values) == (
        "UNSUPPORTED_OPERATION",
        {"op": "set_cell", "format": "pptx"},
    )


def test_deleting_a_slide_removes_its_notes_and_its_only_picture_and_leaves_valid_links(
    deck, tmp_path
):
    """Parts only the deleted slide used go with it; shared ones stay."""
    report, out = run(deck, tmp_path, {"op": "delete_slide", "slide": 2})

    assert report.saved
    after = parts(out)
    assert slide_texts(out) == [
        slide_texts(deck)[0],
        "Chart caption",
        "Closing slide | Page 4",
    ]
    assert "ppt/media/image1.png" not in after
    assert not any(name.startswith("ppt/notesSlides/notesSlide2") for name in after)
    check_deck(out)


def test_deleting_a_slide_others_link_to_drops_those_links(deck, tmp_path):
    """A link to a deleted slide would leave an r:id without its relationship."""
    report, out = run(deck, tmp_path, {"op": "delete_slide", "slide": 1})

    assert report.saved
    assert slide_texts(out)[0] == "Back to the summary"
    assert b"hlinkClick" not in parts(out)["ppt/slides/slide2.xml"]
    check_deck(out)


def test_the_last_slide_cannot_be_deleted(deck, tmp_path):
    """PowerPoint cannot open a deck without slides."""
    report, _ = run(
        deck, tmp_path, *({"op": "delete_slide", "slide": n} for n in (1, 2, 3, 4))
    )

    assert report.outcomes[3].code == "LAST_SLIDE"
    assert not report.saved


def test_a_duplicate_lands_after_its_slide_with_its_own_notes_and_chart(deck, tmp_path):
    """A copy owns its notes and chart, so editing one never changes the other."""
    report, out = run(
        deck,
        tmp_path,
        {"op": "duplicate_slide", "slide": 1},
        {"op": "duplicate_slide", "slide": 3},
    )

    assert report.saved
    texts = slide_texts(out)
    original = slide_texts(deck)
    assert texts == [
        original[0],
        original[0],
        original[1],
        original[2],
        original[2],
        original[3],
    ]
    copied = Presentation(str(out)).slides
    assert copied[1].notes_slide.notes_text_frame.text == "Speak about revenue first"
    assert copied[1].notes_slide.part is not copied[0].notes_slide.part
    chart_parts = {
        shape.chart.part.partname
        for shape in (*copied[3].shapes, *copied[4].shapes)
        if shape.has_chart
    }
    assert len(chart_parts) == 2
    check_deck(out)


def test_slide_numbers_refer_to_the_deck_as_sent(deck, tmp_path):
    """The model numbers slides as it read them, whatever earlier operations did."""
    report, out = run(
        deck,
        tmp_path,
        {"op": "duplicate_slide", "slide": 1},
        {"op": "delete_slide", "slide": 3},
        {
            "op": "replace_text",
            "slide": 2,
            "quote": "Back to the summary",
            "text": "Summary",
        },
        {"op": "replace_text", "slide": 1, "quote": "fast", "text": "quickly"},
    )

    assert report.saved, report.as_text()
    texts = slide_texts(out)
    assert texts[0].startswith("Quarterly revenue grew quickly")
    assert texts[1].startswith("Quarterly revenue grew fast")
    assert texts[2:] == ["Summary", "Closing slide | Page 4"]
    check_deck(out)


def test_a_deleted_slide_cannot_be_edited_later_in_the_call(deck, tmp_path):
    """An edit to a slide deleted earlier in the call is refused, not applied elsewhere."""
    report, _ = run(
        deck,
        tmp_path,
        {"op": "delete_slide", "slide": 2},
        {"op": "replace_text", "slide": 2, "quote": "Back", "text": "x"},
    )

    assert (report.outcomes[1].code, report.outcomes[1].values) == (
        "SLIDE_NOT_FOUND",
        {"slide": 2, "count": 4},
    )


def test_replacing_a_quote_with_itself_changes_nothing(deck, tmp_path):
    """No copy is saved when no text would change."""
    report, out = run(
        deck,
        tmp_path,
        {
            "op": "replace_text",
            "slide": 3,
            "quote": "Chart caption",
            "text": "Chart caption",
        },
    )

    assert not report.saved
    assert "NOTHING_CHANGED" in codes(report)
    assert not out.exists()


def test_sections_follow_deleted_and_duplicated_slides(deck, tmp_path):
    """PowerPoint repairs a deck whose sections name missing slides."""
    from tests.unit.artifacts.revised_copies.office_fixtures import P, rewrite_part

    p14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"

    def add_sections(data: bytes) -> bytes:
        root = etree.fromstring(data, PARSER)
        ids = [e.get("id") for e in root.iter(f"{{{P}}}sldId")]
        ext_list = root.find(f"{{{P}}}extLst")
        if ext_list is None:
            ext_list = etree.SubElement(root, f"{{{P}}}extLst")
        ext = etree.SubElement(
            ext_list, f"{{{P}}}ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
        )
        sections = etree.SubElement(ext, f"{{{p14}}}sectionLst")
        for name, members in (("Intro", ids[:2]), ("Rest", ids[2:])):
            section = etree.SubElement(
                sections,
                f"{{{p14}}}section",
                name=name,
                id=f"{{00000000-0000-0000-0000-00000000000{len(name)}}}",
            )
            listed = etree.SubElement(section, f"{{{p14}}}sldIdLst")
            for member in members:
                etree.SubElement(listed, f"{{{p14}}}sldId", id=member)
        return etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )

    rewrite_part(deck, "ppt/presentation.xml", add_sections)

    report, out = run(
        deck,
        tmp_path,
        {"op": "delete_slide", "slide": 2},
        {"op": "duplicate_slide", "slide": 3},
    )

    assert report.saved
    root = etree.fromstring(parts(out)["ppt/presentation.xml"], PARSER)
    slides = [e.get("id") for e in root.iter(f"{{{P}}}sldId")]
    sectioned = [e.get("id") for e in root.iter(f"{{{p14}}}sldId")]
    assert sectioned == slides
    check_deck(out)


def test_text_xml_cannot_carry_is_refused(deck, tmp_path):
    """Control characters would make the slide unreadable XML."""
    report, _ = run(
        deck,
        tmp_path,
        {"op": "replace_text", "slide": 3, "quote": "Chart", "text": "Bell\x07"},
    )

    assert (report.outcomes[0].code, report.outcomes[0].values) == (
        "BAD_OPERATION",
        {"field": "text"},
    )


def test_a_file_that_is_not_a_deck_is_refused(tmp_path):
    """A workbook named .pptx is refused before any operation runs."""
    import openpyxl

    from modules.artifacts.revised_copies.engines.package import PackageRefusedError

    path = tmp_path / "book.pptx"
    openpyxl.Workbook().save(path)

    with pytest.raises(PackageRefusedError) as refusal:
        run(path, tmp_path, {"op": "delete_slide", "slide": 1})
    assert refusal.value.code == "NOT_A_PRESENTATION"
