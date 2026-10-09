"""Real workbooks and decks for the Excel and PowerPoint engine tests, plus
the structural checks that stand in for opening them in Office."""

from __future__ import annotations

import posixpath
import re
import struct
import zipfile
import zlib
from collections.abc import Callable
from pathlib import Path

import openpyxl
import xlsxwriter
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"

PARSER = etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True)


def tiny_png() -> bytes:
    """A 2x2 red PNG, so the media parts have real bytes to compare."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 2 for _ in range(2))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def rewrite_part(path: Path, name: str, change: Callable[[bytes], bytes]) -> None:
    """Rewrites one zip entry in place, the others copied as they are."""
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in entries:
            target.writestr(info, change(data) if info.filename == name else data)


def add_part(path: Path, name: str, data: bytes) -> None:
    """Adds one zip entry."""
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as target:
        target.writestr(name, data)


def parts(path: Path) -> dict[str, bytes]:
    """Every part's bytes by name, in zip order."""
    with zipfile.ZipFile(path) as package:
        return {
            info.filename: package.read(info.filename) for info in package.infolist()
        }


# --- workbooks -------------------------------------------------------------


def build_workbook(path: Path, *, image: Path) -> Path:
    """Data sheet with shared strings, a SUM, a native chart, a text-box shape,
    a picture, a merge, an array formula and a defined name; a second sheet
    with a table; then shared formulas, an inline string and a calc chain
    written into the XML, since xlsxwriter writes none of them."""
    image.write_bytes(tiny_png())
    book = xlsxwriter.Workbook(str(path))
    data = book.add_worksheet("Data")
    data.write_row("A1", ["Region", "Sales", "Bonus"])
    for row, (region, sales) in enumerate(
        [("North", 10), ("South", 20), ("East", 30), ("West", 40), ("Online", 50)],
        start=1,
    ):
        data.write_string(row, 0, region)
        data.write_number(row, 1, sales)
        data.write_formula(row, 2, f"=B{row + 1}*2", None, sales * 2)
    data.write_string("A7", "Total")
    data.write_formula("B7", "=SUM(B2:B6)", None, 150)
    data.merge_range("E1:F1", "Merged title")
    data.write_array_formula("H2:H3", "{=B2:B3*10}", None, 100)
    chart = book.add_chart({"type": "column"})
    chart.add_series({"categories": "=Data!$A$2:$A$6", "values": "=Data!$B$2:$B$6"})
    data.insert_chart("J2", chart)
    data.insert_textbox("J20", "A note shape")
    data.insert_image("N2", str(image))
    book.define_name("Total", "=Data!$B$7")

    table = book.add_worksheet("Table")
    table.add_table(
        "A1:B3",
        {
            "columns": [{"header": "Item"}, {"header": "Cost"}],
            "data": [["Pen", 2], ["Pad", 3]],
        },
    )
    book.close()

    rewrite_part(path, "xl/worksheets/sheet1.xml", _shared_formulas_and_inline_string)
    _add_calc_chain(path)
    return path


def _shared_formulas_and_inline_string(data: bytes) -> bytes:
    root = etree.fromstring(data, PARSER)
    cells = {c.get("r"): c for c in root.iter(f"{{{S}}}c")}
    for ref in ("C2", "C3", "C4", "C5", "C6"):
        formula = cells[ref].find(f"{{{S}}}f")
        formula.attrib.clear()
        formula.set("t", "shared")
        formula.set("si", "0")
        formula.text = "B2*2" if ref == "C2" else None
    cells["C2"].find(f"{{{S}}}f").set("ref", "C2:C6")
    note = cells["A7"]
    for child in list(note):
        note.remove(child)
    note.set("t", "inlineStr")
    inline = etree.SubElement(note, f"{{{S}}}is")
    etree.SubElement(inline, f"{{{S}}}t").text = "Total"
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _add_calc_chain(path: Path) -> None:
    chain = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<c r="C2" i="1"/><c r="C3"/><c r="C4"/><c r="C5"/><c r="C6"/><c r="B7"/>'
        b'<c r="H2" a="1"/></calcChain>'
    )
    add_part(path, "xl/calcChain.xml", chain)

    def rel(data: bytes) -> bytes:
        return data.replace(
            b"</Relationships>",
            b'<Relationship Id="rIdChain" '
            b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" '
            b'Target="calcChain.xml"/></Relationships>',
        )

    def override(data: bytes) -> bytes:
        return data.replace(
            b"</Types>",
            b'<Override PartName="/xl/calcChain.xml" ContentType='
            b'"application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>',
        )

    rewrite_part(path, "xl/_rels/workbook.xml.rels", rel)
    rewrite_part(path, "[Content_Types].xml", override)


def build_openpyxl_workbook(path: Path) -> Path:
    """openpyxl's own writer: shared strings, styles, a defined name, gaps in rows."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet["A1"] = "Item"
    sheet["B1"] = "Amount"
    sheet["A2"] = "Rent"
    sheet["B2"] = 1200
    sheet["B2"].number_format = "#,##0.00"
    sheet["A5"] = "Total"
    sheet["B5"] = "=SUM(B2:B4)"
    book.defined_names["Rent"] = openpyxl.workbook.defined_name.DefinedName(
        "Rent", attr_text="Budget!$B$2"
    )
    book.save(path)
    return path


def check_workbook(path: Path) -> None:
    """What Excel would complain about, without Excel: every part parses, every
    relationship resolves, rows and cells are in order with matching row
    numbers, shared-string indices exist, and openpyxl reads it."""
    contents = parts(path)
    _check_package(contents)
    strings = 0
    if "xl/sharedStrings.xml" in contents:
        strings = len(
            etree.fromstring(contents["xl/sharedStrings.xml"], PARSER).findall(
                f"{{{S}}}si"
            )
        )
    for name, data in contents.items():
        if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
            continue
        root = etree.fromstring(data, PARSER)
        last_row = 0
        for row in root.iter(f"{{{S}}}row"):
            number = int(row.get("r"))
            assert number > last_row, f"{name}: row {number} out of order"
            last_row = number
            last_column = 0
            for cell in row.findall(f"{{{S}}}c"):
                column, cell_row = _split(cell.get("r"))
                assert cell_row == number, f"{name}: {cell.get('r')} in row {number}"
                assert column > last_column, f"{name}: {cell.get('r')} out of order"
                last_column = column
                if cell.get("t") == "s":
                    assert int(cell.findtext(f"{{{S}}}v")) < strings
                tags = [etree.QName(child).localname for child in cell]
                order = ["f", "v", "is", "extLst"]
                assert tags == sorted(tags, key=order.index), (
                    f"{name}: {cell.get('r')} {tags}"
                )
    openpyxl.load_workbook(path).close()


def _split(ref: str) -> tuple[int, int]:
    letters, digits = re.fullmatch(r"([A-Z]+)(\d+)", ref).groups()
    column = 0
    for letter in letters:
        column = column * 26 + ord(letter) - 64
    return column, int(digits)


# --- decks ------------------------------------------------------------------


def build_deck(path: Path, *, image: Path) -> Path:
    """Four slides: mixed-format runs, a group, a table and notes on slide 1;
    a picture, a link to slide 1 and notes on slide 2; a chart on slide 3;
    a plain slide 4 without notes."""
    image.write_bytes(tiny_png())
    deck = Presentation()
    blank = deck.slide_layouts[6]
    title_only = deck.slide_layouts[5]

    first = deck.slides.add_slide(title_only)
    paragraph = first.shapes.title.text_frame.paragraphs[0]
    for text, bold, italic in (
        ("Quarterly ", True, False),
        ("revenue grew", False, True),
        (" fast", False, False),
    ):
        run = paragraph.add_run()
        run.text = text
        run.font.bold = bold
        run.font.italic = italic
    group = first.shapes.add_group_shape()
    grouped = group.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    grouped.text_frame.text = "Inside the group"
    second_paragraph = grouped.text_frame.add_paragraph()
    second_paragraph.text = "Second line here"
    table = first.shapes.add_table(
        2, 2, Inches(1), Inches(4), Inches(4), Inches(1)
    ).table
    table.cell(0, 0).text = "Owner"
    table.cell(1, 1).text = "Due Friday"
    first.notes_slide.notes_text_frame.text = "Speak about revenue first"

    second = deck.slides.add_slide(blank)
    box = second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Back to the summary"
    box.click_action.target_slide = first
    second.shapes.add_picture(str(image), Inches(5), Inches(1), Inches(1), Inches(1))
    second.notes_slide.notes_text_frame.text = "Pause for questions"

    third = deck.slides.add_slide(blank)
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2"]
    chart_data.add_series("Revenue", (1.0, 2.0))
    third.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    )
    caption = third.shapes.add_textbox(Inches(1), Inches(5.5), Inches(4), Inches(1))
    caption.text_frame.text = "Chart caption"
    caption.text_frame.paragraphs[0].runs[0].font.size = Pt(20)

    fourth = deck.slides.add_slide(blank)
    fourth.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text_frame.text = "Closing slide"
    footer = fourth.shapes.add_textbox(
        Inches(1), Inches(6), Inches(2), Inches(1)
    ).text_frame.paragraphs[0]
    footer.add_run().text = "Page "
    field = etree.SubElement(
        footer._p,
        f"{{{A}}}fld",
        id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}",
        type="slidenum",
    )
    etree.SubElement(field, f"{{{A}}}t").text = "4"
    deck.save(path)
    return path


def check_deck(path: Path) -> None:
    """Every part parses, every relationship resolves, every r:id a slide
    uses has a relationship, slide ids are unique, and python-pptx opens it."""
    contents = parts(path)
    _check_package(contents)
    presentation = etree.fromstring(contents["ppt/presentation.xml"], PARSER)
    ids = [element.get("id") for element in presentation.iter(f"{{{P}}}sldId")]
    assert len(ids) == len(set(ids))
    for name, data in contents.items():
        if (
            not name.endswith(".xml")
            or "/_rels/" in name
            or name == "[Content_Types].xml"
        ):
            continue
        rels_name = _rels_name(name)
        known = set()
        if rels_name in contents:
            known = {
                rel.get("Id") for rel in etree.fromstring(contents[rels_name], PARSER)
            }
        root = etree.fromstring(data, PARSER)
        for element in root.iter():
            for attribute, value in element.attrib.items():
                if attribute.startswith(f"{{{R}}}") and attribute != f"{{{R}}}version":
                    assert value in known, (
                        f"{name}: {attribute}={value} has no relationship"
                    )
    Presentation(str(path))


def slide_texts(path: Path) -> list[str]:
    """Each slide's shape texts joined with " | ", groups and tables included."""
    deck = Presentation(str(path))
    texts = []
    for slide in deck.slides:
        chunks = []
        for shape in slide.shapes:
            chunks.extend(_shape_texts(shape))
        texts.append(" | ".join(chunks))
    return texts


def _shape_texts(shape) -> list[str]:
    if shape.shape_type is not None and shape.shape_type == 6:  # group
        return [text for child in shape.shapes for text in _shape_texts(child)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return [
            cell.text for row in shape.table.rows for cell in row.cells if cell.text
        ]
    if shape.has_text_frame and shape.text_frame.text:
        return [shape.text_frame.text]
    return []


# --- shared ------------------------------------------------------------------


def _rels_name(name: str) -> str:
    folder, base = posixpath.split(name)
    return posixpath.join(folder, "_rels", base + ".rels")


def _check_package(contents: dict[str, bytes]) -> None:
    types = etree.fromstring(contents["[Content_Types].xml"], PARSER)
    defaults = {d.get("Extension").lower() for d in types.findall(f"{{{CT}}}Default")}
    overrides = {
        o.get("PartName").lstrip("/") for o in types.findall(f"{{{CT}}}Override")
    }
    for override in overrides:
        assert override in contents, f"override for missing part {override}"
    for name, data in contents.items():
        if name == "[Content_Types].xml" or name.endswith("/"):
            continue
        extension = name.rsplit(".", 1)[-1].lower()
        assert name in overrides or extension in defaults, f"{name} has no content type"
        if name.endswith((".xml", ".rels")):
            etree.fromstring(data, PARSER)
        if not name.endswith(".rels"):
            continue
        source_folder = posixpath.dirname(posixpath.dirname(name))
        for rel in etree.fromstring(data, PARSER).findall(f"{{{PR}}}Relationship"):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target")
            resolved = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join(source_folder, target))
            )
            assert resolved in contents, (
                f"{name}: {rel.get('Id')} -> missing {resolved}"
            )
