"""Opening an Office file read-only, changing parts of it and saving a copy."""

import zipfile
from pathlib import Path

import docx
import pytest
from lxml import etree

from modules.artifacts.revised_copies.engines import package as package_module
from modules.artifacts.revised_copies.engines.package import (
    PackageRefusedError,
    open_package,
    package_problems,
)

pytestmark = pytest.mark.unit

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
COMMENTS_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
)
COMMENTS_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
)


def _word_file(tmp_path: Path) -> Path:
    document = docx.Document()
    document.add_paragraph("Hello there.")
    path = tmp_path / "in.docx"
    document.save(str(path))
    return path


def _zip(tmp_path: Path, parts: dict[str, bytes], name: str = "x.docx") -> Path:
    path = tmp_path / name
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for part, data in parts.items():
            archive.writestr(part, data)
    return path


CONTENT_TYPES = (
    b'<?xml version="1.0" encoding="UTF-8"?>'
    b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    b'<Default Extension="xml" ContentType="application/xml"/></Types>'
)


def test_saving_untouched_keeps_every_part_byte_identical_in_order(tmp_path):
    """A copy with nothing changed has the same parts, order and compression."""
    source = _word_file(tmp_path)
    out = tmp_path / "out.docx"

    open_package(source).save(out)

    with zipfile.ZipFile(source) as a, zipfile.ZipFile(out) as b:
        assert [i.filename for i in a.infolist()] == [i.filename for i in b.infolist()]
        for left, right in zip(a.infolist(), b.infolist(), strict=True):
            assert a.read(left) == b.read(right)
            assert left.compress_type == right.compress_type


def test_set_xml_keeps_the_declaration_and_other_parts(tmp_path):
    """A changed part keeps its standalone declaration; the rest stay as they were."""
    source = _word_file(tmp_path)
    package = open_package(source)
    root = package.xml("word/document.xml")
    root.find(f".//{{{W}}}t").text = "Changed."

    package.set_xml("word/document.xml", root)
    out = tmp_path / "out.docx"
    package.save(out)

    saved = open_package(out)
    assert saved.read("word/document.xml").startswith(
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    )
    assert docx.Document(str(out)).paragraphs[0].text == "Changed."
    assert saved.read("word/styles.xml") == open_package(source).read("word/styles.xml")
    assert open_package(source).read("word/document.xml") != saved.read(
        "word/document.xml"
    )


def test_add_part_registers_its_content_type_and_relationship(tmp_path):
    """A new part gets an override and a relationship with a fresh id."""
    package = open_package(_word_file(tmp_path))
    root = etree.Element(f"{{{W}}}comments", nsmap={"w": W})

    rid = package.add_part(
        "word/comments.xml", root, COMMENTS_TYPE, "word/document.xml", COMMENTS_REL
    )

    out = tmp_path / "out.docx"
    package.save(out)
    saved = open_package(out)
    assert saved.content_type("word/comments.xml") == COMMENTS_TYPE
    assert saved.related("word/document.xml", COMMENTS_REL) == ["word/comments.xml"]
    rels = saved.xml("word/_rels/document.xml.rels")
    assert [r.get("Id") for r in rels].count(rid) == 1
    assert package_problems(saved) == []
    docx.Document(str(out))


def test_remove_drops_the_part_its_relationships_and_override(tmp_path):
    """No relationship or override is left pointing at a removed part."""
    package = open_package(_word_file(tmp_path))
    package.add_part(
        "word/comments.xml",
        etree.Element(f"{{{W}}}comments", nsmap={"w": W}),
        COMMENTS_TYPE,
        "word/document.xml",
        COMMENTS_REL,
    )

    package.remove("word/comments.xml")

    assert "word/comments.xml" not in package.names
    assert package.related("word/document.xml", COMMENTS_REL) == []
    assert package.content_type("word/comments.xml") == "application/xml"
    assert package_problems(package) == []


def test_a_part_with_a_doctype_is_refused(tmp_path):
    """Entity tricks never reach the engine: a DOCTYPE anywhere refuses the file."""
    evil = (
        b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY x SYSTEM "file:///c:/windows/win.ini">]>'
        b"<d>&x;</d>"
    )
    path = _zip(tmp_path, {"[Content_Types].xml": CONTENT_TYPES, "a.xml": evil})

    with pytest.raises(PackageRefusedError) as refused:
        open_package(path).xml("a.xml")

    assert refused.value.code == "UNSAFE_XML"
    assert refused.value.values == {"part": "a.xml", "reason": "doctype"}


def test_the_structural_check_refuses_a_doctype_in_any_part(tmp_path):
    """A DOCTYPE in a part no engine edits still refuses the file."""
    path = _zip(
        tmp_path,
        {
            "[Content_Types].xml": CONTENT_TYPES,
            "customXml/item1.xml": b'<!DOCTYPE a [<!ENTITY e "x">]><a>&e;</a>',
        },
    )

    with pytest.raises(PackageRefusedError) as refused:
        package_problems(open_package(path))

    assert refused.value.code == "UNSAFE_XML"


def test_a_part_bigger_than_the_cap_is_refused_while_reading(tmp_path, monkeypatch):
    """Decompressed bytes are counted as they stream, whatever the header says."""
    monkeypatch.setattr(package_module, "PART_LIMIT", 1024)
    path = _zip(
        tmp_path,
        {
            "[Content_Types].xml": CONTENT_TYPES,
            "big.xml": b"<a>" + b" " * 4096 + b"</a>",
        },
    )

    with pytest.raises(PackageRefusedError) as refused:
        open_package(path).read("big.xml")

    assert refused.value.code == "PART_TOO_LARGE"
    assert refused.value.values["part"] == "big.xml"


def test_the_total_of_all_parts_is_capped(tmp_path, monkeypatch):
    """Many small parts cannot add up past the package's total."""
    monkeypatch.setattr(package_module, "TOTAL_LIMIT", 2048)
    parts = {"[Content_Types].xml": CONTENT_TYPES}
    parts.update({f"p{i}.xml": b"<a>" + b" " * 900 + b"</a>" for i in range(4)})
    path = _zip(tmp_path, parts)

    package = open_package(path)
    with pytest.raises(PackageRefusedError) as refused:
        for name in package.names:
            package.read(name)

    assert refused.value.code == "PART_TOO_LARGE"


def test_too_many_entries_are_refused(tmp_path, monkeypatch):
    """An archive of countless tiny entries is refused before any is read."""
    monkeypatch.setattr(package_module, "ENTRY_LIMIT", 3)
    parts = {"[Content_Types].xml": CONTENT_TYPES}
    parts.update({f"p{i}.xml": b"<a/>" for i in range(4)})

    with pytest.raises(PackageRefusedError) as refused:
        open_package(_zip(tmp_path, parts))

    assert refused.value.code == "PART_TOO_LARGE"


def test_a_file_that_is_not_a_zip_is_not_a_package(tmp_path):
    """Plain bytes are refused with NOT_A_PACKAGE."""
    path = tmp_path / "x.docx"
    path.write_bytes(b"hello")

    with pytest.raises(PackageRefusedError) as refused:
        open_package(path)

    assert refused.value.code == "NOT_A_PACKAGE"


def test_a_zip_without_content_types_is_not_a_package(tmp_path):
    """A zip is an Office package only with [Content_Types].xml."""
    with pytest.raises(PackageRefusedError) as refused:
        open_package(_zip(tmp_path, {"a.xml": b"<a/>"}))

    assert refused.value.code == "NOT_A_PACKAGE"


def test_a_password_protected_file_is_encrypted(tmp_path):
    """Office saves an encrypted file as an OLE container, not a zip."""
    path = tmp_path / "x.docx"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 512)

    with pytest.raises(PackageRefusedError) as refused:
        open_package(path)

    assert refused.value.code == "ENCRYPTED"


def test_problems_name_a_missing_target_and_a_part_without_type(tmp_path):
    """The structural check finds dangling relationships and untyped parts."""
    content_types = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        b'<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        b'<Override PartName="/a.xml" ContentType="application/xml"/></Types>'
    )
    rels = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" Type="t" Target="a.xml"/>'
        b'<Relationship Id="rId2" Type="t" Target="missing.xml"/>'
        b'<Relationship Id="rId3" Type="t" Target="https://example.com" TargetMode="External"/>'
        b"</Relationships>"
    )
    path = _zip(
        tmp_path,
        {
            "[Content_Types].xml": content_types,
            "_rels/.rels": rels,
            "a.xml": b"<a/>",
            "b.bin": b"x",
            "c.xml": b"<c>",
        },
    )

    problems = package_problems(open_package(path))

    assert any("missing.xml" in p for p in problems)
    assert any("b.bin" in p for p in problems)
    assert any("c.xml" in p for p in problems)
    assert not any("example.com" in p for p in problems)
