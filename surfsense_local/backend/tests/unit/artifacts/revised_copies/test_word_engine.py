"""The Word engine: every edit a tracked change, comments real, the user's file untouched."""

import hashlib
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import docx
import pytest
from lxml import etree

from modules.artifacts.revised_copies.engines import word
from modules.artifacts.revised_copies.engines.package import PackageRefusedError
from tests.unit.artifacts.revised_copies.word_fixtures import (
    COMMENTS_TYPE,
    anchored_text,
    comments_of,
    deleted,
    docx_with_body,
    inserted,
    para,
    plain_text,
    revision_authors,
    rich_docx,
    run,
    structure_problems,
    text_now,
    w,
)

pytestmark = pytest.mark.unit

WHEN = datetime(2026, 10, 6, 9, 30, tzinfo=UTC)
NBSP = chr(0xA0)


def apply(path: Path, tmp_path: Path, *operations: dict, partial: bool = False):
    """Run the Word engine on a copy, with a fixed date."""
    out = tmp_path / "out.docx"
    report = word.apply(path, list(operations), out, partial=partial, date=WHEN)
    return report, out


def accepted(path: Path, tmp_path: Path, **options) -> list[str]:
    """Accept every change into a new file and read its paragraphs."""
    out = tmp_path / "accepted.docx"
    word.accept_all(path, out, **options)
    assert structure_problems(out) == []
    return plain_text(out)


def rejected(path: Path, tmp_path: Path) -> list[str]:
    """Reject every change into a new file and read its paragraphs."""
    out = tmp_path / "rejected.docx"
    word.reject_all(path, out)
    assert structure_problems(out) == []
    return plain_text(out)


def body_of(path: Path) -> etree._Element:
    """The body element of a saved file."""
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("word/document.xml")).find(w("body"))


def changes(path: Path, tag: str) -> list[str]:
    """The text inside each w:ins or w:del of the body, in order."""
    text_tag = w("delText") if tag == "del" else w("t")
    return [
        "".join(t.text or "" for t in el.iter(text_tag))
        for el in body_of(path).iter(w(tag))
        if el.getparent().tag != w("rPr")
    ]


def codes(report) -> list[str | None]:
    """Each outcome's code, None when applied."""
    return [o.code for o in report.outcomes]


# Replacing text


def test_a_replacement_is_a_tracked_change_by_surfsense(tmp_path):
    """Text is deleted and inserted as w:del and w:ins with author and date."""
    source = docx_with_body(tmp_path, para(run("The Supplier delivers goods.")))

    report, out = apply(
        source, tmp_path, {"op": "replace_text", "quote": "Supplier", "text": "Vendor"}
    )

    assert report.saved and report.applied == 1 and codes(report) == [None]
    assert text_now(out) == ["The Vendor delivers goods."]
    assert changes(out, "del") == ["Supplier"]
    assert changes(out, "ins") == ["Vendor"]
    for el in body_of(out).iter(w("ins"), w("del")):
        assert el.get(w("author")) == "SurfSense"
        assert el.get(w("date")) == "2026-10-06T09:30:00Z"
    assert structure_problems(out) == []
    assert accepted(out, tmp_path) == ["The Vendor delivers goods."]
    assert rejected(out, tmp_path) == ["The Supplier delivers goods."]


def test_only_the_words_that_differ_are_marked(tmp_path):
    """Replacing "within 30 days" with "within 45 days" strikes only 30."""
    source = rich_docx(tmp_path)

    report, out = apply(
        source,
        tmp_path,
        {
            "op": "replace_text",
            "quote": "due within 30 days",
            "text": "due within 45 days",
        },
    )

    assert report.saved
    assert changes(out, "del") == ["30"]
    assert changes(out, "ins") == ["45"]


def test_a_replacement_that_only_adds_words_inserts_without_deleting(tmp_path):
    """New words in the middle of a quote become one insertion."""
    source = docx_with_body(tmp_path, para(run("Payment to the Seller.")))

    report, out = apply(
        source,
        tmp_path,
        {
            "op": "replace_text",
            "quote": "the Seller",
            "text": "the Buyer and the Seller",
        },
    )

    assert report.saved
    assert changes(out, "del") == []
    assert text_now(out) == ["Payment to the Buyer and the Seller."]
    assert rejected(out, tmp_path) == ["Payment to the Seller."]


def test_empty_text_deletes_the_quote(tmp_path):
    """An empty replacement is a tracked deletion and nothing else."""
    source = docx_with_body(tmp_path, para(run("Pay fully and promptly.")))

    report, out = apply(
        source, tmp_path, {"op": "replace_text", "quote": "fully and ", "text": ""}
    )

    assert report.saved
    assert changes(out, "ins") == []
    assert accepted(out, tmp_path) == ["Pay promptly."]
    assert rejected(out, tmp_path) == ["Pay fully and promptly."]


def test_a_quote_across_differently_formatted_runs_keeps_the_first_runs_format(
    tmp_path,
):
    """The quote may cross runs; the inserted run copies the first quoted run's rPr."""
    source = docx_with_body(
        tmp_path,
        para(
            run("The "),
            run("Sup", bold=True),
            run("plier shall", italic=True),
            run(" deliver."),
        ),
    )

    report, out = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "Supplier shall", "text": "Vendor will"},
    )

    assert report.saved
    insertion = next(body_of(out).iter(w("ins")))
    assert insertion.find(f"{w('r')}/{w('rPr')}/{w('b')}") is not None
    assert accepted(out, tmp_path) == ["The Vendor will deliver."]
    assert rejected(out, tmp_path) == ["The Supplier shall deliver."]
    reverted = docx.Document(str(tmp_path / "rejected.docx")).paragraphs[0]
    assert [
        (r.text, bool(r.bold), bool(r.italic)) for r in reverted.runs if r.text
    ] == [
        ("The ", False, False),
        ("Sup", True, False),
        ("plier shall", False, True),
        (" deliver.", False, False),
    ]


def test_tabs_and_no_break_spaces_match_a_plain_space(tmp_path):
    """Whitespace runs of any kind equal one space in a quote."""
    source = docx_with_body(
        tmp_path,
        para(
            f'<w:r><w:t>Net</w:t><w:tab/><w:t xml:space="preserve">30{NBSP} days</w:t></w:r>'
        ),
    )

    report, out = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "Net 30 days", "text": "Net 60 days"},
    )

    assert report.saved
    assert changes(out, "del") == ["30"]
    assert accepted(out, tmp_path) == [f"Net\t60{NBSP} days"]


def test_a_cell_and_a_content_control_are_in_scope(tmp_path):
    """Table cells and content controls are edited like body text."""
    source = docx_with_body(
        tmp_path,
        "<w:tbl><w:tr><w:tc>"
        + para(run("Fee: 100"))
        + "</w:tc></w:tr></w:tbl>"
        + para(
            run("Client: "),
            "<w:sdt><w:sdtPr><w:alias w:val='Client'/></w:sdtPr><w:sdtContent>"
            + run("Acme Ltd")
            + "</w:sdtContent></w:sdt>",
        ),
    )

    report, out = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "Fee: 100", "text": "Fee: 120"},
        {"op": "replace_text", "quote": "Acme Ltd", "text": "Acme Group"},
    )

    assert report.saved and report.applied == 2
    assert accepted(out, tmp_path) == ["Fee: 120", "Client: Acme Group"]
    assert body_of(out).find(f".//{w('sdtContent')}/{w('ins')}") is not None


def test_text_inside_a_hyperlink_stays_inside_the_link(tmp_path):
    """An edit within a link's text keeps the link around the new words."""
    source = docx_with_body(
        tmp_path,
        para(
            run("See "),
            '<w:hyperlink r:id="rIdLink">' + run("the old terms") + "</w:hyperlink>",
            run("."),
        ),
    )

    report, out = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "old terms", "text": "new terms"},
    )

    assert report.saved
    link = body_of(out).find(f".//{w('hyperlink')}")
    assert link.find(w("ins")) is not None and link.find(w("del")) is not None
    assert accepted(out, tmp_path) == ["See the new terms."]
    assert structure_problems(out) == []


def test_a_later_operation_sees_the_text_as_earlier_ones_left_it(tmp_path):
    """Editing words SurfSense just inserted nests the deletion in that insertion."""
    source = docx_with_body(tmp_path, para(run("Notice period: one month.")))

    report, out = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "one month", "text": "two months"},
        {"op": "replace_text", "quote": "two months", "text": "three months"},
    )

    assert report.saved and report.applied == 2
    assert body_of(out).find(f".//{w('ins')}/{w('del')}") is not None
    assert text_now(out) == ["Notice period: three months."]
    assert accepted(out, tmp_path) == ["Notice period: three months."]
    assert rejected(out, tmp_path) == ["Notice period: one month."]
    assert structure_problems(out) == []


# Other authors' revisions and comments


def _reviewed(tmp_path: Path) -> Path:
    return docx_with_body(
        tmp_path,
        para(
            run("The term is "),
            inserted(7, "Alice", run("two")),
            deleted(8, "Alice", "one"),
            run(" years."),
        )
        + para(
            '<w:commentRangeStart w:id="3"/>',
            run("Governing law is England."),
            '<w:commentRangeEnd w:id="3"/>',
            '<w:r><w:commentReference w:id="3"/></w:r>',
        ),
        comments=[(3, "Alice", "Check this.")],
    )


def test_quotes_read_another_authors_insertions_and_skip_their_deletions(tmp_path):
    """The quote matches the text as it reads now; Alice's markup stays hers."""
    source = _reviewed(tmp_path)

    report, out = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "two years", "text": "three years"},
    )

    assert report.saved, report.as_text()
    assert text_now(out)[0] == "The term is three years."
    assert revision_authors(out)["Alice"] == 2
    assert accepted(out, tmp_path)[0] == "The term is three years."
    assert rejected(out, tmp_path)[0] == "The term is one years."
    assert structure_problems(out) == []


def test_a_comment_joins_existing_comments_without_renumbering_them(tmp_path):
    """Alice's comment keeps id 3; SurfSense's gets an id unused anywhere in the file."""
    source = _reviewed(tmp_path)

    report, out = apply(
        source,
        tmp_path,
        {"op": "add_comment", "quote": "Governing law", "text": "Should be Delaware."},
    )

    assert report.saved
    new_id = report.outcomes[0].values["comment_id"]
    found = comments_of(out)
    assert found["3"] == ("Alice", "Check this.")
    assert found[new_id] == ("SurfSense", "Should be Delaware.")
    assert new_id not in {"3", "7", "8"}
    assert anchored_text(out, new_id) == "Governing law"
    assert structure_problems(out) == []


def test_a_comment_creates_the_comments_part_when_the_file_has_none(tmp_path):
    """The part, its content type and its relationship are added together."""
    source = docx_with_body(tmp_path, para(run("Liability is unlimited.")))

    report, out = apply(
        source,
        tmp_path,
        {
            "op": "replace_text",
            "quote": "unlimited",
            "text": "capped at the fees paid",
            "comment": "Market standard cap.",
        },
    )

    assert report.saved
    comment_id = report.outcomes[0].values["comment_id"]
    assert comments_of(out) == {comment_id: ("SurfSense", "Market standard cap.")}
    with zipfile.ZipFile(out) as archive:
        assert COMMENTS_TYPE.encode() in archive.read("[Content_Types].xml")
    assert anchored_text(out, comment_id) == "capped at the fees paid"
    assert structure_problems(out) == []


def test_external_drops_internal_comments_and_keeps_the_rest(tmp_path):
    """The copy a counterparty gets has no internal comment, range or reference."""
    source = _reviewed(tmp_path)
    report, out = apply(
        source,
        tmp_path,
        {
            "op": "add_comment",
            "quote": "Governing law",
            "text": "Our fallback is NY.",
            "internal": True,
        },
        {"op": "add_comment", "quote": "The term", "text": "Please confirm."},
    )
    internal = report.outcomes[0].values["comment_id"]
    shared = report.outcomes[1].values["comment_id"]

    external = tmp_path / "external.docx"
    word.external(out, external, [internal])

    found = comments_of(external)
    assert set(found) == {"3", shared}
    with zipfile.ZipFile(external) as archive:
        document = archive.read("word/document.xml").decode()
    assert f'w:id="{internal}"' not in document
    assert text_now(external) == text_now(out)
    assert structure_problems(external) == []


def test_accept_all_without_comments_leaves_a_clean_file(tmp_path):
    """The clean download: every change accepted, every comment and its part gone."""
    source = _reviewed(tmp_path)
    _, out = apply(
        source,
        tmp_path,
        {
            "op": "replace_text",
            "quote": "England",
            "text": "Delaware",
            "comment": "Why",
        },
    )

    clean = tmp_path / "clean.docx"
    word.accept_all(out, clean, drop_comments=True)

    assert plain_text(clean) == ["The term is two years.", "Governing law is Delaware."]
    assert comments_of(clean) == {}
    with zipfile.ZipFile(clean) as archive:
        assert "word/comments.xml" not in archive.namelist()
        assert b"comment" not in archive.read("word/document.xml")
    assert structure_problems(clean) == []
    assert word.counts(clean) == word.RevisionCounts(changes=0, comments=0)


def test_reject_all_undoes_paragraph_and_formatting_revisions(tmp_path):
    """Inserted paragraph marks, deleted ones and format changes are all undone."""
    mark = '<w:pPr><w:rPr><w:{kind} w:id="{id}" w:author="Bob" w:date="2026-01-01T00:00:00Z"/></w:rPr></w:pPr>'
    source = docx_with_body(
        tmp_path,
        para(inserted(20, "Bob", run("Split ")), ppr=mark.format(kind="ins", id=21))
        + para(run("here."))
        + para(run("Joined "), ppr=mark.format(kind="del", id=22))
        + para(run("text."))
        + para(
            '<w:r><w:rPr><w:b/><w:rPrChange w:id="23" w:author="Bob" w:date="2026-01-01T00:00:00Z">'
            "<w:rPr/></w:rPrChange></w:rPr><w:t>Bolded</w:t></w:r>"
        ),
    )

    assert rejected(source, tmp_path) == ["here.", "Joined ", "text.", "Bolded"]
    assert accepted(source, tmp_path) == ["Split ", "here.", "Joined text.", "Bolded"]
    plain = docx.Document(str(tmp_path / "rejected.docx")).paragraphs[3].runs[0]
    assert not plain.bold
    bold = docx.Document(str(tmp_path / "accepted.docx")).paragraphs[3].runs[0]
    assert bold.bold


# Paragraphs


def test_inserted_paragraphs_follow_the_anchor_and_copy_its_list_numbering(tmp_path):
    """A paragraph after a list item is the next list item; its mark is tracked."""
    source = rich_docx(tmp_path)
    original = plain_text(source)

    report, out = apply(
        source,
        tmp_path,
        {
            "op": "insert_paragraphs",
            "quote": "Second obligation",
            "text": "Third obligation of the parties.\nFourth obligation.",
        },
    )

    assert report.saved
    paragraphs = docx.Document(str(out)).paragraphs
    styles = [p.style.name for p in paragraphs]
    after = text_now(out)
    third = after.index("Third obligation of the parties.")
    assert after[third + 1] == "Fourth obligation."
    assert after[third - 1] == "Second obligation of the parties."
    assert styles[third] == "List Number" and styles[third + 1] == "List Number"
    marks = body_of(out).findall(f".//{w('pPr')}/{w('rPr')}/{w('ins')}")
    assert len(marks) == 2
    assert accepted(out, tmp_path) == after
    assert rejected(out, tmp_path) == original
    assert structure_problems(out) == []


def test_inserted_paragraphs_take_a_named_style(tmp_path):
    """style accepts a style's name or id."""
    source = rich_docx(tmp_path)

    report, out = apply(
        source,
        tmp_path,
        {
            "op": "insert_paragraphs",
            "quote": "between Acme",
            "text": "Definitions",
            "style": "Heading 2",
        },
        {
            "op": "insert_paragraphs",
            "quote": "Closing words",
            "text": "Signatures",
            "style": "Heading2",
        },
    )

    assert report.saved
    accepted(out, tmp_path)
    styles = {
        p.text: p.style.name
        for p in docx.Document(str(tmp_path / "accepted.docx")).paragraphs
    }
    assert styles["Definitions"] == "Heading 2"
    assert styles["Signatures"] == "Heading 2"


def test_an_unknown_style_is_refused(tmp_path):
    """STYLE_NOT_FOUND names the style; nothing is saved."""
    source = rich_docx(tmp_path)

    report, out = apply(
        source,
        tmp_path,
        {
            "op": "insert_paragraphs",
            "quote": "Closing words",
            "text": "x",
            "style": "Fancy",
        },
    )

    assert not report.saved and not out.exists()
    assert report.outcomes[0].code == "STYLE_NOT_FOUND"
    assert report.outcomes[0].values == {"style": "Fancy"}


def test_deleting_paragraphs_through_a_quote_removes_them_on_accept(tmp_path):
    """Text and paragraph marks are tracked deletions; reject brings them back."""
    source = rich_docx(tmp_path)
    original = plain_text(source)

    report, out = apply(
        source,
        tmp_path,
        {
            "op": "delete_paragraphs",
            "quote": "First obligation",
            "through": "Second obligation",
            "comment": "Not needed.",
        },
    )

    assert report.saved, report.as_text()
    expected = [
        line
        for line in original
        if line
        not in ("First obligation of the parties.", "Second obligation of the parties.")
    ]
    assert accepted(out, tmp_path) == expected
    assert rejected(out, tmp_path) == original
    assert structure_problems(out) == []


def test_deleting_the_last_paragraph_keeps_an_empty_one(tmp_path):
    """The final paragraph mark of the body cannot be deleted, so it stays empty."""
    source = docx_with_body(tmp_path, para(run("Keep.")) + para(run("Drop this.")))

    report, out = apply(source, tmp_path, {"op": "delete_paragraphs", "quote": "Drop"})

    assert report.saved
    assert accepted(out, tmp_path) == ["Keep.", ""]
    assert rejected(out, tmp_path) == ["Keep.", "Drop this."]


def test_through_before_the_quote_is_refused(tmp_path):
    """THROUGH_BEFORE_QUOTE when the last paragraph comes first."""
    source = rich_docx(tmp_path)

    report, _ = apply(
        source,
        tmp_path,
        {
            "op": "delete_paragraphs",
            "quote": "Second obligation",
            "through": "First obligation",
        },
    )

    assert codes(report) == ["THROUGH_BEFORE_QUOTE"]


def test_a_paragraph_range_across_a_table_is_refused(tmp_path):
    """RANGE_NOT_CONTIGUOUS when a table lies between the two quotes."""
    source = rich_docx(tmp_path)

    report, _ = apply(
        source,
        tmp_path,
        {
            "op": "delete_paragraphs",
            "quote": "Second obligation",
            "through": "Payment is due",
        },
    )

    assert codes(report) == ["RANGE_NOT_CONTIGUOUS"]


# Refusals


def test_a_quote_found_twice_refuses_the_whole_batch(tmp_path):
    """QUOTE_AMBIGUOUS with its count; the others are skipped and nothing is written."""
    source = docx_with_body(
        tmp_path,
        para(run("Acme pays.")) + para(run("Acme pays.")) + para(run("Beta waits.")),
    )

    report, out = apply(
        source,
        tmp_path,
        {"op": "add_comment", "quote": "Beta waits", "text": "ok"},
        {"op": "replace_text", "quote": "Acme pays", "text": "Beta pays"},
    )

    assert not report.saved and not out.exists()
    assert [o.status for o in report.outcomes] == ["skipped", "refused"]
    assert codes(report) == ["BATCH_REFUSED", "QUOTE_AMBIGUOUS"]
    assert report.outcomes[0].values == {"index": 1}
    assert report.outcomes[1].values == {"quote": "Acme pays", "count": 2}
    assert report.output_sha256 is None


def test_partial_keeps_the_operations_that_apply(tmp_path):
    """With partial, a refused operation leaves the others applied and saved."""
    source = docx_with_body(tmp_path, para(run("Alpha beta.")))

    report, out = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "missing", "text": "x"},
        {"op": "replace_text", "quote": "beta", "text": "gamma"},
        partial=True,
    )

    assert report.saved and out.exists()
    assert codes(report) == ["QUOTE_NOT_FOUND", None]
    assert text_now(out) == ["Alpha gamma."]


@pytest.mark.parametrize(
    ("operation", "code", "values"),
    [
        (
            {"op": "replace_text", "quote": "nowhere", "text": "x"},
            "QUOTE_NOT_FOUND",
            {"quote": "nowhere"},
        ),
        (
            {"op": "replace_text", "quote": "first. Second", "text": "x"},
            "QUOTE_CROSSES_PARAGRAPHS",
            {"quote": "first. Second"},
        ),
        (
            {"op": "replace_text", "quote": "Page 3", "text": "x"},
            "QUOTE_IN_FIELD",
            {"quote": "Page 3"},
        ),
        (
            {"op": "replace_text", "quote": "Ref A", "text": "x"},
            "QUOTE_IN_FIELD",
            {"quote": "Ref A"},
        ),
        ({"op": "replace_text", "text": "x"}, "BAD_OPERATION", {"field": "quote"}),
        ({"op": "replace_text", "quote": "first"}, "BAD_OPERATION", {"field": "text"}),
        (
            {"op": "add_comment", "quote": "first", "text": ""},
            "BAD_OPERATION",
            {"field": "text"},
        ),
        (
            {"op": "add_comment", "quote": "first", "text": "x", "internal": "yes"},
            "BAD_OPERATION",
            {"field": "internal"},
        ),
        ({"quote": "first"}, "BAD_OPERATION", {"field": "op"}),
        (
            {"op": "set_cell", "sheet": "A", "cell": "B2", "value": 1},
            "UNSUPPORTED_OPERATION",
            {"op": "set_cell", "format": "docx"},
        ),
        (
            {"op": "replace_text", "quote": "first", "text": "a\nb"},
            "NEWLINE_IN_TEXT",
            {},
        ),
        (
            {"op": "replace_text", "quote": "first", "text": "a\x0bb"},
            "BAD_OPERATION",
            {"field": "text"},
        ),
        (
            {"op": "insert_paragraphs", "quote": "first", "text": "a\x01"},
            "BAD_OPERATION",
            {"field": "text"},
        ),
        (
            {"op": "add_comment", "quote": "first", "text": "a\x0c"},
            "BAD_OPERATION",
            {"field": "text"},
        ),
        (
            {"op": "replace_text", "quote": "first", "text": "x", "comment": "\x02"},
            "BAD_OPERATION",
            {"field": "comment"},
        ),
        (
            {"op": "replace_text", "quote": "The first", "text": "The  first"},
            "SAME_TEXT",
            {"quote": "The first"},
        ),
    ],
)
def test_each_refusal_has_its_code(tmp_path, operation, code, values):
    """Every refusal names its code and the values the model needs to fix it."""
    field = (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        + run("3")
        + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    source = docx_with_body(
        tmp_path,
        para(run("The first."))
        + para(run("Second one."))
        + para(run("Page "), field)
        + para('<w:fldSimple w:instr=" REF a ">' + run("Ref A") + "</w:fldSimple>"),
    )

    report, out = apply(source, tmp_path, operation)

    assert not report.saved and not out.exists()
    assert report.outcomes[0].code == code
    assert report.outcomes[0].values == values
    assert report.outcomes[0].message


def test_text_outside_a_field_is_editable_and_the_field_survives(tmp_path):
    """A quote next to a field edits around it; the field's runs are untouched."""
    field = (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        + run("3")
        + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    source = docx_with_body(tmp_path, para(run("Page "), field, run(" of the deed")))

    report, out = apply(
        source, tmp_path, {"op": "replace_text", "quote": "deed", "text": "contract"}
    )

    assert report.saved
    assert accepted(out, tmp_path) == ["Page 3 of the contract"]
    assert len(body_of(out).findall(f".//{w('fldChar')}")) == 3


def test_headers_footers_and_untouched_parts_stay_byte_identical(tmp_path):
    """Only the document part changes; headers, footers and styles are copied as they were."""
    source = rich_docx(tmp_path)

    _, out = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "Acme and Beta", "text": "Acme and Gamma"},
    )
    header_only, _ = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "Confidential header", "text": "x"},
    )

    assert header_only.outcomes[0].code == "QUOTE_NOT_FOUND"

    with zipfile.ZipFile(source) as a, zipfile.ZipFile(out) as b:
        for name in a.namelist():
            if name != "word/document.xml":
                assert a.read(name) == b.read(name), name


def test_the_users_file_is_never_written(tmp_path):
    """The original's bytes are the same after any edit."""
    source = rich_docx(tmp_path)
    before = source.read_bytes()

    report, _ = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "Closing words", "text": "Final words"},
        {"op": "add_comment", "quote": "Final words", "text": "Fine."},
    )

    assert report.saved
    assert source.read_bytes() == before
    assert report.input_sha256 == hashlib.sha256(before).hexdigest()
    assert (
        report.output_sha256
        == hashlib.sha256((tmp_path / "out.docx").read_bytes()).hexdigest()
    )


def test_the_report_reads_for_the_model_and_stores_as_metadata(tmp_path):
    """as_text lists each operation; as_metadata has the schema the version stores."""
    source = docx_with_body(tmp_path, para(run("Alpha beta.")))

    report, _ = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "beta", "text": "gamma"},
        {"op": "add_comment", "quote": "Alpha", "text": "Hm"},
    )

    assert report.as_text().splitlines() == [
        "#0 replace_text: applied.",
        "#1 add_comment: applied.",
    ]
    metadata = report.as_metadata()
    assert metadata["schema"] == "revise-report/1"
    assert metadata["applied"] == 2 and metadata["refused"] == 0 and metadata["saved"]
    assert (
        metadata["ops"][1]["values"]["comment_id"]
        == report.outcomes[1].values["comment_id"]
    )


def test_no_operations_is_nothing_changed(tmp_path):
    """Zero applied operations save nothing and say so."""
    source = docx_with_body(tmp_path, para(run("Alpha.")))

    report, out = apply(source, tmp_path)

    assert not report.saved and not out.exists()
    assert [n.code for n in report.must_tell_user] == ["NOTHING_CHANGED"]


def test_counts_tally_body_changes_and_comments(tmp_path):
    """counts reports the body's revisions and the comment total."""
    source = _reviewed(tmp_path)

    assert word.counts(source) == word.RevisionCounts(changes=2, comments=1)
    _, out = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "England", "text": "Wales", "comment": "x"},
    )
    assert word.counts(out) == word.RevisionCounts(changes=4, comments=2)


def test_a_document_with_a_doctype_is_refused(tmp_path):
    """No entity is ever expanded from a customer's file."""
    source = docx_with_body(tmp_path, para(run("Alpha.")))
    with zipfile.ZipFile(source) as archive:
        parts = {i.filename: archive.read(i) for i in archive.infolist()}
    parts["word/document.xml"] = parts["word/document.xml"].replace(
        b"<w:document", b'<!DOCTYPE x [<!ENTITY e "boom">]><w:document', 1
    )
    evil = tmp_path / "evil.docx"
    with zipfile.ZipFile(evil, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    with pytest.raises(PackageRefusedError) as refused:
        word.apply(
            evil,
            [{"op": "add_comment", "quote": "Alpha", "text": "x"}],
            tmp_path / "o.docx",
        )

    assert refused.value.code == "UNSAFE_XML"


def test_editing_inside_another_authors_insertion_splits_it(tmp_path):
    """New words land between the halves of Alice's insertion; her text stays hers."""
    source = docx_with_body(
        tmp_path,
        para(
            run("The ball is "),
            inserted(5, "Alice", run("big red round")),
            run(" today."),
        ),
    )

    report, out = apply(
        source,
        tmp_path,
        {
            "op": "replace_text",
            "quote": "big red round",
            "text": "big blue round",
            "comment": "Colour",
        },
    )

    assert report.saved
    assert text_now(out) == ["The ball is big blue round today."]
    assert anchored_text(out, report.outcomes[0].values["comment_id"]) == "blue"
    assert revision_authors(out) == {"Alice": 3, "SurfSense": 2}
    assert structure_problems(out) == []
    assert accepted(out, tmp_path) == ["The ball is big blue round today."]
    assert rejected(out, tmp_path) == ["The ball is  today."]


def test_every_kind_of_edit_together_accepts_to_the_intent_and_rejects_to_the_original(
    tmp_path,
):
    """A batch of all four operations: accept gives the intended text, reject the original."""
    source = rich_docx(tmp_path)
    original = plain_text(source)

    report, out = apply(
        source,
        tmp_path,
        {
            "op": "replace_text",
            "quote": "between Acme and Beta",
            "text": "between Acme Inc. and Beta LLC",
        },
        {
            "op": "insert_paragraphs",
            "quote": "between Acme",
            "text": "Recitals follow.\n\nWhereas both agree.",
        },
        {"op": "delete_paragraphs", "quote": "Second obligation"},
        {"op": "replace_text", "quote": "Twelve months", "text": "Twenty-four months"},
        {
            "op": "add_comment",
            "quote": "Closing words",
            "text": "Add signature block.",
            "internal": True,
        },
        {"op": "replace_text", "quote": "Payment is due", "text": "Payment falls due"},
    )

    assert report.saved and report.applied == 6, report.as_text()
    assert structure_problems(out) == []
    assert accepted(out, tmp_path) == [
        "Master Services Agreement",
        "This agreement is between Acme Inc. and Beta LLC.",
        "Recitals follow.",
        "",
        "Whereas both agree.",
        "First obligation of the parties.",
        "Fee",
        "100 dollars",
        "Term",
        "Twenty-four months",
        "Payment falls due within 30 days of invoice.",
        "Closing words.",
    ]
    assert rejected(out, tmp_path) == original
    counts = word.counts(out)
    assert counts.comments == 1 and counts.changes > 0


def test_moves_and_row_insertions_by_others_are_decided_too(tmp_path):
    """Accept and reject cover moved text and tracked table rows, not only w:ins and w:del."""
    move = '<w:{kind} w:id="{id}" w:author="Bob" w:date="2026-01-01T00:00:00Z">{body}</w:{kind}>'
    row = (
        '<w:tr><w:trPr><w:ins w:id="40" w:author="Bob" w:date="2026-01-01T00:00:00Z"/></w:trPr>'
        "<w:tc>" + para(inserted(41, "Bob", run("New row"))) + "</w:tc></w:tr>"
    )
    source = docx_with_body(
        tmp_path,
        para(
            run("Start "),
            move.format(
                kind="moveFrom", id=30, body="<w:r><w:delText>moved</w:delText></w:r>"
            ),
        )
        + para(run("End "), move.format(kind="moveTo", id=31, body=run("moved")))
        + "<w:tbl><w:tr><w:tc>"
        + para(run("Old row"))
        + "</w:tc></w:tr>"
        + row
        + "</w:tbl>"
        + para(run("After.")),
    )

    assert accepted(source, tmp_path) == [
        "Start ",
        "End moved",
        "Old row",
        "New row",
        "After.",
    ]
    assert rejected(source, tmp_path) == ["Start moved", "End ", "Old row", "After."]


def test_words_added_before_a_quote_take_the_quotes_format(tmp_path):
    """An insertion at the quote's start joins the quote's first run, not the text before."""
    source = docx_with_body(
        tmp_path, para(run("Paid by "), run("Acme", bold=True), run("."))
    )

    report, out = apply(
        source, tmp_path, {"op": "replace_text", "quote": "Acme", "text": "the Acme"}
    )

    assert report.saved
    assert changes(out, "del") == [] and changes(out, "ins") == ["the "]
    insertion = next(body_of(out).iter(w("ins")))
    assert insertion.find(f"{w('r')}/{w('rPr')}/{w('b')}") is not None
    assert accepted(out, tmp_path) == ["Paid by the Acme."]


def test_a_refused_batch_reads_as_one_line_per_operation(tmp_path):
    """as_text names the refusal with its code and each skipped operation."""
    source = docx_with_body(tmp_path, para(run("Alpha beta.")))

    report, _ = apply(
        source,
        tmp_path,
        {"op": "replace_text", "quote": "beta", "text": "gamma"},
        {"op": "add_comment", "quote": "delta", "text": "x"},
    )

    lines = report.as_text().splitlines()
    assert lines[0].startswith("#0 replace_text skipped (BATCH_REFUSED): ")
    assert lines[1].startswith("#1 add_comment refused (QUOTE_NOT_FOUND): ")
    assert report.as_metadata()["ops"][1]["values"] == {"quote": "delta"}


def test_paragraphs_added_after_the_last_one_go_away_on_reject(tmp_path):
    """Rejecting an insertion at the end of the body leaves no empty paragraph behind."""
    source = docx_with_body(tmp_path, para(run("Only paragraph.")))

    report, out = apply(
        source,
        tmp_path,
        {"op": "insert_paragraphs", "quote": "Only", "text": "Added one.\nAdded two."},
    )

    assert report.saved
    assert accepted(out, tmp_path) == ["Only paragraph.", "Added one.", "Added two."]
    assert rejected(out, tmp_path) == ["Only paragraph."]


def test_deleting_a_paragraph_with_a_text_box_leaves_the_box_text_alone(tmp_path):
    """A text box is its own story: its text is not turned into deleted text."""
    box = (
        '<w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox>'
        "<w:txbxContent>" + para(run("Boxed words")) + "</w:txbxContent>"
        "</v:textbox></v:shape></w:pict></w:r>"
    )
    source = docx_with_body(
        tmp_path, para(run("Lead "), box, run(" tail.")) + para(run("Next."))
    )

    report, out = apply(source, tmp_path, {"op": "delete_paragraphs", "quote": "Lead"})

    assert report.saved, report.as_text()
    assert structure_problems(out) == []
    assert rejected(out, tmp_path) == ["Lead  tail.", "Next."]
