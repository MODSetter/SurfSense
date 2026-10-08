import hashlib
from pathlib import Path

import openpyxl
import pytest
from lxml import etree

from modules.artifacts.revised_copies.engines import excel
from tests.unit.artifacts.revised_copies.office_fixtures import (
    CT,
    PARSER,
    PR,
    S,
    build_openpyxl_workbook,
    build_workbook,
    check_workbook,
    parts,
)

pytestmark = pytest.mark.unit

UNTOUCHED = (
    "xl/charts/chart1.xml",
    "xl/drawings/drawing1.xml",
    "xl/drawings/_rels/drawing1.xml.rels",
    "xl/media/image1.png",
    "xl/sharedStrings.xml",
    "xl/styles.xml",
    "xl/worksheets/sheet2.xml",
    "xl/tables/table1.xml",
    "xl/theme/theme1.xml",
)


@pytest.fixture
def book(tmp_path: Path) -> Path:
    """A real workbook with every part an edit must leave alone."""
    return build_workbook(tmp_path / "Sales.xlsx", image=tmp_path / "logo.png")


def run(book: Path, tmp_path: Path, *operations: dict):
    """Applies the operations to a copy next to the test's files."""
    out = tmp_path / "out.xlsx"
    return excel.apply(book, list(operations), out), out


def cell(path: Path, sheet: str, ref: str) -> etree._Element | None:
    """The cell element as written, or None."""
    index = {"Data": 1, "Table": 2}[sheet]
    root = etree.fromstring(parts(path)[f"xl/worksheets/sheet{index}.xml"], PARSER)
    return next((c for c in root.iter(f"{{{S}}}c") if c.get("r") == ref), None)


def codes(report) -> set[str]:
    """Every code in the report, from outcomes and notices."""
    return {o.code for o in report.outcomes if o.code} | {
        n.code for n in report.must_tell_user
    }


def test_a_value_edit_writes_only_the_sheet_and_leaves_charts_shapes_and_media_byte_identical(
    book, tmp_path
):
    """openpyxl's save drops shapes; editing the sheet XML must leave every other part's bytes and the user's file alone."""
    before = parts(book)
    digest = hashlib.sha256(book.read_bytes()).hexdigest()

    report, out = run(
        book, tmp_path, {"op": "set_cell", "sheet": "Data", "cell": "B3", "value": 25}
    )

    assert report.saved and report.applied == 1
    after = parts(out)
    for name in UNTOUCHED:
        assert after[name] == before[name], name
    assert list(after) == [n for n in before if n in after]
    assert hashlib.sha256(book.read_bytes()).hexdigest() == digest
    check_workbook(out)
    assert openpyxl.load_workbook(out)["Data"]["B3"].value == 25


def test_a_string_goes_in_inline_so_shared_strings_stay_as_they_were_and_the_cell_keeps_its_style(
    book, tmp_path
):
    """Inline strings keep the shared-string table consistent without rewriting it."""
    styled = cell(book, "Data", "A2")
    report, out = run(
        book,
        tmp_path,
        {"op": "set_cell", "sheet": "Data", "cell": "A2", "value": "Northwest"},
    )

    assert report.saved
    written = cell(out, "Data", "A2")
    assert written.get("t") == "inlineStr"
    assert written.get("s") == styled.get("s")
    assert parts(out)["xl/sharedStrings.xml"] == parts(book)["xl/sharedStrings.xml"]
    check_workbook(out)
    assert openpyxl.load_workbook(out)["Data"]["A2"].value == "Northwest"


def test_a_written_formula_has_no_cached_value_and_the_workbook_recalculates_on_open(
    book, tmp_path
):
    """A stale cached value would show a wrong total until Excel recalculates."""
    report, out = run(
        book,
        tmp_path,
        {"op": "set_cell", "sheet": "Data", "cell": "B8", "formula": "=AVERAGE(B2:B6)"},
    )

    assert report.saved
    written = cell(out, "Data", "B8")
    assert written.findtext(f"{{{S}}}f") == "AVERAGE(B2:B6)"
    assert written.find(f"{{{S}}}v") is None
    calc = etree.fromstring(parts(out)["xl/workbook.xml"], PARSER).find(
        f"{{{S}}}calcPr"
    )
    assert calc.get("fullCalcOnLoad") == "1"
    notice = next(
        n for n in report.must_tell_user if n.code == "FORMULAS_NOT_RECALCULATED"
    )
    assert notice.values["count"] >= 1
    check_workbook(out)
    assert openpyxl.load_workbook(out)["Data"]["B8"].value == "=AVERAGE(B2:B6)"


def test_removing_a_formula_removes_the_calc_chain_with_its_relationship_and_override(
    book, tmp_path
):
    """A calc chain naming a cell without a formula makes Excel repair the file."""
    report, out = run(
        book, tmp_path, {"op": "set_cell", "sheet": "Data", "cell": "B7", "value": 0}
    )

    assert report.saved
    after = parts(out)
    assert "xl/calcChain.xml" not in after
    assert b"calcChain" not in after["xl/_rels/workbook.xml.rels"]
    assert b"calcChain" not in after["[Content_Types].xml"]
    check_workbook(out)


def test_replacing_a_value_keeps_the_calc_chain(book, tmp_path):
    """The calc chain only goes when the set of formula cells changes."""
    report, out = run(
        book, tmp_path, {"op": "set_cell", "sheet": "Data", "cell": "B2", "value": 11}
    )

    assert report.saved
    assert parts(out)["xl/calcChain.xml"] == parts(book)["xl/calcChain.xml"]


def test_missing_rows_and_cells_are_inserted_in_order(book, tmp_path):
    """Excel rejects rows or cells out of order; the dimension grows to cover them."""
    report, out = run(
        book,
        tmp_path,
        {"op": "set_cell", "sheet": "Data", "cell": "D3", "value": "between"},
        {"op": "set_cell", "sheet": "Data", "cell": "A12", "value": "below"},
        {"op": "set_cell", "sheet": "Data", "cell": "AA1", "value": True},
    )

    assert report.saved and report.applied == 3
    check_workbook(out)
    sheet = openpyxl.load_workbook(out)["Data"]
    assert (sheet["D3"].value, sheet["A12"].value, sheet["AA1"].value) == (
        "between",
        "below",
        True,
    )
    root = etree.fromstring(parts(out)["xl/worksheets/sheet1.xml"], PARSER)
    assert root.find(f"{{{S}}}dimension").get("ref") == "A1:AA12"


def test_a_range_is_written_row_by_row_and_null_clears(book, tmp_path):
    """A range's rows map onto the grid in order, and null empties a cell."""
    report, out = run(
        book,
        tmp_path,
        {
            "op": "set_range",
            "sheet": "Data",
            "range": "A2:B3",
            "values": [["N", 1.5], [None, False]],
        },
    )

    assert report.saved
    check_workbook(out)
    sheet = openpyxl.load_workbook(out)["Data"]
    assert [[c.value for c in row] for row in sheet["A2:B3"]] == [
        ["N", 1.5],
        [None, False],
    ]


def test_text_xml_cannot_carry_is_escaped_as_excel_reads_it(book, tmp_path):
    """Control characters and literal _xHHHH_ survive as ST_Xstring escapes, and the engine reads its own escapes back."""
    text = "  tab	here and _x0041_ literal "
    report, out = run(
        book, tmp_path, {"op": "set_cell", "sheet": "Data", "cell": "D2", "value": text}
    )

    assert report.saved
    check_workbook(out)
    written = cell(out, "Data", "D2").find(f"{{{S}}}is/{{{S}}}t")
    assert written.text == "  tab	here_x0001_ and _x005F_x0041_ literal "
    assert written.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
    again = excel.apply(
        out,
        [{"op": "set_cell", "sheet": "Data", "cell": "D2", "value": text}],
        tmp_path / "again.xlsx",
    )
    assert "NOTHING_CHANGED" in codes(again)


def test_a_shared_formula_dependent_can_be_overwritten(book, tmp_path):
    """Only the shared formula's master is protected; its dependents keep their formulas."""
    report, out = run(
        book, tmp_path, {"op": "set_cell", "sheet": "Data", "cell": "C4", "value": 7}
    )

    assert report.saved
    check_workbook(out)
    assert openpyxl.load_workbook(out)["Data"]["C4"].value == 7
    assert openpyxl.load_workbook(out)["Data"]["C5"].value == "=B5*2"


@pytest.mark.parametrize(
    ("operation", "code", "values"),
    [
        (
            {"op": "set_cell", "sheet": "Summary", "cell": "A1", "value": 1},
            "SHEET_NOT_FOUND",
            {"sheet": "Summary", "sheets": "Data, Table"},
        ),
        (
            {"op": "set_cell", "sheet": "Data", "cell": "B0", "value": 1},
            "BAD_CELL",
            {"cell": "B0"},
        ),
        (
            {"op": "set_cell", "sheet": "Data", "cell": "XFE1", "value": 1},
            "BAD_CELL",
            {"cell": "XFE1"},
        ),
        (
            {"op": "set_range", "sheet": "Data", "range": "C4:A2", "values": [[1]]},
            "BAD_RANGE",
            {"range": "C4:A2"},
        ),
        (
            {"op": "set_cell", "sheet": "Data", "cell": "F1", "value": "x"},
            "CELL_IN_MERGE",
            {"cell": "F1", "merge": "E1:F1"},
        ),
        (
            {"op": "set_cell", "sheet": "Data", "cell": "H3", "value": 1},
            "CELL_IN_ARRAY_FORMULA",
            {"cell": "H3"},
        ),
        (
            {"op": "set_cell", "sheet": "Data", "cell": "C2", "value": 1},
            "SHARED_FORMULA_MASTER",
            {"cell": "C2"},
        ),
        (
            {"op": "set_range", "sheet": "Data", "range": "A2:B3", "values": [[1, 2]]},
            "RANGE_SHAPE_MISMATCH",
            {"rows": 2, "columns": 2},
        ),
        (
            {"op": "set_cell", "sheet": "Table", "cell": "B1", "value": "Price"},
            "CELL_IN_TABLE_HEADER",
            {"cell": "B1", "table": "Table1"},
        ),
    ],
)
def test_refusals_name_the_problem_and_save_nothing(
    book, tmp_path, operation, code, values
):
    """One refused operation saves nothing, and the refusal carries a code and the values to act on."""
    report, out = run(
        book,
        tmp_path,
        {"op": "set_cell", "sheet": "Data", "cell": "B2", "value": 3},
        operation,
    )

    assert not report.saved
    assert not out.exists()
    refused = report.outcomes[1]
    assert (refused.status, refused.code, refused.values) == ("refused", code, values)
    assert report.outcomes[0].status == "skipped"
    assert report.outcomes[0].code == "BATCH_REFUSED"


@pytest.mark.parametrize(
    ("operation", "field"),
    [
        ({"op": "set_cell", "sheet": "Data", "cell": "B2"}, "value"),
        (
            {
                "op": "set_cell",
                "sheet": "Data",
                "cell": "B2",
                "value": 1,
                "formula": "=1",
            },
            "formula",
        ),
        ({"op": "set_cell", "cell": "B2", "value": 1}, "sheet"),
        ({"op": "set_cell", "sheet": "Data", "value": 1}, "cell"),
        ({"op": "set_cell", "sheet": "Data", "cell": "B2", "value": [1]}, "value"),
        (
            {"op": "set_range", "sheet": "Data", "range": "A1:A2", "values": "1,2"},
            "values",
        ),
        ({"op": "set_cell", "sheet": "Data", "cell": "B2", "formula": "="}, "formula"),
    ],
)
def test_malformed_operations_name_the_field(book, tmp_path, operation, field):
    """The model learns which field to fix."""
    report, _ = run(book, tmp_path, operation)

    assert report.outcomes[0].code == "BAD_OPERATION"
    assert report.outcomes[0].values == {"field": field}


def test_a_word_operation_is_unsupported_for_a_workbook(book, tmp_path):
    """Operations of another format are refused, not ignored."""
    report, _ = run(
        book, tmp_path, {"op": "replace_text", "quote": "North", "text": "South"}
    )

    assert report.outcomes[0].code == "UNSUPPORTED_OPERATION"
    assert report.outcomes[0].values == {"op": "replace_text", "format": "xlsx"}


def test_writing_what_is_already_there_changes_nothing(book, tmp_path):
    """No copy is saved when no cell would change."""
    report, out = run(
        book, tmp_path, {"op": "set_cell", "sheet": "Data", "cell": "B2", "value": 10}
    )

    assert not report.saved
    assert "NOTHING_CHANGED" in codes(report)
    assert not out.exists()


def test_an_openpyxl_workbook_keeps_its_defined_names_and_formats(tmp_path):
    """Workbooks from another writer, with gaps in their rows, edit as cleanly."""
    book = build_openpyxl_workbook(tmp_path / "Budget.xlsx")

    report, out = run(
        book,
        tmp_path,
        {
            "op": "set_range",
            "sheet": "Budget",
            "range": "A3:B3",
            "values": [["Power", 90]],
        },
    )

    assert report.saved
    check_workbook(out)
    loaded = openpyxl.load_workbook(out)
    assert loaded["Budget"]["B3"].value == 90
    assert loaded["Budget"]["B2"].number_format == "#,##0.00"
    assert "Rent" in loaded.defined_names
    assert parts(out)["xl/styles.xml"] == parts(book)["xl/styles.xml"]


def test_the_report_and_relationships_survive_a_macro_free_round_trip(book, tmp_path):
    """The report's hashes name the files it read and wrote."""
    report, out = run(
        book, tmp_path, {"op": "set_cell", "sheet": "Data", "cell": "B5", "value": 41}
    )

    assert report.input_sha256 == hashlib.sha256(book.read_bytes()).hexdigest()
    assert report.output_sha256 == hashlib.sha256(out.read_bytes()).hexdigest()
    types = etree.fromstring(parts(out)["[Content_Types].xml"], PARSER)
    assert types.tag == f"{{{CT}}}Types"
    rels = etree.fromstring(parts(out)["xl/worksheets/_rels/sheet1.xml.rels"], PARSER)
    assert all(r.tag == f"{{{PR}}}Relationship" for r in rels)


def test_a_macro_workbook_keeps_its_macros_byte_for_byte(tmp_path):
    """An .xlsm keeps its VBA project and its macro-enabled content type."""
    import xlsxwriter

    macros = tmp_path / "vbaProject.bin"
    macros.write_bytes(bytes(range(256)) * 8)
    path = tmp_path / "Macros.xlsm"
    book = xlsxwriter.Workbook(str(path))
    book.add_worksheet("Sheet1").write("A1", 1)
    book.add_vba_project(str(macros))
    book.close()

    report, out = run(
        path, tmp_path, {"op": "set_cell", "sheet": "Sheet1", "cell": "A1", "value": 2}
    )

    assert report.saved
    assert parts(out)["xl/vbaProject.bin"] == parts(path)["xl/vbaProject.bin"]
    assert b"macroEnabled" in parts(out)["[Content_Types].xml"]
    check_workbook(out)


def test_a_file_that_is_not_a_workbook_is_refused(tmp_path):
    """A deck named .xlsx is refused before any operation runs."""
    from pptx import Presentation

    from modules.artifacts.revised_copies.engines.package import PackageRefusedError

    path = tmp_path / "deck.xlsx"
    Presentation().save(path)

    with pytest.raises(PackageRefusedError) as refusal:
        run(
            path,
            tmp_path,
            {"op": "set_cell", "sheet": "Sheet1", "cell": "A1", "value": 2},
        )
    assert refusal.value.code == "NOT_A_WORKBOOK"


def test_sheet_names_match_regardless_of_case_as_in_excel(book, tmp_path):
    """Excel will not hold two sheets differing only in case, so "data" can only mean Data."""
    report, out = run(
        book, tmp_path, {"op": "set_cell", "sheet": "data", "cell": "B3", "value": 21}
    )

    assert report.saved
    assert openpyxl.load_workbook(out)["Data"]["B3"].value == 21
