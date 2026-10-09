"""What a live case reads from a Word file the agent made, checked on files made here."""

import hashlib
import io

import docx
import pytest
from PIL import Image

from tests.live.word_file import WordFile

pytestmark = pytest.mark.unit


def _png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (40, 20), "teal").save(out, format="PNG")
    return out.getvalue()


def _saved(document: docx.Document) -> bytes:
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_it_reads_headings_tables_and_pictures_in_document_order() -> None:
    """A restructure is checked by where the blocks sit, so order is kept."""
    made = docx.Document()
    made.add_heading("Memo", level=0)
    made.add_paragraph("To: the board")
    table = made.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Option"
    table.rows[1].cells[0].text = "Keep"
    made.add_heading("1. Background", level=1)
    made.add_paragraph("We run three depots.")
    made.add_heading("Detail", level=2)
    made.add_picture(io.BytesIO(_png()))

    word = WordFile(_saved(made))

    assert [b.kind for b in word.blocks] == [
        "title",
        "paragraph",
        "table",
        "heading",
        "paragraph",
        "heading",
        "paragraph",
    ]
    assert [(h.level, h.text) for h in word.headings] == [
        (1, "1. Background"),
        (2, "Detail"),
    ]
    assert word.tables == [[["Option", ""], ["Keep", ""]]]
    assert word.pictures == {hashlib.sha256(_png()).hexdigest()}


def test_a_heading_is_numbered_by_its_text_or_by_word_numbering() -> None:
    """The agent may type the numbers or use Word's numbering; both count."""
    made = docx.Document()
    made.add_heading("2. Options", level=1)
    made.add_heading("Recommendation", level=1)
    listed = made.add_heading("Costs", level=1)
    properties = listed._p.get_or_add_pPr()
    numbering = properties._add_numPr()
    numbering.get_or_add_numId().val = 1
    numbering.get_or_add_ilvl().val = 0

    word = WordFile(_saved(made))

    assert [h.numbered for h in word.headings] == [True, False, True]
