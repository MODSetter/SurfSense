"""Live case 5: a Word report charted from a spreadsheet's yearly numbers, then the chart changed and a paragraph added.

The chart's numbers exist only in the workbook, which ingest reads as a table.
"""

import io
import re

import openpyxl
import pytest

from tests.live.live_agent import LiveAgent
from tests.live.turn_renders import (
    assert_next_version,
    assert_pages_checked,
    last_version,
)
from tests.live.word_file import WordFile

pytestmark = pytest.mark.live

CASE = "spreadsheet-report"
TURNS = (
    "Write a short Word report on Nordlys Energy's growth from the spreadsheet, "
    "with a bar chart of yearly revenue.",
    "Switch the chart to a line chart and add a short trend paragraph.",
)
# Year: members, revenue in NOK thousand, energy sold in MWh.
YEARS = {
    2019: (1_840, 3_120, 9_450),
    2020: (2_015, 3_385, 10_120),
    2021: (2_390, 3_910, 11_870),
    2022: (2_960, 4_870, 13_340),
    2023: (3_310, 5_240, 14_010),
    2024: (3_585, 5_615, 14_960),
    2025: (3_870, 6_080, 15_720),
}
REVENUE = [revenue for _, revenue, _ in YEARS.values()]


async def test_the_agent_charts_the_workbook_then_turns_the_chart_into_a_line(
    live: LiveAgent,
) -> None:
    """v1 bars the workbook's revenue; v2 draws it as a line, as a new picture, beside a new paragraph."""
    workbook = await live.upload(
        "Nordlys members and revenue 2019-2025.xlsx", _workbook()
    )
    await live.wait_ready(workbook)
    thread = await live.thread()

    first = await live.turn(thread, TURNS[0])
    draft = await last_version(live, first, "turn 1")
    assert draft.format == "docx", f"turn 1 made {draft}, not a Word file"
    await assert_pages_checked(live, first, "turn 1")
    draft_script = live.spec(draft.id).get("text", "")
    assert _keeps_revenue(draft_script), (
        f"v{draft.number}'s script does not chart the workbook's revenue {REVENUE}"
    )
    assert re.search(r"\.barh?\(", draft_script), f"v{draft.number} draws no bars"
    before = WordFile(await live.file(draft.id))
    assert before.pictures, f"v{draft.number} holds no chart"

    second = await live.turn(thread, TURNS[1])
    edit = await last_version(live, second, "turn 2")
    assert_next_version(draft, edit, "turn 2")
    await assert_pages_checked(live, second, "turn 2")
    edit_script = live.spec(edit.id).get("text", "")
    assert re.search(r"\.plot\(", edit_script), f"v{edit.number} draws no line"
    assert not re.search(r"\.barh?\(", edit_script), f"v{edit.number} still has bars"
    assert _keeps_revenue(edit_script), (
        f"v{edit.number}'s line does not keep the revenue {REVENUE}"
    )
    after = WordFile(await live.file(edit.id))
    assert after.pictures - before.pictures, f"v{edit.number} kept the bar chart"
    added = [
        p
        for p in after.paragraphs
        if p not in before.paragraphs and len(p.split()) >= 15
    ]
    assert added, f"v{edit.number} adds no trend paragraph"


def _keeps_revenue(script: str) -> bool:
    """Every year's revenue is in the script, in NOK thousand or in NOK million."""
    flat = re.sub(r"(?<=\d)[_,](?=\d{3})", "", script)
    written = [float(n) for n in re.findall(r"\d+(?:\.\d+)?", flat)]
    return all(
        any(abs(n - r) < 0.5 or abs(n - r / 1000) < 0.006 for n in written)
        for r in REVENUE
    )


def _workbook() -> bytes:
    """The co-operative's yearly figures, one row per year, as the board keeps them."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Yearly figures"
    sheet.append(["Year", "Members", "Revenue (NOK thousand)", "Energy sold (MWh)"])
    for year, row in YEARS.items():
        sheet.append([year, *row])
    notes = book.create_sheet("Notes")
    notes.append(["Nordlys Energy is a solar co-operative in Trondelag, founded 2017."])
    notes.append(["2022: the Melhus roof array (2.1 MW) came online in March."])
    notes.append(["2024: membership fee rose from 450 to 600 NOK a year."])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
