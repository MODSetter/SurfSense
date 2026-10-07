"""Live case 6: a three-section Word memo, then restructured: two sections merged, a summary table on top, headings numbered.

An edit that moves and joins blocks, where the demo's edit only adds them.
"""

import io

import docx
import pytest

from tests.live.live_agent import LiveAgent
from tests.live.turn_renders import (
    assert_next_version,
    assert_pages_checked,
    last_version,
)
from tests.live.word_file import WordFile

pytestmark = pytest.mark.live

CASE = "memo-restructure"
TURNS = (
    "Write an internal memo as a Word document from these notes, with three "
    "sections: Background, Options and Recommendation.",
    "Merge sections 2 and 3, add a summary table at the top and number the headings.",
)


async def test_the_agent_restructures_a_memo_it_wrote(live: LiveAgent) -> None:
    """v1 has three sections; v2, its next version, has a table above two numbered sections."""
    notes = await live.upload("WMS selection notes.docx", _selection_notes())
    await live.wait_ready(notes)
    thread = await live.thread()

    first = await live.turn(thread, TURNS[0])
    draft = await last_version(live, first, "turn 1")
    assert draft.format == "docx", f"turn 1 made {draft}, not a Word file"
    await assert_pages_checked(live, first, "turn 1")
    memo = WordFile(await live.file(draft.id))
    assert len(_sections(memo)) == 3, (
        f"v{draft.number} sections: {[h.text for h in _sections(memo)]}"
    )

    second = await live.turn(thread, TURNS[1])
    edit = await last_version(live, second, "turn 2")
    assert_next_version(draft, edit, "turn 2")
    await assert_pages_checked(live, second, "turn 2")
    memo = WordFile(await live.file(edit.id))
    kinds = [b.kind for b in memo.blocks]
    assert "table" in kinds, f"v{edit.number} has no table"
    top_table = kinds.index("table")
    sections = _sections(memo)
    below = [h for h in sections if memo.blocks.index(h) > top_table]
    assert len(below) == 2, (
        f"v{edit.number}: the sections under the table are "
        f"{[h.text for h in below]}, not two"
    )
    assert "background" in below[0].text.lower(), (
        f"v{edit.number}: the summary table is not above the first section: "
        f"{[(b.kind, b.text[:40]) for b in memo.blocks]}"
    )
    assert all(h.numbered for h in sections), (
        f"v{edit.number} headings not all numbered: {[h.text for h in sections]}"
    )


def _sections(memo: WordFile):
    """The headings at the Background section's level: a memo may head its title as a heading too."""
    levels = [h.level for h in memo.headings if "background" in h.text.lower()]
    return [h for h in memo.headings if levels and h.level == levels[0]]


def _selection_notes() -> bytes:
    """Notes from choosing a warehouse management system: the situation, three options and our pick."""
    notes = docx.Document()
    notes.add_heading("WMS selection: working notes, September 2026", 1)
    notes.add_paragraph(
        "Our two warehouses (Vantaa 14,000 m2, Turku 6,500 m2) run Stockpilot 4, "
        "installed 2014. The vendor ends support on 30 June 2027. Pick errors are "
        "at 1.8% of order lines and a stock count takes both sites offline for two "
        "days each quarter. Peak season (November-December) is when it hurts most."
    )
    notes.add_heading("Options we looked at", 2)
    for point in (
        "A. Keep Stockpilot and pay for extended support: 38,000 EUR a year, no "
        "new features, the hardware still ages.",
        "B. Lagerline Cloud WMS: 64,000 EUR to implement, 2,900 EUR a month, "
        "scanner app included, live in about five months.",
        "C. Build our own on top of the ERP: about 210,000 EUR and a year of two "
        "developers, nothing live before peak 2027.",
    ):
        notes.add_paragraph(point, style="List Bullet")
    notes.add_heading("Where we landed", 2)
    notes.add_paragraph(
        "Option B. It is live before support ends, cuts pick errors (Lagerline's "
        "reference customers report under 0.5%), and allows cycle counts without "
        "closing a site. Start in January 2027 so the cut-over misses peak. Needs "
        "sign-off from the CFO, Elina Koski, by 15 November."
    )
    out = io.BytesIO()
    notes.save(out)
    return out.getvalue()
