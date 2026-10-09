"""What a live case reads from a revised copy, checked on files made here."""

import io
import zipfile

import docx
import openpyxl
import pytest

from tests.live.revised_files import RedlinedWord, cell_number

pytestmark = pytest.mark.unit

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
STAMP = 'w:author="{author}" w:date="2026-10-06T10:00:00Z"'


def _word(body: str, comments: str = "") -> bytes:
    """python-docx's template with this body and, when given, a comments part."""
    out = io.BytesIO()
    docx.Document().save(out)
    with zipfile.ZipFile(io.BytesIO(out.getvalue())) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    parts["word/document.xml"] = (
        f'<w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>'
    ).encode()
    if comments:
        parts["word/comments.xml"] = (
            f'<w:comments xmlns:w="{W}">{comments}</w:comments>'
        ).encode()
    made = io.BytesIO()
    with zipfile.ZipFile(made, "w") as package:
        for name, data in parts.items():
            package.writestr(name, data)
    return made.getvalue()


def test_it_reads_each_authors_insertions_deletions_and_comments() -> None:
    """A redline is checked by who made each change and what it says."""
    ours = STAMP.format(author="SurfSense")
    theirs = STAMP.format(author="Counsel")
    body = (
        "<w:p><w:r><w:t>Payable within </w:t></w:r>"
        f'<w:del w:id="1" {ours}><w:r><w:delText>60</w:delText></w:r></w:del>'
        f'<w:ins w:id="2" {ours}><w:r><w:t>30</w:t></w:r></w:ins>'
        "<w:r><w:t> days.</w:t></w:r></w:p>"
        f'<w:p><w:ins w:id="3" {theirs}><w:r><w:t>Added by counsel.</w:t></w:r></w:ins></w:p>'
        "<w:p/>"
    )
    comments = (
        f'<w:comment w:id="0" {ours}><w:p><w:r><w:t>Shorter terms </w:t></w:r>'
        "<w:r><w:t>help cash flow.</w:t></w:r></w:p></w:comment>"
    )

    word = RedlinedWord(_word(body, comments))

    assert [(c.author, c.text) for c in word.insertions] == [
        ("SurfSense", "30"),
        ("Counsel", "Added by counsel."),
    ]
    assert [(c.author, c.text) for c in word.deletions] == [("SurfSense", "60")]
    assert [(c.author, c.text) for c in word.comments] == [
        ("SurfSense", "Shorter terms help cash flow.")
    ]
    assert word.paragraphs == ["Payable within 30 days.", "Added by counsel."]


def test_a_file_with_no_changes_or_comments_reads_as_its_paragraphs() -> None:
    """A clean or rejected copy has no markup left to read."""
    word = RedlinedWord(_word("<w:p><w:r><w:t>Plain.</w:t></w:r></w:p>"))

    assert (word.insertions, word.deletions, word.comments) == ([], [], [])
    assert word.paragraphs == ["Plain."]


def test_a_cell_reads_as_its_number_whether_written_as_a_value_or_a_formula() -> None:
    """The agent may raise a line as 52800, =48000*1.1 or a sum over cells."""
    book = openpyxl.Workbook()
    sheet = book.active
    for row, value in enumerate((100, 200, "=48000*1.1", "=SUM(B1:B3)+B1"), start=1):
        sheet[f"B{row}"] = value
    sheet["C1"] = "=b1 + B2 * 2"

    assert cell_number(sheet, "B1") == 100
    assert cell_number(sheet, "B3") == pytest.approx(52800)
    assert cell_number(sheet, "B4") == pytest.approx(53200)
    assert cell_number(sheet, "C1") == 500


def test_a_formula_it_cannot_read_is_refused_by_name() -> None:
    """A function the reader does not know fails the case loudly, never as a wrong number."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["A1"] = "=AVERAGE(B1:B2)"

    with pytest.raises(ValueError, match="AVERAGE"):
        cell_number(sheet, "A1")
