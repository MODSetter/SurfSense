"""Live case 10: the user's own budget workbook with a chart, revised: one line raised by 10% and the total fixed.

The agent edits cells in a copy; the chart and the user's file stay as they were.
"""

import io
import zipfile

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference

from tests.live.live_agent import LiveAgent
from tests.live.revised_files import cell_number
from tests.live.revised_versions import last_revised, source_bytes

pytestmark = pytest.mark.live

CASE = "budget-revision"
TURN = (
    "In my 2027 budget workbook, raise the Marketing line by 10% and fix the "
    "total: it does not add up all the lines."
)
# Line: 2027 budget in EUR. Marketing is row 4; the total in row 8 leaves out Training.
LINES = {
    "Salaries": 410_000,
    "Rent": 64_000,
    "Marketing": 48_000,
    "Travel": 22_000,
    "Software": 18_500,
    "Training": 9_500,
}
RAISED = 48_000 * 1.1
TOTAL = sum(LINES.values()) - LINES["Marketing"] + RAISED


async def test_the_agent_revises_a_copy_of_the_users_budget(live: LiveAgent) -> None:
    """The copy holds the raised line and a total of every line; its chart is the source's, byte for byte."""
    original = _budget()
    budget = await live.upload("Kestrel Studio budget 2027.xlsx", original)
    await live.wait_ready(budget)
    thread = await live.thread()

    frames = await live.turn(thread, TURN)
    copy = await last_revised(live, frames, "turn 1")
    assert copy["format"] == "xlsx", f"made {copy['format']}, not an Excel copy"
    assert copy["revision"]["derived_from_document_id"] == budget
    revised = await live.file(copy["id"])

    charts = _charts(original)
    assert charts, "the source has no chart part"
    assert _charts(revised) == charts, "the chart parts changed"

    sheet = openpyxl.load_workbook(io.BytesIO(revised))["Budget"]
    assert sheet["A4"].value == "Marketing", "the lines moved"
    assert cell_number(sheet, "B4") == pytest.approx(RAISED), (
        f"Marketing holds {sheet['B4'].value!r}"
    )
    assert cell_number(sheet, "B8") == pytest.approx(TOTAL), (
        f"the total holds {sheet['B8'].value!r}"
    )
    for row, line in enumerate(LINES, start=2):
        if line != "Marketing":
            assert sheet[f"B{row}"].value == LINES[line], f"{line} changed"

    assert source_bytes(live, budget) == original, "the user's workbook was changed"


def _charts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return {
            name: package.read(name)
            for name in package.namelist()
            if name.startswith("xl/charts/")
        }


def _budget() -> bytes:
    """A small studio's budget, its total one row short, charted beside the lines."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Line", "2027 budget (EUR)"])
    for line, amount in LINES.items():
        sheet.append([line, amount])
    sheet.append(["Total", "=SUM(B2:B6)"])
    chart = BarChart()
    chart.title = "2027 budget by line"
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=7), titles_from_data=True
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=7))
    sheet.add_chart(chart, "D2")
    notes = book.create_sheet("Notes")
    notes.append(["Approved by the partners on 14 September 2026."])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
