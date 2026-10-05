"""Live case 4: a PDF from the first turn, then squeezed to one page with its totals in bold.

The demo starts in Word; this starts in ReportLab and edits a layout property.
"""

import io
import re

import docx
import pytest

from tests.live.live_agent import LiveAgent, answer
from tests.live.pdf_file import PdfFile
from tests.live.turn_renders import (
    assert_next_version,
    assert_pages_checked,
    last_version,
)

pytestmark = pytest.mark.live

CASE = "pdf-brief"
TURNS = (
    "Write a client brief for Aurora Dental Group as a PDF from these sources, "
    "with a table of the work packages and their costs.",
    "Make it fit on one page and bold the totals.",
)
WORK_PACKAGES = {
    "Discovery workshops": 6_500,
    "Booking data migration": 14_800,
    "SMS reminder integration": 8_200,
    "Staff training in three clinics": 4_900,
    "Project management": 5_400,
}
ONE_OFF_TOTAL = sum(WORK_PACKAGES.values())  # 39,800
_CLAIMS_NO_FORMAT = re.compile(
    r"(didn't|did not|don't|not) (name|specify|say|give|ask for) (a|the|which) format",
    re.IGNORECASE,
)


async def test_the_agent_drafts_a_pdf_brief_then_fits_it_to_one_page(
    live: LiveAgent,
) -> None:
    """v1 is a PDF; v2 is its next version, one page, with the one-off total set in bold."""
    scope = await live.upload("Aurora scope call.docx", _scope_call())
    background = await live.note("Aurora Dental Group background", _BACKGROUND)
    await live.wait_ready(scope, background)
    thread = await live.thread()

    first = await live.turn(thread, TURNS[0])
    draft = await last_version(live, first, "turn 1")
    assert draft.format == "pdf", f"turn 1 made {draft}, not a PDF"
    said = answer(first)
    assert not _CLAIMS_NO_FORMAT.search(said), (
        f"turn 1 says no format was named, though the user asked for a PDF: {said}"
    )
    await assert_pages_checked(live, first, "turn 1")

    second = await live.turn(thread, TURNS[1])
    edit = await last_version(live, second, "turn 2")
    assert_next_version(draft, edit, "turn 2")
    await assert_pages_checked(live, second, "turn 2")

    pdf = PdfFile(await live.file(edit.id))
    assert pdf.pages == 1, f"v{edit.number} has {pdf.pages} pages"
    bold = [_digits(run) for run in pdf.bold_text]
    assert any(str(ONE_OFF_TOTAL) in run for run in bold), (
        f"v{edit.number} does not set the {ONE_OFF_TOTAL:,} total in bold: "
        f"{pdf.bold_text}"
    )


def _digits(text: str) -> str:
    """'39,800 EUR' and 'EUR 39 800' both read as 39800."""
    return re.sub(r"(?<=\d)[,.\s](?=\d{3})", "", text)


_BACKGROUND = (
    "Aurora Dental Group runs three clinics in Tampere, Lahti and Jyvaskyla, with "
    "41 dentists and hygienists and about 28,000 active patients. Bookings live "
    "in an old on-premise system, DentBook 6, whose vendor ends support in March "
    "2027. Patients book by phone; the front desks spend about 30 hours a week "
    "on calls that are only reminders or reschedules. Missed appointments run at "
    "9% of all bookings, which the finance lead, Sanna Virtanen, puts at 310,000 "
    "EUR of lost chair time a year. Aurora wants online booking, SMS reminders "
    "and one patient record across the three clinics before DentBook's support "
    "ends. Their CEO, Juha Lehtonen, signs anything over 25,000 EUR."
)


def _scope_call() -> bytes:
    """Our notes from the scope call: the work packages, their prices and the dates."""
    notes = docx.Document()
    notes.add_heading("Scope call: Aurora Dental Group, 29 September 2026", 1)
    notes.add_paragraph(
        "Attendees: Sanna Virtanen (finance lead, Aurora), Mikko Salo (IT, Aurora), "
        "Priya Nair (engagement lead, Kestrel Data Works)."
    )
    notes.add_heading("Work packages and prices we agreed to quote", 2)
    descriptions = {
        "Discovery workshops": "two days per clinic, map booking flows and data",
        "Booking data migration": "move 28,000 patient records out of DentBook 6",
        "SMS reminder integration": "48-hour and 2-hour reminders with reply-to-cancel",
        "Staff training in three clinics": "front desk and clinicians, half a day each",
        "Project management": "weekly status, steering group every two weeks",
    }
    for name, price in WORK_PACKAGES.items():
        notes.add_paragraph(
            f"{name}: {price:,} EUR, {descriptions[name]}.", style="List Bullet"
        )
    notes.add_paragraph(
        "Hosting and support after go-live: 1,250 EUR per month, billed quarterly."
    )
    notes.add_heading("Timeline", 2)
    notes.add_paragraph(
        "Discovery in November 2026, migration and SMS build in December and "
        "January, training in the second half of January, go-live on 1 February "
        "2027, well before DentBook support ends in March."
    )
    notes.add_heading("Risks they raised", 2)
    for point in (
        "Duplicate patient records across clinics, maybe 6% of the total.",
        "Front desk staff worry the SMS replies will flood them.",
        "Lahti clinic has weak broadband; it needs a backup line.",
    ):
        notes.add_paragraph(point, style="List Bullet")
    out = io.BytesIO()
    notes.save(out)
    return out.getvalue()
