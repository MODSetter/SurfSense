"""Live case: with Office support on, a Word summary converted to PDF, then a workbook whose total LibreOffice computed.

Office support runs this machine's LibreOffice through the test seam. The quote
gives quantities and unit prices only, so no total is there to copy.
"""

import io
from typing import Any

import openpyxl
import pytest

from modules.artifacts.models import Artifact
from shared.db import create_session_factory
from tests.installed_office import installed_office_on
from tests.live.live_agent import LiveAgent, steps
from tests.live.pdf_file import PdfFile
from tests.live.turn_renders import (
    Version,
    assert_pages_checked,
    last_version,
    pages_left_unread,
    previews_of,
    sees_pages,
)

pytestmark = pytest.mark.live

CASE = "office-conversion"
TURNS = (
    "Write a one-page Word summary of the Fjellstue catering quote for the team.",
    "Convert it to PDF so I can email it.",
    "Now make an Excel workbook of the quote's line items: quantity, unit price "
    "and amount for each, and a total row that adds up the amounts with a formula.",
)
CONVERT = "surfsense_convert_document"
# Item: quantity, unit price in NOK excluding VAT.
LINES = {
    "Lunch buffet (2 days)": (46, 285),
    "Three-course dinner": (23, 640),
    "Coffee and fruit breaks": (92, 75),
    "Allergy and vegetarian surcharge": (9, 60),
    "Meeting room hire (per day)": (2, 3_400),
    "Minibus from Oppdal station (per trip)": (2, 1_850),
}
TOTAL = sum(quantity * price for quantity, price in LINES.values())


@pytest.fixture
def office(monkeypatch: pytest.MonkeyPatch, no_office_support: None) -> None:
    """Turned on after the suite's default of off."""
    installed_office_on(monkeypatch)


async def test_the_agent_converts_its_word_file_and_lets_libreoffice_total_the_workbook(
    office: None, live: LiveAgent
) -> None:
    """The PDF is a new document with pages; the workbook's saved total is the quote's sum."""
    quote = await live.note("Fjellstue catering quote, November offsite", _quote())
    await live.wait_ready(quote)
    thread = await live.thread()

    first = await live.turn(thread, TURNS[0])
    draft = await last_version(live, first, "turn 1")
    assert draft.format == "docx", f"turn 1 made {draft}, not a Word file"
    await assert_pages_checked(live, first, "turn 1")

    second = await live.turn(thread, TURNS[1])
    converted = [s for s in steps(second, CONVERT) if s["status"] == "completed"]
    assert converted, f"turn 2 did not call {CONVERT}"
    call = converted[-1]
    assert call["input"].get("artifact_id") == draft.id, (
        f"turn 2 converted {call['input']}, not artifact {draft.id}"
    )
    assert call.get("artifact"), f"turn 2's conversion made no PDF: {call}"
    pdf = await _version(live, call["artifact"]["id"])
    assert pdf.format == "pdf" and pdf.root != draft.root, (
        f"turn 2 made {pdf}, not a new PDF beside {draft}"
    )
    pages = PdfFile(await live.file(pdf.id))
    assert pages.pages >= 1, f"{pdf} has no pages"
    assert "Fjellstue" in pages.text, f"{pdf} does not hold the summary"
    if sees_pages(live):
        unread = pages_left_unread(
            second, pdf, previews_of(live, pdf), live.proxy.exchanges
        )
        assert not unread, f"turn 2: page(s) {unread} of {pdf} never reached the agent"

    third = await live.turn(thread, TURNS[2])
    book = await last_version(live, third, "turn 3")
    assert book.format == "xlsx", f"turn 3 made {book}, not a workbook"
    recalculation = _metadata(live, book.id).get("recalculation") or {}
    assert str(recalculation.get("by") or "").startswith("LibreOffice"), (
        f"{book} was not recalculated: {recalculation}"
    )
    data = await live.file(book.id)
    totals = _formula_values(data)
    assert any(abs(_number(v) - TOTAL) < 0.01 for v in totals.values()), (
        f"no formula in {book} saves the total {TOTAL}: {totals}"
    )


async def _version(live: LiveAgent, artifact_id: int) -> Version:
    """The ready version an artifact id names."""
    (made,) = [a for a in await live.artifacts() if a["id"] == artifact_id]
    assert made["status"] == "ready" and made["version"], f"{made} is not ready"
    return Version(
        made["id"],
        made["format"],
        made["version"]["root_id"],
        made["version"]["number"],
    )


def _metadata(live: LiveAgent, artifact_id: int) -> dict[str, Any]:
    with create_session_factory(live.engine)() as session:
        artifact = session.get(Artifact, artifact_id)
        assert artifact is not None
        return dict(artifact.artifact_metadata or {})


def _formula_values(data: bytes) -> dict[str, Any]:
    """Each formula, by sheet and cell, with the value the file saves for it."""
    formulas = openpyxl.load_workbook(io.BytesIO(data))
    saved = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    return {
        f"{sheet.title}!{cell.coordinate} {cell.value}": saved[sheet.title][
            cell.coordinate
        ].value
        for sheet in formulas.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }


def _number(value: Any) -> float:
    return float(value) if isinstance(value, int | float) else float("nan")


def _quote() -> str:
    """The caterer's quote as the office manager pasted it: quantities and prices, no total."""
    lines = "\n".join(
        f"- {item}: {quantity} x {price:,} NOK"
        for item, (quantity, price) in LINES.items()
    )
    return (
        "Quote from Fjellstue Kitchen, Oppdal, for the Kestrel Data Works team "
        "offsite on 12-13 November 2026, 23 people. Prices are per unit, "
        "excluding 25% VAT.\n\n"
        f"{lines}\n\n"
        "Valid until 15 October 2026. A 30% deposit is due on booking; the rest "
        "within 14 days of the event. Cancellation after 1 November is charged "
        "in full. Contact: Ingrid Solbakken, bookings@fjellstue.example."
    )
