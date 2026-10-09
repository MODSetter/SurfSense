"""Live case 9: the user's own Word contract, redlined: payment term shortened, liability cap tightened, a comment on why.

The agent revises a copy as tracked changes; the user's file is never written.
"""

import io
import re

import docx
import pytest

from tests.live.live_agent import LiveAgent
from tests.live.revised_files import RedlinedWord
from tests.live.revised_versions import (
    decided,
    download,
    last_revised,
    source_bytes,
)

pytestmark = pytest.mark.live

CASE = "contract-redline"
TURN = (
    "Redline my services agreement: change the payment term from 60 to 30 days, "
    "tighten the liability cap from 24 to 12 months of fees, and add a comment "
    "explaining why."
)
PAYMENT = "Each invoice is payable within sixty (60) days of the date of receipt"
CAP = (
    "the Supplier's total liability under this Agreement shall not exceed the "
    "fees paid by the Customer in the twenty-four (24) months"
)


async def test_the_agent_redlines_a_copy_of_the_users_contract(live: LiveAgent) -> None:
    """A revised copy holds SurfSense's tracked changes and a comment; accepting gives the new terms, rejecting the old, and the user's file is untouched."""
    original = _agreement()
    contract = await live.upload("Harbour Logistics services agreement.docx", original)
    await live.wait_ready(contract)
    thread = await live.thread()

    frames = await live.turn(thread, TURN)
    copy = await last_revised(live, frames, "turn 1")
    assert copy["format"] == "docx", f"made {copy['format']}, not a Word copy"
    assert copy["revision"]["derived_from_document_id"] == contract
    redline = RedlinedWord(await live.file(copy["id"]))
    authors = {c.author for c in redline.insertions + redline.deletions}
    assert redline.insertions and redline.deletions, (
        f"v{copy['version']['number']} has {len(redline.insertions)} insertions and "
        f"{len(redline.deletions)} deletions"
    )
    assert authors == {"SurfSense"}, f"changes by {authors}"
    assert any(c.author == "SurfSense" and c.text.strip() for c in redline.comments), (
        f"no comment by SurfSense: {redline.comments}"
    )

    accepted = RedlinedWord(await download(live, copy["id"], "clean"))
    assert not (accepted.insertions or accepted.deletions or accepted.comments)
    payment = _paragraph(accepted, "payable")
    assert _says(payment, "30", "thirty") and not _says(payment, "60", "sixty"), payment
    cap = _paragraph(accepted, "total liability")
    assert _says(cap, "12", "twelve") and not _says(cap, "24", "twenty-four"), cap

    rejected = RedlinedWord(
        await live.file((await decided(live, copy["id"], "reject-all"))["id"])
    )
    assert not (rejected.insertions or rejected.deletions)
    assert rejected.paragraphs == RedlinedWord(original).paragraphs

    assert source_bytes(live, contract) == original, "the user's contract was changed"


def _paragraph(word: RedlinedWord, words: str) -> str:
    found = [p for p in word.paragraphs if words in p.lower()]
    assert len(found) == 1, f"paragraphs with {words!r}: {found}"
    return found[0]


def _says(text: str, number: str, spelled: str) -> bool:
    return bool(re.search(rf"\b{number}\b", text)) or spelled in text.lower()


def _agreement() -> bytes:
    """A short services agreement as a customer's lawyer would send it."""
    agreement = docx.Document()
    agreement.add_heading("Logistics Services Agreement", 0)
    agreement.add_paragraph(
        "This Agreement is made on 1 September 2026 between Harbour Logistics Oy "
        "(the Supplier) and Pellervo Foods Ltd (the Customer)."
    )
    agreement.add_heading("1. Services", 1)
    agreement.add_paragraph(
        "The Supplier shall store, pick and deliver the Customer's chilled goods "
        "from its Kotka warehouse to the Customer's stores in southern Finland, "
        "as set out in Schedule 1."
    )
    agreement.add_heading("2. Fees and payment", 1)
    agreement.add_paragraph(
        "The Supplier shall invoice the fees in Schedule 2 monthly in arrears. "
        f"{PAYMENT}. Late payments bear interest at the rate set by the Finnish "
        "Interest Act."
    )
    agreement.add_heading("3. Liability", 1)
    agreement.add_paragraph(
        f"Except for death, personal injury, fraud or wilful misconduct, {CAP} "
        "before the event giving rise to the claim."
    )
    agreement.add_paragraph(
        "Neither party is liable for loss of profit, loss of business or any "
        "indirect or consequential loss."
    )
    agreement.add_heading("4. Term", 1)
    agreement.add_paragraph(
        "This Agreement runs for three years from its date and renews for one "
        "year at a time unless either party gives six months' written notice."
    )
    out = io.BytesIO()
    agreement.save(out)
    return out.getvalue()
