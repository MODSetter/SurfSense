"""The frozen worker runs document scripts with every library a script is told it has,
and decodes a recording in its decode child.

A script is not analysed when the worker is frozen, so whatever it imports, or
a library imports for it by name, ships only if the spec says so; the decode
child's miniaudio loads a cffi backend by name too.

The frozen binary comes from tests/packaging/frozen_backend.py: the release
build's, or one frozen from the spec.
"""

import io
import sys
from pathlib import Path

import docx
import openpyxl
import pypdfium2
import pytest

from tests.packaging.frozen_backend import frozen_backend
from worker.audio_decode.run import decode_in_child
from worker.document_script.analysis_run import AnalysisInput, run_analysis_script
from worker.document_script.run import run_document_script

pytestmark = pytest.mark.packaging

BACKEND = Path(__file__).resolve().parents[2]

WORD_WITH_A_CHART = """\
import os
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm

figure, axes = plt.subplots()
axes.bar(["2024", "2025"], [3, 5])
figure.savefig("chart.png")
document = Document()
document.add_heading("Yearly costs", level=1)
document.add_picture("chart.png", width=Cm(12))
document.save(os.environ["OUTPUT_PATH"])
"""

# savefig picks its canvas by format, importing that backend by name.
PDF_OF_VECTOR_CHARTS = """\
import os
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages

figure, axes = plt.subplots()
axes.plot([1, 2, 3], [2, 4, 3])
figure.savefig("chart.svg")
figure.savefig("chart.pdf")
with PdfPages(os.environ["OUTPUT_PATH"]) as pages:
    pages.savefig(figure)
"""


@pytest.fixture(scope="module")
def frozen_worker(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """The backend frozen from bundling/backend.spec, as the installer ships it."""
    return frozen_backend(tmp_path_factory.mktemp("frozen-backend"))


@pytest.fixture
def as_frozen(frozen_worker: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """The runner starts the frozen binary in script mode, as a packaged app does."""
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "executable", str(frozen_worker))


@pytest.mark.usefixtures("as_frozen")
def test_a_frozen_worker_places_a_matplotlib_chart_in_a_word_file() -> None:
    """matplotlib draws on Agg with no display, and python-docx places the PNG."""
    result = run_document_script(
        WORD_WITH_A_CHART, output_name="document.docx", images={}
    )

    assert result.ok, result.traceback_tail
    assert len(docx.Document(io.BytesIO(result.output or b"")).inline_shapes) == 1


@pytest.mark.usefixtures("as_frozen")
def test_a_frozen_worker_saves_vector_charts_and_matplotlib_pdfs() -> None:
    """A chart saved as SVG or PDF, or a PDF made with PdfPages, works when frozen."""
    result = run_document_script(
        PDF_OF_VECTOR_CHARTS, output_name="document.pdf", images={}
    )

    assert result.ok, result.traceback_tail
    pdf = pypdfium2.PdfDocument(result.output)
    try:
        assert len(pdf) == 1
    finally:
        pdf.close()


DECK_AND_WORKBOOKS = """\
import os
import openpyxl
import xlsxwriter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

book = xlsxwriter.Workbook("costs.xlsx")
book.add_worksheet("Costs").write_formula(0, 0, "=1+1")
book.close()
assert openpyxl.load_workbook("costs.xlsx")["Costs"]["A1"].value == "=1+1"

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[5])
data = CategoryChartData()
data.categories = ["2024", "2025"]
data.add_series("Cost", (3, 5))
slide.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data
)
deck.save(os.environ["OUTPUT_PATH"])
"""


@pytest.mark.usefixtures("as_frozen")
def test_a_frozen_worker_writes_decks_and_workbooks() -> None:
    """python-pptx's templates, xlsxwriter and openpyxl all ship in the frozen worker."""
    result = run_document_script(
        DECK_AND_WORKBOOKS, output_name="document.pptx", images={}
    )

    assert result.ok, result.traceback_tail
    assert (result.output or b"").startswith(b"PK")


ANALYSIS_OF_A_WORKBOOK = """\
import os
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

stock = pd.read_excel(os.path.join(os.environ["INPUT_DIR"], "Stock.xlsx"))
print(int(np.sum(stock["units"])))
stock.to_csv(os.path.join(os.environ["OUTPUT_DIR"], "stock.csv"), index=False)
stock.plot(x="product", y="units", kind="bar")
plt.savefig(os.path.join(os.environ["OUTPUT_DIR"], "stock.png"))
"""


@pytest.mark.usefixtures("as_frozen")
def test_a_frozen_worker_analyses_a_workbook_with_pandas(tmp_path: Path) -> None:
    """pandas reads .xlsx through openpyxl and plots through matplotlib, all by name."""
    book = openpyxl.Workbook()
    book.active.append(["product", "units"])
    book.active.append(["Crate", 12])
    book.save(tmp_path / "Stock.xlsx")
    source = AnalysisInput(tmp_path / "Stock.xlsx", "Stock.xlsx", 1, "Stock")

    run = run_analysis_script(ANALYSIS_OF_A_WORKBOOK, [source], tmp_path / "kept")

    assert run.ok, run.traceback_tail
    assert run.stdout == "12\n"
    assert run.kept == ("stock.csv", "stock.png")


@pytest.mark.usefixtures("as_frozen")
def test_a_frozen_worker_decodes_a_recording_in_its_child(tmp_path: Path) -> None:
    """worker --decode-audio reads an MP3 to 16 kHz mono, as ingest starts it."""
    tone = BACKEND / "tests" / "fixtures" / "audio" / "tone.mp3"
    out = tmp_path / "audio16k.raw"

    outcome = decode_in_child(tone, out, between_polls=lambda: None)

    assert outcome.kind == "decoded", outcome
    assert out.stat().st_size == pytest.approx(16000 * 2, rel=0.05)
