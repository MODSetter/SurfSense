"""The Word builder: python-docx's default template and its built-in styles (decision 12)."""

from collections.abc import Mapping
from io import BytesIO
from pathlib import Path

import docx
from docx.document import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, RGBColor
from docx.text.paragraph import Paragraph as WordParagraph
from PIL import Image

from worker.studio.office.markdown.blocks import (
    Block,
    ChartBlock,
    Code,
    Figure,
    Heading,
    ListBlock,
    Paragraph,
    Quote,
    Rule,
    Spans,
    Table,
)
from worker.studio.office.markdown.chart import draw_chart
from worker.studio.office.markdown.links import safe_href
from worker.studio.office.markdown.parse import parse_markdown

# The template's list styles stop at a third level.
_LIST_DEPTH = 3
_CODE_FONT = "Courier New"
_LINK_BLUE = RGBColor(0x05, 0x63, 0xC1)
_SCREEN_DPI = 96


def markdown_to_word(markdown: str, figures: Mapping[str, Path]) -> bytes:
    """The .docx a Markdown spec describes; `figures` maps a figure name to its PNG."""
    document = docx.Document()
    _Writer(document, figures).blocks(parse_markdown(markdown), depth=0)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _Writer:
    def __init__(self, document: Document, figures: Mapping[str, Path]) -> None:
        self._document = document
        self._figures = figures
        section = document.sections[0]
        self._width = Emu(
            section.page_width - section.left_margin - section.right_margin
        )

    def blocks(self, blocks: list[Block] | tuple[Block, ...], depth: int) -> None:
        for block in blocks:
            self.block(block, depth)

    def block(self, block: Block, depth: int, style: str | None = None) -> None:
        match block:
            case Heading(level, spans):
                self._spans(self._document.add_heading("", min(level, 9)), spans)
            case Paragraph(spans):
                self._spans(self._document.add_paragraph(style=style), spans)
            case ListBlock():
                self._list(block, depth + 1)
            case Table():
                self._table(block)
            case Figure():
                self._figure(block)
            case ChartBlock():
                self._chart(block)
            case Code(text):
                run = self._document.add_paragraph(style="No Spacing").add_run(text)
                run.font.name = _CODE_FONT
            case Quote(inner):
                for child in inner:
                    self.block(child, depth, style="Quote")
            case Rule():
                _bottom_border(self._document.add_paragraph())

    def _list(self, block: ListBlock, depth: int) -> None:
        level = min(depth, _LIST_DEPTH)
        kind = "List Number" if block.ordered else "List Bullet"
        first = kind if level == 1 else f"{kind} {level}"
        rest = "List Continue" if level == 1 else f"List Continue {level}"
        for item in block.items:
            styled = False
            for child in item:
                if isinstance(child, Paragraph):
                    self.block(child, depth, style=rest if styled else first)
                    styled = True
                else:
                    self.block(child, depth)

    def _table(self, table: Table) -> None:
        columns = max(len(row) for row in (table.header, *table.rows))
        if columns == 0:
            return
        grid = self._document.add_table(rows=0, cols=columns)
        grid.style = "Table Grid"
        rows = [table.header, *table.rows] if table.header else list(table.rows)
        for index, row in enumerate(rows):
            cells = grid.add_row().cells
            for column, spans in enumerate(row):
                runs = self._spans(cells[column].paragraphs[0], spans)
                if index == 0 and table.header:
                    for run in runs:
                        run.bold = True

    def _figure(self, figure: Figure) -> None:
        path = self._figures.get(figure.name) if figure.name else None
        if path is None:
            self._document.add_paragraph().add_run(figure.caption).italic = True
            return
        self._picture(path.read_bytes())
        self._document.add_paragraph(figure.caption, style="Caption")

    def _chart(self, block: ChartBlock) -> None:
        if block.chart is not None:
            self._picture(draw_chart(block.chart))
            return
        self._document.add_paragraph().add_run(block.note).italic = True
        if block.fallback is not None:
            self._table(block.fallback)

    def _picture(self, png: bytes) -> None:
        """At its own size, or the page's width when it is wider."""
        with Image.open(BytesIO(png)) as image:
            dpi = image.info.get("dpi", (_SCREEN_DPI, _SCREEN_DPI))[0] or _SCREEN_DPI
            natural = Emu(int(image.width / float(dpi) * 914400))
        self._document.add_picture(BytesIO(png), width=min(natural, self._width))

    def _spans(self, paragraph: WordParagraph, spans: Spans) -> list:
        runs = []
        for span in spans:
            href = safe_href(span.href)
            run = paragraph.add_run(span.text)
            run.bold = span.bold or None
            run.italic = span.italic or None
            run.font.strike = span.strike or None
            if span.code:
                run.font.name = _CODE_FONT
            if href is not None:
                _link(paragraph, run, href)
            runs.append(run)
        return runs


def _link(paragraph: WordParagraph, run: object, href: str) -> None:
    """Move the run inside a w:hyperlink: python-docx has no call that writes one."""
    relation = paragraph.part.relate_to(
        href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run.font.color.rgb = _LINK_BLUE  # type: ignore[attr-defined]
    run.font.underline = True  # type: ignore[attr-defined]
    hyperlink.append(run._r)  # type: ignore[attr-defined]
    paragraph._p.append(hyperlink)


def _bottom_border(paragraph: WordParagraph) -> None:
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (
        ("val", "single"),
        ("sz", "6"),
        ("space", "1"),
        ("color", "auto"),
    ):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)
