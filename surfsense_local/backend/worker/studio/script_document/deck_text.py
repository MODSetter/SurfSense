"""The text of a PowerPoint deck a script wrote: each slide's title, text, tables and notes in order."""

from collections.abc import Iterable, Iterator
from io import BytesIO

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from worker.studio.script_document.extracted_text import UnreadableDocumentError


def deck_text(data: bytes) -> str:
    """One block per slide, headed `## Slide <n>: <title>`, in slide order."""
    try:
        deck = pptx.Presentation(BytesIO(data))
    # A file that is not a PowerPoint package fails as a zip, package, key or value error.
    except Exception as error:
        raise UnreadableDocumentError from error
    return "\n\n".join(
        _slide_text(number, slide) for number, slide in enumerate(deck.slides, 1)
    )


def _slide_text(number: int, slide: Slide) -> str:
    title_shape = slide.shapes.title
    title = _one_line(title_shape.text_frame.text) if title_shape is not None else ""
    lines = [f"## Slide {number}: {title}" if title else f"## Slide {number}"]
    lines += _shape_lines(
        shape
        for shape in slide.shapes
        if title_shape is None or shape.shape_id != title_shape.shape_id
    )
    notes = _notes(slide)
    if notes:
        lines.append(f"Notes: {notes}")
    return "\n".join(lines)


def _shape_lines(shapes: Iterable[BaseShape]) -> Iterator[str]:
    """A shape's own lines; a group's shapes in their order; a picture or chart as a marker."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_lines(shape.shapes)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = (cell for cell in row.cells if not cell.is_spanned)
                yield " | ".join(_one_line(cell.text) for cell in cells)
        elif getattr(shape, "has_chart", False):
            chart = shape.chart
            title = (
                _one_line(chart.chart_title.text_frame.text) if chart.has_title else ""
            )
            yield f"[chart: {title}]" if title else "[chart]"
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield "[picture]"
        elif shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = _one_line(paragraph.text)
                if text:
                    yield text


def _notes(slide: Slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return _one_line(frame.text) if frame is not None else ""


def _one_line(text: str) -> str:
    """Line breaks inside a cell or title would read as the next item."""
    return " ".join(text.split())
