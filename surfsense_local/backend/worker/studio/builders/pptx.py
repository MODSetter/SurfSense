from io import BytesIO

from worker.studio.builders.types import Builder, Built, Source
from worker.studio.builders.util import as_list, as_text, parse_json, slug

MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

_SCHEMA = (
    'Return only JSON, no prose: {"title": str, "slides": '
    '[{"title": str, "bullets": [str]}]}.'
)


def prompt(_sources: list[Source], user_prompt: str | None) -> str:
    focus = f" Emphasise: {user_prompt}." if user_prompt else ""
    return (
        "Turn the sources below into a slide deck, using their facts only. Keep "
        "each bullet short." + focus + " " + _SCHEMA
    )


def build(raw: str, _sources: list[Source]) -> Built:
    spec = parse_json(raw)
    title = as_text(spec.get("title")) or "Presentation"

    from pptx import Presentation

    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    lines = [f"# {title}"]

    for slide in as_list(spec.get("slides")):
        if not isinstance(slide, dict):
            continue
        heading = as_text(slide.get("title")) or "Slide"
        added = deck.slides.add_slide(deck.slide_layouts[1])
        added.shapes.title.text = heading
        lines.append(f"\n## {heading}")

        body = added.placeholders[1].text_frame
        first = True
        for bullet in as_list(slide.get("bullets")):
            text = as_text(bullet)
            if not text:
                continue
            paragraph = body.paragraphs[0] if first else body.add_paragraph()
            paragraph.text = text
            lines.append(f"- {text}")
            first = False

    buffer = BytesIO()
    deck.save(buffer)
    return Built(
        title=title,
        markdown="\n".join(lines),
        primary=buffer.getvalue(),
        primary_mime=MIME,
        primary_filename=f"{slug(title, 'deck')}.pptx",
    )


pptx = Builder(key="pptx", prompt=prompt, build=build)
