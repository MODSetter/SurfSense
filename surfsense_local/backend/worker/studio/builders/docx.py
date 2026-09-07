from io import BytesIO

from worker.studio.builders.types import Builder, Built, Source
from worker.studio.builders.util import as_list, as_text, parse_json, slug

MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_SCHEMA = (
    'Return only JSON, no prose: {"title": str, "sections": '
    '[{"heading": str, "paragraphs": [str]}]}.'
)


def prompt(_sources: list[Source], user_prompt: str | None) -> str:
    focus = f" Emphasise: {user_prompt}." if user_prompt else ""
    return (
        "Write a structured document from the sources below, using their facts "
        "only." + focus + " " + _SCHEMA
    )


def build(raw: str, _sources: list[Source]) -> Built:
    spec = parse_json(raw)
    title = as_text(spec.get("title")) or "Document"

    from docx import Document as Docx

    document = Docx()
    document.add_heading(title, level=0)
    lines = [f"# {title}"]

    for section in as_list(spec.get("sections")):
        if not isinstance(section, dict):
            continue
        heading = as_text(section.get("heading"))
        if heading:
            document.add_heading(heading, level=1)
            lines.append(f"\n## {heading}")
        for paragraph in as_list(section.get("paragraphs")):
            text = as_text(paragraph)
            if text:
                document.add_paragraph(text)
                lines.append(text)

    buffer = BytesIO()
    document.save(buffer)
    return Built(
        title=title,
        markdown="\n".join(lines),
        primary=buffer.getvalue(),
        primary_mime=MIME,
        primary_filename=f"{slug(title, 'document')}.docx",
    )


docx = Builder(key="docx", prompt=prompt, build=build)
